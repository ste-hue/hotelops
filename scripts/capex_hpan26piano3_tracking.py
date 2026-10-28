"""Costruisce HPAN26PIANO3_Tracking.xlsx — cockpit CapEx del terzo piano (skill progetto-capex).

    python scripts/capex_hpan26piano3_tracking.py              # costruisce in output/
    python scripts/capex_hpan26piano3_tracking.py --upload     # e aggiorna il file su Drive IN PLACE

Perimetro: solo Piano Terzo, 20 camere (deciso 04/09/2026). Committente INTUR.
Convenzione importi: IMPONIBILE (IVA esclusa), come il budget banca.
Sorgenti (nostro Drive, cartella 1RO88…): Preventivi_PIANO3/, FATTURE_PIANO3/, 06_contrattualistica/.
I documenti sono dichiarati qui sotto (DOCS): un evento nuovo = una riga nuova, poi rilancia.
L'upload riscrive sempre lo stesso fileID: il link non cambia mai, niente _v2.
Non lanciarlo mentre il file è aperto in Excel.
"""

import io
import os
import sys
from datetime import date

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

PROJECT = "HPAN26PIANO3"
TRACKING_FILE_ID = "1dV0WZLJru2N7PSBqE5IBCqqw5vNqBIYw"
WRITER = "stefano@panoramagroup.it"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
FILE = "https://drive.google.com/file/d/{}/view"
FOLDER = "https://drive.google.com/drive/folders/{}"
AGGIORNATO = date(2026, 10, 28)

# Budget: richiesta a MPS (BP Solo Terzo MASTER 11/10, foglio Impieghi e fonti), netto IVA.
BUDGET = [
    ("Opere edili + impianti (mutuo edilizio SAL)", 1_458_742.00),
    ("Fit-out e FF&E 20 camere (finanziamento MCC)", 557_518.50),
]
BUDGET_INTERNO = 1_688_894.50
BUDGET_FONTE_ID = "17j8aKnoHNExa9QYyZBDiN87xXGnm8vSV"

# Fornitori: F-code locale al progetto; l'identità è nel registry globale.
FORNITORI = [
    (
        "F001",
        "A.M.C.N. Costruzioni S.r.l.",
        "OPERE",
        "Opere edili 20 camere P3 (a misura)",
    ),
    (
        "F002",
        "S.T.E. S.r.l. (Aniello Esposito)",
        "OPERE",
        "Impianti elettrici e speciali P3",
    ),
    (
        "F003",
        "Santelia (Luigi Santelia)",
        "OPERE",
        "Impianti idrico-sanitario, clima, antincendio",
    ),
    ("F004", "Progetto Clima", "OPERE", "VRF clima (preventivo di luglio)"),
    ("F005", "KONE S.p.A.", "OPERE", "2 ascensori lato Amalfi + montacarichi"),
    ("F006", "OTIS Servizi S.r.l.", "OPERE", "2 ascensori (alternativa a KONE)"),
    ("F007", "Centrogest", "FIT-OUT", "Casseforti 21 pz"),
    ("F008", "ALDEN", "SPESE TECNICHE", "Studio di fattibilità bando Green Tour"),
]

# Documenti. tipo: PREV / CONTR / FATT / STORICO (non conta) / RIF (riferimento, non conta).
# importo = imponibile. Una revisione sostituisce la precedente: la vecchia diventa STORICO.
DOCS = [
    (
        date(2026, 8, 29),
        "F001",
        "STORICO",
        "Offerta PIANI1-3",
        0,
        "1OeqWb8A1hkmZo6F_riX2BMZ6RsBqIAai",
        "Offerta su P1+P3 (33 camere): superata dal contratto Rev09",
    ),
    (
        date(2026, 10, 15),
        "F001",
        "CONTR",
        "PG-LEG-APP-01 Rev9",
        482_000.00,
        "18ipcwesHipk3RUdb4rkA18HyYEXwyVDc",
        "Firmato. A misura, incl. 15.000 sicurezza. Reverse charge. Anticipo 15%, SAL ogni 60k, ritenuta 10%. Fine 20/03/2027",
    ),
    (
        date(2026, 10, 20),
        "F001",
        "STORICO",
        "Proforma FPR 92/26",
        0,
        "1n8nDIp4dewJuzHWtxq2qlx2UHcO7sEk3",
        "Proforma acconto, sostituita dalla fattura",
    ),
    (
        date(2026, 10, 21),
        "F001",
        "STORICO",
        "Proforma FPR 92/26 REV",
        0,
        "1ANkxrdPi-bWIDkJp8m0yAM620q6YSC3Z",
        "Proforma acconto rev., sostituita dalla fattura",
    ),
    (
        date(2026, 10, 27),
        "F001",
        "FATT",
        "FPR 92/26",
        72_300.00,
        "17NkjBFoKc08JfiUaUyxQ2xD0oTiPbPRl",
        "Acconto 15% sul contratto. Reverse charge",
    ),
    (
        date(2026, 10, 6),
        "F002",
        "PREV",
        "Computo-offerta P3",
        162_656.18,
        "12iqglUOFrSyYZOtBY8fYzpSU91vEzXbI",
        "Firmata solo da STE (concorrente): non è un ordine",
    ),
    (
        date(2026, 9, 7),
        "F003",
        "PREV",
        "0066-26",
        198_000.00,
        "1fZjbeC9YvKdtwmBDAlagQw9I3Y0DCMPV",
        "21 camere, a corpo",
    ),
    (
        date(2026, 10, 16),
        "F003",
        "RIF",
        "Computo HP impianti meccanici",
        141_496.32,
        "1cOJQmoM_-opVTZreK9-P_v-TiTJhxqYk",
        "Computo Hospitality Project: da chiarire se sostituisce il preventivo 0066-26",
    ),
    (
        date(2026, 7, 16),
        "F004",
        "PREV",
        "VRF 6191257459",
        94_425.00,
        "17h7v8Kx87QABxdiTBA5yGjAq8BAHRF5y",
        "Prima della decisione di perimetro (04/09): verificare camere e sovrapposizione con Santelia",
    ),
    (
        date(2026, 10, 15),
        "F005",
        "PREV",
        "T-0010076131 v3",
        161_018.00,
        "111MV7HEINk44cCZqqz2TP2lOxn9XSJir",
        "2 ascensori lato Amalfi. Prezzo fisso fino al 30/06/2027",
    ),
    (
        date(2026, 10, 15),
        "F005",
        "PREV",
        "T-0010066110 v4",
        82_017.00,
        "1o0cQtpJhlZGtRd6OJ7UeLZQCrCL7dLif",
        "Montacarichi",
    ),
    (
        date(2026, 10, 14),
        "F006",
        "PREV",
        "ITK62585/01",
        103_400.00,
        "1uvUaZnjcOk2oWpxHtaHhHm2_6hD75mGX",
        "2 ascensori Gen3 Edge: alternativa all'offerta KONE T-0010076131",
    ),
    (
        date(2026, 9, 14),
        "F007",
        "PREV",
        "959",
        2_144.00,
        "1qetrunGikpgEwPjBHBwMYwFOYEBcC0Kj",
        "21 pezzi",
    ),
    (
        date(2026, 10, 15),
        "F008",
        "FATT",
        "267",
        5_100.00,
        "1Hq1iT12R5pmH2ykAWGbNLso32-OXyVva",
        "CUP J95J26000200008",
    ),
    (
        date(2026, 9, 23),
        "",
        "RIF",
        "CME Pisacane solo terzo piano",
        482_450.46,
        "1bbIFHciy7H54xYP5-IvIoQwDkUL0Szph",
        "Computo metrico DL: base del contratto AMCN, non è un costo",
    ),
]

OOS = [
    (
        "EdilContract S.r.l.",
        "Preventivo 05/09/2026",
        date(2026, 9, 5),
        966.00,
        "Intestato e firmato da ORTI, 2 tavoli + 2 sedie (Studio Bacci): non sono camere P3",
        "Arredi aree comuni ORTI",
        "1NIGrvpzCps-iNAXT7UAIYS3r7tusCasv",
    ),
    (
        "Da attribuire",
        "Livraison IT",
        None,
        None,
        "Documento arrivato nel progetto, non pertinente",
        "Da attribuire",
        "1sWHMWj0TQHz0ugsNBzIwEAnNLSorAX5W",
    ),
]

# ---------- stile ----------
BLU, GIALLO = "1F3864", "FFF2CC"
EUR = '#,##0.00\\ "€"'


def header(ws, row, cols, widths):
    for c, (name, w) in enumerate(zip(cols, widths, strict=True), start=1):
        cell = ws.cell(row=row, column=c, value=name)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=BLU)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = w
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def link(cell, url, text="Apri →"):
    cell.value = text
    cell.hyperlink = url
    cell.font = Font(color="0563C1", underline="single")


# ---------- calcoli per fornitore ----------
def per_fornitore(cod):
    docs = [d for d in DOCS if d[1] == cod]
    prev = sum(d[4] for d in docs if d[2] == "PREV")
    contr = sum(d[4] for d in docs if d[2] == "CONTR")
    fatt = sum(d[4] for d in docs if d[2] == "FATT")
    return {
        "n_prev": sum(d[2] == "PREV" for d in docs),
        "prev": prev,
        "n_contr": sum(d[2] == "CONTR" for d in docs),
        "contr": contr,
        "n_fatt": sum(d[2] == "FATT" for d in docs),
        "fatt": fatt,
    }


def status(k):
    impegno = k["contr"] > 0
    if not impegno and k["fatt"] == 0:
        return "SOLO PREVENTIVO" if k["prev"] else "SCOPERTO"
    if impegno and k["fatt"] == 0:
        return "SOLO IMPEGNO"
    if not impegno:
        return "SOLO COMPETENZA"
    return "COMPLETO"


# Pagato: dalle partite Esolver. Finché l'export non è caricato resta None (= non verificato).
PAGATO = {}

wb = openpyxl.Workbook()

# ══════════════ Summary ══════════════
ws = wb.active
ws.title = "Summary"
ws.column_dimensions["A"].width = 52
ws.column_dimensions["B"].width = 20
ws.column_dimensions["C"].width = 70
r = 1
ws.cell(row=r, column=1, value=f"{PROJECT} — COCKPIT ECONOMICO").font = Font(
    bold=True, size=14, color=BLU
)
r += 1
ws.cell(
    row=r,
    column=1,
    value="Ristrutturazione 20 camere Piano Terzo · Hotel Panorama, Maiori (SA)",
).font = Font(italic=True)
r += 2
for k, v, n in [
    (
        "Committente",
        "INTUR S.r.l.",
        "UNA sola società: tutte le fatture vanno intestate qui",
    ),
    (
        "Perimetro",
        "Solo Piano Terzo, 20 camere",
        "deciso 04/09/2026, confermato da Hospitality Project 07/09",
    ),
    (
        "Lavori",
        "19/10/2026 → 14/04/2027",
        "uscita impresa edile 20/03/2027 (contratto AMCN art. 12)",
    ),
    (
        "Convenzione importi",
        "IMPONIBILE (IVA esclusa)",
        "come il budget banca; AMCN è in reverse charge",
    ),
    ("PM / DL", "Hospitality Project / Ing. Pisacane", ""),
    (
        "Ultimo aggiornamento",
        AGGIORNATO.strftime("%d/%m/%Y"),
        "scripts/capex_hpan26piano3_tracking.py",
    ),
]:
    ws.cell(row=r, column=1, value=k).font = Font(bold=True)
    ws.cell(row=r, column=2, value=v)
    ws.cell(row=r, column=3, value=n).font = Font(italic=True, color="808080")
    r += 1

r += 1
ws.cell(row=r, column=1, value="BUDGET — richiesta a MPS (netto IVA)").font = Font(
    bold=True, color=BLU
)
link(
    ws.cell(row=r, column=3),
    FILE.format(BUDGET_FONTE_ID),
    "Fonte: BP Solo Terzo MASTER 11/10 →",
)
r += 1
for label, v in BUDGET:
    ws.cell(row=r, column=1, value=label)
    ws.cell(row=r, column=2, value=v).number_format = EUR
    r += 1
budget_tot = sum(v for _, v in BUDGET)
ws.cell(row=r, column=1, value="TOTALE BUDGET").font = Font(bold=True)
c = ws.cell(row=r, column=2, value=budget_tot)
c.number_format, c.font = EUR, Font(bold=True)
r += 1
ws.cell(row=r, column=1, value="Tetto interno (nota)")
ws.cell(row=r, column=2, value=BUDGET_INTERNO).number_format = EUR
ws.cell(
    row=r,
    column=3,
    value="stesso BP, versione interna: imprevisti 5%, zero contingency",
).font = Font(italic=True, color="808080")
r += 2

kk = {f[0]: per_fornitore(f[0]) for f in FORNITORI}
impegno = sum(k["contr"] for k in kk.values())
competenza = sum(k["fatt"] for k in kk.values())
pagato_noto = sum(v for v in PAGATO.values() if v is not None)
prev_aperti = sum(k["prev"] for k in kk.values() if k["contr"] == 0)

ws.cell(row=r, column=1, value="LE 3 DIMENSIONI").font = Font(bold=True, color=BLU)
r += 1
for label, v, n in [
    (
        "IMPEGNO — contratti e ordini firmati",
        impegno,
        "solo AMCN: gli altri sono ancora preventivi",
    ),
    ("COMPETENZA — fatturato", competenza, ""),
    (
        "CASSA — pagato (partite Esolver)",
        pagato_noto if PAGATO else None,
        "n.d.: serve l'export «Situazione partite fornitori» INTUR",
    ),
    ("Residuo budget dopo l'impegno", budget_tot - impegno, ""),
    (
        "Preventivi aperti, non impegnati",
        prev_aperti,
        "somma lorda: KONE e OTIS sugli ascensori sono alternative, non vanno sommate",
    ),
]:
    ws.cell(row=r, column=1, value=label)
    c = ws.cell(row=r, column=2, value=v if v is not None else "n.d.")
    c.number_format = EUR
    ws.cell(row=r, column=3, value=n).font = Font(italic=True, color="808080")
    r += 1

# ══════════════ Tracking_Fornitori ══════════════
ws = wb.create_sheet("Tracking_Fornitori")
COLS = [
    "Cod",
    "Ragione Sociale",
    "Macro",
    "Scope",
    "# Prev",
    "€ Prev",
    "# Contr",
    "€ Contr",
    "# Fatt",
    "€ Fatt",
    "Totale €",
    "Status",
    "Cartella F* (Drive)",
]
header(ws, 1, COLS, [7, 32, 15, 44, 7, 15, 8, 15, 7, 15, 15, 17, 16])
for i, (cod, nome, macro, scope) in enumerate(FORNITORI, start=2):
    k = kk[cod]
    impegno_f = k["contr"] or k["prev"]
    vals = [
        cod,
        nome,
        macro,
        scope,
        k["n_prev"],
        k["prev"],
        k["n_contr"],
        k["contr"],
        k["n_fatt"],
        k["fatt"],
        max(impegno_f, k["fatt"]),
        status(k),
    ]
    for c, v in enumerate(vals, start=1):
        cell = ws.cell(row=i, column=c, value=v)
        if c in (6, 8, 10, 11):
            cell.number_format = EUR
    ws.cell(row=i, column=13, value="(nasce col primo ordine)")

# ══════════════ Documents ══════════════
ws = wb.create_sheet("Documents")
COLS = [
    "Data",
    "Fornitore",
    "Cod F*",
    "Tipo",
    "Numero Doc",
    "€ Imponibile",
    "Conta?",
    "Link Drive",
    "Note",
]
header(ws, 1, COLS, [11, 30, 7, 9, 28, 15, 8, 10, 70])
nomi = {f[0]: f[1] for f in FORNITORI}
for i, (d, cod, tipo, num, imp, fid, note) in enumerate(
    sorted(DOCS, key=lambda x: (x[1] or "Z", x[0])), start=2
):
    vals = [
        d,
        nomi.get(cod, "Ing. Pisacane (DL)"),
        cod,
        tipo,
        num,
        imp or None,
        "sì" if tipo in ("PREV", "CONTR", "FATT") else "no",
    ]
    for c, v in enumerate(vals, start=1):
        cell = ws.cell(row=i, column=c, value=v)
        if c == 1:
            cell.number_format = "dd/mm/yyyy"
        if c == 6:
            cell.number_format = EUR
    link(ws.cell(row=i, column=8), FILE.format(fid))
    ws.cell(row=i, column=9, value=note).alignment = Alignment(wrap_text=True)

# ══════════════ Chiusura_Fornitori ══════════════
ws = wb.create_sheet("Chiusura_Fornitori")
COLS = [
    "Cod",
    "Fornitore",
    "Macro",
    "Scope",
    "€ Prev/Contratto",
    "€ Fatturato",
    "€ Pagato",
    "€ Da Pagare",
    "€ Da Fatturare",
    "Stato Operativo",
    "Lavoro Finito?",
    "Ultima Evidenza",
    "Chi Lo Sa",
    "Prossima Azione",
    "Note",
]
header(ws, 1, COLS, [7, 30, 15, 36, 16, 15, 13, 14, 15, 18, 10, 13, 12, 46, 40])
AZIONI = {
    "F001": (
        "Verificare il bonifico dell'acconto 72.300 (contratto: entro 10 gg dalla consegna cantiere del 19/10)",
        "Prossimo incasso AMCN: SAL a 60k, certificato DL",
    ),
    "F002": ("Decidere e firmare l'ordine (offerta 162.656 vs 180.000 nel BP)", ""),
    "F003": (
        "Chiarire se il computo HP 141.496 del 16/10 sostituisce il preventivo 198.000",
        "",
    ),
    "F004": (
        "Verificare se superato: preventivo di luglio, scope sovrapposto a Santelia",
        "",
    ),
    "F005": (
        "Scegliere KONE o OTIS per i 2 ascensori; montacarichi solo KONE",
        "Ascensori: verificare dove stanno nel BP (non nei 848k opere)",
    ),
    "F006": (
        "Scegliere KONE o OTIS per i 2 ascensori",
        "Alternativa a KONE T-0010076131 (161.018)",
    ),
    "F007": ("Ordinare se confermato", ""),
    "F008": ("Verificare il pagamento", "Spesa del bando Green Tour, non lavori"),
}
for i, (cod, nome, macro, scope) in enumerate(FORNITORI, start=2):
    k = kk[cod]
    base = k["contr"] or k["prev"]
    pagato = PAGATO.get(cod)
    da_pagare = k["fatt"] - pagato if pagato is not None else None
    da_fatt = max(base - k["fatt"], 0)
    if k["contr"] == 0 and k["fatt"] == 0:
        stato = "PREVENTIVO"
    elif k["fatt"] and da_fatt:
        stato = "LAVORI IN CORSO"
    elif da_fatt:
        stato = "DA FATTURARE"
    elif pagato is None or da_pagare:
        stato = "DA PAGARE"
    else:
        stato = "CHIUSO"
    ultima = max(d[0] for d in DOCS if d[1] == cod)
    azione, note = AZIONI[cod]
    vals = [
        cod,
        nome,
        macro,
        scope,
        base,
        k["fatt"],
        pagato if pagato is not None else "n.d.",
        da_pagare if da_pagare is not None else "n.d.",
        da_fatt,
        stato,
        None,
        ultima,
        None,
        azione,
        note,
    ]
    for c, v in enumerate(vals, start=1):
        cell = ws.cell(row=i, column=c, value=v)
        if c in (5, 6, 7, 8, 9):
            cell.number_format = EUR
        if c == 12:
            cell.number_format = "dd/mm/yyyy"
        if c in (14, 15):
            cell.alignment = Alignment(wrap_text=True)
    if stato != "PREVENTIVO":
        ws.cell(row=i, column=10).fill = PatternFill("solid", fgColor=GIALLO)

# ══════════════ OOS_Out_Of_Scope ══════════════
ws = wb.create_sheet("OOS_Out_Of_Scope")
COLS = [
    "Fornitore",
    "Documento",
    "Data",
    "€ Imponibile",
    "Perché è fuori perimetro",
    "Dove va invece",
    "Link Drive",
]
header(ws, 1, COLS, [24, 24, 11, 14, 60, 26, 10])
for i, (forn, doc, d, imp, perche, dove, fid) in enumerate(OOS, start=2):
    for c, v in enumerate([forn, doc, d, imp, perche, dove], start=1):
        cell = ws.cell(row=i, column=c, value=v)
        if c == 3:
            cell.number_format = "dd/mm/yyyy"
        if c == 4:
            cell.number_format = EUR
    link(ws.cell(row=i, column=7), FILE.format(fid))

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(REPO, "output", f"{PROJECT}_Tracking.xlsx")
os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)
print("SAVED", OUT)

if "--upload" in sys.argv:
    from googleapiclient.http import MediaIoBaseUpload

    from workspace.drive import get_drive_writer

    svc = get_drive_writer(WRITER)
    with open(OUT, "rb") as fh:
        media = MediaIoBaseUpload(io.BytesIO(fh.read()), mimetype=XLSX, resumable=False)
    meta = (
        svc.files()
        .update(
            fileId=TRACKING_FILE_ID,
            media_body=media,
            fields="id,name,modifiedTime",
            supportsAllDrives=True,
        )
        .execute()
    )
    print("AGGIORNATO IN PLACE:", meta)
else:
    print("\n(nessun upload — rilancia con --upload per aggiornare il file su Drive)")
