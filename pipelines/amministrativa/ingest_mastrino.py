#!/usr/bin/env python3
"""
Mastrino Consolidato + Categorie Conti loader.

Loads from 'Costi Ricavi 2025-2026 Budget.xlsx':
  - CATEGORIE sheet → d_categorie_conti (WRITE_TRUNCATE, dimension)
  - mastrino_consolidato sheet (leaf rows only) → f_mastrino_consolidato (WRITE_TRUNCATE, 2025 actuals)

BU normalization: Hotel→HOTEL, Angelina→RESIDENCE, CVM→CVM, Spiaggia→LIDO, Affitti→HQ

Usage:
    python -m pipelines.amministrativa.ingest_mastrino \\
        --file "/path/to/Costi Ricavi 2025-2026 Budget.xlsx" \\
        --dry-run
    python -m pipelines.amministrativa.ingest_mastrino \\
        --file "/path/to/Costi Ricavi 2025-2026 Budget.xlsx"
"""

import argparse
import logging
import sys
from pathlib import Path

try:
    import openpyxl
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

try:
    from google.cloud import bigquery
    HAS_BQ = True
except ImportError:
    HAS_BQ = False

BQ_PROJECT = "hotelops-suite"
BQ_CATEGORIE_TABLE = f"{BQ_PROJECT}.hotelops.d_categorie_conti"
BQ_MASTRINO_TABLE = f"{BQ_PROJECT}.hotelops.f_mastrino_consolidato"

DEFAULT_SOURCE = Path(
    "/Users/stefanodellapietra/Desktop/WORK/artifacts/Costi Ricavi 2025-2026 Budget.xlsx"
)

BU_MAP = {
    "Hotel": "HOTEL",
    "Angelina": "RESIDENCE",
    "CVM": "CVM",
    "Spiaggia": "LIDO",
    "Affitti": "HQ",
}

BQ_SCHEMA_CATEGORIE = [
    bigquery.SchemaField("codice_conto", "STRING", mode="REQUIRED"),
    bigquery.SchemaField("descrizione", "STRING"),
    bigquery.SchemaField("tipo_costo", "STRING"),
    bigquery.SchemaField("categoria_ce", "STRING"),
    bigquery.SchemaField("anno_inizio", "INTEGER"),
    bigquery.SchemaField("anno_fine", "INTEGER"),
] if HAS_BQ else []

BQ_SCHEMA_MASTRINO = [
    bigquery.SchemaField("societa_id", "STRING", mode="REQUIRED"),
    bigquery.SchemaField("anno", "INTEGER", mode="REQUIRED"),
    bigquery.SchemaField("mese", "INTEGER", mode="REQUIRED"),
    bigquery.SchemaField("codice_conto", "STRING"),
    bigquery.SchemaField("descrizione", "STRING"),
    bigquery.SchemaField("importo", "FLOAT64"),
    bigquery.SchemaField("tipo_costo", "STRING"),
    bigquery.SchemaField("business_unit_id", "STRING"),
    bigquery.SchemaField("categoria", "STRING"),
    bigquery.SchemaField("categoria_ce", "STRING"),
] if HAS_BQ else []


def setup_logger() -> logging.Logger:
    log = logging.getLogger("ingest_mastrino")
    log.setLevel(logging.DEBUG)
    h = logging.StreamHandler(sys.stdout)
    h.setFormatter(logging.Formatter("%(levelname)s %(message)s"))
    log.addHandler(h)
    return log


def parse_categorie(ws, logger: logging.Logger) -> list[dict]:
    rows = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            continue  # header
        conto = row[0]
        if not conto:
            continue
        rows.append({
            "codice_conto": str(conto).strip(),
            "descrizione": str(row[1]).strip() if row[1] else None,
            "tipo_costo": str(row[2]).strip() if row[2] else None,
            "categoria_ce": str(row[3]).strip() if row[3] else None,
            "anno_inizio": int(row[4]) if isinstance(row[4], (int, float)) else None,
            "anno_fine": int(row[5]) if isinstance(row[5], (int, float)) else None,
        })
    logger.info(f"  Categorie: {len(rows)} conti")
    return rows


def parse_mastrino(ws, logger: logging.Logger) -> list[dict]:
    rows = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            continue  # header
        if row[0] is None:
            continue
        livello = row[8]
        if livello != 3:
            continue  # leaf rows only

        bu_raw = row[10]
        business_unit_id = BU_MAP.get(bu_raw) if bu_raw else None

        rows.append({
            "societa_id": str(row[0]).strip(),
            "anno": int(row[14]) if isinstance(row[14], (int, float)) else None,
            "mese": int(row[1]) if isinstance(row[1], (int, float)) else None,
            "codice_conto": str(row[4]).strip() if row[4] else None,
            "descrizione": str(row[6]).strip() if row[6] else None,
            "importo": float(row[7]) if isinstance(row[7], (int, float)) else 0.0,
            "tipo_costo": str(row[9]).strip() if row[9] else None,
            "business_unit_id": business_unit_id,
            "categoria": str(row[11]).strip() if row[11] else None,
            "categoria_ce": str(row[12]).strip() if row[12] else None,
        })

    logger.info(f"  Mastrino: {len(rows)} righe (leaf)")
    return rows


def load_to_bq(rows: list[dict], table_id: str, schema, bq_client, logger: logging.Logger):
    if not rows:
        logger.warning(f"Nessuna riga da caricare in {table_id}")
        return
    job_config = bigquery.LoadJobConfig(
        schema=schema,
        write_disposition=bigquery.WriteDisposition.WRITE_TRUNCATE,
    )
    job = bq_client.load_table_from_json(rows, table_id, job_config=job_config)
    job.result()
    if job.errors:
        logger.error(f"BQ errors: {job.errors}")
        sys.exit(1)
    logger.info(f"  {table_id}: {len(rows)} righe (WRITE_TRUNCATE)")


def main():
    parser = argparse.ArgumentParser(description="Load Mastrino Consolidato + Categorie → BigQuery")
    parser.add_argument(
        "--file",
        default=str(DEFAULT_SOURCE),
        help=f"Path to Costi Ricavi xlsx (default: {DEFAULT_SOURCE})",
    )
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()

    logger = setup_logger()
    filepath = Path(args.file)
    if not filepath.exists():
        logger.error(f"File non trovato: {filepath}")
        sys.exit(1)

    if not HAS_OPENPYXL:
        logger.error("openpyxl non installato")
        sys.exit(1)

    wb = openpyxl.load_workbook(str(filepath), read_only=True, data_only=True)

    categorie_rows = parse_categorie(wb["CATEGORIE"], logger)
    mastrino_rows = parse_mastrino(wb["mastrino_consolidato"], logger)
    wb.close()

    # Summary
    from collections import defaultdict
    by_cat = defaultdict(float)
    for r in mastrino_rows:
        by_cat[(r["societa_id"], r["categoria"])] += r["importo"] or 0
    for (soc, cat), tot in sorted(by_cat.items()):
        logger.info(f"  {soc} | {cat}: {tot:>15,.0f}")

    if args.dry_run:
        logger.info("DRY RUN — nessuna scrittura")
        return

    if not HAS_BQ:
        logger.error("google-cloud-bigquery non installato")
        sys.exit(1)

    bq_client = bigquery.Client(project=BQ_PROJECT)
    load_to_bq(categorie_rows, BQ_CATEGORIE_TABLE, BQ_SCHEMA_CATEGORIE, bq_client, logger)
    load_to_bq(mastrino_rows, BQ_MASTRINO_TABLE, BQ_SCHEMA_MASTRINO, bq_client, logger)

    logger.info("DONE")


if __name__ == "__main__":
    main()
