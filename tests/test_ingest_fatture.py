"""Tests for ingest/flussi/ingest_fatture.py (Lista fatture Esolver dettagliata).

Fixtures sintetiche costruite in tmp_path: replicano i due layout reali
(acquisti senza header / vendite con header) senza dati personali.
"""

from __future__ import annotations

import logging
from pathlib import Path

import openpyxl
import pytest

from ingest.flussi.ingest_fatture import detect_tipo_registro, parse_file

LOGGER = logging.getLogger("test_ingest_fatture")


def _make_acquisti_xlsx(path: Path, rows: list[dict]) -> Path:
    """Layout ACQUISTO: no header, doc=5, des_conto=44, imp=56, conto=57, iva=58, tot=68."""
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in rows:
        row = [""] * 78
        row[1] = "ORTI S.R.L."
        row[2] = "11/06/26 - 15.08"  # timestamp di stampa (volatile)
        row[5] = r["doc"]
        row[6] = "Acquisti"
        row[7] = "1"
        row[8] = "716 - Fattura da fornitore"
        row[9] = r.get("num_doc_orig", "X123 - 01/01/25")
        row[10] = r.get("cod_clifor", 999)
        row[11] = r.get("rag_sociale", "FORNITORE TEST")
        row[44] = r.get("des_conto", "Spese varie")
        row[56] = r["imponibile"]
        row[57] = r["cod_conto"]
        row[58] = r.get("cod_iva", "22")
        row[68] = r.get("tot_documento", r["imponibile"] * 1.22)
        ws.append(row)
    wb.save(path)
    return path


def _make_vendite_xlsx(path: Path, rows: list[dict]) -> Path:
    """Layout VENDITA: header row, doc=5, des_conto=55, imp=68, conto=69, iva=70, tot=89."""
    wb = openpyxl.Workbook()
    ws = wb.active
    header = [""] * 97
    header[5] = "RifRegistrazione"
    header[6] = "TipoRegistoIvaDecod"
    header[69] = "ContoCodConto"
    ws.append(header)
    for r in rows:
        row = [""] * 97
        row[1] = "ORTI S.R.L."
        row[5] = r["doc"]
        row[6] = "Vendite"
        row[7] = "5"
        row[8] = "721 - Fattura cliente"
        row[9] = r.get("cod_clifor", 92)
        row[10] = r.get("rag_sociale", "CLIENTE TEST")
        row[55] = r.get("des_conto", "Ricavi test")
        row[68] = r["imponibile"]
        row[69] = r["cod_conto"]
        row[70] = r.get("cod_iva", "22")
        row[89] = r.get("tot_documento", r["imponibile"] * 1.22)
        ws.append(row)
    wb.save(path)
    return path


class TestDetectTipoRegistro:
    def test_acquisti(self, tmp_path):
        f = _make_acquisti_xlsx(
            tmp_path / "a.xlsx",
            [
                {
                    "doc": "FT n. 1 del 04/01/25",
                    "imponibile": 100.0,
                    "cod_conto": "570101",
                }
            ],
        )
        wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
        rows = list(wb.worksheets[0].iter_rows(values_only=True))
        assert detect_tipo_registro(rows) == "ACQUISTO"

    def test_vendite_with_header(self, tmp_path):
        f = _make_vendite_xlsx(
            tmp_path / "v.xlsx",
            [
                {
                    "doc": "FT n. 1  del 07/01/25",
                    "imponibile": 100.0,
                    "cod_conto": "479502",
                }
            ],
        )
        wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
        rows = list(wb.worksheets[0].iter_rows(values_only=True))
        assert detect_tipo_registro(rows) == "VENDITA"

    def test_unrecognized_raises(self):
        with pytest.raises(ValueError):
            detect_tipo_registro([("x",) * 10])


class TestParseAcquisti:
    def test_basic_row(self, tmp_path):
        f = _make_acquisti_xlsx(
            tmp_path / "a.xlsx",
            [
                {
                    "doc": "FT n. 1 del 04/01/25",
                    "imponibile": 757.79,
                    "cod_conto": "65030599",
                }
            ],
        )
        rows = parse_file(f, "ORTI", LOGGER)
        assert len(rows) == 1
        r = rows[0]
        assert r["tipo_registro"] == "ACQUISTO"
        assert r["cod_conto"] == "65030599"
        assert r["imponibile"] == 757.79
        assert r["data_registrazione"] == "2025-01-04"
        assert r["anno"] == 2025 and r["mese"] == 1
        assert r["sigla_doc"] == "FT"
        assert r["num_registrazione"] == 1

    def test_nc_negates_imponibile(self, tmp_path):
        f = _make_acquisti_xlsx(
            tmp_path / "a.xlsx",
            [
                {
                    "doc": "NC n. 17 del 05/03/25",
                    "imponibile": 2400.0,
                    "cod_conto": "55070101",
                }
            ],
        )
        rows = parse_file(f, "ORTI", LOGGER)
        assert rows[0]["imponibile"] == -2400.0

    def test_multi_riga_stesso_documento(self, tmp_path):
        f = _make_acquisti_xlsx(
            tmp_path / "a.xlsx",
            [
                {
                    "doc": "FT n. 4 del 13/01/25",
                    "imponibile": 15093.43,
                    "cod_conto": "610107",
                    "cod_iva": "22",
                },
                {
                    "doc": "FT n. 4 del 13/01/25",
                    "imponibile": 104.7,
                    "cod_conto": "610107",
                    "cod_iva": "ES16",
                },
            ],
        )
        rows = parse_file(f, "ORTI", LOGGER)
        assert len(rows) == 2
        assert rows[0]["hash_riga"] != rows[1]["hash_riga"]

    def test_righe_identiche_disambiguate(self, tmp_path):
        same = {
            "doc": "FT n. 9 del 01/02/25",
            "imponibile": 50.0,
            "cod_conto": "570101",
            "cod_iva": "22",
        }
        f = _make_acquisti_xlsx(tmp_path / "a.xlsx", [same, dict(same)])
        rows = parse_file(f, "ORTI", LOGGER)
        assert len(rows) == 2
        assert rows[0]["hash_riga"] != rows[1]["hash_riga"]

    def test_hash_stabile_tra_export(self, tmp_path):
        """Stessa riga logica in due file con timestamp stampa diversi → stesso hash."""
        row = {
            "doc": "FT n. 1 del 04/01/25",
            "imponibile": 757.79,
            "cod_conto": "65030599",
        }
        f1 = _make_acquisti_xlsx(tmp_path / "a1.xlsx", [row])
        f2 = _make_acquisti_xlsx(tmp_path / "a2.xlsx", [row])
        wb = openpyxl.load_workbook(f2)
        wb.active.cell(row=1, column=3).value = "12/07/26 - 09.00"  # altro timestamp
        wb.save(f2)
        h1 = parse_file(f1, "ORTI", LOGGER)[0]["hash_riga"]
        h2 = parse_file(f2, "ORTI", LOGGER)[0]["hash_riga"]
        assert h1 == h2

    def test_sigle_composte_ue_reverse_charge(self, tmp_path):
        """FT-UE, AFT-UE, FT-RC (reverse charge / acquisti UE) sono righe valide."""
        f = _make_acquisti_xlsx(
            tmp_path / "a.xlsx",
            [
                {
                    "doc": "FT-UE n. 2 del 04/04/24",
                    "imponibile": 29.0,
                    "cod_conto": "570101",
                },
                {
                    "doc": "AFT-UE n. 8 del 17/05/24",
                    "imponibile": 100.0,
                    "cod_conto": "570101",
                },
                {
                    "doc": "FT-RC n. 17 del 15/03/24",
                    "imponibile": 550.0,
                    "cod_conto": "570101",
                },
                {
                    "doc": "NC-UE n. 3 del 20/06/24",
                    "imponibile": 10.0,
                    "cod_conto": "570101",
                },
            ],
        )
        rows = parse_file(f, "INTUR", LOGGER)
        assert [r["sigla_doc"] for r in rows] == ["FT-UE", "AFT-UE", "FT-RC", "NC-UE"]
        assert rows[3]["imponibile"] == -10.0  # NC-UE negata

    def test_skip_righe_senza_doc_pattern(self, tmp_path):
        f = _make_acquisti_xlsx(
            tmp_path / "a.xlsx",
            [
                {
                    "doc": "FT n. 1 del 04/01/25",
                    "imponibile": 10.0,
                    "cod_conto": "570101",
                }
            ],
        )
        wb = openpyxl.load_workbook(f)
        footer = [""] * 78
        footer[5] = "Totale generale"
        wb.active.append(footer)
        wb.save(f)
        rows = parse_file(f, "ORTI", LOGGER)
        assert len(rows) == 1


class TestParseVendite:
    def test_basic_row_double_space(self, tmp_path):
        """Il riferimento vendite ha doppio spazio: 'FT n. 1  del 07/01/25'."""
        f = _make_vendite_xlsx(
            tmp_path / "v.xlsx",
            [
                {
                    "doc": "FT n. 1  del 07/01/25",
                    "imponibile": 21000.0,
                    "cod_conto": "479502",
                }
            ],
        )
        rows = parse_file(f, "ORTI", LOGGER)
        assert len(rows) == 1
        r = rows[0]
        assert r["tipo_registro"] == "VENDITA"
        assert r["cod_conto"] == "479502"
        assert r["imponibile"] == 21000.0
        assert r["data_registrazione"] == "2025-01-07"

    def test_header_row_skipped(self, tmp_path):
        f = _make_vendite_xlsx(
            tmp_path / "v.xlsx",
            [
                {
                    "doc": "FT n. 1  del 07/01/25",
                    "imponibile": 100.0,
                    "cod_conto": "479502",
                }
            ],
        )
        rows = parse_file(f, "ORTI", LOGGER)
        assert len(rows) == 1  # header non produce righe
