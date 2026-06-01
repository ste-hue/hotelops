from io import BytesIO
from pathlib import Path

import openpyxl
import pytest

from verticals.condges.pf_rotate.step0_normalize_intur import normalize_intur


@pytest.fixture
def intur_pf_bytes() -> bytes:
    """Mini INTUR PF con col A 'sporca': importi, RID, codici."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ut = wb.create_sheet("Utenze")
    ut["A2"] = "CODICE"
    ut["D2"] = "APRILE"
    ut["E2"] = "MAGGIO"
    ut["B3"] = "Utenze"
    # 3 righe: codice mancante / RID / importo
    ut["A4"] = None
    ut["B4"] = "Energia elettrica"
    ut["A5"] = "RID"
    ut["B5"] = "Vodafone"
    ut["A6"] = 14400.0  # importo arretrato
    ut["B6"] = "Fastweb SPA"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture
def fornitori_csv(tmp_path: Path) -> Path:
    p = tmp_path / "d_fornitori.csv"
    p.write_text(
        "codice_fornitore,nome_esolver,nome_pf,voce_id,is_intercompany,is_excluded,exclude_reason,societa_id\n"
        "100,Energia elettrica,Energia elettrica,USCITE_UTENZE,False,False,,INTUR\n"
        "44,Vodafone,Vodafone,USCITE_UTENZE,False,False,,INTUR\n"
        "1592,FASTWEB SPA,Fastweb,USCITE_UTENZE,False,False,,INTUR\n"
    )
    return p


def test_normalize_overwrites_col_a_with_codici(intur_pf_bytes, fornitori_csv):
    out_bytes, overwrites, unmapped = normalize_intur(
        pf_bytes=intur_pf_bytes,
        fornitori_csv=fornitori_csv,
    )
    wb = openpyxl.load_workbook(BytesIO(out_bytes))
    ut = wb["Utenze"]
    assert ut["A4"].value == 100  # Energia elettrica → 100
    assert ut["A5"].value == 44  # Vodafone → 44
    assert ut["A6"].value == 1592  # Fastweb (importo overwritten)

    # overwrites: la sola sovrascrittura di un valore non-None va loggata
    refs = [(o["foglio"], o["riga"], o["vecchio"]) for o in overwrites]
    assert ("Utenze", 5, "RID") in refs
    assert ("Utenze", 6, 14400.0) in refs
    assert unmapped == []
