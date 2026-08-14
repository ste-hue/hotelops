import argparse
import re
import zipfile
from datetime import date
from io import BytesIO

import openpyxl
import pytest

from verticals.condges.pf_rotate.cli_handler import (
    _parse_exclude,
    _resolve_data_saldo,
    add_subparser,
)
from verticals.condges.pf_rotate.excel_model import find_month_periods, periodo


def _inject_cached_xml_values(
    xlsx_bytes: bytes, sheet_xml: str, cell_values: dict[str, float]
) -> bytes:
    """Scrive un valore <v> cached dentro una cella-formula, direttamente
    nell'XML del foglio — simula un file REALMENTE aperto/salvato in Excel
    (openpyxl non calcola formule: un round-trip via openpyxl le lascerebbe
    sempre senza <v>, vedi cli_handler._handle_extend). Usato per costruire un
    fixture su cui il gemello data_only=True caricato dal disco (post-fix)
    trova valori veri, non None.
    """
    with zipfile.ZipFile(BytesIO(xlsx_bytes)) as zin:
        contents = {n: zin.read(n) for n in zin.namelist()}
    xml = contents[sheet_xml].decode()
    for addr, val in cell_values.items():
        pat = re.compile(rf'(<c r="{re.escape(addr)}"[^>]*>)(.*?)(</c>)')
        m = pat.search(xml)
        assert m, f"cella {addr} non trovata in {sheet_xml}"
        inner = m.group(2)
        assert "<f>" in inner, f"cella {addr} non ha una formula: {inner}"
        # <v/> self-closing (Python <=3.13) o <v></v> (3.14+): entrambe le forme
        # vanno rimosse, altrimenti resta un <v> vuoto PRIMA del nostro e openpyxl
        # legge quello — il cached value risulta assente solo su certe versioni.
        inner = re.sub(r"<v\s*/>|<v>.*?</v>", "", inner) + f"<v>{val}</v>"
        xml = xml[: m.start()] + m.group(1) + inner + m.group(3) + xml[m.end() :]
    contents[sheet_xml] = xml.encode()
    out = BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for n, data in contents.items():
            zout.writestr(n, data)
    return out.getvalue()


def test_resolve_data_saldo_explicit_wins():
    """--data-saldo esplicito è override: vince sull'anno/mese derivato."""
    assert _resolve_data_saldo(date(2026, 4, 30), 2026, 5) == date(2026, 4, 30)


def test_resolve_data_saldo_derives_last_day_of_month():
    assert _resolve_data_saldo(None, 2026, 5) == date(2026, 5, 31)


def test_resolve_data_saldo_february_non_leap():
    assert _resolve_data_saldo(None, 2026, 2) == date(2026, 2, 28)


def test_resolve_data_saldo_february_leap():
    assert _resolve_data_saldo(None, 2024, 2) == date(2024, 2, 29)


def test_parse_exclude_comma_separated():
    assert _parse_exclude(["264,48"]) == {264, 48}


def test_parse_exclude_repeatable_and_empty():
    assert _parse_exclude(["264", "48"]) == {264, 48}
    assert _parse_exclude([]) == set()


def test_pf_extend_cli_scrive_file_esteso(tmp_path, minimal_pf_orti_bytes):
    pf_path = tmp_path / "ORTI_PF.xlsx"
    pf_path.write_bytes(minimal_pf_orti_bytes)
    out_dir = tmp_path / "out"

    parser = argparse.ArgumentParser()
    sub = parser.add_subparsers(dest="command")
    add_subparser(sub)
    args = parser.parse_args(
        ["pf-extend", "--pf", str(pf_path), "--to", "2027-06", "--out", str(out_dir)]
    )

    rc = args.func(args)
    assert rc == 0

    outputs = list(out_dir.glob("*_extended_2027-06_*.xlsx"))
    assert len(outputs) == 1
    wb = openpyxl.load_workbook(outputs[0])
    cols = find_month_periods(wb["Piano Finanziario"])
    assert max(cols) == periodo(2027, 6)


def test_pf_extend_cli_trim_before_wiring_no_op(tmp_path, minimal_pf_orti_bytes):
    """--trim-before con cutoff = primo mese già esistente: nessuna colonna da
    eliminare, trim_before torna [] senza toccare i valori cached (nessuna
    freeze necessaria) — verifica solo che l'argomento sia collegato a
    trim_before DOPO extend_to, senza esercitare il path di congelamento."""
    pf_path = tmp_path / "ORTI_PF.xlsx"
    pf_path.write_bytes(minimal_pf_orti_bytes)
    out_dir = tmp_path / "out"

    parser = argparse.ArgumentParser()
    sub = parser.add_subparsers(dest="command")
    add_subparser(sub)
    args = parser.parse_args(
        [
            "pf-extend",
            "--pf",
            str(pf_path),
            "--to",
            "2027-06",
            "--trim-before",
            "2026-04",  # = primo mese del fixture: no-op
            "--out",
            str(out_dir),
        ]
    )

    rc = args.func(args)
    assert rc == 0

    outputs = list(out_dir.glob("*_extended_2027-06_*.xlsx"))
    assert len(outputs) == 1
    wb = openpyxl.load_workbook(outputs[0])
    cols = find_month_periods(wb["Piano Finanziario"])
    assert min(cols) == periodo(2026, 4)  # nulla eliminato
    assert max(cols) == periodo(2027, 6)


def test_pf_extend_cli_trim_before_fail_loud_senza_cache(
    tmp_path, minimal_pf_orti_bytes
):
    """Il fixture è puro-openpyxl (mai aperto in Excel): estendere e potare
    nella STESSA invocazione, senza un giro per Excel in mezzo, non ha un
    valore calcolato da congelare per la cascata della nuova prima colonna —
    la CLI deve fallire rumorosamente (ValueError), mai scrivere un file
    silenziosamente rotto."""
    pf_path = tmp_path / "ORTI_PF.xlsx"
    pf_path.write_bytes(minimal_pf_orti_bytes)
    out_dir = tmp_path / "out"

    parser = argparse.ArgumentParser()
    sub = parser.add_subparsers(dest="command")
    add_subparser(sub)
    args = parser.parse_args(
        [
            "pf-extend",
            "--pf",
            str(pf_path),
            "--to",
            "2027-06",
            "--trim-before",
            "2027-01",  # elimina il vecchio storico -> serve un cached value
            "--out",
            str(out_dir),
        ]
    )

    with pytest.raises(ValueError, match="apri e salva il file in Excel"):
        args.func(args)

    # l'input su disco non è mai stato toccato
    wb_in = openpyxl.load_workbook(pf_path)
    cols_in = find_month_periods(wb_in["Piano Finanziario"])
    assert max(cols_in) == periodo(2026, 12)


def test_pf_extend_cli_trim_before_con_cache_reale_riesce(
    tmp_path, minimal_pf_orti_bytes
):
    """Con un file che HA valori cached veri (iniettati via XML, non via
    openpyxl round-trip: vedi _inject_cached_xml_values) per le formule che
    trim_before deve congelare, il trim riesce end-to-end dalla CLI — prova
    che wb_values ora legge dal disco (--pf), non da un re-save che le
    spoglia sempre. Nessuna estensione in gioco (--to = massimo già presente,
    no-op): isola la wiring del freeze dal caso extend+trim in un colpo solo
    (già coperto, sempre fail-loud, dal test _fail_loud_senza_cache sopra).

    Celle iniettate (determinato empiricamente lanciando il trim e osservando
    quali celle sollevano 'valore assente', stesso approccio di
    _inject_cached_freeze_values in test_pf_rotate_extend.py):
      G4  = cascata SALDO MESE PRECEDENTE della prima superstite (AGOSTO),
            che referenzia F37 (LUGLIO, eliminato);
      L12 = TOTALI/TOTALE ENTRATE, SUM(C12:K12) copre lo storico eliminato;
      L27 = TOTALI/TOTALE USCITE, idem.
    """
    raw = _inject_cached_xml_values(
        minimal_pf_orti_bytes,
        "xl/worksheets/sheet1.xml",
        {"G4": 100_000.0, "L12": 900.0, "L27": 2700.0},
    )
    pf_path = tmp_path / "ORTI_PF.xlsx"
    pf_path.write_bytes(raw)
    out_dir = tmp_path / "out"

    parser = argparse.ArgumentParser()
    sub = parser.add_subparsers(dest="command")
    add_subparser(sub)
    args = parser.parse_args(
        [
            "pf-extend",
            "--pf",
            str(pf_path),
            "--to",
            "2026-12",  # già il massimo del fixture: extend_to è un no-op
            "--trim-before",
            "2026-08",  # elimina APR-LUG, AGOSTO diventa la prima superstite
            "--out",
            str(out_dir),
        ]
    )

    rc = args.func(args)
    assert rc == 0

    outputs = list(out_dir.glob("*_extended_2026-12_*.xlsx"))
    assert len(outputs) == 1
    wb = openpyxl.load_workbook(outputs[0])
    cols = find_month_periods(wb["Piano Finanziario"])
    assert min(cols) == periodo(2026, 8)  # APR-LUG potati
    assert max(cols) == periodo(2026, 12)
    first = cols[periodo(2026, 8)]
    v = wb["Piano Finanziario"].cell(4, first).value
    assert v == 100_000.0  # il valore cached iniettato, non più una formula


def test_pf_extend_cli_no_op_non_conta_skip_come_colonne_aggiunte(
    tmp_path, minimal_pf_orti_bytes
):
    """Su un file già al target, extend_to ritorna comunque righe 'skip: ...'
    per ogni foglio di servizio (DA MAPPARE/ESCLUSI) — senza filtro, il CLI
    le contava come 'Colonne aggiunte' e scriveva un output anche a fronte di
    zero estensioni reali. No trim richiesto: no-op vero deve stampare
    'Già coperto' e non scrivere nessun file."""
    wb = openpyxl.load_workbook(BytesIO(minimal_pf_orti_bytes))
    wb.create_sheet("DA MAPPARE")
    wb.create_sheet("ESCLUSI")
    buf = BytesIO()
    wb.save(buf)
    pf_path = tmp_path / "ORTI_PF.xlsx"
    pf_path.write_bytes(buf.getvalue())
    out_dir = tmp_path / "out"

    parser = argparse.ArgumentParser()
    sub = parser.add_subparsers(dest="command")
    add_subparser(sub)
    args = parser.parse_args(
        [
            "pf-extend",
            "--pf",
            str(pf_path),
            "--to",
            "2026-12",  # già il massimo del fixture
            "--out",
            str(out_dir),
        ]
    )

    rc = args.func(args)
    assert rc == 0
    assert not out_dir.exists()  # no-op vero: nessun file scritto
