from datetime import date
from io import BytesIO
from unittest.mock import MagicMock

import openpyxl
import pytest

from verticals.condges.pf_rotate.excel_model import find_layout
from verticals.condges.pf_rotate.step1_saldi import (
    SaldiIncompletiError,
    banche_mancanti,
    fetch_saldi_da_bq,
    preflight_saldi,
    write_saldi_banca,
)


def test_write_saldi_writes_to_cutover_column(minimal_pf_orti_bytes):
    """APRILE chiuso → col C. Saldi vanno in C32/C33."""
    wb = openpyxl.load_workbook(BytesIO(minimal_pf_orti_bytes), data_only=False)
    pf = wb["Piano Finanziario"]
    # pre-stato fixture: C32=251897.54, C33=87439.92, C4=339337.46
    # Aggiorna a un nuovo cutover ipotetico (es. 31/05):
    write_saldi_banca(
        wb,
        mese_chiuso=4,  # ancora APRILE (smoke: identica sovrascrittura)
        data_saldo=date(2026, 4, 30),
        saldi={"MPS": 300000.0, "Intesa": 50000.0},
    )
    assert pf["C32"].value == 300000.0
    assert pf["C33"].value == 50000.0
    assert pf["C4"].value == 350000.0  # hardcoded = somma
    # C35 resta formula
    assert pf["C35"].value == "=SUM(C32:C34)"
    # data scritta in C31
    assert pf["C31"].value == "30/04/2026"


def test_write_saldi_partial_only_known_banks(minimal_pf_orti_bytes):
    """Se passi solo MPS, Intesa resta com'era e il totale C4 riflette la somma osservata."""
    import datetime

    wb = openpyxl.load_workbook(BytesIO(minimal_pf_orti_bytes), data_only=False)
    pf = wb["Piano Finanziario"]
    pre_c33 = pf["C33"].value  # 87439.92

    write_saldi_banca(
        wb,
        mese_chiuso=4,
        data_saldo=datetime.date(2026, 4, 30),
        saldi={"MPS": 200000.0},
    )
    assert pf["C32"].value == 200000.0
    assert pf["C33"].value == pre_c33  # invariato
    assert pf["C4"].value == 200000.0 + pre_c33


def _build_fixed_snapshot_wb() -> openpyxl.Workbook:
    """Minimal INTUR-style workbook: C1='DATA RILEVAZ', month headers shifted to D+."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    pf = wb.create_sheet("Piano Finanziario")
    pf["A1"] = "INTUR"
    pf["C1"] = "DATA RILEVAZ"
    # row 2: month headers in D..M (E = APRILE)
    pf["C2"] = None
    mesi = [
        "MARZO",
        "APRILE",
        "MAGGIO",
        "GIUGNO",
        "LUGLIO",
        "AGOSTO",
        "SETTEMBRE",
        "OTTOBRE",
        "NOVEMBRE",
        "DICEMBRE",
    ]
    for i, m in enumerate(mesi):
        pf.cell(2, 4 + i, m)  # cols D(4)..M(13)

    # Anchors for find_layout
    pf["A3"] = "SALDO MESE PRECEDENTE"
    pf["A5"] = "Fitto Hotel"
    pf["A12"] = "TOTALE ENTRATE"
    pf["A14"] = "Salari e Stipendi"
    pf["A27"] = "TOTALE USCITE"
    pf["A29"] = "CASH FLOW"
    pf["A31"] = "Saldo MPS"
    pf["A32"] = "Saldo Intesa"
    pf["A33"] = "Saldo BCP"
    pf["A35"] = "TOTALE BANCHE"
    pf["A38"] = "Saldo di Periodo"
    return wb


def test_write_saldi_banca_fixed_snapshot_writes_to_col_c():
    """INTUR-style: saldi banche scritti in col C, NON nella col del mese chiuso."""
    wb = _build_fixed_snapshot_wb()
    pf = wb["Piano Finanziario"]

    out = write_saldi_banca(
        wb,
        mese_chiuso=4,  # APRILE → col E (5); ma saldi vanno in C
        data_saldo=date(2026, 4, 30),
        saldi={"MPS": 100, "Intesa": 200, "BCP": 50},
    )

    assert pf["C31"].value == 100
    assert pf["C32"].value == 200
    assert pf["C33"].value == 50
    assert pf["C2"].value == "30/04/2026"
    # Mese chiuso col E NON deve essere toccato sulle righe saldi
    assert pf["E31"].value is None
    assert pf["E32"].value is None
    assert pf["E33"].value is None
    # saldo_iniziale_row (E3) NON deve essere toccato in fixed-snapshot
    assert pf["E3"].value is None
    assert out["snapshot_kind"] == "fixed-snapshot"
    assert out["col"] == "C"


def test_find_layout_detects_fixed_snapshot():
    """find_layout setta snapshot_kind a 'fixed-snapshot' se C1 contiene 'DATA RILEVAZ'."""
    wb1 = _build_fixed_snapshot_wb()
    assert find_layout(wb1).snapshot_kind == "fixed-snapshot"

    # Second fixture: C1=None (ORTI-style) → month-closed
    wb2 = openpyxl.Workbook()
    wb2.remove(wb2.active)
    pf2 = wb2.create_sheet("Piano Finanziario")
    pf2["A1"] = "ORTI"
    # C1 lasciato None
    pf2.cell(2, 3, "APRILE")
    pf2["A4"] = "SALDO MESE PRECEDENTE"
    pf2["A6"] = "Entrate Hotel"
    pf2["A12"] = "TOTALE ENTRATE"
    pf2["A14"] = "Utenze"
    pf2["A27"] = "TOTALE USCITE"
    pf2["A29"] = "CASH FLOW"
    pf2["A32"] = "Saldo MPS"
    pf2["A35"] = "TOTALE BANCHE"
    pf2["A37"] = "Saldo di Periodo"
    assert find_layout(wb2).snapshot_kind == "month-closed"


def test_fetch_saldi_da_bq_query_columns(monkeypatch):
    """La query usa 'data_riferimento' + 'saldo_eur' (non 'data_saldo' / 'saldo')."""
    captured_sql: dict[str, str] = {}

    # Fake rows: oggetti con attributi banca_id + saldo_eur
    fake_row_mps = MagicMock()
    fake_row_mps.banca_id = "MPS"
    fake_row_mps.saldo_eur = 1234.56
    fake_row_intesa = MagicMock()
    fake_row_intesa.banca_id = "INTESA"
    fake_row_intesa.saldo_eur = 789.01

    fake_job = MagicMock()
    fake_job.result.return_value = [fake_row_mps, fake_row_intesa]

    fake_client = MagicMock()

    def _query(sql, **kwargs):
        captured_sql["sql"] = sql
        return fake_job

    fake_client.query.side_effect = _query

    monkeypatch.setattr(
        "core.bq.client.get_client",
        lambda: fake_client,
    )

    out = fetch_saldi_da_bq("ORTI", date(2026, 4, 30))

    sql = captured_sql["sql"]
    assert "data_riferimento" in sql
    assert "saldo_eur" in sql
    assert "data_saldo" not in sql
    assert "SELECT banca_id, saldo " not in sql
    # ORDER BY banca_id: ordine deterministico → niente heisenbug nel match (vedi
    # test_write_saldi_no_mps_kross_collision).
    assert "ORDER BY banca_id" in sql
    assert out == {"MPS": 1234.56, "INTESA": 789.01}


def _build_month_closed_wb_with_kross() -> openpyxl.Workbook:
    """ORTI-style (month-closed) con DUE righe MPS: 'Saldo MPS' + 'Saldo MPS Kross'."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    pf = wb.create_sheet("Piano Finanziario")
    pf["A1"] = "ORTI"
    # C1 lasciato None → month-closed; APRILE in col C (mese_chiuso=4)
    pf.cell(2, 3, "APRILE")
    pf["A4"] = "SALDO MESE PRECEDENTE"
    pf["A6"] = "Entrate Hotel"
    pf["A12"] = "TOTALE ENTRATE"
    pf["A14"] = "Utenze"
    pf["A27"] = "TOTALE USCITE"
    pf["A29"] = "CASH FLOW"
    pf["A32"] = "Saldo MPS"
    pf["A33"] = "Saldo MPS Kross"
    pf["A35"] = "TOTALE BANCHE"
    pf.cell(35, 3, "=SUM(C32:C34)")
    pf["A37"] = "Saldo di Periodo"
    return wb


def test_write_saldi_no_mps_kross_collision():
    """'Saldo MPS' → MPS, 'Saldo MPS Kross' → MPS_KROSS. Niente cross-match.

    Col fuzzy-overlap precedente, 'Saldo MPS Kross' (parole {mps,kross}) matchava
    la chiave 'MPS' ({mps}) e si prendeva il saldo sbagliato — dipendente dall'ordine
    del dict, quindi heisenbug.
    """
    wb = _build_month_closed_wb_with_kross()
    pf = wb["Piano Finanziario"]
    write_saldi_banca(
        wb,
        mese_chiuso=4,
        data_saldo=date(2026, 4, 30),
        saldi={"MPS": 100000.0, "MPS_KROSS": 999.0},
    )
    assert pf["C32"].value == 100000.0  # Saldo MPS
    assert pf["C33"].value == 999.0  # Saldo MPS Kross


# ── Preflight gate O-A ──────────────────────────────────────────────────────


def test_banche_mancanti_intur_missing_sella():
    assert banche_mancanti("INTUR", {"MPS": 1, "INTESA": 2}) == ["SELLA"]


def test_banche_mancanti_case_insensitive_and_kross_optional():
    # Kross è opzionale → ORTI completo con solo INTESA+MPS; chiavi case-insensitive.
    assert banche_mancanti("ORTI", {"intesa": 1, "mps": 2}) == []


def test_preflight_saldi_hard_fail_actionable():
    with pytest.raises(SaldiIncompletiError) as ei:
        preflight_saldi("INTUR", {"MPS": 1, "INTESA": 2}, date(2026, 5, 31))
    msg = str(ei.value)
    assert "INTUR/SELLA/2026-05-31" in msg  # tupla mancante esplicita
    assert "saldi-ufficiali" in msg  # punta al comando di cattura (O-C)


def test_preflight_saldi_allow_partial_returns_missing_no_raise():
    assert preflight_saldi(
        "INTUR", {"MPS": 1, "INTESA": 2}, date(2026, 5, 31), allow_partial=True
    ) == ["SELLA"]


def test_preflight_saldi_complete_ok():
    assert (
        preflight_saldi("INTUR", {"MPS": 1, "INTESA": 2, "SELLA": 3}, date(2026, 5, 31))
        == []
    )


def test_write_saldi_fixed_snapshot_non_tocca_r4_mesi(intur_pf_bytes):
    """INTUR fixed-snapshot: saldi solo in colonna C, r4 dei mesi mai toccato."""
    wb = openpyxl.load_workbook(BytesIO(intur_pf_bytes), data_only=False)
    pf = wb["Piano Finanziario"]
    before = [pf.cell(4, c).value for c in range(3, 12)]
    write_saldi_banca(
        wb,
        mese_chiuso=4,
        data_saldo=date(2026, 4, 30),
        saldi={"Sella": 10000.0, "MPS": 20000.0, "Intesa": 3000.0},
    )
    after = [pf.cell(4, c).value for c in range(3, 12)]
    assert before == after
