"""Test del tool di validazione semantica candidato vs ground truth.

Workbook sintetici minimali (in memoria) — NON tocca i file reali.
Si testa il CLASSIFICATORE: è la cosa di cui ci fidiamo.

Layout foglio voce (sia candidato che ground truth, layout standard):
  r1  : titolo voce (col B)
  r2  : header — col A "Cod", col B "Fornitore / Voce", col 3..14 = GEN..DIC
  r3  : TOTALE (formula, ignorata)
  r5  : "PARTITE APERTE (motore)"
  r6..: blocco A — col A codice, col B nome, importi
  rK  : "PREVISIONI (manuale)"
  rK+1..: blocco B
  ultima riga blocco B: "RETTIFICA PARTITE/PREVISIONI"
"""

from __future__ import annotations

from io import BytesIO

import openpyxl

from verticals.condges.pf_generator.costanti import MESI, col_mese
from verticals.condges.pf_generator.valida import confronta_pf


def _make_wb(sheets: dict[str, list[tuple]]) -> bytes:
    """Costruisce un workbook xlsx in memoria.

    ``sheets``: {nome_foglio: [ (codice|None, nome, {mese:int -> valore}), ... ]}
    Le righe sono scritte a partire da r6 (blocco A).
    """
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for nome_foglio, righe in sheets.items():
        ws = wb.create_sheet(nome_foglio)
        ws.cell(row=1, column=2, value=nome_foglio)
        ws.cell(row=2, column=1, value="Cod")
        ws.cell(row=2, column=2, value="Fornitore / Voce")
        for m, nome_mese in enumerate(MESI, start=1):
            ws.cell(row=2, column=col_mese(m), value=nome_mese)
        ws.cell(row=3, column=2, value="TOTALE")
        ws.cell(row=5, column=2, value="PARTITE APERTE (motore)")
        r = 6
        for codice, nome, mesi in righe:
            if codice is not None:
                ws.cell(row=r, column=1, value=int(codice))
            ws.cell(row=r, column=2, value=nome)
            for mese, val in mesi.items():
                ws.cell(row=r, column=col_mese(mese), value=val)
            r += 1
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_uguale():
    """Stessi fornitori, stessi importi → nessuna differenza, UGUALE."""
    sheets = {
        "Materie Prime e Consumo": [
            (412, "Alberto D'Urso", {5: 876.42}),
            (19, "AMAZON EU", {5: 591.42, 6: 100.0}),
        ]
    }
    b = _make_wb(sheets)
    rep = confronta_pf(b, b, primo_mese_aperto=5)
    assert rep["verdetto"] == "UGUALE", rep
    assert rep["errori"] == []


def test_importo_spostato():
    """Stesso fornitore, stesso totale, ma mese diverso → ERRORE."""
    cand = _make_wb({"Materie Prime e Consumo": [(412, "Alberto D'Urso", {5: 876.42})]})
    rosa = _make_wb({"Materie Prime e Consumo": [(412, "Alberto D'Urso", {6: 876.42})]})
    rep = confronta_pf(cand, rosa, primo_mese_aperto=5)
    classi = {e["classe"] for e in rep["errori"]}
    assert "IMPORTO_SPOSTATO" in classi, rep
    assert rep["verdetto"] == "PEGGIORE"


def test_previsione_aggiunta_da_rosa():
    """Riga in Rosa, assente nel candidato, NON in scad_codici → ATTESA."""
    cand = _make_wb({"Consulenze": [(100, "Studio X", {5: 1000.0})]})
    rosa = _make_wb(
        {
            "Consulenze": [
                (100, "Studio X", {5: 1000.0}),
                (999, "Polizza Q3", {6: 9000.0}),
            ]
        }
    )
    rep = confronta_pf(cand, rosa, primo_mese_aperto=5, scad_codici={100})
    classi_att = {a["classe"] for a in rep["attese"]}
    assert "PREVISIONE_AGGIUNTA_DA_ROSA" in classi_att, rep
    assert rep["errori"] == [], rep


def test_partita_persa():
    """Codice con partita aperta in export, in Rosa ma assente nel candidato → ERRORE."""
    cand = _make_wb({"Materie Prime e Consumo": [(412, "Alberto D'Urso", {5: 876.42})]})
    rosa = _make_wb(
        {
            "Materie Prime e Consumo": [
                (412, "Alberto D'Urso", {5: 876.42}),
                (500, "Fornitore Perso", {5: 1234.0}),
            ]
        }
    )
    rep = confronta_pf(cand, rosa, primo_mese_aperto=5, scad_codici={412, 500})
    classi = {e["classe"] for e in rep["errori"]}
    assert "PARTITA_PERSA" in classi, rep
    assert rep["verdetto"] == "PEGGIORE"


def test_doppio_conteggio_rosa_candidato_pulito():
    """Rosa ha l'importo duplicato (2x), candidato pulito → il candidato è MIGLIORE.

    La differenza è un errore di Rosa che il motore non commette.
    """
    cand = _make_wb({"Materie Prime e Consumo": [(412, "Alberto D'Urso", {5: 876.42})]})
    rosa = _make_wb(
        {"Materie Prime e Consumo": [(412, "Alberto D'Urso", {5: 1752.84})]}
    )
    rep = confronta_pf(cand, rosa, primo_mese_aperto=5, scad_codici={412})
    # candidato ≈ 1/2 di Rosa → doppio conteggio di Rosa, candidato corretto
    classi_att = {a["classe"] for a in rep["attese"]}
    assert "DOPPIO_CONTEGGIO" in classi_att, rep
    assert rep["errori"] == [], rep
    assert rep["verdetto"] == "MIGLIORE", rep


def test_rettifica_only():
    """Candidato ha la riga RETTIFICA (strutturale), Rosa no → ATTESA, non errore."""
    cand = _make_wb(
        {
            "Materie Prime e Consumo": [
                (412, "Alberto D'Urso", {5: 876.42}),
                (None, "RETTIFICA PARTITE/PREVISIONI", {5: -200.0}),
            ]
        }
    )
    rosa = _make_wb({"Materie Prime e Consumo": [(412, "Alberto D'Urso", {5: 876.42})]})
    rep = confronta_pf(cand, rosa, primo_mese_aperto=5, scad_codici={412})
    classi_att = {a["classe"] for a in rep["attese"]}
    assert "RETTIFICA" in classi_att, rep
    assert rep["errori"] == [], rep
    # la rettifica entra nel totale candidato → quadratura riflette -200
    q = rep["quadrature"][5]
    assert abs(q["candidato"] - (876.42 - 200.0)) < 0.01, q


def test_fornitore_sbagliato():
    """Stesso codice, nome materialmente diverso → ERRORE FORNITORE_SBAGLIATO."""
    cand = _make_wb(
        {
            "Materie Prime e Consumo": [
                (412, "Fornitore Totalmente Diverso", {5: 876.42})
            ]
        }
    )
    rosa = _make_wb({"Materie Prime e Consumo": [(412, "Alberto D'Urso", {5: 876.42})]})
    rep = confronta_pf(cand, rosa, primo_mese_aperto=5, scad_codici={412})
    classi = {e["classe"] for e in rep["errori"]}
    assert "FORNITORE_SBAGLIATO" in classi, rep
