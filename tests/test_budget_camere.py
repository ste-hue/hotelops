"""Budget camere per driver: modello, fonti, foglio. Dati SINTETICI."""

import calendar
from collections import Counter
from datetime import date

import pytest

from verticals.condges.budget_camere import modello as m

CAMERE = [
    {"room_id": "H101", "cod_camera": "DSTA", "tipologia": "Standard"},
    {"room_id": "H102", "cod_camera": "TSTA", "tipologia": "Standard"},
    {"room_id": "H301", "cod_camera": "DSTA", "tipologia": "Standard"},
    {"room_id": "H302", "cod_camera": "DSUP", "tipologia": "Superior"},
    {"room_id": "H201", "cod_camera": "DDLX", "tipologia": "Deluxe"},
]
PROGETTO = {"H301": "Deluxe"}
INV26 = {"Standard": 3, "Superior": 1, "Deluxe": 1}
INV27 = {"Standard": 2, "Superior": 1, "Deluxe": 2}
MAGGIO_GIUGNO = ((5, 1), (6, 30))


def _giorni(mese):
    return [
        date(2026, mese, g) for g in range(1, calendar.monthrange(2026, mese)[1] + 1)
    ]


def _categorie(caricato=date(2026, 6, 10)):
    return [
        {"data": d, "codice": cod, "notti": n, "ricavo": n * p, "caricato": caricato}
        for d in _giorni(5) + _giorni(6)
        for cod, n, p in (("DSTA", 2, 100.0), ("TSTA", 1, 100.0), ("DSUP", 1, 200.0))
    ]


def _pms():
    return (
        [{"data": d, "notti": 4, "ricavo": 520.0} for d in _giorni(5)]
        + [{"data": d, "notti": 4, "ricavo": 520.0} for d in _giorni(6)[:9]]
        + [{"data": d, "notti": 4, "ricavo": None} for d in _giorni(6)[9:]]
    )


def test_giorni_apertura_2027():
    g = m.giorni_apertura(2027, m.APERTURA_2027, m.CHIUSURA_2027)
    assert g == {4: 11, 5: 31, 6: 30, 7: 31, 8: 31, 9: 30, 10: 20}
    assert sum(g.values()) == 184


def test_calendario_estremi_inclusi():
    a, c = m.APERTURA_2027, m.CHIUSURA_2027
    assert not m.in_calendario(date(2026, 4, 19), a, c)
    assert m.in_calendario(date(2026, 4, 20), a, c)
    assert m.in_calendario(date(2026, 10, 20), a, c)
    assert not m.in_calendario(date(2026, 10, 21), a, c)


def test_inventario_con_e_senza_progetto():
    assert m.inventario(CAMERE) == INV26
    assert m.inventario(CAMERE, PROGETTO) == INV27


def test_inventario_rifiuta_camera_o_categoria_sconosciuta():
    with pytest.raises(ValueError, match="H999"):
        m.inventario(CAMERE, {"H999": "Deluxe"})
    with pytest.raises(ValueError, match="Imperial"):
        m.inventario(CAMERE, {"H301": "Imperial"})


def test_progetto_terzo_piano():
    p = m.carica_progetto()
    assert len(p) == 20
    assert Counter(p.values()) == {
        "Deluxe": 9,
        "Classic": 6,
        "Executive": 3,
        "Suite": 2,
    }


def test_notti_2027_conserva_il_totale():
    base = {"Standard": 100, "Superior": 50, "Deluxe": 31}
    out = m.notti_2027(base, INV26, INV27)
    assert out == {"Standard": 67, "Superior": 51, "Deluxe": 63}
    assert sum(out.values()) == sum(base.values())


def test_effetti_somma_esatta():
    e = m.effetti(
        [
            {"notti_base": 100, "prezzo_base": 100.0, "notti": 90, "prezzo": 110.0},
            {"notti_base": 50, "prezzo_base": 200.0, "notti": 70, "prezzo": 200.0},
        ]
    )
    assert e["base"] == pytest.approx(20000)
    assert e["occupazione"] == pytest.approx(1333.3333, abs=1e-3)
    assert e["mix"] == pytest.approx(1666.6667, abs=1e-3)
    assert e["prezzo"] == pytest.approx(900)
    assert e["base"] + e["occupazione"] + e["mix"] + e["prezzo"] == pytest.approx(
        e["budget"]
    )


def test_effetti_categoria_senza_base():
    e = m.effetti(
        [{"notti_base": 0, "prezzo_base": None, "notti": 10, "prezzo": 300.0}]
    )
    assert e["base"] == 0
    assert e["base"] + e["occupazione"] + e["mix"] + e["prezzo"] == pytest.approx(3000)


def test_costruisci_righe_e_mesi():
    dati = m.costruisci(CAMERE, _categorie(), _pms(), PROGETTO, 2026, *MAGGIO_GIUGNO)
    assert len(dati["righe"]) == 2 * 3
    r = {(x["mese"], x["categoria"]): x for x in dati["righe"]}
    # DSTA + TSTA si fondono in Standard
    assert r[(5, "Standard")]["notti_base"] == 93
    assert r[(5, "Standard")]["prezzo_base"] == pytest.approx(100.0)
    assert r[(5, "Standard")]["notti_2027"] == 83
    assert r[(5, "Superior")]["notti_2027"] == 41
    # categoria senza vendite: prezzo vuoto, non zero
    assert r[(5, "Deluxe")]["notti_base"] == 0
    assert r[(5, "Deluxe")]["prezzo_base"] is None
    assert r[(5, "Deluxe")]["camere_2027"] == 2

    mesi = {x["mese"]: x for x in dati["mesi"]}
    assert mesi[5]["stato"] == "osservato"
    assert mesi[5]["raccordo"] == pytest.approx(1.04)
    assert mesi[5]["giorni_2027"] == 31
    assert mesi[5]["notti_reali"] == 124
    assert mesi[5]["ricavo_calendario"] == pytest.approx(31 * 520.0)
    # giugno: file per tipologia del 10/6 → stima, raccordo vuoto
    assert mesi[6]["stato"] == "stima"
    assert mesi[6]["raccordo"] is None


def test_raccordo_vuoto_se_01room_incompleto():
    # file per tipologia esportato dopo fine giugno, ma 01ROOM caricato solo per 9 giorni
    dati = m.costruisci(
        CAMERE, _categorie(date(2026, 7, 5)), _pms(), PROGETTO, 2026, *MAGGIO_GIUGNO
    )
    giugno = next(x for x in dati["mesi"] if x["mese"] == 6)
    assert giugno["stato"] == "stima"
    assert giugno["raccordo"] is None


def test_tipologia_sconosciuta_con_notti_esplode():
    cat = _categorie() + [
        {
            "data": date(2026, 5, 3),
            "codice": "XYZ",
            "notti": 2,
            "ricavo": 300.0,
            "caricato": date(2026, 6, 10),
        }
    ]
    with pytest.raises(ValueError, match="XYZ"):
        m.costruisci(CAMERE, cat, _pms(), PROGETTO, 2026, *MAGGIO_GIUGNO)


def test_tipologia_sconosciuta_senza_notti_e_fuori_calendario_ignorate():
    cat = _categorie() + [
        {
            "data": date(2026, 5, 3),
            "codice": "DEP",
            "notti": 0,
            "ricavo": 0.0,
            "caricato": date(2026, 6, 10),
        },
        {
            "data": date(2026, 7, 3),
            "codice": "DSUP",
            "notti": 9,
            "ricavo": 900.0,
            "caricato": date(2026, 6, 10),
        },
    ]
    dati = m.costruisci(CAMERE, cat, _pms(), PROGETTO, 2026, *MAGGIO_GIUGNO)
    assert sum(x["notti_base"] for x in dati["righe"]) == 4 * 61


from verticals.condges.budget_camere import fonti as f  # noqa: E402


def test_fonti_interrogano_le_tabelle_giuste():
    viste = []

    def finta(sql):
        viste.append(sql)
        return [{"ok": 1}]

    assert f.leggi_camere(finta) == [{"ok": 1}]
    assert f.leggi_categorie(2026, finta) == [{"ok": 1}]
    assert f.leggi_pms(2026, finta) == [{"ok": 1}]
    camere, categorie, pms = viste
    assert "d_camere" in camere and "'HOTEL'" in camere
    assert "f_bookings_tipologia" in categorie and "ricavo_camera" in categorie
    assert "2026" in categorie
    assert "f_pms_statistiche" in pms and "camere_vendute" in pms
    assert "f_produzione_pms" in pms and "'01ROOM'" in pms
    assert "revenue_room" not in pms
    # importo_imponibile è NUMERIC (Decimal in Python): il modello lavora in float
    assert "CAST(SUM(importo_imponibile) AS FLOAT64)" in pms


def test_leggi_categorie_senza_righe_esplode():
    with pytest.raises(ValueError, match="f_bookings_tipologia"):
        f.leggi_categorie(2026, lambda sql: [])


from openpyxl import load_workbook  # noqa: E402

from verticals.condges.budget_camere import foglio  # noqa: E402


@pytest.fixture
def libro(tmp_path):
    dati = m.costruisci(CAMERE, _categorie(), _pms(), PROGETTO, 2026, *MAGGIO_GIUGNO)
    return load_workbook(foglio.scrivi(dati, tmp_path / "budget.xlsx"))


def test_foglio_prezzi(libro):
    ws = libro["Prezzi"]
    assert ws.max_row == 1 + 6
    riga = {ws.cell(1, c).value: ws.cell(2, c).value for c in range(1, 17)}
    # prima riga: maggio, Deluxe (categorie in ordine alfabetico), nessuna vendita
    assert (riga["Mese"], riga["Categoria"]) == (5, "Deluxe")
    assert riga["Prezzo medio 2026"] is None
    assert riga["Prezzo 2027"] == '=IF(H2="","",H2*(1+J2))'
    assert riga["Ricavo base"] == '=IF(H2="",0,G2*H2)'
    assert riga["Ricavo 2027 a prezzi 2026"] == '=IF(H2="",0,I2*H2)'
    assert riga["Ricavo 2027"] == '=IF(K2="",0,I2*K2)'
    assert riga["Aumento %"] == 0
    # seconda riga: maggio, Standard
    assert ws["C3"].value == "Standard"
    assert ws["G3"].value == 93
    assert ws["H3"].value == pytest.approx(100.0)
    assert ws["I3"].value == 83
    # le celle di input sono colorate, quelle calcolate no
    assert ws["J3"].fill.fgColor.rgb.endswith("FFF2CC")
    assert not ws["K3"].fill.fgColor.rgb.endswith("FFF2CC")


def test_foglio_mesi(libro):
    ws = libro["Mesi"]
    assert [ws["A2"].value, ws["A3"].value, ws["A4"].value] == [5, 6, "Totale"]
    assert ws["C2"].value == "osservato" and ws["C3"].value == "stima"
    assert ws["I2"].value == "=H2-G2"
    assert ws["J2"].value == "=SUMIF(Prezzi!$A$2:$A$7,A2,Prezzi!$G$2:$G$7)"
    assert ws["N2"].value == "=IF(J2=0,0,(L2-J2)*K2/J2)"
    assert ws["O2"].value == (
        "=SUMIF(Prezzi!$A$2:$A$7,A2,Prezzi!$M$2:$M$7)-IF(J2=0,0,L2*K2/J2)"
    )
    assert ws["P2"].value == "=M2-SUMIF(Prezzi!$A$2:$A$7,A2,Prezzi!$M$2:$M$7)"
    assert ws["Q2"].value == pytest.approx(1.04)
    assert ws["Q3"].value is None  # giugno è una stima: raccordo vuoto, non zero
    assert ws["R2"].value == '=IF(Q2="","n.d.",M2*Q2)'
    assert ws["R4"].value == '=IF(COUNTBLANK(Q2:Q3)>0,"n.d.",SUM(R2:R3))'
    assert ws["S2"].value == "=L2/(5*E2)"
    assert ws["M4"].value == "=SUM(M2:M3)"


def test_foglio_leggimi(libro):
    testo = " ".join(
        str(r[0].value) for r in libro["Leggimi"].iter_rows() if r[0].value
    )
    assert "1 maggio" in testo and "30 giugno" in testo
    assert "giugno" in testo and "stima" in testo


from datetime import datetime  # noqa: E402

from verticals.condges.budget_camere import base  # noqa: E402


def _finta(sql):
    if "d_camere" in sql:
        return CAMERE
    if "f_bookings_tipologia" in sql:
        return _categorie()
    return _pms()


def test_run_scrive_il_foglio_e_riepiloga(tmp_path, capsys, monkeypatch):
    monkeypatch.setattr(m, "carica_progetto", lambda: PROGETTO)
    path = base.run(
        tmp_path, *MAGGIO_GIUGNO, query=_finta, adesso=datetime(2026, 10, 5, 9, 30)
    )
    assert path == tmp_path / "budget_camere_HOTEL_2027_base_20261005-0930.xlsx"
    assert load_workbook(path).sheetnames == ["Leggimi", "Prezzi", "Mesi"]
    out = capsys.readouterr().out
    assert "maggio" in out and "osservato" in out
    assert "giugno" in out and "stima" in out
    assert str(path) in out


def test_giorno():
    assert base.giorno("04-20") == (4, 20)
    with pytest.raises(ValueError):
        base.giorno("20 aprile")
