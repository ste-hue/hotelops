"""Estrazione blocco previsioni + consuntivi dal PF precedente."""

from __future__ import annotations

from io import BytesIO

import openpyxl

from verticals.condges.pf_generator.previsioni import estrai_entrate, estrai_previsioni


def _pf_precedente_bytes() -> bytes:
    """PF legacy minimale: foglio Materie Prime (nome col typo storico),
    1 riga partite-like (cod fornitore), 1 riga previsione (cod conto PF),
    1 riga senza codice, 1 riga RETTIFICA da ignorare."""
    wb = openpyxl.Workbook()
    wb.active.title = "Piano Finanziario"
    ws = wb.create_sheet("Materie Prime-Consumo ")
    mesi = [
        "GENNAIO",
        "FEBBRAIO",
        "MARZO",
        "APRILE",
        "MAGGIO",
        "GIUGNO",
        "LUGLIO",
        "AGOSTO",
        "SETTEMBRE",
        "OTTOBRE",
        "NOVEMBRE",
        "DICEMBRE",
    ]
    for i, m in enumerate(mesi):
        ws.cell(row=2, column=8 + i, value=m)
    # riga fornitore con partite correnti (cod 92) -> NON previsione
    ws.cell(row=5, column=1, value=92)
    ws.cell(row=5, column=2, value="Amalfi sei esse")
    ws.cell(row=5, column=11, value=350.0)  # APRILE (consuntivo, mese chiuso)
    ws.cell(row=5, column=13, value=999.0)  # GIUGNO (scrittura motore vecchia)
    # riga previsione ricorrente (cod 1076 NON nelle partite correnti)
    ws.cell(row=6, column=1, value=1076)
    ws.cell(row=6, column=2, value="Noleggio Tesla")
    for col in range(12, 20):  # MAG..DIC
        ws.cell(row=6, column=col, value=1030.64)
    # riga manuale senza codice
    ws.cell(row=7, column=2, value="Scorte extra stagione")
    ws.cell(row=7, column=14, value=500.0)  # LUGLIO
    # riga rettifica di un run precedente: MAI estratta
    ws.cell(row=8, column=2, value="RETTIFICA PARTITE/PREVISIONI")
    ws.cell(row=8, column=13, value=-1030.64)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _pf_riepilogo_bytes(primo_mese_aperto: int = 6) -> bytes:
    """PF riepilogo minimale con foglio 'Piano Finanziario'.

    Layout:
        row 2: header mesi (GEN..DIC)
        row 3: SALDO MESE PRECED (da ignorare)
        row 4: Entrate Hotel      giu=150000
        row 5: Entrate CVM        giu=5000, lug=8000
        row 6: TOTALE ENTRATE     (formula/numero)
        row 7: BANCHE (da ignorare)
        row 8: USCITE header (da ignorare)
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Piano Finanziario"
    mesi = [
        "GENNAIO",
        "FEBBRAIO",
        "MARZO",
        "APRILE",
        "MAGGIO",
        "GIUGNO",
        "LUGLIO",
        "AGOSTO",
        "SETTEMBRE",
        "OTTOBRE",
        "NOVEMBRE",
        "DICEMBRE",
    ]
    for i, m in enumerate(mesi):
        ws.cell(row=2, column=3 + i, value=m)  # GEN=col3 .. DIC=col14
    # col di giugno = col3 + 5 = 8, luglio = 9
    col_giu = 3 + 5  # 8
    col_lug = 3 + 6  # 9

    ws.cell(row=3, column=2, value="SALDO MESE PRECED")
    ws.cell(row=3, column=col_giu, value=120000.0)

    ws.cell(row=4, column=2, value="Entrate Hotel")
    ws.cell(row=4, column=col_giu, value=150000.0)

    ws.cell(row=5, column=2, value="Entrate CVM")
    ws.cell(row=5, column=col_giu, value=5000.0)
    ws.cell(row=5, column=col_lug, value=8000.0)

    ws.cell(row=6, column=2, value="TOTALE ENTRATE")
    ws.cell(row=6, column=col_giu, value=155000.0)

    ws.cell(row=7, column=2, value="BANCHE ATTIVE")
    ws.cell(row=7, column=col_giu, value=1000.0)

    ws.cell(row=8, column=2, value="USCITE")
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


class TestEstraiEntrate:
    def test_estrae_entrate_esclude_saldo_totale(self):
        out = estrai_entrate(
            _pf_riepilogo_bytes(),
            primo_mese_aperto=6,
        )
        nomi = [r["nome"] for r in out]
        assert "Entrate Hotel" in nomi
        assert "Entrate CVM" in nomi
        # SALDO e TOTALE e BANCHE esclusi
        assert not any("SALDO" in n.upper() for n in nomi)
        assert not any("TOTALE" in n.upper() for n in nomi)
        assert not any("BANCHE" in n.upper() for n in nomi)

    def test_solo_mesi_aperti(self):
        out = estrai_entrate(_pf_riepilogo_bytes(), primo_mese_aperto=6)
        hotel = next(r for r in out if r["nome"] == "Entrate Hotel")
        cvm = next(r for r in out if r["nome"] == "Entrate CVM")
        # Hotel ha solo giugno
        assert set(hotel["mesi"].keys()) == {6}
        assert hotel["mesi"][6] == 150000.0
        # CVM ha giugno e luglio
        assert set(cvm["mesi"].keys()) == {6, 7}
        assert cvm["mesi"][6] == 5000.0
        assert cvm["mesi"][7] == 8000.0

    def test_primo_mese_aperto_filtra_passato(self):
        # con primo_mese_aperto=7, giugno non deve apparire
        out = estrai_entrate(_pf_riepilogo_bytes(), primo_mese_aperto=7)
        nomi_con_valori = [r["nome"] for r in out]
        # Entrate Hotel aveva solo giugno (< 7) → non inclusa
        assert "Entrate Hotel" not in nomi_con_valori
        # CVM aveva anche luglio → inclusa, solo mese 7
        cvm = next((r for r in out if r["nome"] == "Entrate CVM"), None)
        assert cvm is not None
        assert set(cvm["mesi"].keys()) == {7}


class TestEstraiPrevisioni:
    def test_estrae_solo_righe_non_partite(self):
        out = estrai_previsioni(
            _pf_precedente_bytes(),
            codici_partite={92},
            primo_mese_aperto=5,
        )
        righe = out["USCITE_MATERIE_PRIME"]
        nomi = [r["nome"] for r in righe]
        assert "Noleggio Tesla" in nomi
        assert "Scorte extra stagione" in nomi
        assert "Amalfi sei esse" not in nomi  # ha partite correnti
        assert all("RETTIFICA" not in n.upper() for n in nomi)

    def test_previsioni_solo_mesi_aperti(self):
        out = estrai_previsioni(
            _pf_precedente_bytes(), codici_partite={92}, primo_mese_aperto=5
        )
        tesla = next(
            r for r in out["USCITE_MATERIE_PRIME"] if r["nome"] == "Noleggio Tesla"
        )
        assert tesla["mesi"] == {m: 1030.64 for m in range(5, 13)}
        assert tesla["codice"] == 1076
