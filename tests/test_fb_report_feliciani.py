"""Test verticals/fb/genera_report_feliciani — logica pura (no BQ)."""

from datetime import date, timedelta

import pandas as pd

from verticals.fb.genera_report_feliciani import build_trend_settimanale, build_workbook


def _riga(d: date, servizio: str, coperti: int, ricavi: float) -> dict:
    return {
        "data": d,
        "servizio": servizio,
        "n_comande": 10,
        "coperti": coperti,
        "ricavi": ricavi,
        "coperto_medio": ricavi / coperti,
        "articoli": coperti * 3,
        "articoli_per_coperto": 3.0,
        "bevande_per_coperto": 1.2,
        "food_per_coperto": 1.8,
        "antipasti_per_coperto": 0.5,
        "primi_per_coperto": 0.7,
        "secondi_per_coperto": 0.4,
        "dessert_per_coperto": 0.2,
        "bottiglie_vino_per_10_coperti": 4.0,
    }


def _df_due_settimane_complete() -> pd.DataFrame:
    # Due settimane ISO complete: lun 06/07→dom 12/07 e lun 13/07→dom 19/07 (2026).
    # Ogni giorno pranzo+cena; settimana 2 ricavi +10%.
    rows = []
    for i in range(7):
        for start, base in [(date(2026, 7, 6), 100.0), (date(2026, 7, 13), 110.0)]:
            d = start + timedelta(days=i)
            rows.append(_riga(d, "pranzo", 4, base))
            rows.append(_riga(d, "cena", 6, base * 2))
    return pd.DataFrame(rows)


def test_trend_settimanale_aggrega_e_calcola_delta():
    wk = build_trend_settimanale(_df_due_settimane_complete())
    assert len(wk) == 2
    recente, precedente = wk.iloc[0], wk.iloc[1]  # più recente in alto
    assert precedente["coperti"] == 70 and precedente["ricavi"] == 2100.0
    assert recente["coperti"] == 70 and recente["ricavi"] == 2310.0
    assert recente["delta_coperti_pct"] == 0.0
    assert abs(recente["delta_ricavi_pct"] - 0.1) < 1e-9
    assert pd.isna(precedente["delta_ricavi_pct"])  # prima settimana: nessun confronto
    assert "(parziale)" not in recente["settimana"]


def test_trend_settimana_parziale_marcata_e_senza_delta():
    df = _df_due_settimane_complete()
    # Terza settimana con solo 2 giorni (lun 20/07, mar 21/07) = parziale
    df = pd.concat(
        [
            df,
            pd.DataFrame(
                [
                    _riga(date(2026, 7, 20), "cena", 6, 200.0),
                    _riga(date(2026, 7, 21), "cena", 6, 200.0),
                ]
            ),
        ],
        ignore_index=True,
    )
    wk = build_trend_settimanale(df)
    assert len(wk) == 3
    parziale = wk.iloc[0]
    assert "(parziale)" in parziale["settimana"]
    assert pd.isna(parziale["delta_ricavi_pct"]) and pd.isna(
        parziale["delta_coperti_pct"]
    )
    # Le settimane complete mantengono il loro confronto
    assert abs(wk.iloc[1]["delta_ricavi_pct"] - 0.1) < 1e-9


def test_trend_prima_settimana_parziale_annulla_delta_successivo():
    df = _df_due_settimane_complete()
    # Settimana zero con 1 solo giorno (dom 05/07) = parziale in testa
    df = pd.concat(
        [pd.DataFrame([_riga(date(2026, 7, 5), "cena", 6, 300.0)]), df],
        ignore_index=True,
    )
    wk = build_trend_settimanale(df)
    assert len(wk) == 3
    assert "(parziale)" in wk.iloc[2]["settimana"]
    # La settimana dopo quella parziale non deve mostrare un Δ% fuorviante
    assert pd.isna(wk.iloc[1]["delta_ricavi_pct"])
    assert abs(wk.iloc[0]["delta_ricavi_pct"] - 0.1) < 1e-9


def test_trend_settimanale_vuoto():
    assert build_trend_settimanale(pd.DataFrame()).empty


def test_build_workbook_8_fogli():
    from verticals.fb.genera_report_feliciani import classifica_quadranti

    df = _df_due_settimane_complete()
    wb = build_workbook(
        df_giorni=df,
        df_trend=build_trend_settimanale(df),
        df_coperti=_df_coperti(),
        df_modello=_df_modello(),
        df_reparti=pd.DataFrame(
            [
                {
                    "anno": 2026,
                    "mese": 6,
                    "reparto": "CUCINA",
                    "outlet": "RISTORANTE",
                    "importo": 12000.0,
                }
            ]
        ),
        df_menu=classifica_quadranti(
            pd.DataFrame(
                [
                    {
                        "piatto": "PD.00004",
                        "descrizione": "GNOCCHI",
                        "tipo": "PRIMI PIATTI",
                        "qty": 100.0,
                        "prezzo_medio_netto": 14.0,
                        "costo_unitario": 4.0,
                        "margine_unitario": 10.0,
                        "margine_totale": 1000.0,
                    }
                ]
            )
        ),
    )
    assert wb.sheetnames == [
        "Cruscotto",
        "Giornaliero Servizio",
        "Coperti Completi",
        "Breakfast",
        "Consumi",
        "Menu Engineering",
        "Definizioni",
        "Trend Settimanale",
    ]
    ws = wb["Giornaliero Servizio"]
    assert ws.max_row == 29  # header + 14 giorni × 2 servizi
    assert ws.cell(row=2, column=2).value == "Pranzo"
    assert ws.cell(row=2, column=5).value == 100.0
    # colonne breakdown accodate nella stessa riga
    assert ws.cell(row=1, column=9).value == "Bevande/Cop"


def test_trend_buco_stagionale_annulla_delta():
    df = _df_due_settimane_complete()
    # Settimana completa 6 settimane dopo (buco stagionale in mezzo)
    df = pd.concat(
        [
            df,
            pd.DataFrame(
                [
                    _riga(date(2026, 8, 31) + timedelta(days=i), "cena", 6, 200.0)
                    for i in range(7)
                ]
            ),
        ],
        ignore_index=True,
    )
    wk = build_trend_settimanale(df)
    assert len(wk) == 3
    dopo_buco = wk.iloc[0]
    assert "(parziale)" not in dopo_buco["settimana"]
    assert pd.isna(dopo_buco["delta_ricavi_pct"])  # confronto oltre il buco: mai
    assert abs(wk.iloc[1]["delta_ricavi_pct"] - 0.1) < 1e-9  # consecutive intatte


def _df_coperti() -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "data_servizio": date(2026, 7, 20),
                "tipo_pasto": "BRK",
                "hotel": 120,
                "residence": 8,
                "cvm": 2,
                "esterni": 0,
                "paganti": 130,
                "dipendenti": 0,
                "courtesy_pm": 3,
                "non_paganti": 3,
                "totale": 133,
            },
            {
                "data_servizio": date(2026, 7, 20),
                "tipo_pasto": "DINNER",
                "hotel": 30,
                "residence": 0,
                "cvm": 0,
                "esterni": 5,
                "paganti": 35,
                "dipendenti": 14,
                "courtesy_pm": 2,
                "non_paganti": 16,
                "totale": 51,
            },
        ]
    )


def test_add_sheet_coperti():
    from openpyxl import Workbook

    from verticals.fb.genera_report_feliciani import add_sheet_coperti

    wb = Workbook()
    add_sheet_coperti(wb, _df_coperti())
    ws = wb["Coperti Completi"]
    assert ws.max_row == 3
    header = [ws.cell(1, c).value for c in range(1, 11)]
    assert header == [
        "Data",
        "Pasto",
        "Hotel",
        "Residence",
        "CVM",
        "Esterni",
        "Paganti",
        "Dipendenti",
        "Courtesy/PM",
        "Totale",
    ]
    assert ws.cell(2, 7).value == 130  # paganti BRK


def _df_modello() -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "anno": 2026,
                "mese": 6,
                "periodo": date(2026, 6, 1),
                "outlet": "RISTORANTE",
                "coperti_paganti": 1000,
                "coperti_non_paganti": 200,
                "ricavo_netto": 40000.0,
                "costo_netto": 12000.0,
                "margine": 28000.0,
                "food_cost_pct": 0.30,
                "ricavo_per_coperto": 40.0,
                "costo_per_coperto": 12.0,
                "margine_per_coperto": 28.0,
                "has_costi": True,
                "has_ricavi": True,
            },
            {
                "anno": 2026,
                "mese": 7,
                "periodo": date(2026, 7, 1),
                "outlet": "RISTORANTE",
                "coperti_paganti": 1100,
                "coperti_non_paganti": 180,
                "ricavo_netto": None,
                "costo_netto": None,
                "margine": None,
                "food_cost_pct": None,
                "ricavo_per_coperto": None,
                "costo_per_coperto": None,
                "margine_per_coperto": None,
                "has_costi": False,
                "has_ricavi": False,
            },
        ]
    )


def test_cruscotto_gating_mese_senza_costi():
    from openpyxl import Workbook

    from verticals.fb.genera_report_feliciani import add_sheet_cruscotto

    wb = Workbook()
    add_sheet_cruscotto(wb, _df_modello())
    ws = wb["Cruscotto"]
    assert ws.max_row == 3
    # mese chiuso: margine/cop presente
    assert ws.cell(2, 8).value == 28.0
    # mese senza costi: margine vuoto, mese marcato
    assert ws.cell(3, 8).value is None
    assert "in corso" in str(ws.cell(3, 1).value)


def test_classifica_quadranti_mediane():
    from verticals.fb.genera_report_feliciani import classifica_quadranti

    df = pd.DataFrame(
        [
            {"piatto": "A", "qty": 100.0, "margine_unitario": 10.0},
            {"piatto": "B", "qty": 100.0, "margine_unitario": 2.0},
            {"piatto": "C", "qty": 10.0, "margine_unitario": 10.0},
            {"piatto": "D", "qty": 10.0, "margine_unitario": 2.0},
        ]
    )
    out = classifica_quadranti(df)
    q = dict(zip(out["piatto"], out["quadrante"]))
    assert q == {"A": "Star", "B": "Cavallo", "C": "Enigma", "D": "Cane"}


def test_mensa_staff_senza_etichetta_in_corso():
    from openpyxl import Workbook

    from verticals.fb.genera_report_feliciani import add_sheet_cruscotto

    righe = _df_modello().to_dict("records")
    righe.append(
        {
            "anno": 2026,
            "mese": 6,
            "periodo": date(2026, 6, 1),
            "outlet": "MENSA_STAFF",
            "coperti_paganti": 0,
            "coperti_non_paganti": 900,
            "ricavo_netto": None,
            "costo_netto": None,
            "margine": None,
            "food_cost_pct": None,
            "ricavo_per_coperto": None,
            "costo_per_coperto": None,
            "margine_per_coperto": None,
            "has_costi": False,
            "has_ricavi": False,
        }
    )
    df = pd.DataFrame(righe)
    wb = Workbook()
    add_sheet_cruscotto(wb, df)
    ws = wb["Cruscotto"]
    labels = {
        ws.cell(r, 2).value: ws.cell(r, 1).value for r in range(2, ws.max_row + 1)
    }
    assert "in corso" not in labels["MENSA_STAFF"]  # costo non separabile ≠ in arrivo
    assert "in corso" in labels["RISTORANTE"] or labels["RISTORANTE"] == "2026-06"
