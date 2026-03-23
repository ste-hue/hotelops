"""Tests for ingest.classify — content-based file classifier.

Tests every detector with synthetic files (no real data needed).
Uses tmp_path for all file I/O.
"""

import csv
import re
from pathlib import Path

import pytest

from ingest.classify import (
    LIFECYCLE_APPEND,
    LIFECYCLE_SNAPSHOT,
    ClassificationResult,
    classify,
    classify_batch,
    detect_accodamenti,
    detect_banca,
    detect_bilancino,
    detect_coperti,
    detect_economato,
    detect_gasparotto,
    detect_movimenti_contabili,
    detect_partite_fornitori,
    detect_piano_finanziario,
    detect_scheda_contabile,
    infer_banca,
    infer_societa,
    infer_societa_from_cc,
    route_file,
)


# ── Helpers ──────────────────────────────────────────────────────────────────


def _write_csv_file(
    path: Path, headers: list[str], rows: list[list[str]], delimiter: str = ","
):
    """Write a simple CSV file."""
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=delimiter)
        w.writerow(headers)
        for row in rows:
            w.writerow(row)


def _write_xlsx_file(path: Path, headers: list[str], rows: list[list], sheet_name: str = "Sheet1"):
    """Write a minimal XLSX file using openpyxl."""
    from openpyxl import Workbook

    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.append(headers)
    for row in rows:
        ws.append(row)
    wb.save(path)


def _write_xlsx_multi_sheet(path: Path, sheets: dict[str, tuple[list, list]]):
    """Write XLSX with multiple named sheets. sheets = {name: (headers, rows)}."""
    from openpyxl import Workbook

    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    first = True
    for name, (headers, rows) in sheets.items():
        if first:
            ws = wb.active
            ws.title = name
            first = False
        else:
            ws = wb.create_sheet(name)
        ws.append(headers)
        for row in rows:
            ws.append(row)
    wb.save(path)


def _write_txt(path: Path, content: str):
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding="latin-1") as f:
        f.write(content)


# ── infer_societa ────────────────────────────────────────────────────────────


class TestInferSocieta:
    def test_from_filename_orti(self):
        assert infer_societa("ORTI_MPS_20260322.xlsx") == "ORTI"

    def test_from_filename_intur(self):
        assert infer_societa("INTUR_Sella_20260322.xlsx") == "INTUR"

    def test_from_parent_dir(self, tmp_path):
        d = tmp_path / "datahub" / "banche" / "ORTI"
        d.mkdir(parents=True)
        f = d / "some_file.xlsx"
        assert infer_societa(f.name, f) == "ORTI"

    def test_unknown_returns_none(self):
        assert infer_societa("random_file.xlsx") is None


# ── infer_banca ──────────────────────────────────────────────────────────────


class TestInferBanca:
    def test_cc_pattern_orti(self):
        assert infer_banca("20260101_Conto_Banca_Cc1_Saldo.csv", societa="ORTI") == "INTESA"

    def test_cc_pattern_intur(self):
        assert infer_banca("20260101_Conto_Banca_Cc4_Saldo.csv", societa="INTUR") == "BCP"

    def test_mps_kross_before_mps(self):
        assert infer_banca("ORTI_KROSS_20260322.xlsx") == "MPS_KROSS"

    def test_mps_name(self):
        assert infer_banca("ORTI_MPS_20260322.xlsx") == "MPS"

    def test_intesa_name(self):
        assert infer_banca("IntesaSanpaolo_export.xlsx") == "INTESA"

    def test_sella_name(self):
        assert infer_banca("BancaSella_export.csv") == "SELLA"

    def test_bcp_name(self):
        assert infer_banca("BCP_movimenti.xlsx") == "BCP"

    def test_unknown_returns_none(self):
        assert infer_banca("random.xlsx") is None


# ── infer_societa_from_cc ────────────────────────────────────────────────────


class TestInferSocietaFromCc:
    def test_cc1_intesa_is_orti(self):
        assert infer_societa_from_cc("20260101_Conto_BancaIntesa_Cc1_Saldo.csv", banca="INTESA") == "ORTI"

    def test_cc4_bcp_is_intur(self):
        assert infer_societa_from_cc("file_Cc4_export.csv", banca="BCP") == "INTUR"

    def test_no_cc_returns_none(self):
        assert infer_societa_from_cc("random.csv", banca="MPS") is None


# ── detect_scheda_contabile ──────────────────────────────────────────────────


class TestDetectSchedaContabile:
    def test_csv_saldo_udc(self, tmp_path):
        f = tmp_path / "ORTI_scheda.csv"
        _write_csv_file(
            f,
            ["Data", "Registr.", "Dare", "Avere", "Saldo in UdC"],
            [["01/01/2026", "R001", "100.00", "0.00", "100.00"]],
            delimiter=";",
        )
        result = detect_scheda_contabile(f)
        assert result is not None
        assert result.category == "scheda_contabile"
        assert result.societa == "ORTI"
        assert result.confidence >= 0.90

    def test_csv_dare_avere_saldo(self, tmp_path):
        f = tmp_path / "INTUR_ledger.csv"
        _write_csv_file(
            f,
            ["Data", "Registr.", "Dare", "Avere", "Saldo"],
            [["01/01/2026", "R001", "0.00", "500.00", "500.00"]],
            delimiter=";",
        )
        result = detect_scheda_contabile(f)
        assert result is not None
        assert result.category == "scheda_contabile"
        assert result.societa == "INTUR"

    def test_csv_esolver_naming(self, tmp_path):
        f = tmp_path / "20260101_Conto_BancaIntesaSanpaolo_Cc1_Saldo51520-51_Esercizio2026.csv"
        _write_csv_file(
            f,
            ["Col1", "Col2", "Col3"],
            [["a", "b", "c"]],
            delimiter=";",
        )
        result = detect_scheda_contabile(f)
        assert result is not None
        assert result.category == "scheda_contabile"
        assert result.banca == "INTESA"
        assert result.societa == "ORTI"  # Cc1 + INTESA → ORTI via reverse lookup

    def test_xlsx_dare_avere_saldo_registr(self, tmp_path):
        f = tmp_path / "ORTI_scheda.xlsx"
        _write_xlsx_file(
            f,
            ["Data", "Registr.", "Dare", "Avere", "Saldo in UdC"],
            [["2026-01-01", "R001", 100.0, 0.0, 100.0]],
        )
        result = detect_scheda_contabile(f)
        assert result is not None
        assert result.category == "scheda_contabile"

    def test_lifecycle_is_append(self, tmp_path):
        f = tmp_path / "ORTI_scheda.csv"
        _write_csv_file(
            f,
            ["Dare", "Avere", "Saldo in UdC"],
            [["100", "0", "100"]],
            delimiter=";",
        )
        result = detect_scheda_contabile(f)
        assert result is not None
        assert result.lifecycle == LIFECYCLE_APPEND

    def test_non_matching_csv_returns_none(self, tmp_path):
        f = tmp_path / "random.csv"
        _write_csv_file(f, ["Name", "Age"], [["John", "30"]])
        assert detect_scheda_contabile(f) is None


# ── detect_movimenti_contabili ───────────────────────────────────────────────


class TestDetectMovimentiContabili:
    def test_listamovcont_filename(self, tmp_path):
        """XLS with LISTAMOVCONT in name should match."""
        f = tmp_path / "ORTI_LISTAMOVCONT.xls"
        # Write a minimal binary that xlrd won't crash on — but we test filename match
        f.write_bytes(b"\xd0\xcf\x11\xe0")  # OLE magic bytes
        result = detect_movimenti_contabili(f)
        assert result is not None
        assert result.category == "movimenti_contabili"
        assert result.societa == "ORTI"
        assert result.confidence >= 0.90

    def test_wrong_extension_returns_none(self, tmp_path):
        f = tmp_path / "LISTAMOVCONT.xlsx"
        f.touch()
        assert detect_movimenti_contabili(f) is None

    def test_non_matching_xls_returns_none(self, tmp_path):
        f = tmp_path / "random.xls"
        f.write_bytes(b"\xd0\xcf\x11\xe0")
        assert detect_movimenti_contabili(f) is None


# ── detect_partite_fornitori ─────────────────────────────────────────────────


class TestDetectPartiteFornitori:
    def test_filename_match(self, tmp_path):
        f = tmp_path / "ORTI_PARTITE_FORNITORI.xlsx"
        _write_xlsx_file(f, ["Col"], [[1]])
        result = detect_partite_fornitori(f)
        assert result is not None
        assert result.category == "partite_fornitori"
        assert result.lifecycle == LIFECYCLE_SNAPSHOT
        assert result.societa == "ORTI"
        assert result.confidence >= 0.90

    def test_content_fornitore_scadenza(self, tmp_path):
        f = tmp_path / "export.xlsx"
        _write_xlsx_file(
            f,
            ["Fornitore", "Scadenza", "Importo", "Residuo"],
            [["Acme Srl", "2026-04-15", 1000, 500]],
        )
        result = detect_partite_fornitori(f)
        assert result is not None
        assert result.category == "partite_fornitori"
        assert result.lifecycle == LIFECYCLE_SNAPSHOT

    def test_content_ragsociale_residuo(self, tmp_path):
        f = tmp_path / "INTUR_situazione.xlsx"
        _write_xlsx_file(
            f,
            ["Cod.", "Rag.Sociale", "Residuo", "Data"],
            [["001", "Fornitore Srl", 2500, "2026-03-01"]],
        )
        result = detect_partite_fornitori(f)
        assert result is not None
        assert result.societa == "INTUR"

    def test_wrong_extension_returns_none(self, tmp_path):
        f = tmp_path / "partite.csv"
        f.touch()
        assert detect_partite_fornitori(f) is None

    def test_non_matching_xlsx_returns_none(self, tmp_path):
        f = tmp_path / "random.xlsx"
        _write_xlsx_file(f, ["Name", "Age"], [["John", 30]])
        assert detect_partite_fornitori(f) is None


# ── detect_bilancino ─────────────────────────────────────────────────────────


class TestDetectBilancino:
    def test_esolver_filename(self, tmp_path):
        f = tmp_path / "Marzo_ESOLVER.xls"
        f.write_bytes(b"\xd0\xcf\x11\xe0")
        result = detect_bilancino(f)
        assert result is not None
        assert result.category == "bilancino"
        assert result.lifecycle == LIFECYCLE_SNAPSHOT

    def test_wrong_extension_returns_none(self, tmp_path):
        f = tmp_path / "ESOLVER.xlsx"
        f.touch()
        assert detect_bilancino(f) is None


# ── detect_gasparotto ────────────────────────────────────────────────────────


class TestDetectGasparotto:
    def test_master_completo_filename(self, tmp_path):
        f = tmp_path / "Master Completo Indici 2025 ORTI SRL_Budget26.xlsx"
        _write_xlsx_file(f, ["Col"], [[1]])
        result = detect_gasparotto(f)
        assert result is not None
        assert result.category == "gasparotto"
        assert result.lifecycle == LIFECYCLE_SNAPSHOT
        assert result.societa == "ORTI"
        assert result.confidence >= 0.90

    def test_budget_ce_sheets(self, tmp_path):
        f = tmp_path / "budget_2026.xlsx"
        _write_xlsx_multi_sheet(f, {
            "Budget": (["Mese", "Importo"], [["Gen", 1000]]),
            "Conto Economico": (["Voce", "Totale"], [["Ricavi", 5000]]),
        })
        result = detect_gasparotto(f)
        assert result is not None
        assert result.category == "gasparotto"
        assert result.confidence >= 0.80

    def test_wrong_extension_returns_none(self, tmp_path):
        f = tmp_path / "Master Completo.xls"
        f.write_bytes(b"\xd0\xcf\x11\xe0")
        assert detect_gasparotto(f) is None


# ── detect_piano_finanziario ─────────────────────────────────────────────────


class TestDetectPianoFinanziario:
    def test_filename_match(self, tmp_path):
        f = tmp_path / "ORTI_Piano_Finanziario_03_mar2026.xlsx"
        _write_xlsx_file(f, ["Voce"], [["Entrate Hotel"]])
        result = detect_piano_finanziario(f)
        assert result is not None
        assert result.category == "piano_finanziario"
        assert result.lifecycle == LIFECYCLE_SNAPSHOT
        assert result.societa == "ORTI"

    def test_sheet_name_match(self, tmp_path):
        f = tmp_path / "financial_plan.xlsx"
        _write_xlsx_multi_sheet(f, {
            "Piano Finanziario": (["Voce", "Gen", "Feb"], [["Hotel", 100, 120]]),
            "Note": (["Nota"], [["test"]]),
        })
        result = detect_piano_finanziario(f)
        assert result is not None
        assert result.category == "piano_finanziario"
        assert result.confidence >= 0.90

    def test_wrong_extension_returns_none(self, tmp_path):
        f = tmp_path / "Piano_Finanziario.csv"
        f.touch()
        assert detect_piano_finanziario(f) is None


# ── detect_banca ─────────────────────────────────────────────────────────────


class TestDetectBanca:
    def test_sella_csv(self, tmp_path):
        f = tmp_path / "INTUR" / "sella_export.csv"
        _write_csv_file(
            f,
            ["Codice identificativo", "Data operazione", "Debito", "Credito", "Saldo"],
            [["TX001", "01/01/2026", "100.00", "", "900.00"]],
        )
        result = detect_banca(f)
        assert result is not None
        assert result.category == "banca"
        assert result.banca == "SELLA"
        assert result.lifecycle == LIFECYCLE_APPEND

    def test_intesa_xlsx(self, tmp_path):
        f = tmp_path / "ORTI_intesa.xlsx"
        _write_xlsx_file(
            f,
            ["Data Contabile", "Data Valuta", "Descrizione", "Dare", "Avere"],
            [["2026-01-01", "2026-01-02", "Bonifico", 0, 500]],
        )
        result = detect_banca(f)
        assert result is not None
        assert result.category == "banca"
        assert result.banca == "INTESA"
        assert result.societa == "ORTI"

    def test_mps_xlsx(self, tmp_path):
        f = tmp_path / "MPS_estratto.xlsx"
        _write_xlsx_file(
            f,
            ["Data Cont.", "Data Val.", "Descrizione", "Importo(€)"],
            [["01/01/2026", "02/01/2026", "Commissioni", -50.0]],
        )
        result = detect_banca(f)
        assert result is not None
        assert result.category == "banca"
        assert result.banca == "MPS"

    def test_non_matching_csv_returns_none(self, tmp_path):
        f = tmp_path / "products.csv"
        _write_csv_file(f, ["Product", "Price"], [["Widget", "9.99"]])
        assert detect_banca(f) is None


# ── detect_accodamenti ───────────────────────────────────────────────────────


class TestDetectAccodamenti:
    def test_hotel_corrispettivi(self, tmp_path):
        f = tmp_path / "H_202603_Corrispettivi.txt"
        _write_txt(f, "1001|2026-03-01|Hotel|150.00\n1002|2026-03-02|Hotel|200.00\n")
        result = detect_accodamenti(f)
        assert result is not None
        assert result.category == "accodamenti"
        assert result.societa == "ORTI"
        assert result.confidence >= 0.90
        assert "hotel" in result.file_type

    def test_residence_fatture(self, tmp_path):
        f = tmp_path / "R_202603_Fatture.txt"
        _write_txt(f, "F001|2026-03-01|Residence|500.00\n")
        result = detect_accodamenti(f)
        assert result is not None
        assert "residence" in result.file_type

    def test_cvm_movimenti(self, tmp_path):
        f = tmp_path / "C_202603_Movimenti.txt"
        _write_txt(f, "M001|2026-03-01|CVM|100.00|Extra\n")
        result = detect_accodamenti(f)
        assert result is not None
        assert "cvm" in result.file_type

    def test_pipe_fallback(self, tmp_path):
        f = tmp_path / "unknown.txt"
        _write_txt(f, "A|B|C|D|E\n1|2|3|4|5\n")
        result = detect_accodamenti(f)
        assert result is not None
        assert result.confidence < 0.70  # lower confidence for fallback

    def test_non_txt_returns_none(self, tmp_path):
        f = tmp_path / "accodamenti.csv"
        f.touch()
        assert detect_accodamenti(f) is None


# ── detect_coperti ───────────────────────────────────────────────────────────


class TestDetectCoperti:
    def test_filename_copert(self, tmp_path):
        f = tmp_path / "COPERTI_Marzo2026.csv"
        _write_csv_file(f, ["Data", "Tot"], [["2026-03-01", "50"]])
        result = detect_coperti(f)
        assert result is not None
        assert result.category == "coperti"
        assert result.lifecycle == LIFECYCLE_APPEND

    def test_meal_columns_csv(self, tmp_path):
        f = tmp_path / "daily_data.csv"
        _write_csv_file(
            f,
            ["Date", "Breakfast", "Lunch", "Dinner"],
            [["2026-03-01", "30", "45", "60"]],
        )
        result = detect_coperti(f)
        assert result is not None
        assert result.category == "coperti"

    def test_scarico_sheets(self, tmp_path):
        f = tmp_path / "covers.xlsx"
        _write_xlsx_multi_sheet(f, {
            "SCARICO_BRK": (["Date", "Count"], [["2026-03-01", 30]]),
            "SCARICO_LUNCH": (["Date", "Count"], [["2026-03-01", 45]]),
        })
        result = detect_coperti(f)
        assert result is not None
        assert result.category == "coperti"

    def test_non_matching_returns_none(self, tmp_path):
        f = tmp_path / "staff.csv"
        _write_csv_file(f, ["Name", "Role"], [["Mario", "Chef"]])
        assert detect_coperti(f) is None


# ── detect_economato ─────────────────────────────────────────────────────────


class TestDetectEconomato:
    def test_filename_consumi(self, tmp_path):
        f = tmp_path / "CONSUMI_Marzo.xlsx"
        _write_xlsx_file(f, ["Codice", "Descrizione", "Quantita"], [["P001", "Pasta", 10]])
        result = detect_economato(f)
        assert result is not None
        assert result.category == "economato"
        assert result.lifecycle == LIFECYCLE_APPEND

    def test_content_codice_quantita(self, tmp_path):
        f = tmp_path / "export_reparto.xlsx"
        _write_xlsx_file(
            f,
            ["Codice", "Descrizione", "Quantita", "Euro"],
            [["M001", "Mozzarella", 20, 80.0]],
        )
        result = detect_economato(f)
        assert result is not None
        assert result.category == "economato"

    def test_wrong_extension_returns_none(self, tmp_path):
        f = tmp_path / "consumi.csv"
        f.touch()
        assert detect_economato(f) is None


# ── classify (master function) ───────────────────────────────────────────────


class TestClassify:
    def test_nonexistent_file(self, tmp_path):
        f = tmp_path / "nonexistent.xlsx"
        result = classify(f)
        assert result.category == "error"
        assert result.confidence == 0.0

    def test_unsupported_extension(self, tmp_path):
        f = tmp_path / "photo.jpg"
        f.touch()
        result = classify(f)
        assert result.category == "unknown"

    def test_unrecognized_content(self, tmp_path):
        f = tmp_path / "mystery.csv"
        _write_csv_file(f, ["Alpha", "Beta", "Gamma"], [["x", "y", "z"]])
        result = classify(f)
        assert result.category == "unknown"

    def test_scheda_contabile_via_classify(self, tmp_path):
        f = tmp_path / "ORTI_scheda.csv"
        _write_csv_file(
            f,
            ["Data", "Registr.", "Dare", "Avere", "Saldo in UdC"],
            [["01/01/2026", "R001", "100", "0", "100"]],
            delimiter=";",
        )
        result = classify(f)
        assert result.category == "scheda_contabile"

    def test_tilde_expansion(self):
        """classify() should handle ~ in paths without crashing."""
        result = classify(Path("~/nonexistent_test_file_xyz.csv"))
        assert result.category == "error"  # File doesn't exist, but no crash


# ── classify_batch ───────────────────────────────────────────────────────────


class TestClassifyBatch:
    def test_multiple_files(self, tmp_path):
        f1 = tmp_path / "H_202603_Corrispettivi.txt"
        _write_txt(f1, "1001|2026-03-01|Hotel|150.00\n")

        f2 = tmp_path / "ORTI_PARTITE_FORNITORI.xlsx"
        _write_xlsx_file(f2, ["Col"], [[1]])

        results = classify_batch([f1, f2])
        assert len(results) == 2
        assert results[0].category == "accodamenti"
        assert results[1].category == "partite_fornitori"


# ── route_file ───────────────────────────────────────────────────────────────


class TestRouteFile:
    def test_dry_run_returns_path(self, tmp_path):
        f = tmp_path / "source.csv"
        f.touch()
        result = ClassificationResult(
            file_path=f,
            file_type="test",
            category="test",
            canonical_name="canonical.csv",
            dest_folder="test_dir",
            confidence=0.9,
        )
        datahub = tmp_path / "datahub"
        dest = route_file(result, datahub, dry_run=True)
        assert dest is not None
        assert "canonical.csv" in str(dest)
        # Should NOT create the file in dry-run
        assert not (datahub / "test_dir" / "canonical.csv").exists()

    def test_actual_copy(self, tmp_path):
        f = tmp_path / "source.csv"
        f.write_text("test content")
        result = ClassificationResult(
            file_path=f,
            file_type="test",
            category="test",
            canonical_name="canonical.csv",
            dest_folder="test_dir",
            confidence=0.9,
        )
        datahub = tmp_path / "datahub"
        dest = route_file(result, datahub, dry_run=False)
        assert dest is not None
        assert dest.exists()
        assert dest.read_text() == "test content"

    def test_avoids_overwrite(self, tmp_path):
        f = tmp_path / "source.csv"
        f.write_text("new content")
        datahub = tmp_path / "datahub"
        dest_dir = datahub / "test_dir"
        dest_dir.mkdir(parents=True)
        (dest_dir / "canonical.csv").write_text("old content")

        result = ClassificationResult(
            file_path=f,
            file_type="test",
            category="test",
            canonical_name="canonical.csv",
            dest_folder="test_dir",
            confidence=0.9,
        )
        dest = route_file(result, datahub, dry_run=False)
        assert dest is not None
        assert dest.name == "canonical_1.csv"
        assert (dest_dir / "canonical.csv").read_text() == "old content"

    def test_unknown_category_returns_none(self, tmp_path):
        f = tmp_path / "source.csv"
        f.touch()
        result = ClassificationResult(
            file_path=f,
            file_type="unknown",
            category="unknown",
            confidence=0.0,
        )
        assert route_file(result, tmp_path / "datahub") is None


# ── Lifecycle assignments ────────────────────────────────────────────────────


class TestLifecycle:
    """Verify APPEND vs SNAPSHOT assignments across all file types."""

    def test_append_types(self, tmp_path):
        """Banca, scheda_contabile, movimenti, accodamenti, coperti, economato = APPEND."""
        # Scheda contabile
        f = tmp_path / "ORTI_scheda.csv"
        _write_csv_file(f, ["Dare", "Avere", "Saldo in UdC"], [["100", "0", "100"]], delimiter=";")
        r = detect_scheda_contabile(f)
        assert r is not None
        assert r.lifecycle == LIFECYCLE_APPEND

        # Accodamenti
        f2 = tmp_path / "H_202603_Corrispettivi.txt"
        _write_txt(f2, "1|2|3|4\n")
        r2 = detect_accodamenti(f2)
        assert r2 is not None
        assert r2.lifecycle == LIFECYCLE_APPEND

    def test_snapshot_types(self, tmp_path):
        """Partite fornitori, bilancino, gasparotto, piano finanziario = SNAPSHOT."""
        # Partite fornitori
        f = tmp_path / "ORTI_PARTITE_FORNITORI.xlsx"
        _write_xlsx_file(f, ["Col"], [[1]])
        r = detect_partite_fornitori(f)
        assert r is not None
        assert r.lifecycle == LIFECYCLE_SNAPSHOT

        # Gasparotto
        f2 = tmp_path / "Master Completo ORTI.xlsx"
        _write_xlsx_file(f2, ["Col"], [[1]])
        r2 = detect_gasparotto(f2)
        assert r2 is not None
        assert r2.lifecycle == LIFECYCLE_SNAPSHOT

        # Piano finanziario
        f3 = tmp_path / "ORTI_Piano_Finanziario.xlsx"
        _write_xlsx_file(f3, ["Voce"], [["Hotel"]])
        r3 = detect_piano_finanziario(f3)
        assert r3 is not None
        assert r3.lifecycle == LIFECYCLE_SNAPSHOT
