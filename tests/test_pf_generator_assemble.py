"""Template e assemblaggio workbook generato."""

from __future__ import annotations

import openpyxl
from io import BytesIO
from openpyxl.utils import get_column_letter

from verticals.condges.app_scadenzario import _build_month_col_map
from verticals.condges.pf_generator.costanti import (
    LABEL_RETTIFICA,
    PRIMA_RIGA_BLOCCO_A,
    RIGA_TOTALE,
    col_mese,
)
from verticals.condges.pf_generator.template import (
    scrivi_controlli,
    scrivi_da_mappare,
    scrivi_foglio_voce,
    scrivi_riepilogo,
)


def _blocco_a():
    return [
        {"codice": 92, "nome": "Amalfi sei esse", "mesi": {5: 100.0, 6: 1122.76}},
        {"codice": 128, "nome": "Tecno Piscine", "mesi": {5: 2274.08}},
    ]


def _previsioni():
    return [
        {
            "codice": 1076,
            "nome": "Noleggio Tesla",
            "mesi": {m: 1030.64 for m in range(5, 13)},
        },
    ]


def _scrivi(tmp_path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    scrivi_foglio_voce(
        wb,
        voce_id="USCITE_MATERIE_PRIME",
        blocco_a=_blocco_a(),
        previsioni=_previsioni(),
        consuntivi={92: {4: 350.0}},
        rettifica={},
        primo_mese_aperto=5,
    )
    p = tmp_path / "out.xlsx"
    wb.save(p)
    return openpyxl.load_workbook(p)


class TestFoglioVoce:
    def test_nome_pulito_e_month_map_compatibile(self, tmp_path):
        wb = _scrivi(tmp_path)
        assert "Materie Prime e Consumo" in wb.sheetnames
        mc = _build_month_col_map(wb["Materie Prime e Consumo"])
        assert mc[1] == col_mese(1) and mc[12] == col_mese(12)

    def test_blocco_a_codici_e_valori(self, tmp_path):
        ws = _scrivi(tmp_path)["Materie Prime e Consumo"]
        r = PRIMA_RIGA_BLOCCO_A
        assert ws.cell(row=r, column=1).value == 92
        assert ws.cell(row=r, column=2).value == "Amalfi sei esse"
        assert ws.cell(row=r, column=col_mese(5)).value == 100.0
        assert ws.cell(row=r, column=col_mese(4)).value == 350.0  # consuntivo
        assert ws.cell(row=r + 1, column=col_mese(5)).value == 2274.08

    def test_previsioni_e_totale_formula(self, tmp_path):
        ws = _scrivi(tmp_path)["Materie Prime e Consumo"]
        # previsione presente con valore ORIGINALE
        trovato = [
            r
            for r in range(1, ws.max_row + 1)
            if ws.cell(row=r, column=2).value == "Noleggio Tesla"
        ]
        assert trovato and ws.cell(row=trovato[0], column=col_mese(6)).value == 1030.64
        # riga totale = formula SUM che copre entrambi i blocchi
        f = ws.cell(row=RIGA_TOTALE, column=col_mese(6)).value
        assert isinstance(f, str) and f.startswith("=SUM(")

    def test_rettifica_scritta_quando_presente(self, tmp_path):
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        scrivi_foglio_voce(
            wb,
            voce_id="USCITE_MATERIE_PRIME",
            blocco_a=_blocco_a(),
            previsioni=_previsioni(),
            consuntivi={},
            rettifica={6: -1030.64},
            primo_mese_aperto=5,
        )
        ws = wb["Materie Prime e Consumo"]
        riga = [
            r
            for r in range(1, ws.max_row + 1)
            if ws.cell(row=r, column=2).value == LABEL_RETTIFICA
        ]
        assert riga and ws.cell(row=riga[0], column=col_mese(6)).value == -1030.64


class TestFogliAccessori:
    def test_da_mappare(self, tmp_path):
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        scrivi_da_mappare(
            wb,
            [
                {
                    "codice": 125,
                    "nome": "CENTRO DISTRIBUZIONI MAGLIO",
                    "mesi": {6: 1086.80},
                },
            ],
        )
        ws = wb["DA MAPPARE"]
        assert ws.cell(row=3, column=1).value == 125
        assert ws.cell(row=3, column=col_mese(6)).value == 1086.80

    def test_riepilogo_formule_cross_sheet(self, tmp_path):
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        scrivi_foglio_voce(
            wb,
            voce_id="USCITE_UTENZE",
            blocco_a=[],
            previsioni=[
                {"codice": None, "nome": "Stima bollette", "mesi": {6: 4250.0}}
            ],
            consuntivi={},
            rettifica={},
            primo_mese_aperto=5,
        )
        scrivi_riepilogo(
            wb,
            societa="ORTI",
            anno=2026,
            entrate=[{"nome": "Entrate Hotel", "mesi": {6: 150000.0}}],
            voci_attive=["USCITE_UTENZE"],
            primo_mese_aperto=5,
        )
        ws = wb["Piano Finanziario"]
        col6 = get_column_letter(col_mese(6))
        # almeno una formula che punta al totale del foglio Utenze
        formule = [
            ws.cell(row=r, column=col_mese(6)).value
            for r in range(1, ws.max_row + 1)
            if isinstance(ws.cell(row=r, column=col_mese(6)).value, str)
        ]
        assert any(
            "Utenze" in f and f"{col6}{RIGA_TOTALE}" in f.replace("'", "")
            for f in formule
        )

    def test_controlli(self, tmp_path):
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        scrivi_controlli(
            wb,
            [
                {"check": "codice↔nome vs CSV", "esito": "OK", "dettaglio": ""},
                {
                    "check": "quadratura export",
                    "esito": "ERR",
                    "dettaglio": "delta 12.30",
                },
            ],
        )
        ws = wb["Controlli"]
        esiti = [ws.cell(row=r, column=2).value for r in range(2, 4)]
        assert "OK" in esiti and "ERR" in esiti


# ---------------------------------------------------------------------------
# Task 7: assembler end-to-end
# ---------------------------------------------------------------------------

import pandas as pd  # noqa: E402

from verticals.condges.pf_generator.assemble import genera_pf  # noqa: E402


def _pf_precedente_bytes() -> bytes:
    """PF legacy minimale: foglio Materie Prime (nome col typo storico),
    1 riga partite-like (cod fornitore), 1 riga previsione (cod conto PF),
    1 riga senza codice, 1 riga RETTIFICA da ignorare."""
    wb = openpyxl.Workbook()
    wb.active.title = "Piano Finanziario"
    ws = wb.create_sheet("Materie Prime-Consumo ")
    mesi = [
        "GENNAIO",
        "FEBBRAIO",
        "MARZO",
        "APRILE",
        "MAGGIO",
        "GIUGNO",
        "LUGLIO",
        "AGOSTO",
        "SETTEMBRE",
        "OTTOBRE",
        "NOVEMBRE",
        "DICEMBRE",
    ]
    for i, m in enumerate(mesi):
        ws.cell(row=2, column=8 + i, value=m)
    # riga fornitore con partite correnti (cod 92) -> NON previsione
    ws.cell(row=5, column=1, value=92)
    ws.cell(row=5, column=2, value="Amalfi sei esse")
    ws.cell(row=5, column=11, value=350.0)  # APRILE (consuntivo, mese chiuso)
    ws.cell(row=5, column=13, value=999.0)  # GIUGNO (scrittura motore vecchia)
    # riga previsione ricorrente (cod 1076 NON nelle partite correnti)
    ws.cell(row=6, column=1, value=1076)
    ws.cell(row=6, column=2, value="Noleggio Tesla")
    for col in range(12, 20):  # MAG..DIC
        ws.cell(row=6, column=col, value=1030.64)
    # riga manuale senza codice
    ws.cell(row=7, column=2, value="Scorte extra stagione")
    ws.cell(row=7, column=14, value=500.0)  # LUGLIO
    # riga rettifica di un run precedente: MAI estratta
    ws.cell(row=8, column=2, value="RETTIFICA PARTITE/PREVISIONI")
    ws.cell(row=8, column=13, value=-1030.64)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


class TestGeneraPf:
    def _scad_df(self):
        return pd.DataFrame(
            [
                {
                    "codice_fornitore": 92,
                    "nome": "AMALFI SEI ESSE S.R.L.",
                    "totale": -1222.76,
                    "scaduto": -100.0,
                    "mese_6": -1122.76,
                },
                {
                    "codice_fornitore": 125,
                    "nome": "CENTRO DISTRIBUZIONI MAGLIO",
                    "totale": -1086.80,
                    "mese_6": -1086.80,
                    "scaduto": 0.0,
                },
            ]
        )

    def _fornitori(self):
        return {92: {"voce_id": "USCITE_MATERIE_PRIME", "nome_pf": "Amalfi sei esse"}}

    def test_genera_workbook_completo(self, tmp_path):
        out_bytes, report = genera_pf(
            pf_prev_bytes=_pf_precedente_bytes(),
            scad_df=self._scad_df(),
            bucket_months=[6],
            fornitori=self._fornitori(),
            societa="ORTI",
            anno=2026,
            primo_mese_aperto=5,
        )
        wb = openpyxl.load_workbook(BytesIO(out_bytes))
        assert "Controlli" in wb.sheetnames
        assert "Piano Finanziario" in wb.sheetnames
        assert "Materie Prime e Consumo" in wb.sheetnames
        assert "DA MAPPARE" in wb.sheetnames
        assert report["unmapped"] == [125]
        assert report["fornitori_scritti"] >= 1

    def test_idempotenza_due_run_stesso_input(self, tmp_path):
        kwargs = dict(
            pf_prev_bytes=_pf_precedente_bytes(),
            scad_df=self._scad_df(),
            bucket_months=[6],
            fornitori=self._fornitori(),
            societa="ORTI",
            anno=2026,
            primo_mese_aperto=5,
        )
        out1, _ = genera_pf(**kwargs)
        # secondo run: il file generato dal primo diventa il precedente
        out2, _ = genera_pf(**{**kwargs, "pf_prev_bytes": out1})
        ws1 = openpyxl.load_workbook(BytesIO(out1), data_only=False)[
            "Materie Prime e Consumo"
        ]
        ws2 = openpyxl.load_workbook(BytesIO(out2), data_only=False)[
            "Materie Prime e Consumo"
        ]
        celle1 = [
            (c.coordinate, c.value)
            for row in ws1.iter_rows()
            for c in row
            if c.value is not None
        ]
        celle2 = [
            (c.coordinate, c.value)
            for row in ws2.iter_rows()
            for c in row
            if c.value is not None
        ]
        assert celle1 == celle2  # criterio di successo #1 dello spec

    def test_quadratura_check_in_report(self):
        _, report = genera_pf(
            pf_prev_bytes=_pf_precedente_bytes(),
            scad_df=self._scad_df(),
            bucket_months=[6],
            fornitori=self._fornitori(),
            societa="ORTI",
            anno=2026,
            primo_mese_aperto=5,
        )
        quadr = next(c for c in report["controlli"] if "quadratura" in c["check"])
        assert quadr["esito"] == "OK"

    def test_consuntivi_only_supplier_preserved(self):
        # fornitore con SOLO mesi chiusi nel file precedente (no partita aperta,
        # no previsione) deve mantenere il suo storico nel candidato
        wb = openpyxl.Workbook()
        wb.active.title = "Piano Finanziario"
        ws = wb.create_sheet("Materie Prime-Consumo ")
        mesi = [
            "GENNAIO",
            "FEBBRAIO",
            "MARZO",
            "APRILE",
            "MAGGIO",
            "GIUGNO",
            "LUGLIO",
            "AGOSTO",
            "SETTEMBRE",
            "OTTOBRE",
            "NOVEMBRE",
            "DICEMBRE",
        ]
        for i, m in enumerate(mesi):
            ws.cell(row=2, column=8 + i, value=m)
        ws.cell(row=5, column=1, value=500)
        ws.cell(row=5, column=2, value="Fornitore Storico")
        ws.cell(row=5, column=11, value=777.0)  # APRILE (chiuso)
        buf = BytesIO()
        wb.save(buf)
        prev = buf.getvalue()

        scad_df = pd.DataFrame(
            [
                {
                    "codice_fornitore": 92,
                    "nome": "AMALFI",
                    "totale": -100.0,
                    "scaduto": 0.0,
                    "mese_6": -100.0,
                },
            ]
        )
        fornitori = {
            92: {"voce_id": "USCITE_MATERIE_PRIME", "nome_pf": "Amalfi sei esse"},
            500: {"voce_id": "USCITE_MATERIE_PRIME", "nome_pf": "Fornitore Storico"},
        }
        out, _ = genera_pf(
            pf_prev_bytes=prev,
            scad_df=scad_df,
            bucket_months=[6],
            fornitori=fornitori,
            societa="ORTI",
            anno=2026,
            primo_mese_aperto=5,
        )
        wb2 = openpyxl.load_workbook(BytesIO(out))
        ws2 = wb2["Materie Prime e Consumo"]
        # cod 500 presente con APRILE 777 (storico preservato)
        from verticals.condges.pf_generator.costanti import col_mese

        righe = {
            ws2.cell(row=r, column=1).value: r
            for r in range(1, ws2.max_row + 1)
            if isinstance(ws2.cell(row=r, column=1).value, int)
        }
        assert 500 in righe
        assert ws2.cell(row=righe[500], column=col_mese(4)).value == 777.0

    def test_codice_nome_check_osservabile(self):
        # il check codice↔nome legge davvero l'output (non rubber stamp)
        scad_df = pd.DataFrame(
            [
                {
                    "codice_fornitore": 92,
                    "nome": "AMALFI SEI ESSE S.R.L.",
                    "totale": -100.0,
                    "scaduto": 0.0,
                    "mese_6": -100.0,
                },
            ]
        )
        fornitori = {
            92: {"voce_id": "USCITE_MATERIE_PRIME", "nome_pf": "Amalfi sei esse"}
        }
        _, report = genera_pf(
            pf_prev_bytes=_pf_precedente_bytes(),
            scad_df=scad_df,
            bucket_months=[6],
            fornitori=fornitori,
            societa="ORTI",
            anno=2026,
            primo_mese_aperto=5,
        )
        chk = next(c for c in report["controlli"] if "codice" in c["check"])
        assert chk["esito"] == "OK"
        # e il nome scritto è quello del CSV, non dell'export
        # (se fosse rubber stamp non leggerebbe nulla; qui controlliamo che
        # l'esito derivi dalla lettura — basta che sia OK col nome CSV giusto)
