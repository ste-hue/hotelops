"""Orizzonte rolling del PF (decisione Stefano 2026-10-28).

Ogni rotation mensile allunga il template da sola: dopo la chiusura di M il
foglio copre sempre i 12 mesi aperti M+1..M+12. Prima l'orizzonte era un'ancora
fissa (giugno anno+1) allungata solo a mano con `pf-extend`: a marzo il PF
vedeva tre mesi.
"""

from datetime import date

import openpyxl
import pandas as pd

from verticals.condges.pf_rotate.excel_model import find_month_periods, periodo
from verticals.condges.pf_rotate.rotate import rotate
from verticals.condges.pf_rotate.step3_scadenzario import UnmappedPolicy

SALDI = {"MPS": 100.0, "Intesa": 50.0, "MPS_KROSS": 0.0, "UNICREDIT": 0.0}


def _fornitori_csv(tmp_path):
    p = tmp_path / "fornitori.csv"
    p.write_text(
        "codice_fornitore,nome_esolver,nome_pf,voce_id,is_intercompany,is_excluded,exclude_reason,societa_id\n"
        "18,Acqua - Ausino,Acqua - Ausino,USCITE_UTENZE,False,False,,ORTI\n",
        encoding="utf-8",
    )
    return p


def _rotate_aprile(pf_bytes, tmp_path, scad_df, bucket):
    pf_path = tmp_path / "in.xlsx"
    pf_path.write_bytes(pf_bytes)
    return rotate(
        pf_path=pf_path,
        scad_df=scad_df,
        bucket_periodi=bucket,
        societa="ORTI",
        periodo_chiuso=periodo(2026, 4),
        data_saldo=date(2026, 4, 30),
        saldi=SALDI,
        fornitori_csv=_fornitori_csv(tmp_path),
        out_dir=tmp_path / "out",
        unmapped_policy=UnmappedPolicy.FAIL,
    )


def test_rotate_allunga_il_pf_a_dodici_mesi_aperti(minimal_pf_orti_bytes, tmp_path):
    """Fixture APR..DIC 2026; chiuso aprile → aperti maggio 2026..aprile 2027."""
    scad_df = pd.DataFrame(
        [
            {
                "codice_fornitore": 18,
                "nome": "Acqua Ausino",
                "totale": -500.0,
                "scaduto": 0.0,
                f"mese_{periodo(2026, 5)}": -500.0,
            }
        ]
    )
    result = _rotate_aprile(
        minimal_pf_orti_bytes, tmp_path, scad_df, [periodo(2026, 5)]
    )

    wb = openpyxl.load_workbook(result.out_path)
    for sheet in ("Piano Finanziario", "Utenze"):
        cols = find_month_periods(wb[sheet])
        assert max(cols) == periodo(2027, 4), sheet
        assert len([p for p in cols if p > periodo(2026, 4)]) == 12, sheet


def test_rotate_scrive_le_scadenze_oltre_il_vecchio_orizzonte(
    minimal_pf_orti_bytes, tmp_path
):
    """Una scadenza a febbraio 2027 trova la sua colonna invece di finire in oltre_orizzonte."""
    p_feb27 = periodo(2027, 2)
    scad_df = pd.DataFrame(
        [
            {
                "codice_fornitore": 18,
                "nome": "Acqua Ausino",
                "totale": -200.0,
                "scaduto": 0.0,
                f"mese_{p_feb27}": -200.0,
            }
        ]
    )
    result = _rotate_aprile(minimal_pf_orti_bytes, tmp_path, scad_df, [p_feb27])

    assert result.scadenzario_summary["oltre_orizzonte"] == {}
    wb = openpyxl.load_workbook(result.out_path)
    ut = wb["Utenze"]
    cols = find_month_periods(ut)
    row = next(r for r in range(1, 200) if "Acqua" in str(ut.cell(r, 2).value or ""))
    assert ut.cell(row, cols[p_feb27]).value == 200.0
