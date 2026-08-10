from io import BytesIO
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

from verticals.condges.pf_rotate.excel_model import periodo
from verticals.condges.pf_rotate.step3_scadenzario import (
    UnmappedFornitoriError,
    UnmappedPolicy,
    apply_scadenzario,
)

P = lambda m: periodo(2026, m)  # noqa: E731 - l'anno dei fixture (minimal_pf_orti_bytes) è 2026


@pytest.fixture
def tmp_fornitori_csv(tmp_path: Path) -> Path:
    p = tmp_path / "d_fornitori.csv"
    p.write_text(
        "codice_fornitore,nome_esolver,nome_pf,voce_id,is_intercompany,is_excluded,exclude_reason,societa_id\n"
        "100,KNOWN SPA,Known,USCITE_UTENZE,False,False,,ORTI\n"
    )
    return p


def _make_scad_df(codici: list[int]) -> tuple[pd.DataFrame, list[int]]:
    return pd.DataFrame(
        [
            {
                "codice_fornitore": c,
                "nome": f"Forn{c}",
                "totale": 1000.0,
                "scaduto": 0.0,
                f"mese_{P(5)}": 1000.0,
                f"mese_{P(6)}": 0.0,
            }
            for c in codici
        ]
    ), [P(5), P(6)]


def test_policy_fail_raises_on_unmapped(minimal_pf_orti_bytes, tmp_fornitori_csv):
    scad_df, buckets = _make_scad_df([100, 999])  # 999 unknown
    with pytest.raises(UnmappedFornitoriError) as exc:
        apply_scadenzario(
            pf_bytes=minimal_pf_orti_bytes,
            scad_df=scad_df,
            bucket_months=buckets,
            societa="ORTI",
            fornitori_csv=tmp_fornitori_csv,
            policy=UnmappedPolicy.FAIL,
        )
    assert 999 in exc.value.codici


def test_skip_writes_da_mappare_and_esclusi(minimal_pf_orti_bytes, tmp_path):
    """SKIP non deve più far sparire i fornitori: i non-mappati vanno in
    DA MAPPARE, gli esclusi (is_excluded) in ESCLUSI — niente perso in silenzio."""
    csv = tmp_path / "d_fornitori.csv"
    csv.write_text(
        "codice_fornitore,nome_esolver,nome_pf,voce_id,is_intercompany,is_excluded,exclude_reason,societa_id\n"
        "100,KNOWN SPA,Known,USCITE_UTENZE,False,False,,ORTI\n"
        "264,PANORAMA COMPANY,Panorama Company,,True,True,intercompany gruppo,ORTI\n"
    )
    # debiti negativi (convenzione reale export partite)
    scad_df = pd.DataFrame(
        [
            {"codice_fornitore": c, "nome": f"Forn{c}", "totale": -1000.0,
             "scaduto": 0.0, f"mese_{P(5)}": -1000.0, f"mese_{P(6)}": 0.0}
            for c in (100, 999, 264)
        ]
    )
    buckets = [P(5), P(6)]
    out_bytes, summary = apply_scadenzario(
        pf_bytes=minimal_pf_orti_bytes,
        scad_df=scad_df,
        bucket_months=buckets,
        societa="ORTI",
        fornitori_csv=csv,
        policy=UnmappedPolicy.SKIP,
        scaduto_month=P(5),
    )
    assert summary["da_mappare"] == [999]
    assert summary["esclusi"] == [264]
    wb = openpyxl.load_workbook(BytesIO(out_bytes))
    assert "DA MAPPARE" in wb.sheetnames
    assert "ESCLUSI" in wb.sheetnames
    dm = "\n".join(
        str(c.value)
        for row in wb["DA MAPPARE"].iter_rows()
        for c in row
        if c.value is not None
    )
    assert "999" in dm
    es = "\n".join(
        str(c.value)
        for row in wb["ESCLUSI"].iter_rows()
        for c in row
        if c.value is not None
    )
    assert "264" in es and "intercompany gruppo" in es


def test_policy_skip_writes_known_only(minimal_pf_orti_bytes, tmp_fornitori_csv):
    scad_df, buckets = _make_scad_df([100, 999])
    out_bytes, summary = apply_scadenzario(
        pf_bytes=minimal_pf_orti_bytes,
        scad_df=scad_df,
        bucket_months=buckets,
        societa="ORTI",
        fornitori_csv=tmp_fornitori_csv,
        policy=UnmappedPolicy.SKIP,
    )
    assert summary["skipped_unmapped"] == [999]
    # Code 100 (Known → USCITE_UTENZE → foglio Utenze) scritto
    wb = openpyxl.load_workbook(BytesIO(out_bytes))
    ut = wb["Utenze"]
    # cerca riga con codice 100
    found = False
    for r in range(4, 30):
        if ut.cell(r, 1).value == 100:
            found = True
            break
    assert found


def test_write_pf_scadenza_oltre_orizzonte_va_nel_secchio(minimal_pf_orti_bytes):
    """Periodo senza colonna nel foglio: non scritto, non sommato altrove, riportato."""
    from verticals.condges.pf_rotate.pf_writer import write_pf

    p_fuori = periodo(2027, 7)  # il fixture arriva a dicembre 2026
    scad_df = pd.DataFrame([{
        "codice_fornitore": 18, "nome": "Acqua Ausino",
        "totale": -500.0, "scaduto": 0.0, f"mese_{p_fuori}": -500.0,
    }])
    out, summary = write_pf(
        pf_bytes=minimal_pf_orti_bytes, scad_df=scad_df,
        bucket_periodi=[p_fuori],
        fornitori_map={18: {"voce_id": "USCITE_UTENZE", "nome_pf": "Acqua - Ausino"}},
        excluded=set(), scaduto_periodo=periodo(2026, 5), clear_codici={18},
    )
    assert summary["oltre_orizzonte"] == {18: -500.0}
