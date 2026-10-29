"""Chiave di riga dell'archivio `f_pf_rotazioni` (core.bq.load.load_pf_rotazioni).

Il foglio porta legittimamente due righe sulla stessa voce nello stesso mese: INTUR
«Fitto Hotel» e «Fitto AR» sono entrambe ENTRATE_AFFITTI_INTUR. Con la chiave senza
`occorrenza` collidevano (3 righe su 22 file, misurato il 28/10/2026). Convenzione dal
29/10/2026: societa|voce|anno|mese|fonte|file_sorgente|occorrenza. Workbook in tmp_path,
nessuna connessione a BigQuery (si prova solo build_rows).
"""

from __future__ import annotations

import hashlib
import logging
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import Workbook

from core.bq.load.load_pf_rotazioni import (
    SQL_RIGHE_CHIAVE_VECCHIA,
    _hash,
    build_rows,
    file_mancanti,
    righe_chiave_vecchia,
    verifica_convenzione,
)

MESI = [
    "GENNAIO",
    "FEBBRAIO",
    "MARZO",
    "APRILE",
    "MAGGIO",
    "GIUGNO",
    "LUGLIO",
    "AGOSTO",
    "SETTEMBRE",
    "OTTOBRE",
    "NOVEMBRE",
    "DICEMBRE",
]
COL_LUG = 3 + 6  # luglio 2026: il caso reale di INTUR_PF_2026-07
VOCE = "ENTRATE_AFFITTI_INTUR"
FITTI = {"Fitto Hotel": {COL_LUG: 488_000.0}, "Fitto AR": {COL_LUG: 60_000.0}}


def _pf(path: Path, righe: dict[str, dict[int, float]]) -> Path:
    """PF INTUR 2026: C1 anno, riga 2 mesi, A2 società, una riga per etichetta dalla 10."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Piano Finanziario"
    ws["C1"] = 2026
    for i, m in enumerate(MESI):
        ws.cell(2, 3 + i, m)
    ws["A2"] = "INTUR"
    for r, (label, vals) in enumerate(righe.items(), start=10):
        ws.cell(r, 1, label)
        for col, v in vals.items():
            ws.cell(r, col, v)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


@pytest.fixture()
def log():
    return logging.getLogger("test_load_pf_rotazioni")


def test_due_righe_sulla_stessa_voce_e_mese_restano_due_con_hash_distinti(
    tmp_path, log
):
    p = _pf(tmp_path / "INTUR_PF_2026-07.xlsx", FITTI)
    rows = build_rows([p], log)
    affitti = [r for r in rows if r["voce_id"] == VOCE and r["mese"] == 7]
    assert [r["importo"] for r in affitti] == [488_000.0, 60_000.0]  # ordine del foglio
    assert affitti[0]["hash_riga"] != affitti[1]["hash_riga"]
    assert affitti[0]["hash_riga"] == _hash(affitti[0], 1)
    assert affitti[1]["hash_riga"] == _hash(affitti[1], 2)
    assert {r["file_sorgente"] for r in affitti} == {"INTUR_PF_2026-07.xlsx"}


def test_nessuna_collisione_e_stessi_hash_a_ogni_ricarico(tmp_path, log):
    p = _pf(tmp_path / "INTUR_PF_2026-07.xlsx", {**FITTI, "Utenze": {COL_LUG: 1_000.0}})
    prima = [r["hash_riga"] for r in build_rows([p], log)]
    seconda = [r["hash_riga"] for r in build_rows([p], log)]
    assert len(prima) == 3 and len(set(prima)) == 3
    assert prima == seconda


def test_lo_stesso_contenuto_in_due_versioni_resta_distinto_per_file(tmp_path, log):
    a = _pf(tmp_path / "INTUR_PF_2026-07.xlsx", FITTI)
    b = _pf(tmp_path / "INTUR_PF_2026-08.xlsx", FITTI)
    rows = build_rows([a, b], log)
    per_file = {
        f: [r["hash_riga"] for r in rows if r["file_sorgente"] == f]
        for f in (a.name, b.name)
    }
    assert len(per_file[a.name]) == len(per_file[b.name]) == 2
    assert not set(per_file[a.name]) & set(per_file[b.name])
    assert len({r["hash_riga"] for r in rows}) == 4


def test_lo_stesso_file_sorgente_due_volte_nella_stessa_corsa_viene_rifiutato(
    tmp_path, log
):
    """Originale e ricalcolato con lo stesso nome: la DELETE per file non proteggerebbe."""
    a = _pf(tmp_path / "originali" / "INTUR_PF_2026-07.xlsx", FITTI)
    b = _pf(tmp_path / "ricalcolati" / "INTUR_PF_2026-07.xlsx", FITTI)
    with pytest.raises(ValueError, match="INTUR_PF_2026-07.xlsx"):
        build_rows([a, b], log)


def test_file_mancanti_elenca_cio_che_la_corsa_non_ricaricherebbe():
    """--ricarico-completo: una tabella, una convenzione. Vuoto solo se la corsa copre tutto."""
    in_tabella = {
        "ORTI_PF_2026-09.xlsx",
        "INTUR_PF_2026-09.xlsx",
        "INTUR_PF_2026-10.xlsx",
    }
    assert file_mancanti(in_tabella, {"INTUR_PF_2026-10.xlsx"}) == [
        "INTUR_PF_2026-09.xlsx",
        "ORTI_PF_2026-09.xlsx",
    ]
    assert file_mancanti(in_tabella, in_tabella | {"ORTI_PF_2026-10.xlsx"}) == []


# ── Chiave vecchia riconoscibile in SQL (coordinatore, 28/10/2026, sulla tabella vera:
#    totale 2194 · chiave_vecchia 2194 · chiave_nuova(occ=1) 0 · file 22 · hash_distinti 2191)
QUERY_MISURATA = (
    "SELECT COUNT(*) AS n, COUNT(DISTINCT file_sorgente) AS n_file "
    "FROM `hotelops-suite.hotelops.f_pf_rotazioni` WHERE hash_riga = TO_HEX(MD5(CONCAT("
    "societa_id,'|',voce_id,'|',CAST(anno AS STRING),'|',CAST(mese AS STRING),'|',fonte,'|',file_sorgente)))"
)


class _ClientFinto:
    def __init__(self, n: int, n_file: int):
        self.n, self.n_file, self.sql = n, n_file, None

    def query(self, sql: str):
        self.sql = sql
        return SimpleNamespace(
            result=lambda: iter([SimpleNamespace(n=self.n, n_file=self.n_file)])
        )


def test_la_query_della_chiave_vecchia_e_quella_misurata_sulla_tabella():
    assert SQL_RIGHE_CHIAVE_VECCHIA == QUERY_MISURATA


def test_la_chiave_vecchia_in_sql_e_quella_python_fino_a_7f17c93_e_non_e_mai_la_nuova():
    r = {
        "societa_id": "INTUR",
        "voce_id": VOCE,
        "anno": 2026,
        "mese": 7,
        "fonte": "PIANO_FINANZIARIO",
        "file_sorgente": "INTUR_PF_2026-07.xlsx",
    }
    vecchia = hashlib.md5(
        b"INTUR|ENTRATE_AFFITTI_INTUR|2026|7|PIANO_FINANZIARIO|INTUR_PF_2026-07.xlsx"
    ).hexdigest()  # = TO_HEX(MD5(CONCAT(...))): stesso ordine, stesso separatore, hex minuscolo
    assert _hash(r, 1) != vecchia  # la chiave nuova porta sempre "|occorrenza" in più


def test_righe_chiave_vecchia_esegue_la_query_misurata():
    c = _ClientFinto(2194, 22)
    assert righe_chiave_vecchia(c) == (2194, 22)
    assert c.sql == QUERY_MISURATA


def test_senza_flag_la_chiave_vecchia_blocca_la_corsa_e_dopo_la_finestra_tace():
    with pytest.raises(
        SystemExit,
        match=r"2194 righe con la chiave vecchia \(22 file\): serve --ricarico-completo",
    ):
        verifica_convenzione(2194, 22, ricarico_completo=False)
    verifica_convenzione(2194, 22, ricarico_completo=True)  # la finestra
    verifica_convenzione(
        0, 0, ricarico_completo=False
    )  # carico mensile dopo la finestra
