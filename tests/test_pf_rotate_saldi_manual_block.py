"""Regression: il blocco saldi MANUALE (col B=data, col C=valori testo) del template
ORTI month-closed deve avanzare al cutover, non restare al mese precedente.

Geometria REALE (≠ fixture minimal): mesi in G..R (aprile=J), blocco manuale in B/C.
Il bug: write_saldi_banca scriveva solo la colonna del mese chiuso (J) e lasciava B32/
C32:C33 fermi al mese prima. Fix: avanzare B/C al cutover, come TESTO (formula-safe),
solo quando C non è la colonna cutover.
"""

from datetime import date

import openpyxl

from verticals.condges.pf_rotate.step1_saldi import write_saldi_banca


def _build_orti_real_geometry() -> openpyxl.Workbook:
    """ORTI month-closed con mesi in G(7)..R(18) e blocco manuale stantio in B/C."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    pf = wb.create_sheet("Piano Finanziario")
    pf["A1"] = "ORTI"  # C1 None → month-closed
    mesi = [
        "GENNAIO", "FEBBRAIO", "MARZO", "APRILE", "MAGGIO", "GIUGNO",
        "LUGLIO", "AGOSTO", "SETTEMBRE", "OTTOBRE", "NOVEMBRE", "DICEMBRE",
    ]
    for i, m in enumerate(mesi):
        pf.cell(2, 7 + i, m)  # G(7)..R(18); APRILE = J(10)
    # anchors per find_layout
    pf["A4"] = "SALDO MESE PRECEDENTE"
    pf["A6"] = "Entrate Hotel"
    pf["A12"] = "TOTALE ENTRATE"
    pf["A14"] = "Utenze"
    pf["A27"] = "TOTALE USCITE"
    pf["A29"] = "CASH FLOW"
    pf["A32"] = "Saldo MPS"
    pf["A33"] = "Saldo Intesa"
    pf["A35"] = "TOTALE BANCHE"
    pf["A37"] = "Saldo di Periodo"
    # blocco manuale STANTIO (marzo) + formula somma su colonna C
    pf["B32"] = date(2026, 3, 31)
    pf["C32"] = "52,799.65 €"
    pf["C33"] = "63,061.22 €"
    pf["C35"] = "=SUM(C30:C34)"
    return wb


def test_manual_block_advances_to_cutover():
    wb = _build_orti_real_geometry()
    pf = wb["Piano Finanziario"]

    write_saldi_banca(
        wb,
        mese_chiuso=4,  # APRILE → colonna J (10)
        data_saldo=date(2026, 4, 30),
        saldi={"MPS": 245171.52, "Intesa": 87439.92},
    )

    # blocco manuale B/C avanzato al cutover, come TESTO
    assert pf["B32"].value == "30/04/2026"
    assert pf["C32"].value == "245,171.52 €"
    assert pf["C33"].value == "87,439.92 €"
    # la formula somma NON è toccata
    assert pf["C35"].value == "=SUM(C30:C34)"
    # il blocco engine (colonna cutover J) ha i numeri reali
    assert pf["J31"].value == "30/04/2026"
    assert pf["J32"].value == 245171.52
    assert pf["J33"].value == 87439.92
    assert pf["J4"].value == 245171.52 + 87439.92  # saldo iniziale hardcoded


def test_manual_block_skips_formula_cells():
    """Se C{r} è una formula, NON va sovrascritta (invariante mai-azzerare-formule)."""
    wb = _build_orti_real_geometry()
    pf = wb["Piano Finanziario"]
    pf["C33"] = "=C32"  # qualcuno ha messo una formula nel blocco manuale

    write_saldi_banca(
        wb,
        mese_chiuso=4,
        data_saldo=date(2026, 4, 30),
        saldi={"MPS": 245171.52, "Intesa": 87439.92},
    )

    assert pf["C32"].value == "245,171.52 €"  # value cell → aggiornata
    assert pf["C33"].value == "=C32"  # formula → preservata
