"""Tests for scadenzario Excel bridge."""
from io import BytesIO

import openpyxl


def _make_sintetica_workbook(rows, header_dates=None):
    """Create a minimal sintetica-style workbook in memory."""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(row=1, column=1, value="Fornitore/Tipo pagamento/Conto")
    ws.cell(row=1, column=2, value="Scadenze - Totale")
    ws.cell(row=1, column=3, value="Scadenze - Scaduto fino al 02/04/2026")
    defaults = header_dates or [
        "Scadenze - In scadenza al 02/05/2026",
        "Scadenze - In scadenza al 02/06/2026",
        "Scadenze - In scadenza al 02/07/2026",
    ]
    for i, h in enumerate(defaults):
        ws.cell(row=1, column=4 + i, value=h)
    for r_idx, row_data in enumerate(rows, start=2):
        ws.cell(row=r_idx, column=1, value=row_data[0])
        if len(row_data) > 1 and row_data[1] is not None:
            ws.cell(row=r_idx, column=2, value=row_data[1])
        if len(row_data) > 2 and row_data[2] is not None:
            ws.cell(row=r_idx, column=3, value=row_data[2])
        for i, val in enumerate(row_data[3:]):
            if val is not None:
                ws.cell(row=r_idx, column=4 + i, value=val)
    return wb


class TestParseSinteticaScadenze:
    def test_parses_codice_and_nome(self, tmp_path):
        from verticals.condges.scadenzario_excel import parse_sintetica_scadenze
        wb = _make_sintetica_workbook([
            ("264 PANORAMA COMPANY S.R.L.", -1000, -500, -500),
        ])
        f = tmp_path / "test.xlsx"
        wb.save(f)
        result, months = parse_sintetica_scadenze(f)
        assert len(result) == 1
        assert result[0]["codice_fornitore"] == 264
        assert result[0]["nome"] == "PANORAMA COMPANY S.R.L."

    def test_parses_buckets_from_headers(self, tmp_path):
        from verticals.condges.scadenzario_excel import parse_sintetica_scadenze
        wb = _make_sintetica_workbook([
            ("1 FORNITORE TEST", -1500, -500, -600, -400),
        ])
        f = tmp_path / "test.xlsx"
        wb.save(f)
        result, months = parse_sintetica_scadenze(f)
        r = result[0]
        assert r["totale"] == -1500
        assert r["scaduto"] == -500
        assert r["buckets"] == {5: -600, 6: -400}
        assert months == [5, 6, 7]

    def test_skips_rows_without_numeric_prefix(self, tmp_path):
        from verticals.condges.scadenzario_excel import parse_sintetica_scadenze
        wb = _make_sintetica_workbook([
            ("264 PANORAMA S.R.L.", -1000, -1000),
            ("Totale", -1000, -1000),
        ])
        f = tmp_path / "test.xlsx"
        wb.save(f)
        result, _ = parse_sintetica_scadenze(f)
        assert len(result) == 1

    def test_handles_positive_amounts(self, tmp_path):
        from verticals.condges.scadenzario_excel import parse_sintetica_scadenze
        wb = _make_sintetica_workbook([
            ("1008 MIELE SPA", 2294.28, 2294.28),
        ])
        f = tmp_path / "test.xlsx"
        wb.save(f)
        result, _ = parse_sintetica_scadenze(f)
        assert result[0]["totale"] == 2294.28


class TestMapToVoci:
    def test_maps_supplier_to_voce(self):
        from verticals.condges.scadenzario_excel import map_to_voci
        partite = [
            {"codice_fornitore": 1, "nome": "LE CROISSANT", "totale": -1000,
             "scaduto": -500, "buckets": {5: -500}},
        ]
        fornitori_map = {1: "USCITE_MATERIE_PRIME"}
        mapped, unmapped = map_to_voci(partite, fornitori_map)
        assert "USCITE_MATERIE_PRIME" in mapped
        assert len(mapped["USCITE_MATERIE_PRIME"]) == 1
        assert unmapped == []

    def test_unmapped_supplier_goes_to_unmapped_list(self):
        from verticals.condges.scadenzario_excel import map_to_voci
        partite = [
            {"codice_fornitore": 9999, "nome": "UNKNOWN", "totale": -100,
             "scaduto": -100, "buckets": {}},
        ]
        fornitori_map = {1: "USCITE_MATERIE_PRIME"}
        mapped, unmapped = map_to_voci(partite, fornitori_map)
        assert len(unmapped) == 1
        assert unmapped[0]["codice_fornitore"] == 9999

    def test_multiple_suppliers_same_voce(self):
        from verticals.condges.scadenzario_excel import map_to_voci
        partite = [
            {"codice_fornitore": 1, "nome": "A", "totale": -500,
             "scaduto": -500, "buckets": {}},
            {"codice_fornitore": 4, "nome": "B", "totale": -300,
             "scaduto": -300, "buckets": {}},
        ]
        fornitori_map = {1: "USCITE_MATERIE_PRIME", 4: "USCITE_MATERIE_PRIME"}
        mapped, unmapped = map_to_voci(partite, fornitori_map)
        assert len(mapped["USCITE_MATERIE_PRIME"]) == 2


class TestLoadFornitoriMap:
    def test_loads_csv(self, tmp_path):
        from verticals.condges.scadenzario_excel import load_fornitori_map
        csv_path = tmp_path / "d_fornitori.csv"
        csv_path.write_text(
            "codice_fornitore,nome_esolver,nome_pf,voce_id,is_intercompany\n"
            "1,LE CROISSANT SRL,Le Croissant,USCITE_MATERIE_PRIME,False\n"
            "264,PANORAMA COMPANY S.R.L.,Fitto,USCITE_CANONE_PASSIVO,True\n"
        )
        result = load_fornitori_map(csv_path=csv_path)
        assert result[1] == "USCITE_MATERIE_PRIME"
        assert result[264] == "USCITE_CANONE_PASSIVO"


class TestGenerateExcel:
    def _sample_data(self):
        mapped = {
            "USCITE_MATERIE_PRIME": [
                {"codice_fornitore": 1, "nome": "LE CROISSANT",
                 "totale": -1500, "scaduto": -1000, "buckets": {5: -500}},
                {"codice_fornitore": 4, "nome": "GIACINTO",
                 "totale": -400, "scaduto": -400, "buckets": {}},
            ],
            "USCITE_UTENZE": [
                {"codice_fornitore": 18, "nome": "AUSINO",
                 "totale": -600, "scaduto": 0, "buckets": {5: -600}},
            ],
        }
        return mapped

    def test_creates_riepilogo_sheet(self, tmp_path):
        from verticals.condges.scadenzario_excel import generate_excel
        out = tmp_path / "test.xlsx"
        generate_excel(self._sample_data(), None, [], out, bucket_months=[5, 6, 7])
        wb = openpyxl.load_workbook(out)
        assert "Riepilogo" in wb.sheetnames

    def test_creates_per_voce_sheets(self, tmp_path):
        from verticals.condges.scadenzario_excel import generate_excel
        out = tmp_path / "test.xlsx"
        generate_excel(self._sample_data(), None, [], out, bucket_months=[5, 6, 7])
        wb = openpyxl.load_workbook(out)
        assert "Materie Prime" in wb.sheetnames
        assert "Utenze" in wb.sheetnames

    def test_riepilogo_has_totals(self, tmp_path):
        from verticals.condges.scadenzario_excel import generate_excel
        out = tmp_path / "test.xlsx"
        generate_excel(self._sample_data(), None, [], out, bucket_months=[5, 6, 7])
        wb = openpyxl.load_workbook(out, data_only=True)
        ws = wb["Riepilogo"]
        values = {}
        for row in ws.iter_rows(min_row=2, values_only=False):
            if row[0].value and "Totale" in str(row[0].value):
                values["totale_col"] = row[1].value
        assert values.get("totale_col") is not None

    def test_unmapped_sheet_created_when_needed(self, tmp_path):
        from verticals.condges.scadenzario_excel import generate_excel
        out = tmp_path / "test.xlsx"
        unmapped = [{"codice_fornitore": 999, "nome": "UNKNOWN",
                     "totale": -100, "scaduto": -100, "buckets": {}}]
        generate_excel({}, None, unmapped, out, bucket_months=[5])
        wb = openpyxl.load_workbook(out)
        assert "DA VERIFICARE" in wb.sheetnames

    def test_no_unmapped_sheet_when_all_mapped(self, tmp_path):
        from verticals.condges.scadenzario_excel import generate_excel
        out = tmp_path / "test.xlsx"
        generate_excel(self._sample_data(), None, [], out, bucket_months=[5])
        wb = openpyxl.load_workbook(out)
        assert "DA VERIFICARE" not in wb.sheetnames


class TestLoadPfForecasts:
    def test_parses_uscite_rows(self, tmp_path):
        """Test parsing Rosa's PF Excel uscite rows."""
        from verticals.condges.scadenzario_excel import load_pf_forecasts
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Piano Finanziario"
        ws.cell(row=2, column=1, value="ORTI")
        # Row 2 has actual calendar month names (like the real PF Excel)
        for i, name in enumerate(["APRILE", "MAGGIO", "GIUGNO"], start=10):
            ws.cell(row=2, column=i, value=name)
        ws.cell(row=16, column=1, value="Materie Prime/Consumo")
        ws.cell(row=16, column=10, value=117460)
        ws.cell(row=16, column=11, value=40075)
        ws.cell(row=16, column=12, value=95000)

        f = tmp_path / "pf.xlsx"
        wb.save(f)

        result = load_pf_forecasts(f)
        assert "USCITE_MATERIE_PRIME" in result
        assert result["USCITE_MATERIE_PRIME"][4] == 117460  # April


# ── Cascade tests ────────────────────────────────────────────────────────────


class TestCascadeScaduto:
    def test_moves_overdue_to_current_month(self):
        from verticals.condges.scadenzario_excel import cascade_scaduto

        suppliers = [
            {
                "codice_fornitore": 1,
                "nome": "ACME",
                "totale": 1500,
                "scaduto": 500,
                "buckets": {5: 600, 6: 400},
            },
            {
                "codice_fornitore": 2,
                "nome": "BETA",
                "totale": 300,
                "scaduto": 300,
                "buckets": {},
            },
        ]
        result = cascade_scaduto(suppliers, current_month=4)

        assert result[0]["buckets"][4] == 500
        assert result[0]["buckets"][5] == 600
        assert result[0]["scaduto"] == 0

        assert result[1]["buckets"][4] == 300
        assert result[1]["scaduto"] == 0

    def test_adds_to_existing_current_month(self):
        from verticals.condges.scadenzario_excel import cascade_scaduto

        suppliers = [
            {
                "codice_fornitore": 1,
                "nome": "ACME",
                "totale": 1000,
                "scaduto": 200,
                "buckets": {4: 300, 5: 500},
            },
        ]
        result = cascade_scaduto(suppliers, current_month=4)
        assert result[0]["buckets"][4] == 500

    def test_zero_scaduto_unchanged(self):
        from verticals.condges.scadenzario_excel import cascade_scaduto

        suppliers = [
            {
                "codice_fornitore": 1,
                "nome": "ACME",
                "totale": 600,
                "scaduto": 0,
                "buckets": {5: 600},
            },
        ]
        result = cascade_scaduto(suppliers, current_month=4)
        assert 4 not in result[0]["buckets"]
        assert result[0]["buckets"][5] == 600


# ── Write-back tests ─────────────────────────────────────────────────────────


def _make_pf_fixture(tmp_path):
    """Create a minimal PF Excel mimicking Rosa's layout."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Piano Finanziario"
    ws.cell(row=1, column=3, value=2025)
    ws.cell(row=1, column=7, value=2026)
    months_2025 = ["SETTEMBRE", "OTTOBRE", "NOVEMBRE", "DICEMBRE"]
    months_2026 = [
        "GENNAIO", "FEBBRAIO", "MARZO", "APRILE", "MAGGIO",
        "GIUGNO", "LUGLIO", "AGOSTO", "SETTEMBRE", "OTTOBRE",
        "NOVEMBRE", "DICEMBRE",
    ]
    for i, m in enumerate(months_2025):
        ws.cell(row=2, column=3 + i, value=m)
    for i, m in enumerate(months_2026):
        ws.cell(row=2, column=7 + i, value=m)
    ws.cell(row=16, column=1, value="Materie Prime/Consumo")

    # Detail sheet "Materie Prime-Consumo " (trailing space)
    ws_mp = wb.create_sheet("Materie Prime-Consumo ")
    for i, m in enumerate(months_2025):
        ws_mp.cell(row=2, column=4 + i, value=m)
    for i, m in enumerate(months_2026):
        ws_mp.cell(row=2, column=8 + i, value=m)
    ws_mp.cell(row=4, column=2, value="Materie Prime e Consumo")
    ws_mp.cell(row=5, column=2, value="Le Croissant srl")
    ws_mp.cell(row=6, column=2, value="Giacinto Di Palma")
    ws_mp.cell(row=7, column=2, value="PREVISIONALE")
    ws_mp.cell(row=7, column=11, value=60000)

    out = tmp_path / "PF_test.xlsx"
    wb.save(out)
    return out


class TestWriteBackToPf:
    def test_places_amounts_correctly(self, tmp_path):
        from verticals.condges.scadenzario_excel import write_back_to_pf

        pf_path = _make_pf_fixture(tmp_path)

        mapped = {
            "USCITE_MATERIE_PRIME": [
                {
                    "codice_fornitore": 1,
                    "nome": "LE CROISSANT SRL",
                    "totale": 14277,
                    "scaduto": 0,
                    "buckets": {4: 14277},
                },
                {
                    "codice_fornitore": 4,
                    "nome": "Giacinto Di Palma",
                    "totale": 390,
                    "scaduto": 0,
                    "buckets": {4: 390},
                },
            ],
        }
        fornitori_map_full = {
            1: {"voce_id": "USCITE_MATERIE_PRIME", "nome_pf": "Le Croissant srl"},
            4: {"voce_id": "USCITE_MATERIE_PRIME", "nome_pf": "Giacinto Di Palma"},
        }

        out, summary = write_back_to_pf(pf_path, mapped, fornitori_map_full)

        wb = openpyxl.load_workbook(out)
        ws_mp = wb["Materie Prime-Consumo "]

        # APRILE = col 11 on detail sheets (col 8=GEN + 3)
        assert ws_mp.cell(row=5, column=11).value == 14277
        assert ws_mp.cell(row=6, column=11).value == 390
        # PREVISIONALE untouched
        assert ws_mp.cell(row=7, column=11).value == 60000

        assert "Materie Prime" in summary
        assert len(summary["Materie Prime"]) == 2

    def test_preserves_existing_data(self, tmp_path):
        from verticals.condges.scadenzario_excel import write_back_to_pf

        pf_path = _make_pf_fixture(tmp_path)

        # Pre-write some data in May column
        wb = openpyxl.load_workbook(pf_path)
        ws_mp = wb["Materie Prime-Consumo "]
        ws_mp.cell(row=5, column=12, value=9999)  # May, Le Croissant
        wb.save(pf_path)

        mapped = {
            "USCITE_MATERIE_PRIME": [
                {
                    "codice_fornitore": 1,
                    "nome": "LE CROISSANT SRL",
                    "totale": 14277,
                    "scaduto": 0,
                    "buckets": {4: 14277},  # only April
                },
            ],
        }
        fornitori_map_full = {
            1: {"voce_id": "USCITE_MATERIE_PRIME", "nome_pf": "Le Croissant srl"},
        }

        out, _ = write_back_to_pf(pf_path, mapped, fornitori_map_full)
        wb = openpyxl.load_workbook(out)
        ws_mp = wb["Materie Prime-Consumo "]

        assert ws_mp.cell(row=5, column=11).value == 14277
        assert ws_mp.cell(row=5, column=12).value == 9999  # preserved

    def test_skips_unmapped_nome_pf(self, tmp_path):
        from verticals.condges.scadenzario_excel import write_back_to_pf

        pf_path = _make_pf_fixture(tmp_path)

        mapped = {
            "USCITE_MATERIE_PRIME": [
                {
                    "codice_fornitore": 999,
                    "nome": "UNKNOWN SUPPLIER",
                    "totale": 5000,
                    "scaduto": 0,
                    "buckets": {4: 5000},
                },
            ],
        }
        # No nome_pf for this supplier
        fornitori_map_full = {
            999: {"voce_id": "USCITE_MATERIE_PRIME", "nome_pf": ""},
        }

        out, summary = write_back_to_pf(pf_path, mapped, fornitori_map_full)
        assert "Materie Prime" not in summary  # nothing written


class TestBuildMonthColMap:
    def test_takes_2026_columns_for_duplicate_months(self):
        from verticals.condges.scadenzario_excel import _build_month_col_map

        wb = openpyxl.Workbook()
        ws = wb.active
        # 2025: C4-C7 = Sep-Dec
        for i, m in enumerate(["SETTEMBRE", "OTTOBRE", "NOVEMBRE", "DICEMBRE"]):
            ws.cell(row=2, column=4 + i, value=m)
        # 2026: C8-C19 = Jan-Dec
        months_2026 = [
            "GENNAIO", "FEBBRAIO", "MARZO", "APRILE", "MAGGIO",
            "GIUGNO", "LUGLIO", "AGOSTO", "SETTEMBRE", "OTTOBRE",
            "NOVEMBRE", "DICEMBRE",
        ]
        for i, m in enumerate(months_2026):
            ws.cell(row=2, column=8 + i, value=m)

        col_map = _build_month_col_map(ws)

        # Sep-Dec should use 2026 columns (rightmost)
        assert col_map[9] == 16
        assert col_map[10] == 17
        assert col_map[11] == 18
        assert col_map[12] == 19
        # Forward months
        assert col_map[4] == 11  # APRILE
        assert col_map[5] == 12  # MAGGIO


def _make_pf_fixture_con_codici(tmp_path):
    """PF minimale con codici fornitore in col A (layout reale dei detail sheet)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Piano Finanziario"

    ws_mp = wb.create_sheet("Materie Prime-Consumo ")
    months_2026 = [
        "GENNAIO", "FEBBRAIO", "MARZO", "APRILE", "MAGGIO",
        "GIUGNO", "LUGLIO", "AGOSTO", "SETTEMBRE", "OTTOBRE",
        "NOVEMBRE", "DICEMBRE",
    ]
    for i, m in enumerate(months_2026):
        ws_mp.cell(row=2, column=8 + i, value=m)
    ws_mp.cell(row=4, column=2, value="Materie Prime e Consumo")
    # riga fornitore mappato con scrittura stantia a MAGGIO (col 12)
    ws_mp.cell(row=5, column=1, value=92)
    ws_mp.cell(row=5, column=2, value="Amalfi sei esse")
    ws_mp.cell(row=5, column=12, value=74.61)
    # riga manuale (codice conto PF, non fornitore) con valore a MAGGIO
    ws_mp.cell(row=6, column=1, value=350301)
    ws_mp.cell(row=6, column=2, value="F24")
    ws_mp.cell(row=6, column=12, value=27000)

    out = tmp_path / "PF_codici.xlsx"
    wb.save(out)
    return out


class TestWritePfRotation:
    def _scad_df(self, **extra_cols):
        import pandas as pd
        base = {
            "codice_fornitore": [92],
            "nome": ["AMALFI SEI ESSE S.R.L."],
            "totale": [-1222.76],
            "scaduto": [-100.0],
        }
        base.update(extra_cols)
        return pd.DataFrame(base)

    def test_scaduto_va_nel_mese_indicato_non_oggi(self, tmp_path):
        from verticals.condges.app_scadenzario import write_pf

        pf_path = _make_pf_fixture_con_codici(tmp_path)
        scad_df = self._scad_df(mese_6=[-1122.76])
        fornitori_map = {
            92: {"voce_id": "USCITE_MATERIE_PRIME", "nome_pf": "Amalfi sei esse"}
        }
        out, _ = write_pf(
            pf_path.read_bytes(), scad_df, [6], fornitori_map, scaduto_month=5
        )
        ws = openpyxl.load_workbook(BytesIO(out))["Materie Prime-Consumo "]
        # scaduto (100) -> MAGGIO (col 12), non nel mese di oggi
        assert ws.cell(row=5, column=12).value == 100.0
        # GIUGNO (col 13) = bucket mese_6
        assert ws.cell(row=5, column=13).value == 1122.76

    def test_pulizia_scritture_stantie_nei_mesi_aperti(self, tmp_path):
        from verticals.condges.app_scadenzario import write_pf

        pf_path = _make_pf_fixture_con_codici(tmp_path)
        # niente scaduto, solo giugno: la cella stantia di maggio deve sparire
        scad_df = self._scad_df(scaduto=[0.0], mese_6=[-1122.76])
        fornitori_map = {
            92: {"voce_id": "USCITE_MATERIE_PRIME", "nome_pf": "Amalfi sei esse"}
        }
        out, _ = write_pf(
            pf_path.read_bytes(), scad_df, [6], fornitori_map,
            scaduto_month=5, clear_codici={92},
        )
        ws = openpyxl.load_workbook(BytesIO(out))["Materie Prime-Consumo "]
        # MAGGIO stantio (74.61) pulito
        assert ws.cell(row=5, column=12).value in (None, 0)
        # GIUGNO scritto
        assert ws.cell(row=5, column=13).value == 1122.76
        # riga manuale F24 (350301, non in clear_codici) intatta
        assert ws.cell(row=6, column=12).value == 27000
