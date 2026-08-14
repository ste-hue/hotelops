"""Tests per parse_scadenze — layout sintetica e dettagliata."""

from __future__ import annotations

from datetime import date, timedelta
from io import BytesIO

import openpyxl
import pytest

from verticals.condges.pf_rotate.excel_model import periodo
from verticals.condges.scadenze_parse import (
    parse_scadenze,
    partite_df_to_scadenzario_data,
)

FUTURE = date.today() + timedelta(days=40)
PAST = date.today() - timedelta(days=40)


def _wb_bytes(rows: list[dict]) -> BytesIO:
    """rows: dict {col_1based: value}."""
    wb = openpyxl.Workbook()
    ws = wb.active
    for r_idx, r in enumerate(rows, 1):
        for col, v in r.items():
            ws.cell(row=r_idx, column=col).value = v
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def test_sintetica_layout():
    # importo = col 26 "Saldo scadenza in UDC" (residuo per-scadenza, signed)
    buf = _wb_bytes(
        [
            {11: 42, 12: "FORNITORE SINT", 26: -100.0, 23: FUTURE},
            {11: 42, 12: "FORNITORE SINT", 26: -50.0, 23: PAST},
        ]
    )
    df, buckets = parse_scadenze(buf)
    assert len(df) == 1
    row = df.iloc[0]
    assert row["codice_fornitore"] == 42
    assert row["totale"] == -150.0
    assert row["scaduto"] == -50.0
    assert buckets == [periodo(FUTURE.year, FUTURE.month)]


def test_sintetica_usa_saldo_scadenza_non_apertura():
    """issue #47: usa col 26 (saldo scadenza) NON col 20 (importo apertura).

    col 20 = fattura intera ripetuta su ogni rata → da ignorare;
    col 26 = residuo della singola scadenza → quello da bucketare.
    """
    buf = _wb_bytes(
        [
            {11: 99, 12: "MIELE", 20: -10857.37, 26: -3619.13, 23: FUTURE},
        ]
    )
    df, _ = parse_scadenze(buf)
    row = df.iloc[0]
    p = periodo(FUTURE.year, FUTURE.month)
    assert row[f"mese_{p}"] == -3619.13  # col 26, non -10857.37
    assert row["totale"] == -3619.13


def test_due_rate_stesso_mese_sommate():
    """Due scadenze della stessa voce nello stesso mese → sommate (caso Miele lug)."""
    p = periodo(FUTURE.year, FUTURE.month)
    buf = _wb_bytes(
        [
            {11: 1008, 12: "MIELE", 26: -3619.12, 23: FUTURE},
            {11: 1008, 12: "MIELE", 26: -439.74, 23: FUTURE},
        ]
    )
    df, _ = parse_scadenze(buf)
    assert round(df.iloc[0][f"mese_{p}"], 2) == -4058.86


def test_dettagliata_layout():
    """Dettagliata: col 23 non-data, scadenza in col 34, residuo in col 37."""
    buf = _wb_bytes(
        [
            {11: 42, 12: "FORNITORE DETT", 23: "Q", 30: -999.0, 34: FUTURE, 37: -200.0},
            {11: 42, 12: "FORNITORE DETT", 23: "Q", 30: -999.0, 34: PAST, 37: -80.0},
        ]
    )
    df, buckets = parse_scadenze(buf)
    assert len(df) == 1
    row = df.iloc[0]
    assert row["totale"] == -280.0  # usa col 37 (residuo), non col 30/20
    assert row["scaduto"] == -80.0
    assert buckets == [periodo(FUTURE.year, FUTURE.month)]


def test_dettagliata_riga_senza_scadenza_skippata():
    buf = _wb_bytes(
        [
            {11: 42, 12: "X", 23: "Q", 34: FUTURE, 37: -10.0},
            {11: 43, 12: "Y", 23: "Q"},  # nessuna data in 23 né 34
        ]
    )
    df, _ = parse_scadenze(buf)
    assert list(df["codice_fornitore"]) == [42]


def test_cutoff_esplicito_decide_scaduto():
    """primo_mese_aperto ancora lo scaduto alla rotation, non a date.today().

    Cutoff (2026, 5): scadenze < maggio 2026 -> scaduto; maggio resta mese_5
    anche se oggi e' giugno o oltre.
    """
    buf = _wb_bytes(
        [
            {11: 92, 12: "AMALFI SEI ESSE", 26: -100.0, 23: date(2026, 4, 30)},
            {11: 92, 12: "AMALFI SEI ESSE", 26: -200.0, 23: date(2026, 5, 31)},
            {11: 92, 12: "AMALFI SEI ESSE", 26: -300.0, 23: date(2026, 6, 30)},
        ]
    )
    df, buckets = parse_scadenze(buf, primo_mese_aperto=(2026, 5))
    row = df.iloc[0]
    p5, p6 = periodo(2026, 5), periodo(2026, 6)
    assert row["scaduto"] == -100.0
    assert row[f"mese_{p5}"] == -200.0
    assert row[f"mese_{p6}"] == -300.0
    assert buckets == [p5, p6]


def test_nome_completo_da_due_colonne():
    """Ragione sociale 1 (col 12) + Ragione sociale 2 (col 13) concatenate."""
    buf = _wb_bytes(
        [
            {11: 417, 12: "CIMINI", 13: "FILOMENA", 26: -51.48, 23: FUTURE},
        ]
    )
    df, _ = parse_scadenze(buf)
    assert df.iloc[0]["nome"] == "CIMINI FILOMENA"


def test_parse_scadenze_separa_gli_anni(tmp_path):
    """Giugno 2026 e giugno 2027 NON si sommano nello stesso bucket."""
    import openpyxl
    from datetime import datetime

    from verticals.condges.pf_rotate.excel_model import periodo

    wb = openpyxl.Workbook()
    ws = wb.active
    # layout sintetica: col 11 codice, col 12 nome, col 23 scadenza, col 26 saldo
    for r, (scad, imp) in enumerate(
        [(datetime(2026, 6, 15), -100.0), (datetime(2027, 6, 15), -900.0)], start=2
    ):
        ws.cell(r, 11, 5555)
        ws.cell(r, 12, "FORNITORE BIENNALE")
        ws.cell(r, 23, scad)
        ws.cell(r, 26, imp)
    p = tmp_path / "scad.xlsx"
    wb.save(p)

    from verticals.condges.scadenze_parse import parse_scadenze

    df, bucket_periodi = parse_scadenze(open(p, "rb"), primo_mese_aperto=(2026, 5))

    p26, p27 = periodo(2026, 6), periodo(2027, 6)
    assert set(bucket_periodi) == {p26, p27}
    row = df.iloc[0]
    assert row[f"mese_{p26}"] == -100.0
    assert row[f"mese_{p27}"] == -900.0  # oggi finirebbe sommato in giugno 2026
    assert row["scaduto"] == 0.0


# ── partite_df_to_scadenzario_data: chiavi PERIODO → mesi calendario nudi ──
#
# ScadenzarioData.totale_per_mese usa la stessa convenzione dei layout
# riepilogo/sintetica (mese calendario 1-12, vedi parse_pf.parse_scadenzario
# e scadenzario_excel.parse_sintetica_scadenze): i consumer (tesoreria.py,
# gen_tesoreria_xlsx.py::project_cashflow) indicizzano per mese 1-12, non per
# periodo ordinale. parse_scadenze produce periodi (anno-aware); qui si
# converte, con guardia anti-collisione multi-anno (stile step3_scadenzario).


def test_partite_df_to_scadenzario_data_usa_mesi_calendario_nudi():
    """totale_per_mese e le colonne per-fornitore usano mese 1-12, non periodo
    (chiave a 5 cifre tipo 24317): un consumer che indicizza per mese nudo
    (project_cashflow, tesoreria.py) deve trovare l'uscita, non zero."""
    buf = _wb_bytes(
        [
            {11: 18, 12: "ACQUA AUSINO", 26: -500.0, 23: date(2026, 6, 15)},
        ]
    )
    df, bucket_periodi = parse_scadenze(buf, primo_mese_aperto=(2026, 5))
    p_giugno = periodo(2026, 6)
    # pre-condizione: parse_scadenze e' periodo-keyed
    assert bucket_periodi == [p_giugno]

    data = partite_df_to_scadenzario_data(df, bucket_periodi)

    assert set(data.totale_per_mese.keys()) == {6}  # mese nudo, non periodo(2026,6)
    assert data.totale_per_mese[6] == -500.0
    entry = data.fornitori[0]
    assert entry["mese_6"] == -500.0
    assert "mese_%d" % p_giugno not in entry  # niente chiave periodo residua


def test_partite_df_to_scadenzario_data_collisione_multianno_esplode():
    """Giugno 2026 e giugno 2027 sono periodi diversi (parse_scadenze li tiene
    separati) ma collassano sullo stesso mese calendario nudo: ScadenzarioData
    ha 12 colonne mese, non anno-aware — sommarli in silenzio nella stessa
    colonna perderebbe la separazione. Deve fallire rumorosamente."""
    buf = _wb_bytes(
        [
            {11: 5555, 12: "FORNITORE BIENNALE", 26: -100.0, 23: date(2026, 6, 15)},
            {11: 5555, 12: "FORNITORE BIENNALE", 26: -900.0, 23: date(2027, 6, 15)},
        ]
    )
    df, bucket_periodi = parse_scadenze(buf, primo_mese_aperto=(2026, 5))
    p26, p27 = periodo(2026, 6), periodo(2027, 6)
    assert set(bucket_periodi) == {p26, p27}

    with pytest.raises(ValueError, match=f"{p26}|{p27}"):
        partite_df_to_scadenzario_data(df, bucket_periodi)


def test_seam_parse_scadenze_a_project_cashflow_uscita_fornitori_non_zero():
    """Test di cucitura: parse_scadenze (xlsx sintetico con date reali) →
    partite_df_to_scadenzario_data → tesoreria.project_cashflow. Prima del
    fix, totale_per_mese era periodo-keyed e project_cashflow (che cerca
    fornitori_per_mese.get(mese, 0.0) con mese 1-12) zerava sempre l'uscita
    fornitori in silenzio — questo test cattura la cucitura end-to-end, non
    solo il singolo modulo."""
    from verticals.condges.cashflow import project_cashflow

    buf = _wb_bytes(
        [
            {11: 18, 12: "ACQUA AUSINO", 26: -500.0, 23: date(2026, 6, 15)},
        ]
    )
    df, bucket_periodi = parse_scadenze(buf, primo_mese_aperto=(2026, 5))
    scad_data = partite_df_to_scadenzario_data(df, bucket_periodi)

    rows = project_cashflow(
        saldo_iniziale=10_000.0,
        mese_inizio=6,
        entrate_per_mese={},
        uscite_per_mese={},
        fornitori_per_mese=scad_data.totale_per_mese,
        mese_fine=6,
    )

    assert len(rows) == 1
    # convenzione parse_scadenze: debiti negativi (col 26 "Saldo scadenza in
    # UDC") — il punto del test e' che NON e' 0.0 (lo zeroing silenzioso
    # pre-fix, chiavi periodo cercate come chiavi mese 1-12).
    assert rows[0].uscite_fornitori == -500.0
