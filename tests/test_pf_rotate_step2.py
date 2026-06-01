from io import BytesIO

import openpyxl

from verticals.condges.pf_rotate.step2_azzera import azzera_mese


def test_azzera_master_entrate_to_none(minimal_pf_orti_bytes):
    wb = openpyxl.load_workbook(BytesIO(minimal_pf_orti_bytes), data_only=False)
    # Pre: fixture mette 100.0 in colonna MAGGIO (col D) per ogni entrata r6-r11
    # ... ma chiudiamo APRILE (col C). Inseriamo prima un valore in APRILE manualmente.
    pf = wb["Piano Finanziario"]
    pf["C6"] = 500.0  # entrata APR hardcoded
    pf["C7"] = 200.0

    changes = azzera_mese(wb, mese_chiuso=4)

    assert pf["C6"].value is None
    assert pf["C7"].value is None
    # MAGGIO (col D) intatto
    assert pf["D6"].value == 100.0
    # formule preservate
    assert pf["C12"].value == "=SUM(C6:C11)"
    assert pf["C29"].value == "=C12-C27"
    assert pf["C37"].value == "=C4+C29"
    # changes registrate
    refs = [(ch.sheet, ch.ref, ch.old, ch.new) for ch in changes]
    assert ("Piano Finanziario", "C6", 500.0, None) in refs
    assert ("Piano Finanziario", "C7", 200.0, None) in refs


def test_azzera_dettaglio_only_values(minimal_pf_orti_bytes):
    wb = openpyxl.load_workbook(BytesIO(minimal_pf_orti_bytes), data_only=False)
    # Fixture: Utenze D5 = 2264.12 (Acqua-Ausino APR)
    ut = wb["Utenze"]
    assert ut["D5"].value == 2264.12  # pre-condition

    azzera_mese(wb, mese_chiuso=4)

    assert ut["D5"].value is None  # value cell azzerata
    # Formula totale riga r3 intatta
    assert ut["D3"].value == "=SUM(D4:D50)"
    # Formula totale colonna C (riga 4: =SUM(D4:L4)) intatta
    assert ut["C4"].value == "=SUM(D4:L4)"


def test_azzera_does_not_touch_other_months(minimal_pf_orti_bytes):
    wb = openpyxl.load_workbook(BytesIO(minimal_pf_orti_bytes), data_only=False)
    ut = wb["Utenze"]
    pre_e4 = ut["E4"].value
    pre_e5 = ut["E5"].value

    azzera_mese(wb, mese_chiuso=4)

    assert ut["E4"].value == pre_e4
    assert ut["E5"].value == pre_e5
