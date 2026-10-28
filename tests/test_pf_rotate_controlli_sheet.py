"""Test per l'avanzamento dei controlli in-foglio al nuovo mese chiuso.

Thread "control-discrepancy": le formule del foglio Controlli sono statiche
(ancorate al mese chiuso di quando furono scritte) e dopo ogni rotation
mostrano ERRORE finti sulla colonna appena chiusa.
"""

import re

from openpyxl import Workbook
from openpyxl.worksheet.formula import ArrayFormula

from verticals.condges.pf_rotate.controlli_sheet import advance_controlli
from verticals.condges.pf_rotate.excel_model import periodo

P = lambda m: periodo(2026, m)  # noqa: E731

MESI = [
    "APRILE",
    "MAGGIO",
    "GIUGNO",
    "LUGLIO",
    "AGOSTO",
    "SETTEMBRE",
    "OTTOBRE",
    "NOVEMBRE",
    "DICEMBRE",
]


def _wb_stale_aprile() -> Workbook:
    """Master minimale: mesi C..K, Controlli ancorati a mese chiuso = APRILE."""
    wb = Workbook()
    wb.remove(wb.active)
    pf = wb.create_sheet("Piano Finanziario")
    pf["A1"] = 2026
    for i, m in enumerate(MESI):
        pf.cell(2, 3 + i, m)  # C..K

    ct = wb.create_sheet("Controlli")
    ct["A4"] = "C4 = Saldo iniziale hardcoded?"
    ct["B4"] = '=IF(_xlfn.ISFORMULA(\'Piano Finanziario\'!C4),"ERRORE","OK")'
    ct["A5"] = "D4 = C37 (catena saldo mese prec)?"
    ct["B5"] = "=IF('Piano Finanziario'!D4='Piano Finanziario'!C37,\"OK\",\"ERRORE\")"
    eqs = ",".join(
        f"'Piano Finanziario'!{chr(68 + i)}4='Piano Finanziario'!{chr(67 + i)}37"
        for i in range(8)  # D4=C37 .. K4=J37
    )
    ct["A6"] = "Catena D4:K4 = mese prec riga 37?"
    ct["B6"] = f'=IF(AND({eqs}),"OK","ERRORE")'
    return wb


def test_advance_sposta_il_confine_al_mese_chiuso():
    wb = _wb_stale_aprile()
    changed = advance_controlli(wb, periodo_chiuso=P(5))

    ct = wb["Controlli"]
    # Check hardcoded: ora copre C4 E D4 (aprile + maggio chiusi)
    assert "ISFORMULA('Piano Finanziario'!C4)" in ct["B4"].value
    assert "ISFORMULA('Piano Finanziario'!D4)" in ct["B4"].value
    # Check apertura: giugno (E) parte dal saldo reale di maggio (D37)
    assert (
        ct["B5"].value
        == "=IF('Piano Finanziario'!E4='Piano Finanziario'!D37,\"OK\",\"ERRORE\")"
    )
    # Catena: parte da E4=D37, arriva a K4=J37, NON contiene più D4=C37
    assert "'Piano Finanziario'!E4='Piano Finanziario'!D37" in ct["B6"].value
    assert "'Piano Finanziario'!K4='Piano Finanziario'!J37" in ct["B6"].value
    assert "'Piano Finanziario'!D4='Piano Finanziario'!C37" not in ct["B6"].value
    # Le etichette seguono il confine
    assert "C4:D4" in ct["A4"].value
    assert "E4" in ct["A5"].value and "D37" in ct["A5"].value
    assert changed  # celle riscritte riportate


def test_advance_idempotente():
    wb = _wb_stale_aprile()
    advance_controlli(wb, periodo_chiuso=P(5))
    snapshot = [
        wb["Controlli"][ref].value for ref in ("A4", "B4", "A5", "B5", "A6", "B6")
    ]
    advance_controlli(wb, periodo_chiuso=P(5))
    assert snapshot == [
        wb["Controlli"][ref].value for ref in ("A4", "B4", "A5", "B5", "A6", "B6")
    ]


def test_advance_noop_senza_foglio_controlli():
    wb = Workbook()
    wb.remove(wb.active)
    pf = wb.create_sheet("Piano Finanziario")
    for i, m in enumerate(MESI):
        pf.cell(2, 3 + i, m)
    assert advance_controlli(wb, periodo_chiuso=P(5)) == []


# ── Layout "matriciale" di Rosa (ORTI_PF_2026-09, 2026-10-28) ────────────────
#
# Il foglio Controlli reale usa ISFORMULA su INTERVALLI dentro SUMPRODUCT/MATCH
# (formule matriciali). In Excel funziona; Google Sheets accetta ISFORMULA solo
# su una cella → #N/A / #VALUE! sui controlli di sezione 1 (segnalato da Stefano
# il 28/10: «saranno sempre excel in drive, dobbiamo risolvere»). La rotation
# deve riscrivere quelle righe in formule esplicite cella-per-cella, valide in
# entrambi.

_SHEETS_UNSAFE = re.compile(r"ISFORMULA\([^)]*:")  # ISFORMULA su un intervallo

MESI_ROSA = [
    "APRILE",
    "MAGGIO",
    "GIUGNO",
    "LUGLIO",
    "AGOSTO",
    "SETTEMBRE",
    "OTTOBRE",
    "NOVEMBRE",
    "DICEMBRE",
    "GENNAIO",
    "FEBBRAIO",
    "MARZO",
    "APRILE",
    "MAGGIO",
    "GIUGNO",
]  # C..Q


def _controlli_rosa(ct) -> None:
    ct["A3"] = "1. SALDO PARTENZA"
    ct["A4"] = (
        "Proiezione contigua: nessun saldo hardcoded dopo il primo mese con formula (D4:Q4)?"
    )
    ct["B4"] = ArrayFormula(
        "B4",
        "=IF(SUMPRODUCT(--NOT(_xlfn.ISFORMULA('Piano Finanziario'!D4:Q4)),"
        '--_xlfn.ISFORMULA(\'Piano Finanziario\'!C4:P4))=0,"OK","ERRORE")',
    )
    ct["A5"] = (
        "Primo mese proiettato (prima formula in riga 4) = riga 37 del mese precedente?"
    )
    ct["B5"] = ArrayFormula(
        "B5",
        "=IF(INDEX('Piano Finanziario'!D4:Q4,MATCH(TRUE,_xlfn.ISFORMULA('Piano Finanziario'!D4:Q4),0))"
        "=INDEX('Piano Finanziario'!C37:P37,MATCH(TRUE,_xlfn.ISFORMULA('Piano Finanziario'!D4:Q4),0)),"
        '"OK","ERRORE")',
    )
    ct["A6"] = "Catena: ogni mese con formula in riga 4 = riga 37 del mese precedente?"
    ct["B6"] = ArrayFormula(
        "B6",
        "=IF(SUMPRODUCT(--_xlfn.ISFORMULA('Piano Finanziario'!D4:Q4),"
        "--(ROUND('Piano Finanziario'!D4:Q4-'Piano Finanziario'!C37:P37,2)<>0))=0,\"OK\",\"ERRORE\")",
    )
    ct["A7"] = "Primo mese proiettato:"
    ct["B7"] = ArrayFormula(
        "B7",
        "=INDEX('Piano Finanziario'!D2:Q2,MATCH(TRUE,_xlfn.ISFORMULA('Piano Finanziario'!D4:Q4),0))",
    )
    ct["A10"] = "Riga 12 = SUM(6:11) per mesi C:Q?"
    ct["B10"] = ArrayFormula(
        "B10",
        "=IF(SUMPRODUCT(--(ROUND('Piano Finanziario'!C12:Q12-'Piano Finanziario'!C6:Q6,2)<>0))=0,\"OK\",\"ERRORE\")",
    )


def _wb_rosa_settembre() -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    pf = wb.create_sheet("Piano Finanziario")
    pf["C1"] = 2026
    pf["L1"] = 2027
    for i, m in enumerate(MESI_ROSA):
        pf.cell(2, 3 + i, m)
    _controlli_rosa(wb.create_sheet("Controlli"))
    return wb


def test_layout_matriciale_diventa_formule_esplicite_valide_in_sheets():
    """Chiuso agosto (G): settembre (H) è il primo mese aperto."""
    wb = _wb_rosa_settembre()
    changed = advance_controlli(wb, periodo_chiuso=periodo(2026, 8))

    ct = wb["Controlli"]
    for ref in ("B4", "B5", "B6", "B7"):
        v = ct[ref].value
        assert isinstance(v, str), ref  # non più ArrayFormula
        assert not _SHEETS_UNSAFE.search(v), (ref, v)
    # 1.1: mesi chiusi hardcoded, mesi aperti formula
    assert "NOT(_xlfn.ISFORMULA('Piano Finanziario'!G4))" in ct["B4"].value
    assert "_xlfn.ISFORMULA('Piano Finanziario'!H4)" in ct["B4"].value
    assert "ISFORMULA('Piano Finanziario'!Q4)" in ct["B4"].value
    # 1.2: primo mese aperto parte dal saldo reale del mese chiuso
    assert (
        ct["B5"].value
        == "=IF('Piano Finanziario'!H4='Piano Finanziario'!G37,\"OK\",\"ERRORE\")"
    )
    # 1.3: catena H4=G37 .. Q4=P37, non tocca il mese chiuso
    assert "'Piano Finanziario'!H4='Piano Finanziario'!G37" in ct["B6"].value
    assert "'Piano Finanziario'!Q4='Piano Finanziario'!P37" in ct["B6"].value
    assert "'Piano Finanziario'!G4='Piano Finanziario'!F37" not in ct["B6"].value
    # etichetta del primo mese proiettato: riferimento diretto
    assert ct["B7"].value == "='Piano Finanziario'!H2"
    # i controlli fuori sezione 1 non si toccano
    assert isinstance(ct["B10"].value, ArrayFormula)
    assert set(changed) == {"B4", "B5", "B6", "B7"}


def test_rotate_estende_prima_di_riscrivere_la_catena(minimal_pf_orti_bytes, tmp_path):
    """La catena dei controlli deve arrivare all'ultima colonna DOPO l'estensione rolling."""
    import io
    from datetime import date

    import openpyxl
    import pandas as pd

    from verticals.condges.pf_rotate.excel_model import find_month_periods
    from verticals.condges.pf_rotate.rotate import rotate
    from verticals.condges.pf_rotate.step3_scadenzario import UnmappedPolicy

    wb = openpyxl.load_workbook(io.BytesIO(minimal_pf_orti_bytes))
    _controlli_rosa(wb.create_sheet("Controlli"))
    pf_path = tmp_path / "in.xlsx"
    wb.save(pf_path)
    forn = tmp_path / "fornitori.csv"
    forn.write_text(
        "codice_fornitore,nome_esolver,nome_pf,voce_id,is_intercompany,is_excluded,exclude_reason,societa_id\n"
        "18,Acqua - Ausino,Acqua - Ausino,USCITE_UTENZE,False,False,,ORTI\n",
        encoding="utf-8",
    )
    result = rotate(
        pf_path=pf_path,
        scad_df=pd.DataFrame(
            [
                {
                    "codice_fornitore": 18,
                    "nome": "Acqua Ausino",
                    "totale": -1.0,
                    "scaduto": 0.0,
                    f"mese_{periodo(2026, 5)}": -1.0,
                }
            ]
        ),
        bucket_periodi=[periodo(2026, 5)],
        societa="ORTI",
        periodo_chiuso=periodo(2026, 4),
        data_saldo=date(2026, 4, 30),
        saldi={"MPS": 100.0, "Intesa": 50.0, "MPS_KROSS": 0.0, "UNICREDIT": 0.0},
        fornitori_csv=forn,
        out_dir=tmp_path / "out",
        unmapped_policy=UnmappedPolicy.FAIL,
    )
    out = openpyxl.load_workbook(result.out_path)
    cols = find_month_periods(out["Piano Finanziario"])
    last_l = openpyxl.utils.get_column_letter(cols[periodo(2027, 4)])
    assert f"'Piano Finanziario'!{last_l}4=" in out["Controlli"]["B6"].value
