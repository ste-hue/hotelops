"""Shared pytest fixtures for hotelops tests."""

import os
from io import BytesIO
from unittest.mock import MagicMock, patch

from google.cloud import bigquery
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
import pytest

from core.config import PROJECT

DATAHUB_FOLDERS = [
    "homebanking/ORTI",
    "homebanking/INTUR",
    "movimenti_contabili/ORTI",
    "movimenti_contabili/INTUR",
    "registro_banca_esolver/ORTI",
    "registro_banca_esolver/INTUR",
    "partite_fornitori/ORTI",
    "partite_fornitori/INTUR",
    "piani_finanziari/ORTI",
    "piani_finanziari/INTUR",
    "accodamenti/ORTI",
    "economato",
    "coperti",
    "bilancino/ORTI",
    "bilancino/INTUR",
    "gasparotto",
    "dimensioni",
    "fatti",
    "meta",
]


@pytest.fixture()
def tmp_datahub(tmp_path):
    """Create a minimal datahub directory structure under tmp_path."""
    for folder in DATAHUB_FOLDERS:
        (tmp_path / folder).mkdir(parents=True, exist_ok=True)
    return tmp_path


@pytest.fixture()
def mock_bq_client():
    """Patch google.cloud.bigquery.Client with a MagicMock.

    Yields the mock client so tests can configure return values.
    """
    with patch("google.cloud.bigquery.Client") as mock_cls:
        client = MagicMock()
        mock_cls.return_value = client
        yield client


@pytest.fixture(scope="session")
def bq_client():
    """Real BigQuery client for invariant tests against live views.

    Skipped when HOTELOPS_SKIP_BQ=1 (CI / offline development).
    Tests using this fixture should also be marked @pytest.mark.bq.
    """
    if os.environ.get("HOTELOPS_SKIP_BQ") == "1":
        pytest.skip("HOTELOPS_SKIP_BQ=1 — skipping live BigQuery test")
    return bigquery.Client(project=PROJECT)


def _build_minimal_pf_orti(mese_chiuso_col: int = 3) -> BytesIO:
    """Construct a minimal ORTI PF xlsx in-memory.

    Layout = osservato su 05_ORTIFinancialPlan2026.xlsx (post-witness 2026-05-19).
    mese_chiuso_col = 3 means APRILE is in col C (the closed/cutover column).
    """
    wb = Workbook()
    wb.remove(wb.active)

    # ── Piano Finanziario master ─────────────────────────────────────────
    pf = wb.create_sheet("Piano Finanziario")
    pf["A1"] = "ORTI S.R.L."
    pf["C1"] = 2026
    mesi = [
        "APRILE",
        "MAGGIO",
        "GIUGNO",
        "LUGLIO",
        "AGOSTO",
        "SETTEMBRE",
        "OTTOBRE",
        "NOVEMBRE",
        "DICEMBRE",
    ]
    for i, m in enumerate(mesi):
        pf.cell(2, 3 + i, m)
    pf["L2"] = "TOTALI"

    pf["A4"] = "SALDO MESE PRECEDENTE"
    pf["C4"] = 339337.46  # hardcoded cutover
    for i in range(1, 9):  # D4..K4 cascade
        col = 3 + i
        prev_col = chr(ord("A") + col - 2)  # col 4 (D) -> prev col 3 (C); ref to col-1
        pf.cell(4, col, f"={prev_col}37")

    # Entrate r6-r11 (hardcoded numeric, mese_chiuso azzerato in step 2)
    entrate_labels = [
        "Entrate Hotel ",
        "Entrate Residence",
        "Entrate CVM",
        "Entrate Supermercato",
        "Rientro Sospesi",
        "Caparre Intur",
    ]
    for i, lbl in enumerate(entrate_labels):
        pf.cell(6 + i, 1, lbl)
        # leave mese_chiuso col empty (closed), put 100 in MAGGIO (col D) for sanity
        pf.cell(6 + i, 4, 100.0)

    pf["A12"] = "TOTALE ENTRATE"
    for c in range(3, 12):
        col_letter = chr(ord("A") + c - 1)
        pf.cell(12, c, f"=SUM({col_letter}6:{col_letter}11)")
    pf["L12"] = "=SUM(C12:K12)"

    # Uscite r14-r23 link a fogli dettaglio (qui semplificato: 2 voci)
    pf["A15"] = "Utenze"
    pf["A16"] = "Materie Prime/Consumo"
    for c in range(3, 12):
        col_letter = chr(ord("A") + c - 1)
        pf.cell(15, c, f"=Utenze!{col_letter}3")
        pf.cell(16, c, f"='Materie Prime-Consumo '!{col_letter}3")

    pf["A27"] = "TOTALE USCITE"
    for c in range(3, 12):
        col_letter = chr(ord("A") + c - 1)
        pf.cell(27, c, f"=SUM({col_letter}14:{col_letter}26)")
    pf["L27"] = "=SUM(C27:K27)"

    pf["A29"] = "Cash Flow"
    for c in range(3, 12):
        col_letter = chr(ord("A") + c - 1)
        pf.cell(29, c, f"={col_letter}12-{col_letter}27")
    pf["L29"] = "=L12-L27"

    # Saldi banca area (col del mese chiuso)
    pf["C31"] = "30/04/2026"
    pf["A32"] = "Saldo MPS"
    pf["C32"] = 251897.54
    pf["A33"] = "Saldo Intesa"
    pf["C33"] = 87439.92
    pf["A35"] = "TOTALE BANCHE"
    pf["C35"] = "=SUM(C32:C34)"

    pf["A37"] = "Saldo di Periodo"
    for c in range(3, 12):
        col_letter = chr(ord("A") + c - 1)
        pf.cell(37, c, f"={col_letter}4+{col_letter}29")
    pf["L37"] = "=K37"

    # ── Foglio Utenze (dettaglio, codice in col A) ──────────────────────
    ut = wb.create_sheet("Utenze")
    ut["D1"] = 2026
    ut["A2"] = "CODICE"
    for i, m in enumerate(mesi):
        ut.cell(2, 4 + i, m)
    ut["B3"] = "Utenze"
    ut["C3"] = "=SUM(D3:L3)"
    for i in range(9):  # D3..L3
        col = 4 + i
        col_letter = chr(ord("A") + col - 1)
        ut.cell(3, col, f"=SUM({col_letter}4:{col_letter}50)")
    # 2 fornitori
    ut["A4"] = 570913
    ut["B4"] = "Energia elettrica"
    ut["C4"] = "=SUM(D4:L4)"
    ut["E4"] = 6687.85  # MAGGIO
    ut["A5"] = 18
    ut["B5"] = "Acqua - Ausino"
    ut["C5"] = "=SUM(D5:L5)"
    ut["D5"] = 2264.12  # APRILE (will be cleared by step 2)
    ut["E5"] = 8143.5

    # ── Foglio Materie Prime (idem) ─────────────────────────────────────
    mp = wb.create_sheet("Materie Prime-Consumo ")
    mp["D1"] = 2026
    mp["A2"] = "CODICE"
    for i, m in enumerate(mesi):
        mp.cell(2, 4 + i, m)
    mp["B3"] = "Materie Prime/Consumo"
    mp["C3"] = "=SUM(D3:L3)"
    for i in range(9):
        col = 4 + i
        col_letter = chr(ord("A") + col - 1)
        mp.cell(3, col, f"=SUM({col_letter}4:{col_letter}50)")
    mp["A4"] = 92
    mp["B4"] = "Amalfi sei esse"
    mp["C4"] = "=SUM(D4:L4)"
    mp["D4"] = 357.86
    mp["E4"] = 357.86

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


@pytest.fixture
def minimal_pf_orti_bytes() -> bytes:
    """ORTI PF minimal: APRILE chiuso, 2 fogli dettaglio (Utenze, Materie Prime)."""
    return _build_minimal_pf_orti().getvalue()


def _build_minimal_pf_intur() -> BytesIO:
    """Construct a minimal INTUR-style PF xlsx in-memory.

    Layout: saldo mese precedente at r3, entrate r5-r10, uscite r14-r22,
    saldi banca r31-r33, saldo proiettato r38. Month cols: MARZO=D, APRILE=E..DIC=M.
    Col C = "DATA RILEVAZ" (not a month).
    """
    wb = Workbook()
    wb.remove(wb.active)

    # ── Piano Finanziario master ─────────────────────────────────────────
    pf = wb.create_sheet("Piano Finanziario")
    pf["A1"] = "INTUR S.R.L."
    pf["D1"] = 2026

    # Row 2: headers. Col C = DATA RILEVAZ (non-month), D=MARZO, E=APRILE, ...
    pf["C2"] = "DATA RILEVAZ"
    mesi = [
        "MARZO",
        "APRILE",
        "MAGGIO",
        "GIUGNO",
        "LUGLIO",
        "AGOSTO",
        "SETTEMBRE",
        "OTTOBRE",
        "NOVEMBRE",
        "DICEMBRE",
    ]
    for i, m in enumerate(mesi):
        pf.cell(2, 4 + i, m)  # cols D(4)..M(13)

    # r3: SALDO MESE PRECEDENTE — D3 hardcoded, E3..M3 cascade from prev R38
    pf["A3"] = "SALDO MESE PRECEDENTE"
    pf["D3"] = 719445.19  # hardcoded cutover (MARZO closed)
    for i in range(1, 10):  # E3..M3 cascade
        col_idx = 4 + i  # E=5, F=6, ...
        prev_col = get_column_letter(col_idx - 1)
        pf.cell(3, col_idx, f"={prev_col}38")

    # r5-r10: Entrate
    entrate_labels = [
        "Fitto Hotel",
        "Fitto AR",
        "Entrate Residence",
        "Ribaltamento Costi",
        "Entrate Farmacia",
        "Entrate Spiaggia",
    ]
    for i, lbl in enumerate(entrate_labels):
        pf.cell(5 + i, 1, lbl)
        # leave MARZO (col D) empty (closed month), put 500 in APRILE (col E)
        pf.cell(5 + i, 5, 500.0)

    pf["A12"] = "TOTALE ENTRATE"
    for c in range(4, 14):  # cols D..M
        cl = get_column_letter(c)
        pf.cell(12, c, f"=SUM({cl}5:{cl}11)")

    # r14-r22: Uscite (9 voci)
    uscite_labels = [
        "Salari e Stipendi",
        "Utenze",
        "Materie Prime",
        "Tasse e Imposte",
        "Mutui e Finanziamenti",
        "Consulenze",
        "Godimento Beni di Terzi",
        "Varie ed Eventuali",
        "Canoni e servizi",
    ]
    for i, lbl in enumerate(uscite_labels):
        pf.cell(14 + i, 1, lbl)
        pf.cell(14 + i, 5, 300.0)  # APRILE

    pf["A27"] = "TOTALE USCITE"
    for c in range(4, 14):
        cl = get_column_letter(c)
        pf.cell(27, c, f"=SUM({cl}14:{cl}26)")

    pf["A29"] = "CASH FLOW"
    for c in range(4, 14):
        cl = get_column_letter(c)
        pf.cell(29, c, f"={cl}12-{cl}27")

    # Saldi banca area
    pf["A31"] = "Saldo Banca Sella"
    pf["D31"] = 35186.34
    pf["A32"] = "Saldo MPS"
    pf["D32"] = 682801.07
    pf["A33"] = "Saldo Intesa"
    pf["D33"] = 1457.78

    pf["A35"] = "TOTALE BANCHE"
    for c in range(4, 14):
        cl = get_column_letter(c)
        pf.cell(35, c, f"=SUM({cl}31:{cl}34)")

    # r38: Saldo di Periodo
    pf["A38"] = "Saldo di Periodo"
    for c in range(4, 14):
        cl = get_column_letter(c)
        pf.cell(38, c, f"={cl}3+({cl}29+{cl}35)")

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


@pytest.fixture
def intur_pf_bytes() -> bytes:
    """INTUR PF minimal: MARZO chiuso, saldo r3, saldi banca r31-33, saldo proiettato r38."""
    return _build_minimal_pf_intur().getvalue()


@pytest.fixture(autouse=True)
def _mai_il_lock_pec_vero(monkeypatch):
    """Rete di sicurezza: nessun test può prendere il lock PEC reale su
    `gs://hotelops-raw`. Un test di `main()` che dimenticasse di fingere il
    lock, girando su una macchina con ADC configurate, scriverebbe in
    produzione — e se crashasse la bloccherebbe per due ore.

    I test del lock passano un client finto: quelli continuano a funzionare.
    """
    from ingest import pec_lock

    reale = pec_lock._blob

    def guardia(bucket, client=None):
        if client is None:
            raise AssertionError(
                "questo test sta per prendere il lock PEC VERO su GCS: "
                "inietta un client finto o usa la fixture `lock_libero`"
            )
        return reale(bucket, client)

    monkeypatch.setattr(pec_lock, "_blob", guardia)
