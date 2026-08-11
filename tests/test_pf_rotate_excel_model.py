from io import BytesIO
import openpyxl
import pytest

from verticals.condges.pf_rotate.excel_model import (
    find_layout,
    find_month_periods,
    is_value_cell,
    is_formula_with_refs,
    resolve_sheet_name,
    find_total_row_and_sum_range,
    periodo,
    periodo_anno_mese,
    require_periodo,
)


def test_find_month_periods_returns_periodo_to_col_map(minimal_pf_orti_bytes):
    wb = openpyxl.load_workbook(BytesIO(minimal_pf_orti_bytes), data_only=False)
    ws = wb["Utenze"]
    cols = find_month_periods(ws, header_row=2)
    assert cols == {
        periodo(2026, 4): 4,  # APRILE -> col D (4)
        periodo(2026, 5): 5,
        periodo(2026, 6): 6,
        periodo(2026, 7): 7,
        periodo(2026, 8): 8,
        periodo(2026, 9): 9,
        periodo(2026, 10): 10,
        periodo(2026, 11): 11,
        periodo(2026, 12): 12,
    }


def test_is_value_cell_pure_number():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = 123.45
    assert is_value_cell(ws["A1"]) is True
    assert is_formula_with_refs(ws["A1"]) is False


def test_is_value_cell_formula_only_constants():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "=110000+160000"
    assert is_value_cell(ws["A1"]) is True
    assert is_formula_with_refs(ws["A1"]) is False


def test_is_formula_with_refs_a1_reference():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "=B5"
    assert is_value_cell(ws["A1"]) is False
    assert is_formula_with_refs(ws["A1"]) is True


def test_is_formula_with_refs_cross_sheet():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "=Utenze!D3"
    assert is_value_cell(ws["A1"]) is False
    assert is_formula_with_refs(ws["A1"]) is True


def test_is_formula_with_refs_sum():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "=SUM(C6:C11)"
    assert is_formula_with_refs(ws["A1"]) is True


def test_is_value_cell_empty_is_false():
    wb = openpyxl.Workbook()
    ws = wb.active
    assert is_value_cell(ws["A1"]) is False  # None


def test_resolve_sheet_name_exact():
    wb = openpyxl.Workbook()
    wb.active.title = "Materie Prime-Consumo "
    wb.create_sheet("Utenze")
    assert (
        resolve_sheet_name(wb, ["Materie Prime-Consumo ", "Materie Prime-Conumo "])
        == "Materie Prime-Consumo "
    )


def test_resolve_sheet_name_typo_fallback():
    wb = openpyxl.Workbook()
    wb.active.title = "Materie Prime-Conumo "  # typo variant
    assert (
        resolve_sheet_name(wb, ["Materie Prime-Consumo ", "Materie Prime-Conumo "])
        == "Materie Prime-Conumo "
    )


def test_resolve_sheet_name_none_when_absent():
    wb = openpyxl.Workbook()
    wb.active.title = "Foo"
    assert resolve_sheet_name(wb, ["Bar", "Baz"]) is None


def test_find_total_row_and_sum_range_utenze(minimal_pf_orti_bytes):
    wb = openpyxl.load_workbook(BytesIO(minimal_pf_orti_bytes), data_only=False)
    ws = wb["Utenze"]
    # In fixture: r3 has =SUM(D4:D50), =SUM(E4:E50), ...
    total_row, sum_start, sum_end = find_total_row_and_sum_range(ws)
    assert total_row == 3
    assert sum_start == 4
    assert sum_end == 50


def test_find_layout_orti_fixture(minimal_pf_orti_bytes):
    wb = openpyxl.load_workbook(BytesIO(minimal_pf_orti_bytes), data_only=False)
    layout = find_layout(wb)
    assert layout.saldo_iniziale_row == 4
    assert layout.totale_entrate_row == 12
    assert layout.totale_uscite_row == 27
    assert layout.cashflow_row == 29
    assert layout.totale_banche_row == 35
    assert layout.saldo_proiettato_row == 37
    assert layout.saldi_banca_rows == [32, 33]
    assert layout.entrate_rows == [6, 7, 8, 9, 10, 11]


def test_find_layout_intur_style_fixture():
    """Build a minimal INTUR-style workbook (saldo at r3, etc.) e verifica layout detection."""
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    pf = wb.create_sheet("Piano Finanziario")
    pf["A1"] = "INTUR"
    # INTUR: saldo at r3 (NOT r4)
    pf["A3"] = "SALDO MESE PRECED."
    pf["A5"] = "Fitto Hotel"
    pf["A6"] = "Fitto AR"
    pf["A7"] = "Entrate Residence"
    pf["A12"] = "TOTALE ENTRATE"
    pf["A14"] = "Salari e Stipendi"
    pf["A15"] = "Utenze"
    pf["A27"] = "TOTALE USCITE"
    pf["A29"] = "CASH FLOW"
    pf["A31"] = "Saldo Banca Sella"
    pf["A32"] = "Saldo MPS"
    pf["A33"] = "Saldo Intesa"
    pf["A35"] = "TOTALE BANCHE"
    pf["A38"] = "Saldo di Periodo/Proiettato"

    layout = find_layout(wb)
    assert layout.saldo_iniziale_row == 3
    assert layout.totale_entrate_row == 12
    assert layout.totale_uscite_row == 27
    assert layout.cashflow_row == 29
    assert layout.totale_banche_row == 35
    assert layout.saldo_proiettato_row == 38
    assert layout.saldi_banca_rows == [31, 32, 33]
    assert layout.entrate_rows == [5, 6, 7]


def test_periodo_roundtrip():
    p = periodo(2026, 6)
    assert p == 2026 * 12 + 5
    assert periodo_anno_mese(p) == (2026, 6)


def test_periodo_successivo_attraversa_l_anno():
    assert periodo(2026, 12) + 1 == periodo(2027, 1)


def test_require_periodo_rifiuta_mese_nudo():
    with pytest.raises(ValueError, match="mese nudo"):
        require_periodo(6)
    with pytest.raises(ValueError, match="mese nudo"):
        require_periodo(2026)  # anche un anno nudo è sospetto
    require_periodo(periodo(2026, 6))  # non solleva


def _ws_multi_anno():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["C1"] = 2026
    for i, nome in enumerate(
        ["APRILE", "MAGGIO", "GIUGNO", "LUGLIO", "AGOSTO", "SETTEMBRE",
         "OTTOBRE", "NOVEMBRE", "DICEMBRE", "GENNAIO", "FEBBRAIO", "MARZO",
         "APRILE", "MAGGIO", "GIUGNO"]
    ):
        ws.cell(2, 3 + i, nome)
    return ws


def test_find_month_periods_wrap_dic_gen():
    from verticals.condges.pf_rotate.excel_model import find_month_periods, periodo
    cols = find_month_periods(_ws_multi_anno())
    assert cols[periodo(2026, 4)] == 3
    assert cols[periodo(2026, 12)] == 11
    assert cols[periodo(2027, 1)] == 12   # wrap: DIC→GEN incrementa l'anno
    assert cols[periodo(2027, 4)] == 15   # APRILE 2027 ≠ APRILE 2026
    assert cols[periodo(2026, 4)] == 3    # ...che resta al suo posto
    assert len(cols) == 15


def test_find_month_periods_senza_anno_esplode():
    from verticals.condges.pf_rotate.excel_model import find_month_periods
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(2, 3, "APRILE")  # nessun anno in riga 1
    with pytest.raises(ValueError, match="[Aa]nno non dichiarato"):
        find_month_periods(ws)
