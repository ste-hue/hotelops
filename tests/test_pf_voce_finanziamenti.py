"""La riga principale "Finanziamento ..." del PF INTUR (es. "Finanziamento MPS 550K ARREDI")
va su ENTRATE_FINANZIAMENTI invece di essere scartata come non mappata."""

from __future__ import annotations

import csv
import logging
from pathlib import Path

import pytest
from openpyxl import Workbook

from ingest.flussi.ingest_piano_finanziario_xlsx import _match_voce, parse_piano_finanziario

MESI = [
    "GENNAIO", "FEBBRAIO", "MARZO", "APRILE", "MAGGIO", "GIUGNO",
    "LUGLIO", "AGOSTO", "SETTEMBRE", "OTTOBRE", "NOVEMBRE", "DICEMBRE",
]


def _pf_intur(tmp_path: Path, etichetta: str, col: int, valore: float) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Piano Finanziario"
    ws["C1"] = 2026
    ws["A2"] = "INTUR"
    for i, m in enumerate(MESI):
        ws.cell(2, 3 + i, m)
    ws["A15"] = etichetta
    ws.cell(15, col, valore)
    p = tmp_path / "INTUR_PF_2026-09.xlsx"
    wb.save(p)
    return p


def test_riga_finanziamento_mappata(tmp_path, caplog):
    p = _pf_intur(tmp_path, "Finanziamento MPS 550K ARREDI", 13, 550_000.0)  # M = novembre
    with caplog.at_level(logging.WARNING):
        rows = parse_piano_finanziario(p, None, logging.getLogger("t"))
    rec = [r for r in rows if r["voce_id"] == "ENTRATE_FINANZIAMENTI"]
    assert [(r["anno"], r["mese"], r["importo"]) for r in rec] == [(2026, 11, pytest.approx(550_000.0))]
    assert "UNMAPPED" not in caplog.text


def test_match_voce_finanziamento():
    assert _match_voce("  Finanziamento MPS 550K ARREDI ") == "ENTRATE_FINANZIAMENTI"
    assert _match_voce("Mutui e Finanziamenti") == "USCITE_MUTUI"


def test_voce_nel_catalogo():
    csv_path = Path(__file__).resolve().parents[1] / "core" / "bq" / "dimensioni" / "d_voci_piano_finanziario.csv"
    with csv_path.open(newline="") as fh:
        voci = {r["voce_id"]: r for r in csv.DictReader(fh)}
    v = voci["ENTRATE_FINANZIAMENTI"]
    assert (v["sezione"], v["categoria"], v["societa_id"], v["fonte"]) == ("ENTRATE", "Finanziamenti", "", "MANUALE")
