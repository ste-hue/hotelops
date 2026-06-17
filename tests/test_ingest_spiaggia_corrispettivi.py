from datetime import date, datetime, timezone
import pytest
from core.schemas import SpiaggiaCorrispettivoRow, make_hash

import openpyxl
from ingest.flussi.ingest_spiaggia_corrispettivi import (
    to_eur, parse_data, iter_day_rows, build_corrispettivo_rows, ingest_file,
)
from datetime import date as _d, datetime as _dt, timezone as _tz

_NOW = _dt(2026, 6, 17, tzinfo=_tz.utc)


def _now():
    return datetime(2026, 6, 17, 12, 0, tzinfo=timezone.utc)


def test_corrispettivo_row_valid():
    r = SpiaggiaCorrispettivoRow(
        societa_id="INTUR",
        business_unit_id="LIDO",
        location_id="LIDO",
        data=date(2026, 6, 16),
        anno=2026,
        mese=6,
        corrispettivo_spiaggia=561.0,
        corrispettivo_bar=517.5,
        corrispettivo_totale=1078.5,
        rt_matricola="RT2CFL018343",
        file_sorgente="registro_2026.xlsx",
        hash_riga="abc",
        data_caricamento=_now(),
    )
    assert r.societa_id == "INTUR"
    assert r.corrispettivo_totale == 1078.5


def test_corrispettivo_row_rejects_bad_societa():
    with pytest.raises(Exception):
        SpiaggiaCorrispettivoRow(
            societa_id="PIPPO", business_unit_id="LIDO",
            data=date(2026, 6, 16), anno=2026, mese=6,
            corrispettivo_spiaggia=1.0, corrispettivo_bar=1.0,
            corrispettivo_totale=2.0,
            file_sorgente="x.xlsx", hash_riga="x", data_caricamento=_now(),
        )


def test_to_eur():
    assert to_eur(" 1,083.00 ") == 1083.0
    assert to_eur(" - ") is None
    assert to_eur("866.00") == 866.0
    assert to_eur(None) is None
    assert to_eur(527.0) == 527.0


def test_parse_data():
    assert parse_data("16-Jun", 2026) == _d(2026, 6, 16)
    assert parse_data("1-Jul", 2026) == _d(2026, 7, 1)
    assert parse_data("Totale mese", 2026) is None
    assert parse_data(None, 2026) is None
    # Bug 1: stale-year override — datetime/date cells from 2021 template must use header anno
    from datetime import datetime as _dtmod
    assert parse_data(_dtmod(2021, 6, 16), 2026) == _d(2026, 6, 16)
    assert parse_data(_d(2021, 6, 5), 2026) == _d(2026, 6, 5)


def _make_registro(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Giugno"
    ws["E1"] = "Giugno"; ws["I1"] = "ANNO:"; ws["J1"] = "2026"
    ws["M1"] = "RT"; ws["N1"] = "RT2CFL018343"
    # header rows then day rows: [n, data, totale, 22%, 10%] starting col A
    data = [
        (6, "6-Jun", "866.00", "339.00", "527.00"),
        (16, "16-Jun", "1,078.50", "561.00", "517.50"),
        (None, "Totale mese", "1,944.50", "900.00", "1,044.50"),
    ]
    r0 = 8
    for i, row in enumerate(data):
        for j, val in enumerate(row):
            ws.cell(row=r0 + i, column=1 + j, value=val)
    p = tmp_path / "registro_2026.xlsx"
    wb.save(p)
    return p


def test_iter_day_rows(tmp_path):
    p = _make_registro(tmp_path)
    wb = openpyxl.load_workbook(p, data_only=True)
    rows = iter_day_rows(wb["Giugno"], 2026)
    assert (_d(2026, 6, 6), 866.0, 339.0, 527.0) in rows
    assert (_d(2026, 6, 16), 1078.5, 561.0, 517.5) in rows
    assert len(rows) == 2  # "Totale mese" escluso


def test_build_rows_and_quality_gate(tmp_path):
    p = _make_registro(tmp_path)
    wb = openpyxl.load_workbook(p, data_only=True)
    rows = build_corrispettivo_rows(wb, "registro_2026.xlsx", "raw-1", _NOW)
    assert len(rows) == 2
    r = [x for x in rows if x.data == _d(2026, 6, 16)][0]
    assert r.corrispettivo_spiaggia == 561.0
    assert r.corrispettivo_bar == 517.5
    assert r.corrispettivo_totale == 1078.5
    assert r.societa_id == "INTUR" and r.business_unit_id == "LIDO" and r.location_id == "LIDO"
    assert r.rt_matricola == "RT2CFL018343"
    assert r.raw_object_id == "raw-1"
    assert r.hash_riga == make_hash("corrispettivi", "INTUR", _d(2026, 6, 16).isoformat())


def test_ingest_file_dry_run(tmp_path):
    p = _make_registro(tmp_path)
    counts = ingest_file(p, raw_object_id=None, dry_run=True)
    assert counts["corrispettivi"] == 2


def _make_registro_datetime(tmp_path):
    """Fixture with stale-year datetime cells and a zero-value day row."""
    from datetime import datetime as _dtmod
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Giugno"
    # Header: ANNO: 2026 in row 2 cols D/E (col 4/5)
    ws.cell(row=2, column=4, value="ANNO:")
    ws.cell(row=2, column=5, value=2026)
    # Day row with real values: datetime cell has stale year 2021
    ws.cell(row=8, column=1, value=_dtmod(2021, 6, 16))   # stale year
    ws.cell(row=8, column=2, value=1078.5)
    ws.cell(row=8, column=3, value=561.0)
    ws.cell(row=8, column=4, value=517.5)
    # Zero-value day row (closed day): should be excluded
    ws.cell(row=9, column=1, value=_dtmod(2021, 6, 17))   # stale year
    ws.cell(row=9, column=2, value=0)
    ws.cell(row=9, column=3, value=0)
    ws.cell(row=9, column=4, value=0)
    p = tmp_path / "registro_dt_2026.xlsx"
    wb.save(p)
    return p


def test_stale_year_and_zero_day_exclusion(tmp_path):
    """Bug 1 + Bug 2: stale datetime year uses header anno; zero-day rows are excluded."""
    p = _make_registro_datetime(tmp_path)
    wb = openpyxl.load_workbook(p, data_only=True)
    rows = build_corrispettivo_rows(wb, "registro_dt_2026.xlsx", "raw-dt", _NOW)
    # Only 1 row (the real-value row); zero row excluded
    assert len(rows) == 1
    r = rows[0]
    # Year must come from header (2026), not stale datetime cell (2021)
    assert r.data == _d(2026, 6, 16)
    assert r.anno == 2026
    assert r.data.year == 2026
    assert r.corrispettivo_spiaggia == 561.0
    assert r.corrispettivo_bar == 517.5
