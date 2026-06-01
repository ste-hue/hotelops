from datetime import date
from io import BytesIO

import openpyxl

from verticals.condges.pf_rotate.step5_controlli import verifica_controlli, CheckOutcome
from verticals.condges.pf_rotate.step1_saldi import write_saldi_banca
from verticals.condges.pf_rotate.step2_azzera import azzera_mese


def test_controlli_su_fixture_minimal_post_step2(minimal_pf_orti_bytes):
    """Post step 2: tutti i controlli devono passare sul fixture minimale."""
    wb = openpyxl.load_workbook(BytesIO(minimal_pf_orti_bytes), data_only=False)
    # Replicate the witness flow: APRILE già "chiuso" nel fixture (entrate vuote),
    # quindi step 2 ulteriore è no-op. Eseguilo lo stesso per simmetria.
    azzera_mese(wb, mese_chiuso=4)

    report = verifica_controlli(wb, mese_chiuso=4)

    assert report.n_err == 0, (
        f"Errori inattesi: {[r for r in report.results if r.outcome == CheckOutcome.ERR]}"
    )


def test_controllo_1_c4_must_be_hardcoded(minimal_pf_orti_bytes):
    """Se C4 fosse formula → check #1 ERR."""
    wb = openpyxl.load_workbook(BytesIO(minimal_pf_orti_bytes), data_only=False)
    wb["Piano Finanziario"]["C4"] = "=C35"  # vandalismo

    report = verifica_controlli(wb, mese_chiuso=4)
    check_1 = next(r for r in report.results if r.check_id == "C1")
    assert check_1.outcome == CheckOutcome.ERR


def test_controllo_3_cascade_d4_eq_c37(minimal_pf_orti_bytes):
    """D4 cascade = C37 + r29 chain. Se cambio D4 a un numero hardcoded → ERR."""
    wb = openpyxl.load_workbook(BytesIO(minimal_pf_orti_bytes), data_only=False)
    wb["Piano Finanziario"]["D4"] = 999999.0  # rompe la cascade

    report = verifica_controlli(wb, mese_chiuso=4)
    check_3 = next(r for r in report.results if r.check_id == "C3")
    assert check_3.outcome == CheckOutcome.ERR


def test_controlli_su_intur_style_fixture(intur_pf_bytes):
    """Verifica engine layout-aware su INTUR-style layout."""
    wb = openpyxl.load_workbook(BytesIO(intur_pf_bytes), data_only=False)
    write_saldi_banca(
        wb,
        mese_chiuso=3,  # MARZO
        data_saldo=date(2026, 3, 31),
        saldi={"MPS": 682801.07, "Intesa": 1457.78, "Sella": 35186.34},
    )
    azzera_mese(wb, mese_chiuso=3)
    report = verifica_controlli(wb, mese_chiuso=3)
    assert report.n_err == 0, (
        f"Errori: {[r for r in report.results if r.outcome == CheckOutcome.ERR]}"
    )


def test_verifica_controlli_uses_mese_chiuso_not_first_col(intur_pf_bytes):
    """Cutover col deve derivare da mese_chiuso, non dalla prima colonna mese.

    Setup: D=MARZO con formula `=C35` (cascata), E=APRILE con 100.0 hardcoded.
    - mese_chiuso=4 → C1 controlla E (hardcoded) → OK.
    - mese_chiuso=3 → C1 controlla D (formula) → ERR.
    """
    wb = openpyxl.load_workbook(BytesIO(intur_pf_bytes), data_only=False)
    pf = wb["Piano Finanziario"]
    pf["D3"] = "=C35"  # cascata legittima per APRILE-chiuso
    pf["E3"] = 100.0  # cutover hardcoded per APRILE-chiuso

    report_apr = verifica_controlli(wb, mese_chiuso=4)
    c1_apr = next(r for r in report_apr.results if r.check_id == "C1")
    assert c1_apr.outcome == CheckOutcome.OK

    report_mar = verifica_controlli(wb, mese_chiuso=3)
    c1_mar = next(r for r in report_mar.results if r.check_id == "C1")
    assert c1_mar.outcome == CheckOutcome.ERR


def test_c1_fixed_snapshot_ok(intur_pf_bytes):
    """Layout fixed-snapshot (INTUR): C1 OK quando C2 popolato + saldi banca in col C hardcoded."""
    wb = openpyxl.load_workbook(BytesIO(intur_pf_bytes), data_only=False)
    pf = wb["Piano Finanziario"]
    pf["C1"] = "DATA RILEVAZ"  # attiva branch fixed-snapshot
    pf["C2"] = "30/04/2026"
    pf["C31"] = 49609.01
    pf["C32"] = 376510.97
    pf["C33"] = 1696.48

    report = verifica_controlli(wb, mese_chiuso=4)
    c1 = next(r for r in report.results if r.check_id == "C1")
    assert c1.outcome == CheckOutcome.OK, f"detail={c1.detail}"
    assert "Snapshot C" in c1.title


def test_c1_fixed_snapshot_err_c2_vuota(intur_pf_bytes):
    """Layout fixed-snapshot: C1 ERR quando C2 (data rilevazione) è vuota."""
    wb = openpyxl.load_workbook(BytesIO(intur_pf_bytes), data_only=False)
    pf = wb["Piano Finanziario"]
    pf["C1"] = "DATA RILEVAZ"
    pf["C2"] = None  # data rilevazione mancante
    pf["C31"] = 49609.01
    pf["C32"] = 376510.97
    pf["C33"] = 1696.48

    report = verifica_controlli(wb, mese_chiuso=4)
    c1 = next(r for r in report.results if r.check_id == "C1")
    assert c1.outcome == CheckOutcome.ERR
    assert "C2 vuota" in c1.detail
