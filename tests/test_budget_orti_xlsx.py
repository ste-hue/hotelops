"""Tests for Budget_ORTI monthly-grid → f_budget_mensile parser."""

from __future__ import annotations

import logging
from pathlib import Path

import openpyxl
import pytest

from ingest.flussi.budget_orti_xlsx import (
    parse_budget_orti_workbook,
    sniff_budget_orti_monthly_format,
)


@pytest.fixture
def logger() -> logging.Logger:
    return logging.getLogger("test_budget_orti_xlsx")


def test_sniff_true_without_budget_sheet():
    assert sniff_budget_orti_monthly_format(["Ricavi", "Fissi", "CE"])


def test_sniff_false_when_budget_present():
    assert not sniff_budget_orti_monthly_format(["Budget", "Ricavi", "Fissi"])


def test_parse_minimal_workbook(tmp_path: Path, logger: logging.Logger):
    path = tmp_path / "Budget_ORTI_2026.xlsx"
    wb = openpyxl.Workbook()
    default = wb.active
    wb.remove(default)

    ws = wb.create_sheet("Ricavi")
    headers = [
        "codice_conto",
        "descrizione",
        "BU",
        "Gen",
        "Feb",
        "Mar",
        "Apr",
        "Mag",
        "Giu",
        "Lug",
        "Ago",
        "Set",
        "Ott",
        "Nov",
        "Dic",
    ]
    for col, h in enumerate(headers, start=1):
        ws.cell(1, col, h)
    ws.append(
        [
            479101,
            "Camera doppia",
            "HOTEL",
            0,
            150.5,
            0,
            0,
            0,
            0,
            0,
            0,
            0,
            0,
            0,
            0,
        ]
    )
    wb.save(path)
    wb.close()

    rows = parse_budget_orti_workbook(path, "ORTI", 2026, logger)
    assert len(rows) == 1
    r = rows[0]
    assert r["codice_conto"] == "47.91.01"
    assert r["mese"] == 2
    assert r["importo"] == pytest.approx(150.5)
    assert r["categoria_ce"] == "Ricavi"
    assert r["tipo_costo"] == "IP"
    assert r["business_unit_id"] == "HOTEL"
    assert r["fonte"] == "GASPAROTTO"
