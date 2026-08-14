"""Test hide_past_columns_rotation: nasconde colonne-mese passate, year-aware."""

import openpyxl

from verticals.condges.pf_rotate.excel_model import periodo
from verticals.condges.skeleton_shift import hide_past_columns_rotation

P = lambda m: periodo(2026, m)  # noqa: E731


def _ws_con_anno(wb, nome: str):
    """Foglio dettaglio-style: D1=anno, mesi APRILE..DICEMBRE in D2..L2."""
    ws = wb.create_sheet(nome)
    ws["D1"] = 2026
    mesi = [
        "APRILE", "MAGGIO", "GIUGNO", "LUGLIO", "AGOSTO",
        "SETTEMBRE", "OTTOBRE", "NOVEMBRE", "DICEMBRE",
    ]
    for i, m in enumerate(mesi):
        ws.cell(2, 4 + i, m)
    return ws


def test_nasconde_periodi_passati_mostra_aperti():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = _ws_con_anno(wb, "Utenze")

    hide_past_columns_rotation(wb, P(6))  # GIUGNO 2026 = primo periodo aperto

    # APRILE(D=4)/MAGGIO(E=5) < GIUGNO → nascoste
    assert ws.column_dimensions["D"].hidden is True
    assert ws.column_dimensions["E"].hidden is True
    # GIUGNO(F=6)..DICEMBRE(L=12) >= GIUGNO → visibili
    for col_letter in ["F", "G", "H", "I", "J", "K", "L"]:
        assert ws.column_dimensions[col_letter].hidden is not True


def test_foglio_senza_anno_non_solleva_e_resta_intoccato():
    """Un foglio con colonne-mese ma SENZA anno dichiarato (es. i residui
    "DA MAPPARE1"/"ESCLUSI1" osservati sui file reali) non è databile: deve
    essere un no-op silenzioso, non un crash."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("DA MAPPARE1")  # titolo NON nella skip-list esatta
    mesi = [
        "GENNAIO", "FEBBRAIO", "MARZO", "APRILE", "MAGGIO", "GIUGNO",
        "LUGLIO", "AGOSTO", "SETTEMBRE", "OTTOBRE", "NOVEMBRE", "DICEMBRE",
    ]
    for i, m in enumerate(mesi):
        ws.cell(2, 3 + i, m)
    # nessun anno numerico in riga 1

    hide_past_columns_rotation(wb, P(6))  # non deve sollevare ValueError

    for c in range(3, 15):
        letter = openpyxl.utils.get_column_letter(c)
        assert ws.column_dimensions[letter].hidden is not True
