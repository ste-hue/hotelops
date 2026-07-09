import io

import openpyxl

from workspace.miners.dossier_index import build_index_xlsx


def _items():
    return [
        {
            "key": "drive:a",
            "source": "drive",
            "name": "verbale.pdf",
            "category": "01_Societario",
            "confidence": 0.9,
            "owner": "x@p.it",
            "holders": ["x@p.it"],
            "link": "http://l",
            "modified": "2026-01-01",
            "mime_type": "application/pdf",
            "size": 1,
        },
        {
            "key": "hash:ff",
            "source": "gmail",
            "name": "f24.pdf",
            "category": "02_Fiscale",
            "confidence": 0.8,
            "owner": "y@p.it",
            "holders": ["y@p.it", "z@p.it"],
            "link": "http://m",
            "modified": "2026-02-02",
            "mime_type": "application/pdf",
            "size": 2,
        },
    ]


def test_build_index_has_rows_and_gaplist():
    data = build_index_xlsx(_items(), inaccessible=["am@panoramagroup.it"])
    wb = openpyxl.load_workbook(io.BytesIO(data))
    idx = wb["Indice"]
    rows = list(idx.iter_rows(values_only=True))
    assert rows[0][0] == "Categoria"
    assert len(rows) == 3  # header + 2 item
    gap = wb["GapList"]
    gap_text = "\n".join(str(c) for row in gap.iter_rows(values_only=True) for c in row)
    assert "am@panoramagroup.it" in gap_text
    assert "statuto" in gap_text.lower()
