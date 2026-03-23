#!/usr/bin/env python3
"""
Scheda contabile Esolver → BigQuery f_saldi_banca_snapshot.

Parses Esolver "Scheda contabile" (account ledger) exports for bank accounts.
Extracts the progressive saldo (running balance) and writes end-of-day
snapshots to f_saldi_banca_snapshot.

Supports TWO export formats:
  1. XLSX — from Esolver GUI export (may have DD/MM date swap bug)
  2. CSV (semicolon-delimited) — from Esolver direct export (dates always correct)

Source (Drive):
    hotelops_datahub/registro_banca_esolver/ORTI/
    hotelops_datahub/registro_banca_esolver/INTUR/

Filename patterns (auto-infer societa + banca):
    XLSX: MPSORTI_schedacontabile.xlsx, ORTI_scheda_contabile_MPS_20260318.xlsx
    CSV:  20260101_Conto_ORTI_BancaMPS_Cc2_Saldo35588-05_Esercizio2026.csv
          20260101_Conto_INTUR_BancaSella_Cc1_Saldo32225-02_Esercizio2026.csv

Columns (same in both CSV and XLSX):
    Data registrazione; Rif. registrazione; Causale contabile;
    Dare in UdC; Avere in UdC; Saldo in UdC; Documento;
    Val.; Dare in valuta; Avere in valuta; Riferimenti IVA; Centro imputazione

Usage:
    python -m ingest.amministrativa.ingest_scheda_contabile \\
        --file /path/to/20260101_Conto_ORTI_BancaMPS_Cc2_Saldo35588-05_Esercizio2026.csv \\
        --societa ORTI --banca MPS

    python -m ingest.amministrativa.ingest_scheda_contabile \\
        --dir /path/to/registro_banca_esolver/ORTI/ --societa ORTI

    python -m ingest.amministrativa.ingest_scheda_contabile \\
        --file export.xlsx --societa ORTI --banca MPS --dry-run

Output BQ: f_saldi_banca_snapshot (DELETE-INSERT per societa+banca+data_snapshot)
"""

from __future__ import annotations

import argparse
import csv
import logging
import re
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

try:
    import openpyxl
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

try:
    from google.cloud import bigquery
    HAS_BQ = True
except ImportError:
    HAS_BQ = False

# ── Config ────────────────────────────────────────────────────────────────────

BQ_PROJECT = "hotelops-suite"
BQ_TABLE = f"{BQ_PROJECT}.hotelops.f_saldi_banca_snapshot"

log = logging.getLogger("ingest.scheda_contabile")
logging.basicConfig(level=logging.INFO, format="%(asctime)s  %(levelname)-5s  %(message)s")

# Known bank name patterns in filenames
BANCA_PATTERNS = {
    "MPS_KROSS": ["KROSS", "MPS_KROSS"],  # Check before MPS
    "MPS": ["MPS", "MONTEPASCHI", "MONTE_PASCHI", "BANCAMPS"],
    "SELLA": ["SELLA", "BANCASELLA"],
    "INTESA": ["INTESA", "SANPAOLO", "INTESASANPAOLO"],
    "BCP": ["BCP", "CREDITO_POPOLARE", "CREDITOPOPOLARE"],
}

# Esolver conto number → banca_id mapping (from ABI config)
ESOLVER_CC_MAP = {
    # ORTI
    ("ORTI", "3"): "MPS",
    ("ORTI", "2"): "MPS_KROSS",
    ("ORTI", "1"): "INTESA",
    # INTUR
    ("INTUR", "2"): "MPS",
    ("INTUR", "3"): "INTESA",
    ("INTUR", "1"): "SELLA",
    ("INTUR", "4"): "BCP",
}


# ── Parsing helpers ───────────────────────────────────────────────────────────

def parse_italian_number(val) -> float | None:
    """Parse Italian number format: '1.234,56' or '1234,56' → 1234.56."""
    if val is None:
        return None
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).strip()
    if not s or s == "-":
        return None
    # Remove dots (thousands), replace comma with dot (decimal)
    s = s.replace(".", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def parse_date(val) -> date | None:
    """Parse date from either datetime object or DD/MM/YYYY string."""
    if val is None:
        return None
    if isinstance(val, datetime):
        return val.date()
    if isinstance(val, date) and not isinstance(val, datetime):
        return val
    s = str(val).strip()
    # Try DD/MM/YYYY
    m = re.match(r"(\d{1,2})/(\d{1,2})/(\d{4})", s)
    if m:
        return date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
    # Try YYYY-MM-DD
    m = re.match(r"(\d{4})-(\d{2})-(\d{2})", s)
    if m:
        return date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    return None


# ── DD/MM swap fix (XLSX only) ────────────────────────────────────────────────

def _fix_excel_date_swap(rows_raw: list[dict]) -> list[dict]:
    """
    Fix DD/MM ↔ MM/DD swap caused by Excel date auto-detection.

    Esolver exports dates as DD/MM/YYYY. When DD ≤ 12, Excel sometimes
    interprets them as MM/DD/YYYY, creating datetime objects with swapped
    day/month. Dates where DD > 12 remain as strings (Excel can't swap them).

    Strategy: collect "safe" dates (from string-parsed dates where DD > 12),
    determine the expected date range, then fix datetime-parsed dates that
    fall outside this range by swapping day ↔ month.
    """
    safe_dates = []
    datetime_indices = []

    for i, row in enumerate(rows_raw):
        d = row["data_registrazione"]
        if row.get("_from_string", False):
            safe_dates.append(d)
        else:
            datetime_indices.append(i)

    if not safe_dates or not datetime_indices:
        return rows_raw

    min_safe = min(safe_dates)
    max_safe = max(safe_dates)
    range_start = min_safe - timedelta(days=15)
    range_end = max_safe + timedelta(days=5)

    fixed_count = 0
    for i in datetime_indices:
        d = rows_raw[i]["data_registrazione"]
        if range_start <= d <= range_end:
            continue
        try:
            swapped = date(d.year, d.day, d.month)
            if range_start <= swapped <= range_end:
                rows_raw[i]["data_registrazione"] = swapped
                fixed_count += 1
        except ValueError:
            pass

    if fixed_count:
        log.info(f"Fixed {fixed_count} DD/MM-swapped dates (Excel auto-detection artifact)")

    return rows_raw


# ── Parsers ───────────────────────────────────────────────────────────────────

def parse_scheda_csv(filepath: Path) -> list[dict]:
    """
    Parse semicolon-delimited CSV from Esolver.

    CSV dates are always DD/MM/YYYY strings — no Excel date swap issue.
    """
    rows = []
    # Try common encodings
    for enc in ["utf-8-sig", "latin-1", "cp1252"]:
        try:
            with open(filepath, "r", encoding=enc) as f:
                content = f.read()
            break
        except UnicodeDecodeError:
            continue
    else:
        log.error(f"Cannot decode {filepath.name}")
        return []

    reader = csv.reader(content.splitlines(), delimiter=";")
    for line_no, fields in enumerate(reader, 1):
        if len(fields) < 6:
            continue

        data = parse_date(fields[0])
        if data is None:
            continue  # Header row or empty

        saldo = parse_italian_number(fields[5])
        if saldo is None:
            continue

        dare = parse_italian_number(fields[3]) or 0.0
        avere = parse_italian_number(fields[4]) or 0.0

        rows.append({
            "data_registrazione": data,
            "rif_registrazione": fields[1].strip() if len(fields) > 1 else "",
            "causale": fields[2].strip() if len(fields) > 2 else "",
            "dare": dare,
            "avere": avere,
            "saldo": saldo,
            "documento": fields[6].strip() if len(fields) > 6 else "",
        })

    return rows


def parse_scheda_xlsx(filepath: Path) -> list[dict]:
    """
    Parse XLSX from Esolver GUI export.

    Applies DD/MM swap fix for dates mangled by Excel.
    """
    if not HAS_OPENPYXL:
        log.error("openpyxl required for XLSX: pip install openpyxl")
        return []

    wb = openpyxl.load_workbook(filepath, data_only=True)
    ws = wb.active
    rows = []

    for r in range(1, ws.max_row + 1):
        data_raw = ws.cell(r, 1).value
        from_string = isinstance(data_raw, str)
        data = parse_date(data_raw)
        if data is None:
            continue

        saldo = parse_italian_number(ws.cell(r, 6).value)
        if saldo is None:
            continue

        dare = parse_italian_number(ws.cell(r, 4).value) or 0.0
        avere = parse_italian_number(ws.cell(r, 5).value) or 0.0

        rows.append({
            "data_registrazione": data,
            "rif_registrazione": str(ws.cell(r, 2).value or "").strip(),
            "causale": str(ws.cell(r, 3).value or "").strip(),
            "dare": dare,
            "avere": avere,
            "saldo": saldo,
            "documento": str(ws.cell(r, 7).value or "").strip(),
            "_from_string": from_string,
        })

    # Fix DD/MM swaps from Excel auto-detection
    rows = _fix_excel_date_swap(rows)

    for row in rows:
        row.pop("_from_string", None)

    return rows


def parse_scheda_contabile(filepath: Path) -> list[dict]:
    """Parse scheda contabile file (auto-detect CSV vs XLSX)."""
    suffix = filepath.suffix.lower()
    if suffix in (".csv", ".tsv"):
        return parse_scheda_csv(filepath)
    elif suffix in (".xlsx", ".xls"):
        return parse_scheda_xlsx(filepath)
    else:
        log.error(f"Unsupported format: {suffix}")
        return []


# ── Inference ─────────────────────────────────────────────────────────────────

def extract_daily_saldi(rows: list[dict]) -> dict[date, float]:
    """
    From parsed rows, extract the end-of-day saldo for each date.

    The saldo column is a running balance — we take the LAST row of each day.
    """
    daily_saldo: dict[date, float] = {}
    for row in rows:
        d = row["data_registrazione"]
        daily_saldo[d] = row["saldo"]  # last write wins = end of day
    return daily_saldo


def infer_banca_from_filename(filename: str, societa: str | None = None) -> str | None:
    """
    Try to infer banca ID from filename.

    Handles two naming patterns:
        1. Legacy: MPSORTI_schedacontabile.xlsx
        2. Esolver CSV: 20260101_Conto_ORTI_BancaMPS_Cc2_Saldo35588-05_Esercizio2026.csv
    """
    upper = filename.upper()

    # Try to extract Cc number from Esolver CSV filename pattern
    # Pattern: _Cc{N}_ where N is the Esolver conto number
    cc_match = re.search(r"_CC(\d+)_", upper)
    if cc_match and societa:
        cc_num = cc_match.group(1)
        key = (societa, cc_num)
        if key in ESOLVER_CC_MAP:
            return ESOLVER_CC_MAP[key]

    # Fall back to name matching
    # Check KROSS first (before MPS, since MPS patterns would also match)
    for banca_id in ["MPS_KROSS", "MPS", "SELLA", "INTESA", "BCP"]:
        for pat in BANCA_PATTERNS[banca_id]:
            if pat in upper:
                return banca_id
    return None


def infer_societa_from_filename(filename: str, banca: str | None = None) -> str | None:
    """
    Infer societa from filename.

    Tries in order:
      1. Explicit ORTI/INTUR in filename
      2. Reverse lookup: Cc# + banca → societa (when combo is unambiguous)
      3. Reverse lookup: Cc# alone (if unique across both società)
    """
    upper = filename.upper()
    if "INTUR" in upper:
        return "INTUR"
    if "ORTI" in upper:
        return "ORTI"

    # Reverse inference from Cc# (same logic as classify.py)
    cc_match = re.search(r"_CC(\d+)_", upper)
    if cc_match:
        cc_num = cc_match.group(1)
        # If banca is known, check (societa, cc) → banca matches
        if banca:
            for (soc, cc), bnk in ESOLVER_CC_MAP.items():
                if cc == cc_num and bnk == banca:
                    return soc
        # Otherwise, check if cc# is unique to one societa
        candidates = {soc for (soc, cc) in ESOLVER_CC_MAP if cc == cc_num}
        if len(candidates) == 1:
            return candidates.pop()

    return None


# ── BigQuery write ────────────────────────────────────────────────────────────

def write_saldi_to_bq(
    saldi: dict[date, float],
    societa_id: str,
    banca_id: str,
    dry_run: bool = False,
) -> int:
    """Write end-of-day saldi to f_saldi_banca_snapshot."""
    if not saldi:
        log.warning("No saldi to write")
        return 0

    if not HAS_BQ:
        log.error("google-cloud-bigquery required")
        return 0

    rows_to_write = []
    for d, saldo in sorted(saldi.items()):
        rows_to_write.append({
            "societa_id": societa_id,
            "banca_id": banca_id,
            "data_snapshot": d.isoformat(),
            "saldo_finale": saldo,
            "file_sorgente": "SCHEDA_CONTABILE",
        })

    if dry_run:
        log.info(f"DRY RUN: {len(rows_to_write)} saldi for {societa_id}/{banca_id}")
        for r in rows_to_write:
            log.info(f"  {r['data_snapshot']}: €{r['saldo_finale']:,.2f}")
        return len(rows_to_write)

    client = bigquery.Client(project=BQ_PROJECT)

    # DELETE existing snapshots for this societa+banca (from SCHEDA_CONTABILE source)
    dates = [r["data_snapshot"] for r in rows_to_write]
    date_list = ", ".join(f"'{d}'" for d in dates)
    delete_sql = f"""
    DELETE FROM `{BQ_TABLE}`
    WHERE societa_id = '{societa_id}'
      AND banca_id = '{banca_id}'
      AND file_sorgente = 'SCHEDA_CONTABILE'
      AND CAST(data_snapshot AS STRING) IN ({date_list})
    """
    client.query(delete_sql).result()
    log.info(f"Deleted existing SCHEDA_CONTABILE snapshots for {societa_id}/{banca_id}")

    # INSERT new rows
    errors = client.insert_rows_json(BQ_TABLE, rows_to_write)
    if errors:
        log.error(f"BQ insert errors: {errors}")
        return 0

    log.info(f"Wrote {len(rows_to_write)} saldi to {BQ_TABLE}")
    return len(rows_to_write)


# ── Main ──────────────────────────────────────────────────────────────────────

def process_file(
    filepath: Path,
    societa_id: str,
    banca_id: str | None,
    dry_run: bool,
) -> int:
    """Process a single scheda contabile file."""
    log.info(f"Processing: {filepath.name}")

    # Infer banca if not provided
    if not banca_id:
        banca_id = infer_banca_from_filename(filepath.name, societa_id)
        if not banca_id:
            log.error(f"Cannot infer banca from filename: {filepath.name}. Use --banca.")
            return 0
        log.info(f"Inferred banca: {banca_id}")

    rows = parse_scheda_contabile(filepath)
    if not rows:
        log.warning(f"No rows parsed from {filepath.name}")
        return 0

    log.info(f"Parsed {len(rows)} movements, date range: "
             f"{rows[0]['data_registrazione']} → {rows[-1]['data_registrazione']}")

    daily_saldi = extract_daily_saldi(rows)
    log.info(f"Extracted {len(daily_saldi)} end-of-day saldi")

    return write_saldi_to_bq(daily_saldi, societa_id, banca_id, dry_run)


def main():
    parser = argparse.ArgumentParser(
        description="Scheda contabile Esolver → f_saldi_banca_snapshot"
    )
    parser.add_argument("--file", "-f", help="Single file to process")
    parser.add_argument("--dir", "-d", help="Directory to scan (e.g. registro_banca_esolver/ORTI/)")
    parser.add_argument("--societa", default=None, choices=["ORTI", "INTUR"],
                        help="Società ID (auto-inferred from filename if omitted)")
    parser.add_argument("--banca", default=None,
                        choices=["MPS", "MPS_KROSS", "SELLA", "INTESA", "BCP"],
                        help="Banca ID (auto-inferred from filename if omitted)")
    parser.add_argument("--dry-run", action="store_true",
                        help="Parse and show results, no BQ writes")
    args = parser.parse_args()

    if not args.file and not args.dir:
        parser.error("Specify --file or --dir")

    total = 0

    if args.file:
        filepath = Path(args.file)
        if not filepath.exists():
            log.error(f"File not found: {filepath}")
            sys.exit(1)
        # Infer banca first (needed for reverse societa lookup from Cc#)
        banca = args.banca or infer_banca_from_filename(filepath.name, args.societa)
        societa = args.societa or infer_societa_from_filename(filepath.name, banca)
        if not societa:
            log.error(f"Cannot infer societa from {filepath.name}. Use --societa.")
            sys.exit(1)
        total = process_file(filepath, societa, banca, args.dry_run)

    elif args.dir:
        dirpath = Path(args.dir)
        if not dirpath.exists():
            log.error(f"Directory not found: {dirpath}")
            sys.exit(1)

        # Find all supported files
        files = (
            sorted(dirpath.glob("*.xlsx"))
            + sorted(dirpath.glob("*.xls"))
            + sorted(dirpath.glob("*.csv"))
            + sorted(dirpath.glob("*.CSV"))
        )
        files = [f for f in files if not f.name.startswith("~$")]
        # Deduplicate (case-insensitive glob on some OS)
        seen = set()
        unique_files = []
        for f in files:
            if f.name.lower() not in seen:
                seen.add(f.name.lower())
                unique_files.append(f)
        files = unique_files

        if not files:
            log.warning(f"No scheda contabile files found in {dirpath}")
            sys.exit(0)

        log.info(f"Found {len(files)} files in {dirpath}")
        for filepath in files:
            banca = args.banca or infer_banca_from_filename(filepath.name, args.societa)
            societa = args.societa or infer_societa_from_filename(filepath.name, banca)
            if not societa:
                # Try inferring from parent directory name
                societa = dirpath.name if dirpath.name in ("ORTI", "INTUR") else None
            if not societa:
                log.warning(f"Skipping {filepath.name}: cannot infer societa")
                continue
            n = process_file(filepath, societa, banca, args.dry_run)
            total += n

    log.info(f"Total saldi written: {total}")


if __name__ == "__main__":
    main()
