#!/usr/bin/env python3
"""
Piano dei conti 2026 loader — from "Costi Ricavi 2025-2026 Budget.xlsx".

Source: "piano dei conti nuovo" sheet (1471 rows, 206 leaf accounts)
Target: d_piano_conti in BigQuery (WRITE_TRUNCATE — dimension table)

This replaces the old pianodeiconti.xlsx from reconciliation_dino.
The 2026 PDC has structural changes: new code families (63.05.xx, 65.90.xx)
replace old 61.xx codes.

Sheet structure (row 1 = header):
  Col A: TIPO (CE/SP)
  Col B: SEZIONE (COSTI/RICAVI/ATTIVO/PASSIVO)
  Col C: Codici standard (not used — same as codice)
  Col D: Level 1 segment (2 digits)
  Col E: Level 2 segment (2 digits)
  Col F: Level 3 segment (2 digits)
  Col G: Level 4 segment (2 digits, optional)
  Col H: Conto (dotted format, formula — may be None in read_only)
  Col I: Descrizione
  Col J+: Importo / notes (not needed)

Usage:
    python -m ingest.amministrativa.ingest_piano_conti_nuovo --dry-run
    python -m ingest.amministrativa.ingest_piano_conti_nuovo \\
        --file "/path/to/Costi Ricavi 2025-2026 Budget.xlsx"
"""

from __future__ import annotations

import argparse
import csv
import logging
import sys
from pathlib import Path

try:
    import openpyxl
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

try:
    from google.cloud import bigquery
    HAS_BQ = True
except ImportError:
    HAS_BQ = False

BQ_PROJECT = "hotelops-suite"
BQ_TABLE = f"{BQ_PROJECT}.hotelops.d_piano_conti"

DEFAULT_FILE = Path(
    "/Users/stefanodellapietra/Desktop/WORK/artifacts/"
    "Costi Ricavi 2025-2026 Budget.xlsx"
)

SHEET_NAMES = ["piano dei conti nuovo", "piano dei conti nuovo "]  # handle trailing space

# Account prefix → business_unit_id (revenue accounts only)
BU_RICAVI_MAP = {
    "47.91": "HOTEL",
    "47.92": "RESIDENCE",
    "47.93": "CVM",
    "47.94": "LIDO",
    "47.95": "HQ",
}


def setup_logger() -> logging.Logger:
    log = logging.getLogger("ingest_piano_conti_nuovo")
    log.setLevel(logging.DEBUG)
    if not log.handlers:
        h = logging.StreamHandler(sys.stdout)
        h.setFormatter(logging.Formatter("%(levelname)s  %(message)s"))
        log.addHandler(h)
    return log


def infer_business_unit(codice: str) -> str | None:
    for prefix, bu in BU_RICAVI_MAP.items():
        if codice.startswith(prefix):
            return bu
    return None


def _build_codice(d: str, e: str, f: str, g: str | None) -> str:
    """Build dotted codice from level segments: DD.EE.FF[.GG]"""
    parts = [d, e, f]
    if g:
        parts.append(g)
    return ".".join(parts)


def parse_piano_conti_nuovo(filepath: Path, logger: logging.Logger) -> list[dict]:
    """Parse 'piano dei conti nuovo' sheet → list of dicts for d_piano_conti."""
    if not HAS_OPENPYXL:
        logger.error("openpyxl non installato: pip install openpyxl")
        sys.exit(1)

    wb = openpyxl.load_workbook(str(filepath), read_only=True, data_only=True)

    sheet_name = None
    for candidate in SHEET_NAMES:
        if candidate in wb.sheetnames:
            sheet_name = candidate
            break
    if sheet_name is None:
        logger.error(f"Foglio 'piano dei conti nuovo' non trovato in {filepath.name}")
        logger.info(f"  Fogli disponibili: {wb.sheetnames}")
        wb.close()
        return []

    ws = wb[sheet_name]
    rows_raw = list(ws.iter_rows(min_row=2, values_only=True))  # Skip header
    wb.close()

    records = []
    skipped = 0
    # Track parent info for hierarchy
    current_tipo = None
    current_sezione = None

    for row in rows_raw:
        if not row or len(row) < 9:
            skipped += 1
            continue

        tipo = str(row[0]).strip() if row[0] else None
        sezione = str(row[1]).strip() if row[1] else None

        # Carry forward tipo/sezione from parent rows
        if tipo and tipo in ("CE", "SP"):
            current_tipo = tipo
        if sezione:
            current_sezione = sezione

        # Level segments
        seg_d = str(row[3]).strip() if row[3] else None
        seg_e = str(row[4]).strip() if row[4] else None
        seg_f = str(row[5]).strip() if row[5] else None
        seg_g = str(row[6]).strip() if row[6] else None

        # Need at least D.E for a valid account
        if not seg_d or not seg_e:
            skipped += 1
            continue

        # Pad segments to 2 digits
        seg_d = seg_d.zfill(2) if seg_d else None
        seg_e = seg_e.zfill(2) if seg_e else None
        if seg_f:
            seg_f = seg_f.zfill(2)
        else:
            skipped += 1
            continue  # Need at least 3 levels for a meaningful account

        if seg_g:
            seg_g = seg_g.zfill(2)

        # Col H might be formula/None in read_only mode; build from segments
        codice = _build_codice(seg_d, seg_e, seg_f, seg_g)

        # Descrizione from col I
        descrizione = str(row[8]).strip() if row[8] else ""
        if not descrizione or descrizione == "Descrizione":
            skipped += 1
            continue

        records.append({
            "codice_conto": codice,
            "descrizione": descrizione,
            "tipo_conto": current_tipo or "",
            "sezione": current_sezione or "",
            "partitario": "",  # Not in this source
            "business_unit_id": infer_business_unit(codice),
        })

    logger.info(f"  Piano conti nuovo: {len(records)} conti parsati (skip {skipped})")

    # Deduplicate by codice_conto (keep first occurrence)
    seen = set()
    deduped = []
    for r in records:
        if r["codice_conto"] not in seen:
            seen.add(r["codice_conto"])
            deduped.append(r)
    if len(deduped) < len(records):
        logger.info(f"  Dedup: {len(records)} → {len(deduped)} conti unici")

    return deduped


def quality_summary(rows: list[dict], logger: logging.Logger) -> None:
    by_tipo: dict[str, int] = {}
    by_sezione: dict[str, int] = {}
    for r in rows:
        t = r["tipo_conto"] or "?"
        s = r["sezione"] or "?"
        by_tipo[t] = by_tipo.get(t, 0) + 1
        by_sezione[s] = by_sezione.get(s, 0) + 1

    logger.info("═" * 50)
    logger.info("  PIANO CONTI NUOVO — SUMMARY")
    logger.info("═" * 50)
    for t, n in sorted(by_tipo.items()):
        logger.info(f"  Tipo {t}: {n:>4d} conti")
    for s, n in sorted(by_sezione.items()):
        logger.info(f"  Sezione {s}: {n:>4d} conti")
    bu_conti = [r for r in rows if r["business_unit_id"]]
    logger.info(f"  Conti con BU (47.9x): {len(bu_conti)}")
    logger.info("═" * 50)


FIELDS = ["codice_conto", "descrizione", "tipo_conto", "sezione",
           "partitario", "business_unit_id"]


def dump_csv(rows: list[dict], path: Path, logger: logging.Logger) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=FIELDS, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)
    logger.info(f"  CSV: {path}  ({len(rows)} righe)")


BQ_SCHEMA = [
    bigquery.SchemaField("codice_conto", "STRING", mode="REQUIRED"),
    bigquery.SchemaField("descrizione", "STRING"),
    bigquery.SchemaField("tipo_conto", "STRING"),
    bigquery.SchemaField("sezione", "STRING"),
    bigquery.SchemaField("partitario", "STRING"),
    bigquery.SchemaField("business_unit_id", "STRING"),
] if HAS_BQ else []


def load_to_bq(rows: list[dict], logger: logging.Logger) -> None:
    if not rows:
        logger.warning("Nessuna riga da caricare")
        return

    bq_client = bigquery.Client(project=BQ_PROJECT)
    job = bq_client.load_table_from_json(
        rows, BQ_TABLE,
        job_config=bigquery.LoadJobConfig(
            schema=BQ_SCHEMA,
            write_disposition=bigquery.WriteDisposition.WRITE_TRUNCATE,
        ),
    )
    job.result()
    if job.errors:
        logger.error(f"BQ errors: {job.errors}")
        sys.exit(1)
    logger.info(f"  {BQ_TABLE}: {len(rows)} righe caricate (WRITE_TRUNCATE)")


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Piano dei conti nuovo 2026 → d_piano_conti"
    )
    parser.add_argument("--file", default=str(DEFAULT_FILE),
                        help="Path al file XLSX con foglio 'piano dei conti nuovo'")
    parser.add_argument("--dry-run", action="store_true",
                        help="Parse + CSV, no BQ write")
    parser.add_argument("--output-dir", default="output",
                        help="Directory per CSV output in dry-run")
    args = parser.parse_args()

    logger = setup_logger()
    filepath = Path(args.file)
    if not filepath.exists():
        logger.error(f"File non trovato: {filepath}")
        sys.exit(1)

    logger.info(f"File: {filepath.name}")
    rows = parse_piano_conti_nuovo(filepath, logger)
    if not rows:
        logger.error("Nessun conto estratto")
        sys.exit(1)

    quality_summary(rows, logger)

    if args.dry_run:
        dump_csv(rows, Path(args.output_dir) / "d_piano_conti_nuovo.csv", logger)
        logger.info("DRY RUN — nessuna scrittura su BQ.")
        return

    if not HAS_BQ:
        logger.error("google-cloud-bigquery non installato")
        sys.exit(1)

    load_to_bq(rows, logger)
    logger.info("✓ DONE")


if __name__ == "__main__":
    main()
