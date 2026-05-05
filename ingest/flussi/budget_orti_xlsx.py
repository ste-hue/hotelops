"""Budget_ORTI_*.xlsx — monthly grid sheets → f_budget_mensile row dicts.

Workbooks like ``Budget_ORTI_2026.xlsx`` expose one sheet per macro-block
(Ricavi, Fissi, Variabili, Personale, Finanziari) with columns
``codice_conto``, ``descrizione``, BU, then Jan–Dec (Italian headers).

This is **not** the Gasparotto Master (no ``Budget`` / ``Conto Economico`` pair).
Parsed rows use the same ``fonte`` as ``ingest_gasparotto`` (default GASPAROTTO)
so ``v_budget_canonical`` precedence rules stay unchanged.
"""

from __future__ import annotations

import logging
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

try:
    import openpyxl
except ImportError:
    openpyxl = None  # type: ignore

FONTE_DEFAULT = "GASPAROTTO"

# Sheet logical name (upper) → (categoria_ce, tipo_costo)
SHEET_META: dict[str, tuple[str, str]] = {
    "RICAVI": ("Ricavi", "IP"),
    "FISSI": ("Costi Produttivi", "F"),
    "VARIABILI": ("Costi Produttivi", "V"),
    "PERSONALE": ("Costo del Personale", "P"),
    "FINANZIARI": ("Oneri Finanziari", "X"),
}

_MONTH_SYNONYMS: list[tuple[int, tuple[str, ...]]] = [
    (1, ("GEN", "GENNAIO", "JAN", "JANUARY")),
    (2, ("FEB", "FEBBRAIO", "FEBRUARY")),
    (3, ("MAR", "MARZO", "MAR", "MARCH")),
    (4, ("APR", "APRILE", "APRIL")),
    (5, ("MAG", "MAGGIO", "MAY")),
    (6, ("GIU", "GIUGNO", "JUN", "JUNE")),
    (7, ("LUG", "LUGLIO", "JUL", "JULY")),
    (8, ("AGO", "AGOSTO", "AUG", "AUGUST")),
    (9, ("SET", "SETT", "SETTEMBRE", "SEP", "SEPT", "SEPTEMBER")),
    (10, ("OTT", "OTTOBRE", "OCT", "OCTOBER")),
    (11, ("NOV", "NOVEMBRE", "NOVEMBER")),
    (12, ("DIC", "DICEMBRE", "DEC", "DECEMBER")),
]

_HEADER_TO_MESE: dict[str, int] = {}
for mese, labels in _MONTH_SYNONYMS:
    for lab in labels:
        _HEADER_TO_MESE[lab.upper()] = mese


def sniff_budget_orti_monthly_format(sheetnames: list[str]) -> bool:
    """True when workbook looks like ORTI monthly export (not Gasparotto Master)."""
    upper = {n.strip().upper() for n in sheetnames}
    if "BUDGET" in upper:
        return False
    return "RICAVI" in upper and "FISSI" in upper


def _norm_header_month(cell_val: Any) -> int | None:
    if cell_val is None:
        return None
    s = str(cell_val).strip().upper()
    if not s:
        return None
    s = re.sub(r"^[\d.]+\s*", "", s)
    s = s.replace("À", "A").replace("È", "E")
    return _HEADER_TO_MESE.get(s)


def _normalize_codice_conto(raw: Any) -> str | None:
    if raw is None:
        return None
    if isinstance(raw, float) and raw == int(raw):
        raw = int(raw)
    if isinstance(raw, int):
        s = str(raw)
    else:
        s = str(raw).strip()
    if not s:
        return None
    if "." in s:
        return s
    digits = "".join(c for c in s if c.isdigit())
    if len(digits) == 6:
        return f"{digits[:2]}.{digits[2:4]}.{digits[4:]}"
    if len(digits) >= 4:
        return s
    return None


def _business_unit(raw: Any) -> str | None:
    if raw is None:
        return None
    s = str(raw).strip().upper()
    if not s or "ANNUO" in s:
        return None
    for bid in ("HOTEL", "RESIDENCE", "CVM", "LIDO", "HQ"):
        if s == bid or bid in s.split():
            return bid
    if "SEDE" in s or "AMM" in s:
        return "HQ"
    return None


def _float_cell(val: Any) -> float:
    if val is None:
        return 0.0
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).strip().replace(".", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return 0.0


def _find_header_row(ws: Any, max_scan: int = 5) -> int | None:
    for r in range(1, min(max_scan, ws.max_row) + 1):
        a = ws.cell(r, 1).value
        if a is None:
            continue
        low = str(a).strip().lower()
        if "codice" in low and "conto" in low:
            return r
    return 1 if ws.max_row else None


def _resolve_sheet(wb: Any, logical_upper: str) -> Any | None:
    for name in wb.sheetnames:
        if name.strip().upper() == logical_upper:
            return wb[name]
    return None


def parse_budget_orti_workbook(
    filepath: Path,
    societa_id: str,
    anno: int,
    logger: logging.Logger,
    *,
    fonte: str = FONTE_DEFAULT,
) -> list[dict]:
    """Parse ORTI Budget_* XLSX with Ricavi/Fissi/… monthly columns."""
    if openpyxl is None:
        logger.error("openpyxl non installato")
        return []

    wb = openpyxl.load_workbook(str(filepath), data_only=True)
    try:
        now = datetime.now(timezone.utc).isoformat()
        records: list[dict] = []

        for sheet_key, (categoria_ce, tipo_costo) in SHEET_META.items():
            ws = _resolve_sheet(wb, sheet_key)
            if ws is None:
                logger.warning(f"  Foglio «{sheet_key.title()}» assente — skip")
                continue

            hr = _find_header_row(ws)
            if hr is None:
                logger.warning(f"  Header non trovato in {sheet_key} — skip")
                continue

            month_cols: dict[int, int] = {}
            for c in range(1, ws.max_column + 1):
                m = _norm_header_month(ws.cell(hr, c).value)
                if m is not None:
                    month_cols[m] = c

            if len(month_cols) < 6:
                logger.warning(
                    f"  {sheet_key}: colonne mese insufficienti "
                    f"({len(month_cols)}), skip foglio"
                )
                continue

            for r in range(hr + 1, ws.max_row + 1):
                cod = _normalize_codice_conto(ws.cell(r, 1).value)
                desc_cell = ws.cell(r, 2).value
                desc = str(desc_cell).strip() if desc_cell else ""
                du = str(desc).upper()
                if not cod:
                    continue
                if not desc or "TOTALE" in du or du.startswith("\\"):
                    continue

                bu = _business_unit(ws.cell(r, 3).value)

                for mese, col in month_cols.items():
                    imp = _float_cell(ws.cell(r, col).value)
                    if imp == 0.0:
                        continue
                    records.append(
                        {
                            "societa_id": societa_id,
                            "anno": anno,
                            "mese": mese,
                            "codice_conto": cod,
                            "descrizione": desc,
                            "tipo_costo": tipo_costo,
                            "categoria_ce": categoria_ce,
                            "business_unit_id": bu,
                            "importo": round(imp, 4),
                            "fonte": fonte,
                            "data_caricamento": now,
                        }
                    )

        conti = len({r["codice_conto"] for r in records})
        logger.info(
            f"  Export ORTI: {conti} conti distinti → {len(records)} righe mensili (≠0)"
        )
        return records
    finally:
        wb.close()
