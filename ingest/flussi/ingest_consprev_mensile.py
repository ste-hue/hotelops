#!/usr/bin/env python3
"""Ingest Power BI "Consuntivo + Previsione" mensile → f_consprev_mensile.

Rollup mensile del forecast PMS per (mese × classe × categoria × addebito):
Mese Cons (consumato), Mese Prev+Cons (consumato + portafoglio), Mese BDG
(budget PMS), Mese AP (anno precedente). Ogni export è una FOTOGRAFIA alla
snapshot_date: le fotografie si accumulano, la SNAPSHOT delete è scoped a
(snapshot_date, business_unit_id).

⚠️ La BU non è nel filename né in una colonna: sta nel FOOTER come negazione
("CodiceHotel is not ANGELINARES or HOMEHOLIDAY" → la struttura è il
complemento del set escluso). Content-only by design.

⚠️ La data dello snapshot NON è nel contenuto: si risolve dall'intake del
raw object (--raw-object-id) o si passa esplicita con --snapshot-date.

Parser_module della source POWERBI_CONSPREV_ORTI_SNAPSHOT:
  python -m ingest.flussi.ingest_consprev_mensile --file X --raw-object-id Y

Usage:
    python -m ingest.flussi.ingest_consprev_mensile --file <xlsx> --raw-object-id <id>
    python -m ingest.flussi.ingest_consprev_mensile --file <xlsx> --snapshot-date 2026-07-11 --dry-run
"""

from __future__ import annotations

import argparse
import logging
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

from core.config import F_CONSPREV_MENSILE
from core.schemas import ConsprevMensileRow, make_hash, validate_batch

log = logging.getLogger("ingest.consprev_mensile")

FONTE = "POWERBI_CONSPREV"

CODICI_HOTEL = {"PANORAMAHT": "HOTEL", "ANGELINARES": "RESIDENCE", "HOMEHOLIDAY": "CVM"}

MESI = {
    m: i
    for i, m in enumerate(
        "gennaio febbraio marzo aprile maggio giugno luglio agosto "
        "settembre ottobre novembre dicembre".split(),
        1,
    )
}


def _num(v) -> float:
    if v is None or v == "":
        return 0.0
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def detect_bu_from_footer(rows: list[tuple]) -> str:
    """BU dal footer 'Applied filters' per NEGAZIONE del set CodiceHotel.

    Il footer dice quali strutture sono ESCLUSE ("CodiceHotel is not X or Y"):
    la BU del file è il complemento. Se la negazione non copre esattamente
    len-1 codici, il file è ambiguo → errore, non default silenzioso.
    """
    footer = ""
    for r in reversed(rows[-5:]):
        if r and r[0] and "Applied filters" in str(r[0]):
            footer = str(r[0])
            break
    if not footer:
        raise ValueError("footer 'Applied filters' non trovato — BU non deducibile")
    esclusi = {c for c in CODICI_HOTEL if c in footer}
    inclusi = set(CODICI_HOTEL) - esclusi
    if len(inclusi) != 1:
        raise ValueError(
            f"BU ambigua dal footer (esclusi={sorted(esclusi)}): {footer[:160]!r}"
        )
    return CODICI_HOTEL[inclusi.pop()]


def parse_xlsx(path: Path) -> dict:
    """Estrae {business_unit_id, righe[]} (header-based)."""
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    header = [str(c).strip().lower() if c is not None else "" for c in rows[0]]
    col = {name: i for i, name in enumerate(header)}

    def idx(*names: str) -> int:
        for n in names:
            if n in col:
                return col[n]
        raise ValueError(f"{path.name}: colonna {names} mancante (header: {header})")

    i_anno, i_mese, i_classe = idx("anno"), idx("mese"), idx("classe")
    i_cat, i_add = idx("categoria"), idx("addebito")
    i_cons = idx("mese cons")
    i_pc = idx("mese prev + cons", "mese prev+cons")
    i_bdg = idx("mese bdg")
    i_ap = idx("mese ap")

    bu = detect_bu_from_footer(rows)

    righe = []
    for r in rows[1:]:
        if not r or len(r) <= i_ap:
            continue
        mese = str(r[i_mese]).strip().lower() if r[i_mese] else ""
        if mese not in MESI:
            continue  # Total / footer / righe non-dato
        righe.append(
            {
                "anno": int(_num(r[i_anno])),
                "mese": MESI[mese],
                "classe": str(r[i_classe]).strip() if r[i_classe] else "",
                "categoria": str(r[i_cat]).strip() if r[i_cat] else None,
                "addebito": str(r[i_add]).strip() if r[i_add] else None,
                "mese_cons": _num(r[i_cons]),
                "mese_prev_cons": _num(r[i_pc]),
                "mese_bdg": _num(r[i_bdg]),
                "mese_ap": _num(r[i_ap]),
            }
        )
    return {"business_unit_id": bu, "righe": righe}


def build_rows(
    parsed: dict, snapshot_date: str, raw_object_id: str | None
) -> list[dict]:
    now = datetime.now(timezone.utc).replace(tzinfo=None).isoformat()
    bu = parsed["business_unit_id"]
    out = []
    for r in parsed["righe"]:
        out.append(
            {
                "snapshot_date": snapshot_date,
                "societa_id": "ORTI",
                "business_unit_id": bu,
                **r,
                "fonte": FONTE,
                "hash_riga": make_hash(
                    snapshot_date,
                    bu,
                    r["anno"],
                    r["mese"],
                    r["classe"],
                    r["categoria"] or "",
                    r["addebito"] or "",
                ),
                "raw_object_id": raw_object_id,
                "data_caricamento": now,
            }
        )
    return out


def _snapshot_date_from_raw(raw_object_id: str) -> str:
    from google.cloud import bigquery

    from core.bq.client import get_client
    from core.config import PROJECT

    client = get_client()
    sql = (
        f"SELECT DATE(intake_at) d FROM `{PROJECT}.hotelops.f_raw_objects` "
        "WHERE raw_object_id = @rid"
    )
    job = client.query(
        sql,
        job_config=bigquery.QueryJobConfig(
            query_parameters=[
                bigquery.ScalarQueryParameter("rid", "STRING", raw_object_id)
            ]
        ),
    )
    res = list(job.result())
    if not res:
        raise ValueError(f"raw_object_id {raw_object_id} non trovato in f_raw_objects")
    return res[0]["d"].isoformat()


def ingest_file(
    path: Path,
    raw_object_id: str | None = None,
    snapshot_date: str | None = None,
    dry_run: bool = False,
) -> int:
    if snapshot_date is None:
        if raw_object_id is None:
            raise ValueError("serve --raw-object-id oppure --snapshot-date")
        snapshot_date = _snapshot_date_from_raw(raw_object_id)

    parsed = parse_xlsx(path)
    rows = build_rows(parsed, snapshot_date, raw_object_id)
    validate_batch(rows, ConsprevMensileRow, context=f"consprev_mensile {path.name}")
    if dry_run:
        log.info(
            "[DRY-RUN] %s → %s snapshot %s: %d righe",
            path.name,
            parsed["business_unit_id"],
            snapshot_date,
            len(rows),
        )
        return len(rows)

    from core.bq.write import bq_write_validated

    pydantic_rows = [ConsprevMensileRow(**r) for r in rows]
    bq_write_validated(
        F_CONSPREV_MENSILE,
        pydantic_rows,
        mode="snapshot",
        natural_key=["snapshot_date", "business_unit_id"],
    )
    log.info(
        "OK %s → %s snapshot %s: %d righe",
        path.name,
        parsed["business_unit_id"],
        snapshot_date,
        len(rows),
    )
    return len(rows)


def main() -> None:
    ap = argparse.ArgumentParser(
        description="Ingest Consuntivo+Previsione mensile → f_consprev_mensile"
    )
    ap.add_argument("--file", required=True, type=Path)
    ap.add_argument(
        "--raw-object-id",
        default=None,
        help="FK a f_raw_objects — da `hotelops promote`",
    )
    ap.add_argument(
        "--snapshot-date", default=None, help="YYYY-MM-DD, alternativa al raw-object-id"
    )
    ap.add_argument(
        "--societa",
        default=None,
        help="Ignorato — sempre ORTI. Accettato da `hotelops promote`.",
    )
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()

    logging.basicConfig(level=logging.INFO, format="%(levelname)s: %(message)s")
    n = ingest_file(
        args.file,
        raw_object_id=args.raw_object_id,
        snapshot_date=args.snapshot_date,
        dry_run=args.dry_run,
    )
    log.info("Totale: %d righe", n)


if __name__ == "__main__":
    main()
