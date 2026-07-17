#!/usr/bin/env python3
"""Ingest HotelCube Power BI Daily Production Report (classe cut) → f_produzione_pms.

One xlsx = one struttura × one or more anni (il filtro Power BI può essere
"Anno is 2026" o "Anno is 2026 or 2025"). Period metadata (CodiceHotel, Anno,
Descrizione) lives in the file's last "Applied filters" cell. Data rows are daily;
columns are revenue classes (01ROOM, 02FB, ...). Unpivot wide→long. L'anno di
ogni riga segue la SUA data — mai il filtro.

Lifecycle: SNAPSHOT, natural_key (business_unit_id, anno). Re-loading a
struttura×anno replaces those rows; un file bi-anno rimpiazza entrambi gli
scope (la DELETE del gate copre i combo distinti presenti nel batch).

Parser_module della source POWERBI_PRODUZIONE_ORTI_SNAPSHOT, invocato da
`hotelops promote` come: python -m ingest.flussi.ingest_produzione_pms --file X --raw-object-id Y

Usage:
    python -m ingest.flussi.ingest_produzione_pms --file <xlsx> --raw-object-id <id>
    python -m ingest.flussi.ingest_produzione_pms --file <xlsx> --dry-run
"""
from __future__ import annotations

import argparse
import logging
import re
from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from core.config import F_PRODUZIONE_PMS
from core.schemas import (
    OPERATIONS_CUTOVER_DATE,
    ProduzioneRow,
    make_hash,
    validate_batch,
)

log = logging.getLogger("ingest.produzione_pms")

HOTEL_TO_BU = {
    "PANORAMAHT": "HOTEL",
    "ANGELINARES": "RESIDENCE",
    "HOMEHOLIDAY": "CVM",
}
CLASSE_RE = re.compile(r"^\d{2}[A-Z]+$")


def parse_applied_filters(text: str) -> tuple[str, set[int]]:
    """Extract (business_unit_id, anni) from the 'Applied filters' cell.

    Il filtro Anno può essere singolo ("Anno is 2026") o multiplo
    ("Anno is 2026 or 2025"). Requires Descrizione is Imponibile.
    """
    if not re.search(r"Descrizione is Imponibile", text):
        raise ValueError(
            f"file non Imponibile (atteso 'Descrizione is Imponibile'): {text[:120]!r}"
        )
    # Export multi-struttura ("CodiceHotel is A, B, or C"): la regex prenderebbe
    # solo il primo nome e l'aggregato finirebbe etichettato su UNA BU — rifiuta.
    if re.search(r"CodiceHotel is \w+\s*(,| or )", text):
        raise ValueError(
            f"CodiceHotel multiplo: serve un export per singola struttura: {text[:160]!r}"
        )
    # Filtro Mese = export parziale: la DELETE dello SNAPSHOT copre (BU, anno)
    # intero e cancellerebbe i mesi fuori filtro — rifiuta.
    if re.search(r"Mese is ", text):
        raise ValueError(
            f"filtro Mese presente: export parziale non ammesso (SNAPSHOT su anno intero): {text[:160]!r}"
        )
    hotel_m = re.search(r"CodiceHotel is (\w+)", text)
    anno_m = re.search(r"Anno is (\d{4}(?: or \d{4})*)", text)
    if not (hotel_m and anno_m):
        raise ValueError(f"Applied filters incompleti: {text[:120]!r}")
    codice = hotel_m.group(1)
    if codice not in HOTEL_TO_BU:
        raise ValueError(f"CodiceHotel sconosciuto: {codice}")
    anni = {int(a) for a in re.findall(r"\d{4}", anno_m.group(1))}
    return HOTEL_TO_BU[codice], anni


def detect_classe_columns(header: tuple) -> dict[int, str]:
    """Map column index → classe code for cells matching ^\\d{2}[A-Z]+$.

    Skips 'Total', 'Classe', blanks. Works regardless of leading blank cols.
    """
    out: dict[int, str] = {}
    for idx, cell in enumerate(header):
        if cell is None:
            continue
        name = str(cell).strip()
        if CLASSE_RE.match(name):
            out[idx] = name
    return out


def parse_xlsx(path: Path) -> tuple[tuple[str, set[int]], list[dict]]:
    """Read a Daily Production Report (classe cut) → ((bu, anni), data_rows).

    data_rows = list of {data: date, classe: str, importo: Decimal}. The 'Total'
    row/column and empty cells are skipped; negatives are kept.
    """
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["Export"] if "Export" in wb.sheetnames else wb.worksheets[0]
    rows = [r for r in ws.iter_rows(values_only=True) if any(c is not None for c in r)]
    wb.close()
    if len(rows) < 3:
        raise ValueError(f"{path.name}: troppe poche righe ({len(rows)})")

    # Applied filters = last row whose first cell starts with 'Applied filters'
    filter_text = next(
        (str(r[0]) for r in reversed(rows)
         if r and r[0] and str(r[0]).startswith("Applied filters")),
        "",
    )
    period = parse_applied_filters(filter_text)

    # Header = first row that yields ≥1 classe column
    classe_cols: dict[int, str] = {}
    for r in rows:
        cand = detect_classe_columns(r)
        if cand:
            classe_cols = cand
            break
    if not classe_cols:
        raise ValueError(f"{path.name}: nessuna colonna-classe rilevata")

    data_rows: list[dict] = []
    for r in rows:
        d = r[0] if r else None
        if not isinstance(d, datetime):
            continue
        for idx, classe in classe_cols.items():
            if idx >= len(r):
                continue
            val = r[idx]
            if val is None or val == "":
                continue
            if not isinstance(val, (int, float)):
                continue
            data_rows.append({
                "data": d.date(),
                "classe": classe,
                "importo": Decimal(str(val)),
            })

    _, anni = period
    fuori = {r["data"].year for r in data_rows} - anni
    if fuori:
        raise ValueError(
            f"{path.name}: righe datate fuori dagli anni del filtro "
            f"({sorted(fuori)} ∉ {sorted(anni)})"
        )
    return period, data_rows


def build_rows(
    period: tuple[str, set[int]],
    raw_rows: list[dict],
    file_name: str,
    raw_object_id: str | None = None,
) -> list[dict]:
    """Turn parsed raw rows into f_produzione_pms dict rows.

    Derives societa_id from data via OPERATIONS_CUTOVER_DATE, anno e mese from
    data (per-riga — un file può coprire più anni),
    and hash_riga = md5(business_unit_id, anno, data, classe).
    """
    bu, _anni = period
    now = datetime.now(timezone.utc)
    out: list[dict] = []
    for r in raw_rows:
        d = r["data"]
        anno = d.year
        societa = "INTUR" if d < OPERATIONS_CUTOVER_DATE else "ORTI"
        out.append({
            "societa_id": societa,
            "business_unit_id": bu,
            "data": d,
            "anno": anno,
            "mese": d.month,
            "classe": r["classe"],
            "importo_imponibile": r["importo"],
            "file_sorgente": file_name,
            "hash_riga": make_hash(bu, str(d.year), d.isoformat(), r["classe"]),
            "raw_object_id": raw_object_id,
            "data_caricamento": now,
        })
    return out


def ingest_file(
    path: Path, raw_object_id: str | None = None, dry_run: bool = False
) -> int:
    """Parse one xlsx and SNAPSHOT-write it to f_produzione_pms. Returns row count.

    raw_object_id is the lineage FK passed by `hotelops promote`; stamped on
    every row. SNAPSHOT scope = (business_unit_id, anno).
    """
    period, raw = parse_xlsx(path)
    rows = build_rows(period, raw, path.name, raw_object_id)
    validate_batch(rows, ProduzioneRow, context=f"produzione_pms {path.name}")
    bu, anni = period
    anni_str = ",".join(str(a) for a in sorted(anni))
    if dry_run:
        log.info("[DRY-RUN] %s → %s %s : %d righe", path.name, bu, anni_str, len(rows))
        return len(rows)

    from core.bq.write import bq_write_validated

    pydantic_rows = [ProduzioneRow(**r) for r in rows]
    bq_write_validated(
        F_PRODUZIONE_PMS,
        pydantic_rows,
        mode="snapshot",
        natural_key=["business_unit_id", "anno"],
    )
    log.info("OK %s → %s %s : %d righe", path.name, bu, anni_str, len(rows))
    return len(rows)


def main() -> None:
    ap = argparse.ArgumentParser(description="Ingest Daily Production Report → f_produzione_pms")
    ap.add_argument("--file", required=True, type=Path, help="xlsx Daily Production Report (classe)")
    ap.add_argument(
        "--raw-object-id",
        default=None,
        help="FK a f_raw_objects — passato da `hotelops promote`",
    )
    ap.add_argument(
        "--societa",
        default=None,
        help="Ignorato — societa derivata dal cutover. Accettato da `hotelops promote`.",
    )
    ap.add_argument("--dry-run", action="store_true", help="parse senza scrivere")
    args = ap.parse_args()

    logging.basicConfig(level=logging.INFO, format="%(levelname)s: %(message)s")
    n = ingest_file(args.file, raw_object_id=args.raw_object_id, dry_run=args.dry_run)
    log.info("Totale: %d righe", n)


if __name__ == "__main__":
    main()
