#!/usr/bin/env python3
"""
Pipeline: Situazione Partite Aperte Fornitori (snapshot).

Reads the Esolver "Situazione partite sintetica per fornitori" Excel export
and loads it into f_partite_aperte_fornitori in BigQuery.

Pattern: SNAPSHOT — DELETE-INSERT per societa_id + data_snapshot.
Every time Rosa exports the report, the new snapshot replaces the previous
one for that date, giving us a point-in-time view of open payables.

This is the THIRD TEMPORAL DIMENSION (IMPEGNO): knowing what's already
committed (invoices received, payment terms agreed) but not yet paid.

Usage:
    python -m ingest.flussi.ingest_partite_aperte --file export.xlsx
    python -m ingest.flussi.ingest_partite_aperte --file export.xlsx --societa INTUR
    python -m ingest.flussi.ingest_partite_aperte --file export.xlsx --dry-run
"""

import argparse
import logging
import re
import sys
from datetime import date, datetime
from pathlib import Path

from google.cloud import bigquery

from core.schemas import PartitaApertaFornitoreRow, validate_batch

BQ_TABLE = "hotelops-suite.hotelops.f_partite_aperte_fornitori"
BQ_PROJECT = "hotelops-suite"

# Payment code mapping (col 21 → col 22 in Esolver)
PAYMENT_CODES = {
    "01": "Rimessa diretta",
    "02": "RI.BA.",
    "03": "SDD",
    "04": "Bonifico SEPA",
    "08": "PagoPA",
    "10": "Carta di credito",
}

# Known intercompany suppliers (fornitore names containing these → flag)
INTERCOMPANY_KEYWORDS = ["PANORAMA COMPANY", "INTUR", "ORTI S.R.L."]

logger = logging.getLogger("partite_aperte")


def parse_situazione_partite(filepath: Path, societa_override: str | None = None) -> list[dict]:
    """Parse Esolver 'Situazione partite sintetica per fornitori' Excel.

    Column mapping (from real Esolver export):
        C1:  Società (e.g. 'ORTI S.R.L.')
        C3:  Data snapshot (date)
        C11: Codice fornitore (int)
        C12: Nome fornitore (part 1)
        C13: Nome fornitore (part 2, optional — for long names)
        C19: Descrizione documento (e.g. 'FT n. EE-FAT202600026496 del 12/03/2026')
        C20: Importo residuo (negative = we owe)
        C21: Codice pagamento (e.g. '04')
        C22: Metodo pagamento (e.g. 'Bonifico SEPA')
        C23: Data scadenza (date)
        C25: Importo assoluto (always positive)
    """
    import openpyxl

    wb = openpyxl.load_workbook(str(filepath), data_only=True)
    ws = wb.active

    # Extract metadata from first row
    societa_raw = ws.cell(row=1, column=1).value or ""
    data_snapshot_raw = ws.cell(row=1, column=3).value

    # Determine societa_id
    if societa_override:
        societa_id = societa_override.upper()
    elif "INTUR" in societa_raw.upper():
        societa_id = "INTUR"
    elif "ORTI" in societa_raw.upper():
        societa_id = "ORTI"
    else:
        raise ValueError(f"Cannot determine società from '{societa_raw}'. Use --societa.")

    # Determine snapshot date
    if isinstance(data_snapshot_raw, datetime):
        data_snapshot = data_snapshot_raw.date()
    elif isinstance(data_snapshot_raw, date):
        data_snapshot = data_snapshot_raw
    else:
        raise ValueError(f"Cannot parse snapshot date from C3: {data_snapshot_raw}")

    logger.info(f"Parsing: {filepath.name}")
    logger.info(f"  Società: {societa_id}, Snapshot: {data_snapshot}")

    rows = []
    skipped = 0

    for row_idx in range(1, ws.max_row + 1):
        desc_raw = ws.cell(row=row_idx, column=19).value
        importo_raw = ws.cell(row=row_idx, column=20).value
        scadenza_raw = ws.cell(row=row_idx, column=23).value

        # Skip rows without a document description or amount
        if not desc_raw or importo_raw is None:
            skipped += 1
            continue

        # Parse fornitore
        codice_fornitore = ws.cell(row=row_idx, column=11).value
        nome_part1 = ws.cell(row=row_idx, column=12).value or ""
        nome_part2 = ws.cell(row=row_idx, column=13).value or ""
        nome_fornitore = f"{nome_part1} {nome_part2}".strip()

        if not codice_fornitore or not nome_fornitore:
            skipped += 1
            continue

        # Parse document description: "FT n. EE-FAT202600026496 del 12/03/2026"
        tipo_doc, num_doc, data_doc = _parse_descrizione(str(desc_raw))

        # Parse scadenza
        if isinstance(scadenza_raw, datetime):
            data_scadenza = scadenza_raw.date()
        elif isinstance(scadenza_raw, date):
            data_scadenza = scadenza_raw
        else:
            # If no scadenza, use document date as fallback
            data_scadenza = data_doc or data_snapshot

        # Payment info
        cod_pag = str(ws.cell(row=row_idx, column=21).value or "")
        metodo_pag = ws.cell(row=row_idx, column=22).value or PAYMENT_CODES.get(cod_pag, "Sconosciuto")

        # Amounts
        importo_residuo = float(importo_raw)
        importo_abs_raw = ws.cell(row=row_idx, column=25).value
        importo_abs = float(importo_abs_raw) if importo_abs_raw else abs(importo_residuo)

        # Intercompany flag
        is_intercompany = any(kw in nome_fornitore.upper() for kw in INTERCOMPANY_KEYWORDS)

        rows.append({
            "societa_id": societa_id,
            "data_snapshot": str(data_snapshot),
            "codice_fornitore": int(codice_fornitore),
            "nome_fornitore": nome_fornitore,
            "tipo_documento": tipo_doc,
            "numero_documento": num_doc,
            "data_documento": str(data_doc) if data_doc else str(data_snapshot),
            "data_scadenza": str(data_scadenza),
            "importo_residuo": importo_residuo,
            "importo_abs": importo_abs,
            "codice_pagamento": cod_pag,
            "metodo_pagamento": metodo_pag,
            "is_intercompany": is_intercompany,
            "file_sorgente": filepath.name,
        })

    logger.info(f"  Parsed: {len(rows)} partite, skipped {skipped} rows")
    return rows


def _parse_descrizione(desc: str) -> tuple[str, str, date | None]:
    """Parse 'FT n. EE-FAT202600026496 del 12/03/2026' into (tipo, numero, data)."""
    tipo = "FT"
    numero = desc
    data_doc = None

    # Extract document type (FT, NC, AFT)
    m_tipo = re.match(r"^(\w+)\s+n\.", desc)
    if m_tipo:
        tipo = m_tipo.group(1).upper()

    # Extract document number
    m_num = re.search(r"n\.\s*(.+?)\s+del\s+", desc)
    if m_num:
        numero = m_num.group(1).strip()

    # Extract document date
    m_date = re.search(r"del\s+(\d{1,2})/(\d{2})/(\d{4})", desc)
    if m_date:
        try:
            data_doc = date(
                int(m_date.group(3)),
                int(m_date.group(2)),
                int(m_date.group(1)),
            )
        except ValueError:
            pass

    return tipo, numero, data_doc


def create_table_if_needed(client: bigquery.Client):
    """Create f_partite_aperte_fornitori table if it doesn't exist."""
    try:
        client.get_table(BQ_TABLE)
        return
    except Exception:
        pass

    schema = [
        bigquery.SchemaField("societa_id", "STRING"),
        bigquery.SchemaField("data_snapshot", "DATE"),
        bigquery.SchemaField("codice_fornitore", "INTEGER"),
        bigquery.SchemaField("nome_fornitore", "STRING"),
        bigquery.SchemaField("tipo_documento", "STRING"),
        bigquery.SchemaField("numero_documento", "STRING"),
        bigquery.SchemaField("data_documento", "DATE"),
        bigquery.SchemaField("data_scadenza", "DATE"),
        bigquery.SchemaField("importo_residuo", "FLOAT"),
        bigquery.SchemaField("importo_abs", "FLOAT"),
        bigquery.SchemaField("codice_pagamento", "STRING"),
        bigquery.SchemaField("metodo_pagamento", "STRING"),
        bigquery.SchemaField("is_intercompany", "BOOLEAN"),
        bigquery.SchemaField("file_sorgente", "STRING"),
    ]
    table = bigquery.Table(BQ_TABLE, schema=schema)
    client.create_table(table)
    logger.info(f"Created table {BQ_TABLE}")


def load_to_bq(client: bigquery.Client, rows: list[dict], dry_run: bool = False):
    """DELETE-INSERT snapshot into BigQuery."""
    if not rows:
        logger.warning("No rows to load")
        return

    societa_id = rows[0]["societa_id"]
    data_snapshot = rows[0]["data_snapshot"]

    if dry_run:
        logger.info(f"DRY RUN: would load {len(rows)} rows for {societa_id} @ {data_snapshot}")
        _print_summary(rows)
        return

    create_table_if_needed(client)

    # DELETE existing snapshot for this societa + date
    delete_sql = f"""
    DELETE FROM `{BQ_TABLE}`
    WHERE societa_id = '{societa_id}'
      AND data_snapshot = DATE('{data_snapshot}')
    """
    client.query(delete_sql).result()
    logger.info(f"  Deleted existing snapshot for {societa_id} @ {data_snapshot}")

    # INSERT new rows
    errors = client.insert_rows_json(BQ_TABLE, rows)
    if errors:
        logger.error(f"  BQ insert errors: {errors[:3]}")
        raise RuntimeError(f"BQ insert failed: {errors[:3]}")

    logger.info(f"  ✓ Loaded {len(rows)} partite aperte for {societa_id} @ {data_snapshot}")
    _print_summary(rows)


def _print_summary(rows: list[dict]):
    """Print a quick summary of what was loaded."""
    from collections import defaultdict

    by_month = defaultdict(float)
    intercompany_total = 0.0
    total = 0.0

    for r in rows:
        importo = r["importo_abs"]
        total += importo
        # Group by scadenza month
        scad = r["data_scadenza"]
        month_key = scad[:7]  # "2026-03"
        by_month[month_key] += importo
        if r.get("is_intercompany"):
            intercompany_total += importo

    logger.info(f"\n  SUMMARY: {len(rows)} partite, totale €{total:,.0f} (di cui intercompany €{intercompany_total:,.0f})")
    logger.info("  Per mese di scadenza:")
    for k in sorted(by_month.keys()):
        logger.info(f"    {k}: €{by_month[k]:>12,.0f}")


def main():
    parser = argparse.ArgumentParser(
        description="Ingest Esolver Situazione Partite Fornitori (snapshot)"
    )
    parser.add_argument("--file", required=True, help="Path to Esolver Excel export")
    parser.add_argument("--societa", help="Override società (ORTI/INTUR)")
    parser.add_argument("--dry-run", action="store_true", help="Parse only, no BQ write")
    args = parser.parse_args()

    from ingest._logging import setup_logging
    setup_logging("partite_aperte", Path(__file__).parent / "logs")

    filepath = Path(args.file)
    if not filepath.exists():
        logger.error(f"File not found: {filepath}")
        sys.exit(1)

    # Parse
    rows = parse_situazione_partite(filepath, societa_override=args.societa)

    if not rows:
        logger.warning("No partite found in file")
        sys.exit(0)

    # Validate
    validate_batch(rows, PartitaApertaFornitoreRow, context="Partite Aperte Fornitori")
    logger.info(f"  ✓ {len(rows)} rows validated")

    # Load
    if args.dry_run:
        load_to_bq(None, rows, dry_run=True)
    else:
        client = bigquery.Client(project=BQ_PROJECT)
        load_to_bq(client, rows)


if __name__ == "__main__":
    main()
