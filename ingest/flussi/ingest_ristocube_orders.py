#!/usr/bin/env python3
"""Ingest RistoCube Orders Report → f_ristocube_orders.

Una riga per item × comanda. I campi della comanda (tavolo, sala, coperti,
segmento, etc.) sono ereditati da ogni item (flatten).

Struttura del file (header "Data" SOLO all'inizio):
  Row 0: header [Data, Orario Apertura, Orario Chiusura, Tavoli, Sala, Comanda, Cop., Totale, Operatore]
  Per ogni comanda:
    Row N+0: valori comanda (col[0]=data dd/mm/yy, col[1]=orario)
    Row N+1: sub-header [None, Modalità Chiusura, Segmento Cliente, Importo, MP, Note Direzione]
    Row N+2: valori sub-header
    Row N+3: items-header [None, None, Menu, Articolo POS, Descrizione, Quantità, ...]
    Row N+4..N+4+k-1: k item rows
    -> direttamente next comanda

Pattern: APPEND + hash_riga dedup. Idempotent re-ingest tramite filter_new_rows_by_hash.

È il parser_module di RISTOCUBE_ORDERS_ORTI_APPEND, invocato da
`hotelops promote` come: python -m ingest.flussi.ingest_ristocube_orders --file X --raw-object-id Y

Usage:
    python -m ingest.flussi.ingest_ristocube_orders --file <xlsx> --raw-object-id <id>
    python -m ingest.flussi.ingest_ristocube_orders --file <xlsx> --dry-run
"""
from __future__ import annotations

import argparse
import logging
import re
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

from core.config import F_RISTOCUBE_ORDERS
from core.schemas import RistocubeOrderRow, make_hash, validate_batch

log = logging.getLogger("ingest.ristocube_orders")

SOCIETA_ID = "ORTI"
DATE_RE = re.compile(r"^\d{2}/\d{2}/\d{2}$")


def _bu_from_sala(sala: str) -> str:
    """Sala raw → business_unit_id. BAR e RIS* → HOTEL (tutto nel hotel)."""
    s = (sala or "").upper()
    if s.startswith("BAR") or s.startswith("RIS"):
        return "HOTEL"
    return "HOTEL"  # default


def _is_comanda_values_row(r: tuple) -> bool:
    """Riga è una comanda values? col[0] = date dd/mm/yy."""
    if not r or len(r) < 8 or r[0] is None:
        return False
    if not isinstance(r[0], str):
        return False
    return bool(DATE_RE.match(r[0]))


def _safe_str(v) -> str | None:
    """Convert to stripped string or None."""
    if v is None:
        return None
    s = str(v).strip()
    return s if s else None


def _safe_float(v) -> float | None:
    if v is None:
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _safe_int(v) -> int | None:
    if v is None:
        return None
    try:
        return int(v)
    except (TypeError, ValueError):
        return None


def parse_xlsx(path: Path, raw_object_id: str | None = None) -> list[dict]:
    """Parse one Orders Report xlsx → list of flat item dicts ready for RistocubeOrderRow.

    One dict per item × comanda. Comanda-level fields are inherited by each item.
    """
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    now = datetime.now(timezone.utc)
    items: list[dict] = []
    i = 0
    n = len(rows)

    while i < n:
        r = rows[i]
        if not _is_comanda_values_row(r):
            i += 1
            continue

        # Parse comanda date
        try:
            d = datetime.strptime(r[0], "%d/%m/%y").date()
        except ValueError:
            i += 1
            continue

        comanda_id_raw = r[5]
        comanda_id = _safe_int(comanda_id_raw)
        if comanda_id is None:
            i += 1
            continue

        sala_raw = _safe_str(r[4]) or ""
        bu = _bu_from_sala(sala_raw)

        comanda: dict = {
            "data": d,
            "anno": d.year,
            "mese": d.month,
            "giorno": d.day,
            "orario_apertura": _safe_str(r[1]),
            "orario_chiusura": _safe_str(r[2]),
            "tavolo": _safe_str(r[3]),
            "sala": sala_raw,
            "comanda_id": comanda_id,
            "coperti_comanda": _safe_int(r[6]),
            "totale_comanda": _safe_float(r[7]),
            "operatore_apertura": _safe_str(r[8]) if len(r) > 8 else None,
            "modalita_chiusura": None,
            "segmento_cliente": None,
            "importo_pagamento": None,
            "mp": None,
            "note_direzione": None,
        }

        # Row i+1 = sub-header, i+2 = sub-values
        if i + 2 < n and rows[i + 1] and rows[i + 1][1] == "Modalità Chiusura Comanda":
            sub = rows[i + 2]
            if sub:
                comanda["modalita_chiusura"] = _safe_str(sub[1])
                comanda["segmento_cliente"] = _safe_str(sub[2])
                comanda["importo_pagamento"] = _safe_float(sub[3])
                comanda["mp"] = _safe_str(sub[4])
                comanda["note_direzione"] = _safe_str(sub[5])
            j = i + 3
        else:
            j = i + 1

        # Row j = items-header (skip if present)
        if j < n and rows[j] and rows[j][2] == "Menu":
            j += 1

        # Items: one row per item until next comanda row
        while j < n:
            ir = rows[j]
            if _is_comanda_values_row(ir):
                break
            if not ir or all(c is None for c in ir):
                j += 1
                continue
            # Item cols: col[3]=codice POS, col[4]=descrizione
            if len(ir) < 10 or ir[3] is None or ir[4] is None:
                j += 1
                continue

            item_codice = _safe_str(ir[3]) or ""
            item_quantita = _safe_float(ir[5])
            item_importo_finale = _safe_float(ir[9])

            # Hash: data|comanda_id|item_codice_pos|item_quantita|item_importo_finale
            hash_riga = make_hash(
                str(d),
                str(comanda_id),
                item_codice,
                str(item_quantita),
                str(item_importo_finale),
            )

            items.append({
                "hash_riga": hash_riga,
                "societa_id": SOCIETA_ID,
                "business_unit_id": bu,
                "data": d,
                "anno": d.year,
                "mese": d.month,
                "giorno": d.day,
                "orario_apertura": comanda["orario_apertura"],
                "orario_chiusura": comanda["orario_chiusura"],
                "tavolo": comanda["tavolo"],
                "sala": sala_raw,
                "comanda_id": comanda_id,
                "coperti_comanda": comanda["coperti_comanda"],
                "totale_comanda": comanda["totale_comanda"],
                "operatore_apertura": comanda["operatore_apertura"],
                "modalita_chiusura": comanda["modalita_chiusura"],
                "segmento_cliente": comanda["segmento_cliente"],
                "importo_pagamento": comanda["importo_pagamento"],
                "mp": comanda["mp"],
                "note_direzione": comanda["note_direzione"],
                "item_menu": _safe_str(ir[2]),
                "item_codice_pos": item_codice,
                "item_descrizione": _safe_str(ir[4]),
                "item_quantita": item_quantita,
                "item_importo_originale": _safe_float(ir[6]),
                "item_sconto_tipo": _safe_str(ir[7]),
                "item_importo_sconto": _safe_float(ir[8]),
                "item_importo_finale": item_importo_finale,
                "file_sorgente": path.name,
                "raw_object_id": raw_object_id,
                "data_caricamento": now,
            })
            j += 1

        i = j

    return items


def ingest_file(
    path: Path, raw_object_id: str | None = None, dry_run: bool = False
) -> int:
    """Parse one xlsx and APPEND-write new rows to f_ristocube_orders. Returns row count written."""
    rows = parse_xlsx(path, raw_object_id=raw_object_id)
    validate_batch(rows, RistocubeOrderRow, context=f"ristocube_orders {path.name}")

    if dry_run:
        log.info("[DRY-RUN] %s : %d righe estratte", path.name, len(rows))
        return len(rows)

    from core.bq.dedup import filter_new_rows_by_hash
    from core.bq.write import bq_write_validated

    new_rows = filter_new_rows_by_hash(F_RISTOCUBE_ORDERS, rows, "hash_riga")
    if not new_rows:
        log.info("Tutte le righe già presenti — niente da scrivere (%s)", path.name)
        return 0

    pydantic_rows = [RistocubeOrderRow(**r) for r in new_rows]
    bq_write_validated(F_RISTOCUBE_ORDERS, pydantic_rows, mode="append")
    log.info("OK %s : %d righe nuove scritte (su %d totali)", path.name, len(new_rows), len(rows))
    return len(new_rows)


def main() -> None:
    ap = argparse.ArgumentParser(description="Ingest RistoCube Orders Report → f_ristocube_orders")
    ap.add_argument("--file", required=True, type=Path, help="xlsx Orders Report RISTOCUBE")
    ap.add_argument(
        "--raw-object-id",
        default=None,
        help="FK a f_raw_objects — passato da `hotelops promote`",
    )
    ap.add_argument(
        "--societa",
        default=None,
        help="Ignorato — sempre ORTI. Accettato da `hotelops promote`.",
    )
    ap.add_argument("--dry-run", action="store_true", help="parse senza scrivere")
    args = ap.parse_args()

    logging.basicConfig(level=logging.INFO, format="%(levelname)s: %(message)s")

    from core.pipeline_run import PipelineRun

    with PipelineRun(
        "ingest_ristocube_orders",
        societa_id=SOCIETA_ID,
        file_sorgente=args.file.name,
    ):
        n = ingest_file(args.file, raw_object_id=args.raw_object_id, dry_run=args.dry_run)
        log.info("Totale: %d righe", n)


if __name__ == "__main__":
    main()
