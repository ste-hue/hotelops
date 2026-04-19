#!/usr/bin/env python3
"""
Scheda contabile Esolver conto 190101 (banca) → f_movimenti_contabili.

Parses Esolver "Scheda contabile" XLSX exports filtered on cod_conto=190101,
extracting row-level movements (not just end-of-day saldo like
ingest_scheda_contabile.py). Fills the gap where ingest_movimenti_contabili
exports don't cover the bank ledger account.

Expected 14 columns:
    Data registrazione | Rif. registrazione | Causale contabile |
    Dare in UdC | Avere in UdC | Saldo in UdC | Documento |
    Partitario | Descrizione partitario | Val. |
    Dare in valuta | Avere in valuta | Riferimenti IVA | Centro imputazione

Usage:
    python -m ingest.flussi.ingest_scheda_190101 \\
        --file ~/Desktop/orti2025.XLSX --societa ORTI

    python -m ingest.flussi.ingest_scheda_190101 \\
        --file ~/Desktop/orti2025.XLSX --societa ORTI --dry-run
"""

from __future__ import annotations

import argparse
import logging
import re
import sys
from datetime import date, datetime
from pathlib import Path

import openpyxl
from google.cloud import bigquery

from core.bq.client import get_client
from core.config import F_MOVIMENTI_CONTABILI
from core.schemas import MovimentoContabileRow, make_hash, validate_batch

BQ_TABLE = F_MOVIMENTI_CONTABILI
COD_CONTO_BANCA = "190101"

EXPECTED_HEADER = [
    "Data registrazione",
    "Rif. registrazione",
    "Causale contabile",
    "Dare in UdC",
    "Avere in UdC",
    "Saldo in UdC",
    "Documento",
    "Partitario",
    "Descrizione partitario",
    "Val.",
    "Dare in valuta",
    "Avere in valuta",
    "Riferimenti IVA",
    "Centro imputazione",
]

PNC_RE = re.compile(r"PNC\s*n\s*(\d+)", re.IGNORECASE)


def setup_logger() -> logging.Logger:
    log = logging.getLogger("ingest_scheda_190101")
    log.setLevel(logging.INFO)
    h = logging.StreamHandler(sys.stdout)
    h.setFormatter(logging.Formatter("%(levelname)s %(message)s"))
    log.addHandler(h)
    return log


def parse_file(path: Path, societa: str, logger: logging.Logger) -> list[dict]:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]

    header = [ws.cell(1, c).value for c in range(1, 15)]
    if header != EXPECTED_HEADER:
        logger.error(f"  Header inatteso in {path.name}: {header}")
        wb.close()
        return []

    rows = []
    today = date.today().isoformat()
    for r in range(2, ws.max_row + 1):
        data_raw = ws.cell(r, 1).value
        if not isinstance(data_raw, datetime):
            continue  # skip "Riporto saldi" / blank rows

        rif_reg = str(ws.cell(r, 2).value or "").strip()
        causale = str(ws.cell(r, 3).value or "").strip()

        # Skip opening balance rows — they're not real movements
        if "ripresa saldi" in causale.lower():
            continue

        dare = float(ws.cell(r, 4).value or 0)
        avere = float(ws.cell(r, 5).value or 0)
        saldo = float(ws.cell(r, 6).value or 0)
        documento = ws.cell(r, 7).value
        partitario_raw = ws.cell(r, 8).value
        descr_part = str(ws.cell(r, 9).value or "").strip()

        if partitario_raw is None:
            continue
        partitario = (
            str(int(partitario_raw))
            if isinstance(partitario_raw, (int, float))
            else str(partitario_raw).strip()
        )

        pnc_match = PNC_RE.search(rif_reg)
        pnc_num = int(pnc_match.group(1)) if pnc_match else None

        data_reg = data_raw.date().isoformat()
        hash_riga = make_hash(
            "SCHEDA_190101",
            societa,
            data_reg,
            rif_reg,
            partitario,
            f"{dare:.2f}",
            f"{avere:.2f}",
            f"{saldo:.2f}",
        )

        rows.append(
            {
                "hash_riga": hash_riga,
                "societa_id": societa,
                "id_documento": pnc_num,
                "num_progr_riga": None,
                "gruppo_doc": "SCHEDA_190101",
                "anno": data_raw.year,
                "mese": data_raw.month,
                "data_registrazione": data_reg,
                "sigla_doc": None,
                "rif_registrazione": rif_reg,
                "num_doc_originale": str(documento).strip() if documento else None,
                "data_originale": None,
                "tipo_documento": None,
                "cod_conto": COD_CONTO_BANCA,
                "cod_partitario": partitario,
                "rag_sociale": descr_part or None,
                "causale_contabile": causale or None,
                "imp_dare": dare,
                "imp_avere": avere,
                "cod_divisione": None,
                "file_sorgente": path.name,
                "data_ingresso": today,
            }
        )

    wb.close()
    logger.info(f"  {path.name}: {len(rows)} righe parsed")
    return rows


def load_existing_hashes(bq_client, societa: str, logger: logging.Logger) -> set[str]:
    q = f"""
    SELECT hash_riga FROM `{BQ_TABLE}`
    WHERE societa_id = @soc AND cod_conto = '{COD_CONTO_BANCA}'
      AND gruppo_doc = 'SCHEDA_190101'
    """
    job = bq_client.query(
        q,
        job_config=bigquery.QueryJobConfig(
            query_parameters=[
                bigquery.ScalarQueryParameter("soc", "STRING", societa),
            ]
        ),
    )
    hashes = {r.hash_riga for r in job.result()}
    logger.info(f"  Hashes esistenti BQ ({societa}, 190101, SCHEDA): {len(hashes)}")
    return hashes


def write_to_bq(rows: list[dict], bq_client, logger: logging.Logger):
    if not rows:
        logger.info("  BQ: nessuna riga da inserire")
        return
    job_config = bigquery.LoadJobConfig(
        write_disposition=bigquery.WriteDisposition.WRITE_APPEND,
    )
    job = bq_client.load_table_from_json(rows, BQ_TABLE, job_config=job_config)
    job.result()
    if job.errors:
        logger.error(f"BQ errors: {job.errors}")
    else:
        logger.info(f"  BQ: {len(rows)} righe → {BQ_TABLE}")


def main():
    ap = argparse.ArgumentParser(
        description="Ingest Esolver scheda contabile 190101 → f_movimenti_contabili"
    )
    ap.add_argument(
        "--file", action="append", required=True, help="XLSX file (repeatable)"
    )
    ap.add_argument("--societa", required=True, choices=["ORTI", "INTUR"])
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()

    logger = setup_logger()

    all_rows: list[dict] = []
    for f in args.file:
        all_rows.extend(parse_file(Path(f), args.societa, logger))

    # In-memory dedup (same file overlap or cross-file)
    seen: set[str] = set()
    uniq: list[dict] = []
    for r in all_rows:
        if r["hash_riga"] in seen:
            continue
        seen.add(r["hash_riga"])
        uniq.append(r)
    logger.info(f"Totale parsed: {len(all_rows)}, unici in-file: {len(uniq)}")

    # I1 gate: validate every row before any BQ interaction
    validate_batch(uniq, MovimentoContabileRow, context="ingest_scheda_190101")
    logger.info(f"Validation OK: {len(uniq)} righe")

    if args.dry_run:
        logger.info("DRY-RUN: stop here")
        return

    bq = get_client()
    existing = load_existing_hashes(bq, args.societa, logger)
    new_rows = [r for r in uniq if r["hash_riga"] not in existing]
    logger.info(f"Nuove: {len(new_rows)}, già presenti: {len(uniq) - len(new_rows)}")
    write_to_bq(new_rows, bq, logger)


if __name__ == "__main__":
    main()
