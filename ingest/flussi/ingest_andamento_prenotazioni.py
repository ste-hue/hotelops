#!/usr/bin/env python3
"""Ingest Power BI "Andamento Prenotazioni" (OTB) → f_prenotazioni_otb.

Portafoglio prenotazioni on-the-books (consumato + soggiorni futuri) per
data soggiorno, con o senza dimensione tipologia. Ogni export è una
FOTOGRAFIA alla snapshot_date: le fotografie si accumulano (booking pace),
la SNAPSHOT delete è scoped a (snapshot_date, business_unit_id, dim_tipologia).

Layout (header-based, il file può avere o non avere la colonna tipologia):
  Giorno | CodiceHotel | [Tipologia | Tipologia Venduta] | Camere |
  Presenze ARB | Importo Lordo | ADR Lordo | Imponibile | ADR Imponibile

- `dim_tipologia` dedotta dall'header (content-only): "Tipologia Venduta" →
  VENDUTA, "Tipologia" → ASSEGNATA, assente → NESSUNA. Mai sommare varianti
  diverse della stessa snapshot_date (stessa fotografia, dimensioni diverse).
- BU dal contenuto (colonna CodiceHotel), non dal filename.
- ⚠️ La data dello snapshot NON è nel contenuto: si risolve dall'intake del
  raw object (--raw-object-id) o si passa esplicita con --snapshot-date.

Parser_module della source POWERBI_ANDAMENTOPRENOTAZIONI_ORTI_SNAPSHOT:
  python -m ingest.flussi.ingest_andamento_prenotazioni --file X --raw-object-id Y

Usage:
    python -m ingest.flussi.ingest_andamento_prenotazioni --file <xlsx> --raw-object-id <id>
    python -m ingest.flussi.ingest_andamento_prenotazioni --file <xlsx> --snapshot-date 2026-07-06 --dry-run
"""

from __future__ import annotations

import argparse
import logging
import re
from datetime import date, datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

from core.config import F_PRENOTAZIONI_OTB
from core.schemas import PrenotazioniOtbRow, make_hash, validate_batch

log = logging.getLogger("ingest.andamento_prenotazioni")

FONTE = "POWERBI_ANDAMENTOPRENOTAZIONI"

CODICE_HOTEL_TO_BU = {
    "PANORAMAHT": "HOTEL",
    "ANGELINARES": "RESIDENCE",
    "HOMEHOLIDAY": "CVM",
}


def _num(v) -> float:
    if v is None or v == "":
        return 0.0
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def _parse_giorno(v) -> date | None:
    """Cella Giorno: datetime oppure stringa 'dd/mm/YYYY ggg'."""
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, str):
        try:
            return datetime.strptime(v.strip()[:10], "%d/%m/%Y").date()
        except ValueError:
            return None
    return None


def valida_basi_export(path: Path) -> tuple[bool, str]:
    """Guardiano basi-omogenee per la foto OTB (upload dalla pagina Revenue).

    Il pace è foto-vs-foto della STESSA base (concept BOOKING_PACE_E_BASI):
    un export coi filtri sbagliati inquinerebbe la serie. Controlla:
    - layout foto (colonne Giorno/CodiceHotel/Imponibile presenti);
    - footer con 'DataPrenotazione is not blank';
    - filtro Anno su ALMENO due anni (caso reale 2026-07-13: 'Anno is 2026'
      escludeva ~2.800 notti prenotate nel 2025).
    Ritorna (ok, motivo) — motivo leggibile quando ok=False.
    """
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = list(wb.active.iter_rows(values_only=True))
    finally:
        wb.close()
    if not rows:
        return False, "file vuoto (export venuto male: rifallo)"

    header = {str(c).strip().lower() for c in rows[0] if c is not None}
    attese = {"giorno", "codicehotel", "camere", "imponibile"}
    if not attese <= header:
        return False, (
            "layout sbagliato: non è l'export 'Andamento Prenotazioni' "
            f"(mancano le colonne {sorted(attese - header)})"
        )

    footer = ""
    for r in reversed(rows[-5:]):
        if r and r[0] and "Applied filters" in str(r[0]):
            footer = str(r[0])
            break
    if not footer:
        return False, "footer 'Applied filters' non trovato: export incompleto"
    if "DataPrenotazione is not blank" not in footer:
        return False, (
            "manca il filtro 'DataPrenotazione is not blank': "
            "base diversa dalla serie storica"
        )
    riga_anno = next(
        (line for line in footer.splitlines() if line.strip().startswith("Anno is")),
        "",
    )
    anni = set(re.findall(r"\b(20\d{2})\b", riga_anno))
    if len(anni) < 2:
        return False, (
            f"filtro Anno su un solo anno ({riga_anno.strip() or 'assente'}): "
            "servono ENTRAMBI gli anni (le prenotazioni per l'estate N "
            "arrivano anche nell'anno N-1) — rifai l'export con Anno = "
            "anno corrente e precedente"
        )
    return True, ""


def parse_xlsx(path: Path) -> dict:
    """Estrae {dim_tipologia, righe[]} dal file (header-based)."""
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    header = [str(c).strip().lower() if c is not None else "" for c in rows[0]]
    col = {name: i for i, name in enumerate(header)}
    if "tipologia venduta" in col:
        dim, tip_idx = "VENDUTA", col["tipologia venduta"]
    elif "tipologia" in col:
        dim, tip_idx = "ASSEGNATA", col["tipologia"]
    else:
        dim, tip_idx = "NESSUNA", None

    def idx(name: str) -> int:
        if name not in col:
            raise ValueError(
                f"{path.name}: colonna '{name}' mancante (header: {header})"
            )
        return col[name]

    i_giorno, i_hotel = idx("giorno"), idx("codicehotel")
    i_cam, i_arb = idx("camere"), idx("presenze arb")
    i_lordo, i_adr_l = idx("importo lordo"), idx("adr lordo")
    i_imp, i_adr_i = idx("imponibile"), idx("adr imponibile")

    righe = []
    for r in rows[1:]:
        if not r or len(r) <= i_imp:
            continue
        d = _parse_giorno(r[i_giorno])
        codice = r[i_hotel]
        if d is None or not codice:
            continue  # Total / footer / righe non-dato
        codice = str(codice).strip()
        if codice not in CODICE_HOTEL_TO_BU:
            raise ValueError(f"{path.name}: CodiceHotel sconosciuto '{codice}'")
        righe.append(
            {
                "data": d,
                "business_unit_id": CODICE_HOTEL_TO_BU[codice],
                "tipologia": str(r[tip_idx]).strip()
                if tip_idx is not None and r[tip_idx]
                else None,
                "camere": int(_num(r[i_cam])),
                "presenze_arb": int(_num(r[i_arb])),
                "importo_lordo": _num(r[i_lordo]),
                "adr_lordo": _num(r[i_adr_l]),
                "imponibile": _num(r[i_imp]),
                "adr_imponibile": _num(r[i_adr_i]),
            }
        )
    return {"dim_tipologia": dim, "righe": righe}


def build_rows(
    parsed: dict, snapshot_date: str, raw_object_id: str | None
) -> list[dict]:
    now = datetime.now(timezone.utc).replace(tzinfo=None).isoformat()
    dim = parsed["dim_tipologia"]
    out = []
    for r in parsed["righe"]:
        d_iso = r["data"].isoformat()
        out.append(
            {
                "snapshot_date": snapshot_date,
                "societa_id": "ORTI",
                "business_unit_id": r["business_unit_id"],
                "data": d_iso,
                "dim_tipologia": dim,
                "tipologia": r["tipologia"],
                "camere": r["camere"],
                "presenze_arb": r["presenze_arb"],
                "importo_lordo": round(r["importo_lordo"], 2),
                "adr_lordo": round(r["adr_lordo"], 2),
                "imponibile": round(r["imponibile"], 2),
                "adr_imponibile": round(r["adr_imponibile"], 2),
                "fonte": FONTE,
                "hash_riga": make_hash(
                    snapshot_date,
                    r["business_unit_id"],
                    d_iso,
                    dim,
                    r["tipologia"] or "",
                ),
                "raw_object_id": raw_object_id,
                "data_caricamento": now,
            }
        )
    return out


def _snapshot_date_from_raw(raw_object_id: str) -> str:
    """La data dello snapshot non è nel file: fa fede l'intake del raw object."""
    from core.bq.client import get_client
    from core.config import PROJECT

    client = get_client()
    sql = (
        f"SELECT DATE(intake_at) d FROM `{PROJECT}.hotelops.f_raw_objects` "
        "WHERE raw_object_id = @rid"
    )
    from google.cloud import bigquery

    job = client.query(
        sql,
        job_config=bigquery.QueryJobConfig(
            query_parameters=[
                bigquery.ScalarQueryParameter("rid", "STRING", raw_object_id)
            ]
        ),
    )
    res = list(job.result())
    if not res:
        raise ValueError(f"raw_object_id {raw_object_id} non trovato in f_raw_objects")
    return res[0]["d"].isoformat()


def ingest_file(
    path: Path,
    raw_object_id: str | None = None,
    snapshot_date: str | None = None,
    dry_run: bool = False,
) -> int:
    if snapshot_date is None:
        if raw_object_id is None:
            raise ValueError("serve --raw-object-id oppure --snapshot-date")
        snapshot_date = _snapshot_date_from_raw(raw_object_id)

    parsed = parse_xlsx(path)
    rows = build_rows(parsed, snapshot_date, raw_object_id)
    validate_batch(
        rows, PrenotazioniOtbRow, context=f"andamento_prenotazioni {path.name}"
    )
    if dry_run:
        log.info(
            "[DRY-RUN] %s → snapshot %s, dim=%s: %d righe",
            path.name,
            snapshot_date,
            parsed["dim_tipologia"],
            len(rows),
        )
        return len(rows)

    from core.bq.write import bq_write_validated

    pydantic_rows = [PrenotazioniOtbRow(**r) for r in rows]
    bq_write_validated(
        F_PRENOTAZIONI_OTB,
        pydantic_rows,
        mode="snapshot",
        natural_key=["snapshot_date", "business_unit_id", "dim_tipologia"],
    )
    log.info(
        "OK %s → snapshot %s, dim=%s: %d righe",
        path.name,
        snapshot_date,
        parsed["dim_tipologia"],
        len(rows),
    )
    return len(rows)


def main() -> None:
    ap = argparse.ArgumentParser(
        description="Ingest Andamento Prenotazioni (OTB) → f_prenotazioni_otb"
    )
    ap.add_argument("--file", required=True, type=Path)
    ap.add_argument(
        "--raw-object-id",
        default=None,
        help="FK a f_raw_objects — da `hotelops promote`",
    )
    ap.add_argument(
        "--snapshot-date", default=None, help="YYYY-MM-DD, alternativa al raw-object-id"
    )
    ap.add_argument(
        "--societa",
        default=None,
        help="Ignorato — sempre ORTI. Accettato da `hotelops promote`.",
    )
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()

    logging.basicConfig(level=logging.INFO, format="%(levelname)s: %(message)s")
    n = ingest_file(
        args.file,
        raw_object_id=args.raw_object_id,
        snapshot_date=args.snapshot_date,
        dry_run=args.dry_run,
    )
    log.info("Totale: %d righe", n)


if __name__ == "__main__":
    main()
