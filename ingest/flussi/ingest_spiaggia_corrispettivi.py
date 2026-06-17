#!/usr/bin/env python3
"""Ingest Registro Corrispettivi Spiaggia INTUR (xlsx) → f_spiaggia_corrispettivi.

1 riga/giorno. Split aliquota: 22% = spiaggia, 10% = bar. Un file = un anno.
Lifecycle SNAPSHOT full-replace per anno (natural_key societa_id+anno).

Parser_module della source INTUR_CORRISPETTIVI_SPIAGGIA_INTUR_SNAPSHOT, invocato
da `hotelops promote`:
    python -m ingest.flussi.ingest_spiaggia_corrispettivi --file X.xlsx --raw-object-id Y
"""
from __future__ import annotations

import argparse
import logging
import re
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any, Optional

import openpyxl

from core.config import F_SPIAGGIA_CORRISPETTIVI
from core.schemas import SpiaggiaCorrispettivoRow, make_hash, validate_batch

log = logging.getLogger("ingest.spiaggia_corrispettivi")

SOCIETA = "INTUR"
BUSINESS_UNIT = "LIDO"
LOCATION = "LIDO"

_MONTHS = {
    "jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "jun": 6,
    "jul": 7, "aug": 8, "sep": 9, "oct": 10, "nov": 11, "dec": 12,
}
_DATE_RE = re.compile(r"^\s*(\d{1,2})-([A-Za-z]{3})\s*$")


def to_eur(v: Any) -> Optional[float]:
    """Cella corrispettivo → float. ' 1,083.00 '→1083.0, ' - '/''→None."""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", "")
    if s in ("", "-"):
        return None
    try:
        return float(s)
    except ValueError:
        return None


def parse_data(cell: Any, anno: int) -> Optional[date]:
    """'16-Jun' + anno → date. Datetime cell → date. Altro → None."""
    if isinstance(cell, datetime):
        return cell.date()
    if isinstance(cell, date):
        return cell
    if cell is None:
        return None
    m = _DATE_RE.match(str(cell))
    if not m:
        return None
    day = int(m.group(1))
    month = _MONTHS.get(m.group(2).lower())
    if not month:
        return None
    try:
        return date(anno, month, day)
    except ValueError:
        return None


def _find_anno(ws, file_name: str) -> int:
    """Anno dall'header sheet (cella accanto a 'ANNO:'); fallback dal filename."""
    for row in ws.iter_rows(min_row=1, max_row=6, values_only=True):
        for i, c in enumerate(row):
            if isinstance(c, str) and c.strip().upper().startswith("ANNO"):
                for nxt in row[i + 1:]:
                    n = to_eur(nxt)
                    if n and 2018 <= int(n) <= 2035:
                        return int(n)
    m = re.search(r"(20\d{2})", file_name)
    if m:
        return int(m.group(1))
    raise ValueError(f"anno non determinabile per sheet {ws.title} / {file_name}")


def _find_rt(ws) -> Optional[str]:
    for row in ws.iter_rows(min_row=1, max_row=6, values_only=True):
        for i, c in enumerate(row):
            if isinstance(c, str) and c.strip() == "RT":
                for nxt in row[i + 1:]:
                    if isinstance(nxt, str) and nxt.strip().startswith("RT"):
                        return nxt.strip()
    return None


def iter_day_rows(ws, anno: int) -> list[tuple]:
    """Righe-dato del mese → [(data, totale, spiaggia22, bar10)].

    Riga-dato = una cella matcha 'D-Mon'; i 3 numerici a seguire (saltando vuoti)
    sono totale, 22% (spiaggia), 10% (bar). Esclude Totale/riporto/riporta.
    """
    out: list[tuple] = []
    for row in ws.iter_rows(values_only=True):
        d = None
        idx = None
        for i, c in enumerate(row):
            d = parse_data(c, anno)
            if d is not None:
                idx = i
                break
        if d is None:
            continue
        nums = [to_eur(c) for c in row[idx + 1:]]
        nums = [n for n in nums if n is not None]
        if len(nums) < 3:
            continue
        totale, spiaggia, bar = nums[0], nums[1], nums[2]
        out.append((d, totale, spiaggia, bar))
    return out


def build_corrispettivo_rows(
    wb, file_name: str, raw_object_id: Optional[str], now: datetime
) -> list[SpiaggiaCorrispettivoRow]:
    out: list[SpiaggiaCorrispettivoRow] = []
    for ws in wb.worksheets:
        try:
            anno = _find_anno(ws, file_name)
        except ValueError:
            continue
        rt = _find_rt(ws)
        for d, totale, spiaggia, bar in iter_day_rows(ws, anno):
            # Quality gate per-riga: totale ~= spiaggia + bar (tolleranza 0.05)
            if abs((spiaggia + bar) - totale) > 0.05:
                log.warning("riga %s: totale %.2f != spiaggia+bar %.2f",
                            d, totale, spiaggia + bar)
            out.append(SpiaggiaCorrispettivoRow(
                societa_id=SOCIETA,
                business_unit_id=BUSINESS_UNIT,
                location_id=LOCATION,
                oggetto_id=None,
                funzione_id=None,
                data=d,
                anno=d.year,
                mese=d.month,
                corrispettivo_spiaggia=spiaggia,
                corrispettivo_bar=bar,
                corrispettivo_totale=totale,
                rt_matricola=rt,
                raw_object_id=raw_object_id,
                file_sorgente=file_name,
                hash_riga=make_hash("corrispettivi", SOCIETA, d.isoformat()),
                data_caricamento=now,
            ))
    return out


def ingest_file(
    path: Path, raw_object_id: Optional[str] = None, dry_run: bool = False
) -> dict[str, int]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    now = datetime.now(timezone.utc)
    rows = build_corrispettivo_rows(wb, path.name, raw_object_id, now)
    wb.close()
    if rows:
        validate_batch(
            [r.model_dump() for r in rows],
            SpiaggiaCorrispettivoRow,
            context=f"spiaggia_corrispettivi {path.name}",
        )
    if dry_run:
        log.info("[DRY-RUN] corrispettivi → %s : %d righe", F_SPIAGGIA_CORRISPETTIVI, len(rows))
        return {"corrispettivi": len(rows)}

    from core.bq.write import bq_write_validated
    bq_write_validated(
        F_SPIAGGIA_CORRISPETTIVI,
        rows,
        mode="snapshot",
        natural_key=["societa_id", "anno"],
    )
    log.info("OK corrispettivi → %s : %d righe (full-replace anno)",
             F_SPIAGGIA_CORRISPETTIVI, len(rows))
    return {"corrispettivi": len(rows)}


def main() -> None:
    ap = argparse.ArgumentParser(description="Ingest Registro Corrispettivi Spiaggia → BQ")
    ap.add_argument("--file", required=True, type=Path)
    ap.add_argument("--raw-object-id", default=None)
    ap.add_argument("--societa", default=None, help="Ignorato — sempre INTUR.")
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()
    logging.basicConfig(level=logging.INFO, format="%(levelname)s: %(message)s")
    counts = ingest_file(args.file, raw_object_id=args.raw_object_id, dry_run=args.dry_run)
    log.info("Totale: %s", counts)


if __name__ == "__main__":
    main()
