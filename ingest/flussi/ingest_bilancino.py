#!/usr/bin/env python3
"""
Bilancio di verifica (trial balance) ingestion pipeline.

Reads XLS exports from Esolver (Bilancio di verifica) into BigQuery f_bilancino.
Only ingests leaf-level accounts (Livello di imputazione = Si).

Usage:
    python -m ingest.flussi.ingest_bilancino \\
        --file /path/to/GENNAIO2026ESOLVER.xls \\
        --societa ORTI \\
        --mese 2026-01
    python -m ingest.flussi.ingest_bilancino \\
        --file /path/to/GENNAIO2026ESOLVER.xls \\
        --societa ORTI \\
        --mese 2026-01 \\
        --dry-run
"""

import argparse
import csv
import hashlib
import logging
import sys
from datetime import date
from pathlib import Path

try:
    import xlrd

    HAS_XLRD = True
except ImportError:
    HAS_XLRD = False

try:
    import openpyxl

    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

try:
    from google.cloud import bigquery

    HAS_BQ = True
except ImportError:
    HAS_BQ = False

from core.bq.client import get_client
from core.config import PROJECT

BQ_TABLE = f"{PROJECT}.hotelops.f_bilancino"

# Account prefix → business_unit_id
# 47.91.* = Hotel, 47.92.* = Residence, 47.93.* = CVM, 47.94.* = Spiaggia/Lido, 47.95.* = Affitti (HQ)
BU_RICAVI_MAP = {
    "47.91": "HOTEL",
    "47.92": "RESIDENCE",
    "47.93": "CVM",
    "47.94": "LIDO",
    "47.95": "HQ",
}

FACT_HEADER = [
    "hash_riga",
    "societa_id",
    "mese",
    "codice_conto",
    "descrizione",
    "tipo_conto",
    "sezione",
    "dare",
    "avere",
    "saldo",
    "business_unit_id",
    "categoria",
    "file_sorgente",
    "data_ingresso",
    "raw_object_id",
]

MESI_IT = {
    "GENNAIO": 1, "FEBBRAIO": 2, "MARZO": 3, "APRILE": 4,
    "MAGGIO": 5, "GIUGNO": 6, "LUGLIO": 7, "AGOSTO": 8,
    "SETTEMBRE": 9, "OTTOBRE": 10, "NOVEMBRE": 11, "DICEMBRE": 12,
}


def infer_mese_from_filename(filepath: Path) -> str | None:
    """Try to parse YYYY-MM from filenames like GENNAIO2026ESOLVER.xls."""
    import re

    name = filepath.stem.upper()
    # Try numeric pattern YYYY-MM or YYYYMM
    m = re.search(r"(\d{4})[-_]?(\d{2})", name)
    if m:
        return f"{m.group(1)}-{m.group(2)}"
    # Try Italian month name
    for mese_name, mese_num in MESI_IT.items():
        if mese_name in name:
            year_m = re.search(r"(\d{4})", name)
            if year_m:
                return f"{year_m.group(1)}-{mese_num:02d}"
    return None


def setup_logger() -> logging.Logger:
    log = logging.getLogger("ingest_bilancino")
    log.setLevel(logging.DEBUG)
    h = logging.StreamHandler(sys.stdout)
    h.setFormatter(logging.Formatter("%(levelname)s %(message)s"))
    log.addHandler(h)
    return log


def infer_business_unit(codice_conto: str, tipo_conto: str) -> str | None:
    """Derive business_unit_id from account code prefix for revenue accounts."""
    for prefix, bu in BU_RICAVI_MAP.items():
        if codice_conto.startswith(prefix):
            return bu
    return None


def infer_categoria(codice_conto: str, tipo_conto: str, sezione: str) -> str:
    """Classify account into a high-level category."""
    if tipo_conto == "SP":
        if sezione in ("Attività", "Attivita"):
            return "ATTIVO"
        if sezione in ("Passività", "Passivita"):
            return "PASSIVO"
        return "PATRIMONIALE"
    # CE accounts
    if codice_conto.startswith("47."):
        return "RICAVI"
    return "COSTI"


def make_hash(societa_id: str, mese: str, codice_conto: str) -> str:
    key = f"{societa_id}|{mese}|{codice_conto}"
    return hashlib.md5(key.encode()).hexdigest()


def _parse_tipo_sezione(raw: str) -> tuple[str, str]:
    """Parse 'Tipo conto/Sezione' column → (tipo_conto, sezione).

    Examples:
        ' Stato Patrimoniale: attivo'  → ('SP', 'Attività')
        ' Stato Patrimoniale: passivo' → ('SP', 'Passività')
        ' Conto Economico: costi'      → ('CE', 'Costi')
        ' Conto Economico: ricavi'     → ('CE', 'Ricavi')
    """
    s = raw.strip().lower()
    if "patrimoniale" in s:
        sezione = "Attività" if "attivo" in s else "Passività"
        return "SP", sezione
    if "economico" in s:
        sezione = "Ricavi" if "ricavi" in s else "Costi"
        return "CE", sezione
    return "CE", "Costi"


def _parse_xls(filepath: Path, logger: logging.Logger) -> list[dict]:
    """Parse old-format XLS (13 columns, Livello di imputazione = Si)."""
    if not HAS_XLRD:
        logger.error("xlrd not installed. Run: pip install xlrd")
        sys.exit(1)
    wb = xlrd.open_workbook(str(filepath))
    ws = wb.sheets()[0]
    logger.info(f"Sheet: {ws.name}, rows: {ws.nrows}")

    rows = []
    for r in range(1, ws.nrows):
        row = ws.row_values(r)
        livello = str(row[10]).strip()
        if livello != "Si":
            continue
        codice = str(row[0]).strip()
        dare = float(row[4]) if row[4] else 0.0
        avere = float(row[5]) if row[5] else 0.0
        saldo = float(row[7]) if row[7] else 0.0
        if dare == 0.0 and avere == 0.0 and saldo == 0.0:
            continue
        rows.append(
            {
                "codice_conto": codice,
                "descrizione": str(row[1]).strip(),
                "tipo_conto": str(row[11]).strip(),
                "sezione": str(row[12]).strip(),
                "dare": dare,
                "avere": avere,
                "saldo": saldo,
            }
        )
    return rows


def _parse_xlsx(filepath: Path, logger: logging.Logger) -> list[dict]:
    """Parse new-format XLSX (5 columns: Tipo conto/Sezione, Conto, Partitari, Descrizione, Importo).

    Leaf detection:
    - SP accounts: Partitari = 'S'
    - CE accounts: no 'S' marker, so a code is leaf if no other code starts with it + '.'
    """
    if not HAS_OPENPYXL:
        logger.error("openpyxl not installed. Run: pip install openpyxl")
        sys.exit(1)
    wb = openpyxl.load_workbook(str(filepath), read_only=True, data_only=True)
    ws = wb.active
    logger.info(f"Sheet: {ws.title}, rows: {ws.max_row}")

    # First pass: collect all entries
    entries = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            continue
        tipo_sez_raw = str(row[0] or "").strip()
        codice_raw = str(row[1] or "").strip()
        partitari = str(row[2] or "").strip()
        descrizione = str(row[3] or "").strip()
        importo = float(row[4]) if row[4] else 0.0

        if not codice_raw:
            continue

        tipo_conto, sezione = _parse_tipo_sezione(tipo_sez_raw)
        entries.append(
            {
                "codice_conto": codice_raw,
                "descrizione": descrizione,
                "tipo_conto": tipo_conto,
                "sezione": sezione,
                "importo": importo,
                "partitari": partitari,
            }
        )
    wb.close()

    # Build set of all codes for leaf detection
    all_codes = {e["codice_conto"] for e in entries}

    def is_leaf(entry: dict) -> bool:
        if entry["tipo_conto"] == "SP":
            return entry["partitari"] == "S"
        # CE: leaf if no other code starts with this code + "."
        code = entry["codice_conto"]
        return not any(c.startswith(code + ".") for c in all_codes if c != code)

    rows = []
    for e in entries:
        if not is_leaf(e):
            continue
        if e["importo"] == 0.0:
            continue
        rows.append(
            {
                "codice_conto": e["codice_conto"],
                "descrizione": e["descrizione"],
                "tipo_conto": e["tipo_conto"],
                "sezione": e["sezione"],
                "dare": 0.0,
                "avere": 0.0,
                "saldo": e["importo"],
            }
        )
    return rows


def parse_bilancino(
    filepath: Path, societa_id: str, mese: str, logger: logging.Logger
) -> list[dict]:
    """Parse Esolver Bilancio di verifica (XLS or XLSX) → list of fact dicts."""
    suffix = filepath.suffix.lower()
    if suffix == ".xlsx":
        parsed = _parse_xlsx(filepath, logger)
    else:
        parsed = _parse_xls(filepath, logger)

    today = date.today().isoformat()
    file_sorgente = filepath.name
    rows = []

    for entry in parsed:
        codice = entry["codice_conto"]
        business_unit_id = infer_business_unit(codice, entry["tipo_conto"])
        categoria = infer_categoria(codice, entry["tipo_conto"], entry["sezione"])
        hash_riga = make_hash(societa_id, mese, codice)

        rows.append(
            {
                "hash_riga": hash_riga,
                "societa_id": societa_id,
                "mese": mese,
                "codice_conto": codice,
                "descrizione": entry["descrizione"],
                "tipo_conto": entry["tipo_conto"],
                "sezione": entry["sezione"],
                "dare": entry["dare"],
                "avere": entry["avere"],
                "saldo": entry["saldo"],
                "business_unit_id": business_unit_id,
                "categoria": categoria,
                "file_sorgente": file_sorgente,
                "data_ingresso": today,
            }
        )

    logger.info(f"Leaf rows with activity: {len(rows)}")
    return rows


def load_hashes_bq(
    bq_client, mese: str, societa_id: str, logger: logging.Logger
) -> set:
    try:
        result = bq_client.query(
            f"SELECT hash_riga FROM `{BQ_TABLE}` WHERE mese = '{mese}' AND societa_id = '{societa_id}'"
        ).result()
        hashes = {row.hash_riga for row in result}
        logger.info(f"Hashes BQ esistenti ({societa_id} {mese}): {len(hashes)}")
        return hashes
    except Exception as e:
        logger.warning(f"Impossibile caricare hashes BQ: {e}")
        return set()


BQ_SCHEMA = (
    [
        bigquery.SchemaField("hash_riga", "STRING", mode="REQUIRED"),
        bigquery.SchemaField("societa_id", "STRING", mode="REQUIRED"),
        bigquery.SchemaField("mese", "STRING", mode="REQUIRED"),
        bigquery.SchemaField("codice_conto", "STRING", mode="REQUIRED"),
        bigquery.SchemaField("descrizione", "STRING"),
        bigquery.SchemaField("tipo_conto", "STRING"),
        bigquery.SchemaField("sezione", "STRING"),
        bigquery.SchemaField("dare", "FLOAT64"),
        bigquery.SchemaField("avere", "FLOAT64"),
        bigquery.SchemaField("saldo", "FLOAT64"),
        bigquery.SchemaField("business_unit_id", "STRING"),
        bigquery.SchemaField("categoria", "STRING"),
        bigquery.SchemaField("file_sorgente", "STRING"),
        bigquery.SchemaField("data_ingresso", "STRING"),
        bigquery.SchemaField("raw_object_id", "STRING"),
    ]
    if HAS_BQ
    else []
)


def write_to_bq(rows: list[dict], bq_client, logger: logging.Logger):
    if not rows:
        return
    job_config = bigquery.LoadJobConfig(
        schema=BQ_SCHEMA,
        write_disposition=bigquery.WriteDisposition.WRITE_APPEND,
    )
    job = bq_client.load_table_from_json(rows, BQ_TABLE, job_config=job_config)
    job.result()
    if job.errors:
        logger.error(f"BQ errors: {job.errors}")
    else:
        logger.info(f"  BQ: {len(rows)} righe inserite in {BQ_TABLE}")


def write_to_csv(rows: list[dict], csv_path: Path, logger: logging.Logger):
    if not rows:
        return
    csv_path.parent.mkdir(parents=True, exist_ok=True)
    existing = set()
    if csv_path.exists():
        with open(csv_path, newline="") as f:
            reader = csv.DictReader(f)
            existing = {r["hash_riga"] for r in reader}

    new_rows = [r for r in rows if r["hash_riga"] not in existing]
    if not new_rows:
        logger.info("  CSV: nessuna nuova riga (già presenti)")
        return

    write_header = not csv_path.exists()
    with open(csv_path, "a", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=FACT_HEADER)
        if write_header:
            writer.writeheader()
        writer.writerows(new_rows)
    logger.info(f"  CSV: {len(new_rows)} righe aggiunte → {csv_path}")


def main():
    parser = argparse.ArgumentParser(
        description="Ingest Esolver Bilancio di verifica → f_bilancino"
    )
    parser.add_argument(
        "--file", required=True, help="Path to XLS file (Bilancio di verifica)"
    )
    parser.add_argument(
        "--societa",
        required=True,
        choices=["INTUR", "ORTI"],
        help="Società (INTUR o ORTI)",
    )
    parser.add_argument(
        "--mese",
        default=None,
        help="Mese contabile (YYYY-MM, es. 2026-01). Auto-inferred from filename if omitted.",
    )
    parser.add_argument(
        "--datahub", help="Path to datahub root (for CSV backup). Optional."
    )
    parser.add_argument("--dry-run", action="store_true", help="Parse only, no write")
    parser.add_argument(
        "--raw-object-id",
        default=None,
        help="FK to f_raw_objects.raw_object_id (stamped on every row — used by promotion path)",
    )
    args = parser.parse_args()

    logger = setup_logger()
    filepath = Path(args.file)
    if not filepath.exists():
        logger.error(f"File non trovato: {filepath}")
        sys.exit(1)

    mese = args.mese or infer_mese_from_filename(filepath)
    if not mese:
        logger.error("--mese richiesto (YYYY-MM) — impossibile inferirlo dal filename")
        sys.exit(1)

    rows = parse_bilancino(filepath, args.societa, mese, logger)

    for r in rows:
        r["raw_object_id"] = args.raw_object_id

    # Print summary
    by_cat = {}
    for r in rows:
        cat = r["categoria"]
        by_cat[cat] = by_cat.get(cat, 0) + 1
    for cat, n in sorted(by_cat.items()):
        logger.info(f"  {cat}: {n} conti")

    # Print revenue breakdown
    ricavi = [r for r in rows if r["categoria"] == "RICAVI"]
    if ricavi:
        logger.info("  === Ricavi ===")
        for r in sorted(ricavi, key=lambda x: x["codice_conto"]):
            bu = r["business_unit_id"] or "?"
            logger.info(
                f"    {r['codice_conto']} [{bu}] {r['descrizione']}: saldo={r['saldo']:,.2f}"
            )

    if args.dry_run:
        logger.info("DRY RUN — nessuna scrittura")
        return

    if not HAS_BQ:
        logger.error(
            "google-cloud-bigquery non installato. Run: pip install google-cloud-bigquery"
        )
        sys.exit(1)

    bq_client = get_client()
    existing_hashes = load_hashes_bq(bq_client, mese, args.societa, logger)
    new_rows = [r for r in rows if r["hash_riga"] not in existing_hashes]
    logger.info(
        f"Nuove righe da inserire: {len(new_rows)} (già presenti: {len(rows) - len(new_rows)})"
    )

    write_to_bq(new_rows, bq_client, logger)

    if args.datahub:
        csv_path = Path(args.datahub) / "fatti" / "f_bilancino.csv"
        write_to_csv(new_rows, csv_path, logger)

    logger.info("DONE")


if __name__ == "__main__":
    main()
