#!/usr/bin/env python3
"""Ingest menu engineering RistoCube → f_menu_engineering.

Una riga per (snapshot_date × sala × piatto). Il file è un CUMULATO senza data
nel contenuto: snapshot_date fa fede dall'intake del raw object (pattern OTB,
cfr. ingest_andamento_prenotazioni._snapshot_date_from_raw). Le foto si
accumulano — mai DELETE della foto precedente.

È il parser_module di POWERBI_MENUENGINEERING_ORTI_SNAPSHOT, invocato da
`hotelops promote` come:
    python -m ingest.flussi.ingest_menu_engineering --file X --societa ORTI --raw-object-id Y

Usage:
    python -m ingest.flussi.ingest_menu_engineering --file <xlsx> --dry-run
"""

from __future__ import annotations

import argparse
import logging
from datetime import date, datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

from core.config import F_MENU_ENGINEERING
from core.schemas import MenuEngineeringRow, make_hash, validate_batch

log = logging.getLogger("ingest.menu_engineering")

SOCIETA_ID = "ORTI"
BUSINESS_UNIT_ID = "HOTEL"


def _s(v) -> str | None:
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def _f(v) -> float | None:
    try:
        return float(v) if v is not None else None
    except (TypeError, ValueError):
        return None


def parse_xlsx(
    path: Path, snapshot_date: date, raw_object_id: str | None
) -> list[dict]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    now = datetime.now(timezone.utc)
    out: list[dict] = []
    for r in rows[1:]:
        if not r or len(r) < 14:
            continue
        sala, piatto = _s(r[2]), _s(r[3])
        if piatto is None or sala in (None, "Total"):
            continue  # subtotali, riga Total finale, footer "Applied filters"
        out.append(
            {
                "hash_riga": make_hash(str(snapshot_date), sala, piatto),
                "societa_id": SOCIETA_ID,
                "business_unit_id": BUSINESS_UNIT_ID,
                "snapshot_date": snapshot_date,
                "sala": sala,
                "piatto": piatto,
                "descrizione": _s(r[4]),
                "tipo": _s(r[1]),
                "m_class": _s(r[0]),
                "prezzo_unitario": _f(r[5]),
                "costo_unitario": _f(r[6]),
                "quantita": _f(r[7]),
                "incidenza_pct": _f(r[8]),
                "costo_totale": _f(r[9]),
                "listino": _f(r[10]),
                "vendita": _f(r[11]),
                "importo_addebitato": _f(r[12]),
                "importo_fatturato": _f(r[13]),
                "file_sorgente": path.name,
                "raw_object_id": raw_object_id,
                "data_caricamento": now,
            }
        )
    return out


def _snapshot_date_from_raw(raw_object_id: str) -> date:
    """La data della foto non è nel file: fa fede l'intake del raw object."""
    from google.cloud import bigquery

    from core.bq.client import get_client
    from core.config import PROJECT

    client = get_client()
    sql = (
        f"SELECT DATE(intake_at) d FROM `{PROJECT}.hotelops.f_raw_objects` "
        "WHERE raw_object_id = @rid"
    )
    job = client.query(
        sql,
        job_config=bigquery.QueryJobConfig(
            query_parameters=[
                bigquery.ScalarQueryParameter("rid", "STRING", raw_object_id)
            ]
        ),
    )
    rows = list(job.result())
    if not rows:
        raise SystemExit(f"raw_object_id {raw_object_id} non trovato in f_raw_objects")
    return rows[0][0]


def ingest_file(path: Path, raw_object_id: str | None, dry_run: bool) -> int:
    snapshot_date = (
        _snapshot_date_from_raw(raw_object_id) if raw_object_id else date.today()
    )
    rows = parse_xlsx(path, snapshot_date=snapshot_date, raw_object_id=raw_object_id)
    validate_batch(rows, MenuEngineeringRow, context=f"menu_engineering {path.name}")
    if dry_run:
        log.info(
            "[DRY-RUN] %s : %d righe (snapshot %s)", path.name, len(rows), snapshot_date
        )
        return len(rows)

    from core.bq.dedup import filter_new_rows_by_hash
    from core.bq.write import bq_write_validated

    new_rows = filter_new_rows_by_hash(F_MENU_ENGINEERING, rows, "hash_riga")
    if not new_rows:
        log.info("Foto già presente — niente da scrivere (%s)", path.name)
        return 0
    bq_write_validated(
        F_MENU_ENGINEERING, [MenuEngineeringRow(**r) for r in new_rows], mode="append"
    )
    log.info("OK %s : %d righe (snapshot %s)", path.name, len(new_rows), snapshot_date)
    return len(new_rows)


def main() -> None:
    ap = argparse.ArgumentParser(
        description="Ingest menu engineering → f_menu_engineering"
    )
    ap.add_argument("--file", required=True, type=Path)
    ap.add_argument("--raw-object-id", default=None)
    ap.add_argument(
        "--societa", default=None, help="Accettato da promote — sempre ORTI"
    )
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()
    logging.basicConfig(level=logging.INFO, format="%(levelname)s: %(message)s")
    if args.societa and args.societa != SOCIETA_ID:
        raise SystemExit(
            f"societa {args.societa} != {SOCIETA_ID}: file menu engineering è ORTI"
        )

    from core.pipeline_run import PipelineRun

    with PipelineRun(
        "ingest_menu_engineering", societa_id=SOCIETA_ID, file_sorgente=args.file.name
    ):
        n = ingest_file(args.file, args.raw_object_id, args.dry_run)
        log.info("Totale: %d righe", n)


if __name__ == "__main__":
    main()
