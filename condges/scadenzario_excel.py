"""Scadenzario → PF Excel Bridge.

Reads the Esolver 'Situazione sintetica scadenze' export, maps suppliers
to Piano Finanziario voci via d_fornitori, and generates an annotated
Excel for Rosa to update her PF.
"""

from __future__ import annotations

import csv
import re
from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill


def parse_sintetica_scadenze(filepath: Path) -> tuple[list[dict], list[int]]:
    """Parse Esolver 'Situazione sintetica scadenze' Excel.

    Returns:
        (suppliers, bucket_months) where suppliers is list of dicts
        {codice_fornitore, nome, totale, scaduto, buckets: {month_int: amount}}
        and bucket_months is the ordered list of month ints from headers.
    """
    wb = openpyxl.load_workbook(str(filepath), data_only=True)
    ws = wb.active

    bucket_months = {}
    for col in range(4, ws.max_column + 1):
        header = ws.cell(row=1, column=col).value
        if not header or "scadenza" not in str(header).lower():
            continue
        if "oltre" in str(header).lower():
            continue  # skip catch-all "oltre" bucket
        m = re.search(r"(\d{2})/(\d{2})/(\d{4})", str(header))
        if m:
            bucket_months[col] = int(m.group(2))

    results = []
    for row_idx in range(2, ws.max_row + 1):
        cell_a = ws.cell(row=row_idx, column=1).value
        if not cell_a:
            continue
        cell_str = str(cell_a).strip()
        m = re.match(r"^(\d+)\s+(.+)$", cell_str)
        if not m:
            continue
        codice = int(m.group(1))
        nome = m.group(2).strip()
        totale = ws.cell(row=row_idx, column=2).value or 0
        scaduto = ws.cell(row=row_idx, column=3).value or 0
        buckets = {}
        for col, month in bucket_months.items():
            val = ws.cell(row=row_idx, column=col).value
            if val:
                buckets[month] = float(val)
        results.append(
            {
                "codice_fornitore": codice,
                "nome": nome,
                "totale": float(totale),
                "scaduto": float(scaduto),
                "buckets": buckets,
            }
        )
    return results, sorted(bucket_months.values())


# ── Mapper ─────────────────────────────────────────────────────────────────

FORNITORI_CSV = (
    Path(__file__).parent.parent / "core" / "bq" / "dimensioni" / "d_fornitori.csv"
)


def load_fornitori_map(csv_path: Path = FORNITORI_CSV) -> dict[int, str]:
    """Load d_fornitori CSV, return {codice_fornitore: voce_id}."""
    result = {}
    with open(csv_path, newline="", encoding="utf-8") as f:
        for row in csv.DictReader(f):
            result[int(row["codice_fornitore"])] = row["voce_id"]
    return result


def map_to_voci(
    partite: list[dict], fornitori_map: dict[int, str]
) -> tuple[dict[str, list[dict]], list[dict]]:
    """Map suppliers to PF voci.

    Returns:
        (mapped, unmapped) where mapped = {voce_id: [supplier_dicts]},
        unmapped = [supplier_dicts without voce match]
    """
    mapped: dict[str, list[dict]] = {}
    unmapped: list[dict] = []

    for p in partite:
        voce = fornitori_map.get(p["codice_fornitore"])
        if voce:
            mapped.setdefault(voce, []).append(p)
        else:
            unmapped.append(p)

    return mapped, unmapped


# ── Cascade ───────────────────────────────────────────────────────────────


def cascade_scaduto(
    suppliers: list[dict], current_month: int | None = None
) -> list[dict]:
    """Move overdue (scaduto) amounts into the current month bucket.

    Mutates suppliers in-place and returns them for chaining.
    """
    if current_month is None:
        current_month = date.today().month

    for s in suppliers:
        if s["scaduto"] and s["scaduto"] != 0:
            s["buckets"][current_month] = (
                s["buckets"].get(current_month, 0) + s["scaduto"]
            )
            s["scaduto"] = 0
    return suppliers


# ── Voce labels ────────────────────────────────────────────────────────────

VOCE_LABELS = {
    "USCITE_MATERIE_PRIME": "Materie Prime",
    "USCITE_UTENZE": "Utenze",
    "USCITE_SALARI": "Salari e Stipendi",
    "USCITE_TASSE": "Tasse e Imposte",
    "USCITE_COMMISSIONI": "Commissioni",
    "USCITE_MUTUI": "Mutui e Finanziamenti",
    "USCITE_CONSULENZE": "Consulenze",
    "USCITE_CANONE_PASSIVO": "Godimento Beni di Terzi",
    "USCITE_VARIE_EXT": "Varie ed Eventuali",
    "USCITE_SERVIZI_PRODUZIONE": "Canoni e servizi",
    "USCITE_MARKETING": "Marketing",
}

MESI_NOMI = [
    "Gen",
    "Feb",
    "Mar",
    "Apr",
    "Mag",
    "Giu",
    "Lug",
    "Ago",
    "Set",
    "Ott",
    "Nov",
    "Dic",
]

RED_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
BOLD = Font(bold=True)
NUM_FMT = "#,##0"


# ── Excel Generator ───────────────────────────────────────────────────────


def generate_excel(
    mapped: dict[str, list[dict]],
    forecasts: dict[str, dict[int, float]] | None,
    unmapped: list[dict],
    output_path: Path,
    bucket_months: list[int] | None = None,
) -> Path:
    """Generate the Excel ponte."""
    if bucket_months is None:
        all_months = set()
        for suppliers in mapped.values():
            for s in suppliers:
                all_months.update(s["buckets"].keys())
        bucket_months = sorted(all_months) if all_months else []

    wb = openpyxl.Workbook()

    # ── Riepilogo sheet ──
    ws = wb.active
    ws.title = "Riepilogo"
    headers = (
        ["Voce PF", "Scaduto"] + [MESI_NOMI[m - 1] for m in bucket_months] + ["Totale"]
    )
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = BOLD

    row_idx = 2
    grand_scaduto = 0
    grand_buckets = {m: 0 for m in bucket_months}
    grand_total = 0

    for voce_id in sorted(mapped.keys()):
        suppliers = mapped[voce_id]
        label = VOCE_LABELS.get(voce_id, voce_id)
        scaduto = sum(s["scaduto"] for s in suppliers)
        totale = sum(s["totale"] for s in suppliers)
        month_sums = {
            m: sum(s["buckets"].get(m, 0) for s in suppliers) for m in bucket_months
        }

        ws.cell(row=row_idx, column=1, value=label)
        cell_s = ws.cell(row=row_idx, column=2, value=round(scaduto))
        cell_s.number_format = NUM_FMT
        if scaduto < 0:
            cell_s.fill = RED_FILL
        for i, m in enumerate(bucket_months):
            c = ws.cell(
                row=row_idx,
                column=3 + i,
                value=round(month_sums[m]) if month_sums[m] else None,
            )
            c.number_format = NUM_FMT
        ws.cell(
            row=row_idx, column=3 + len(bucket_months), value=round(totale)
        ).number_format = NUM_FMT

        grand_scaduto += scaduto
        for m in bucket_months:
            grand_buckets[m] += month_sums[m]
        grand_total += totale
        row_idx += 1

    # Totals row
    ws.cell(row=row_idx, column=1, value="Totale").font = BOLD
    ws.cell(row=row_idx, column=2, value=round(grand_scaduto)).font = BOLD
    ws.cell(row=row_idx, column=2).number_format = NUM_FMT
    for i, m in enumerate(bucket_months):
        c = ws.cell(
            row=row_idx,
            column=3 + i,
            value=round(grand_buckets[m]) if grand_buckets[m] else None,
        )
        c.font = BOLD
        c.number_format = NUM_FMT
    ws.cell(
        row=row_idx, column=3 + len(bucket_months), value=round(grand_total)
    ).font = BOLD
    ws.cell(row=row_idx, column=3 + len(bucket_months)).number_format = NUM_FMT

    # ── Per-voce sheets ──
    for voce_id in sorted(mapped.keys()):
        suppliers = mapped[voce_id]
        label = VOCE_LABELS.get(voce_id, voce_id)
        ws_v = wb.create_sheet(title=label[:31])

        headers_v = (
            ["Fornitore", "Cod.", "Scaduto"]
            + [MESI_NOMI[m - 1] for m in bucket_months]
            + ["Totale"]
        )
        for c, h in enumerate(headers_v, 1):
            ws_v.cell(row=1, column=c, value=h).font = BOLD

        r = 2
        for s in sorted(suppliers, key=lambda x: x["totale"]):
            ws_v.cell(row=r, column=1, value=s["nome"])
            ws_v.cell(row=r, column=2, value=s["codice_fornitore"])
            cell_s = ws_v.cell(row=r, column=3, value=round(s["scaduto"]))
            cell_s.number_format = NUM_FMT
            if s["scaduto"] < 0:
                cell_s.fill = RED_FILL
            for i, m in enumerate(bucket_months):
                val = s["buckets"].get(m)
                c = ws_v.cell(row=r, column=4 + i, value=round(val) if val else None)
                c.number_format = NUM_FMT
            ws_v.cell(
                row=r, column=4 + len(bucket_months), value=round(s["totale"])
            ).number_format = NUM_FMT
            r += 1

        # Totals row
        r += 1
        ws_v.cell(row=r, column=1, value="Totale scadenzario").font = BOLD
        ws_v.cell(
            row=r, column=3, value=round(sum(s["scaduto"] for s in suppliers))
        ).font = BOLD
        ws_v.cell(row=r, column=3).number_format = NUM_FMT
        for i, m in enumerate(bucket_months):
            val = sum(s["buckets"].get(m, 0) for s in suppliers)
            c = ws_v.cell(row=r, column=4 + i, value=round(val) if val else None)
            c.font = BOLD
            c.number_format = NUM_FMT
        ws_v.cell(
            row=r,
            column=4 + len(bucket_months),
            value=round(sum(s["totale"] for s in suppliers)),
        ).font = BOLD

        # Gap analysis rows (if forecasts provided)
        if forecasts and voce_id in forecasts:
            r += 1
            ws_v.cell(row=r, column=1, value="Rosa prevede").font = BOLD
            for i, m in enumerate(bucket_months):
                val = forecasts[voce_id].get(m)
                if val:
                    ws_v.cell(
                        row=r, column=4 + i, value=round(val)
                    ).number_format = NUM_FMT

            r += 1
            ws_v.cell(row=r, column=1, value="Gap (non fatturato)").font = BOLD
            for i, m in enumerate(bucket_months):
                forecast_val = forecasts[voce_id].get(m, 0)
                scad_val = sum(s["buckets"].get(m, 0) for s in suppliers)
                if forecast_val:
                    gap = forecast_val - scad_val
                    ws_v.cell(
                        row=r, column=4 + i, value=round(gap)
                    ).number_format = NUM_FMT

    # ── DA VERIFICARE sheet ──
    if unmapped:
        ws_u = wb.create_sheet(title="DA VERIFICARE")
        for c, h in enumerate(["Fornitore", "Cod.", "Totale", "Scaduto"], 1):
            ws_u.cell(row=1, column=c, value=h).font = BOLD
        for r, u in enumerate(unmapped, 2):
            ws_u.cell(row=r, column=1, value=u["nome"])
            ws_u.cell(row=r, column=2, value=u["codice_fornitore"])
            ws_u.cell(row=r, column=3, value=round(u["totale"])).number_format = NUM_FMT
            ws_u.cell(
                row=r, column=4, value=round(u["scaduto"])
            ).number_format = NUM_FMT

    wb.save(output_path)
    return output_path


# ── PF Excel Forecast Parser ──────────────────────────────────────────────

PF_LABEL_TO_VOCE = {
    "salari e stipendi": "USCITE_SALARI",
    "utenze": "USCITE_UTENZE",
    "materie prime/consumo": "USCITE_MATERIE_PRIME",
    "materie prime e consumo": "USCITE_MATERIE_PRIME",
    "tasse e imposte": "USCITE_TASSE",
    "commissioni portali": "USCITE_COMMISSIONI",
    "mutui e finaziamenti": "USCITE_MUTUI",
    "mutui e finanziamenti": "USCITE_MUTUI",
    "consulenze": "USCITE_CONSULENZE",
    "godimento beni di terzi": "USCITE_CANONE_PASSIVO",
    "varie ed eventuali": "USCITE_VARIE_EXT",
    "canoni e servizi": "USCITE_SERVIZI_PRODUZIONE",
}

MONTH_NAMES = {
    "gennaio": 1,
    "febbraio": 2,
    "marzo": 3,
    "aprile": 4,
    "maggio": 5,
    "giugno": 6,
    "luglio": 7,
    "agosto": 8,
    "settembre": 9,
    "ottobre": 10,
    "novembre": 11,
    "dicembre": 12,
}


def load_pf_forecasts(filepath: Path) -> dict[str, dict[int, float]]:
    """Parse Rosa's PF Excel, extract uscite forecasts per voce per month.

    Returns: {voce_id: {month_int: amount}}
    """
    wb = openpyxl.load_workbook(str(filepath), data_only=True)
    ws = wb["Piano Finanziario"]

    # Row 2 has actual calendar month names (APRILE, MAGGIO, etc.)
    # Row 3 has fiscal-year counter months — don't use it
    col_to_month: dict[int, int] = {}
    for col in range(3, ws.max_column + 1):
        val = ws.cell(row=2, column=col).value
        if val:
            month = MONTH_NAMES.get(str(val).strip().lower())
            if month:
                col_to_month[col] = month

    result: dict[str, dict[int, float]] = {}
    for row_idx in range(14, 28):
        label_raw = ws.cell(row=row_idx, column=1).value
        if not label_raw:
            continue
        label = str(label_raw).strip().lower()
        voce_id = PF_LABEL_TO_VOCE.get(label)
        if not voce_id:
            continue

        months: dict[int, float] = {}
        for col, month in col_to_month.items():
            val = ws.cell(row=row_idx, column=col).value
            if val and isinstance(val, (int, float)) and val != 0:
                months[month] = float(val)
        if months:
            result[voce_id] = months

    return result


# ── BQ Fallback ───────────────────────────────────────────────────────────


def load_from_bq(bq_client, societa: str) -> list[dict]:
    """Fallback: load from f_partite_aperte_fornitori, aggregate like sintetica."""
    from google.cloud import bigquery as bq_mod

    from core.config import F_PARTITE_APERTE_FORNITORI

    sql = f"""
    SELECT
        codice_fornitore,
        nome_fornitore AS nome,
        ROUND(SUM(importo_residuo), 2) AS totale,
        ROUND(SUM(CASE WHEN data_scadenza < CURRENT_DATE() THEN importo_residuo ELSE 0 END), 2) AS scaduto,
        EXTRACT(MONTH FROM data_scadenza) AS mese,
        ROUND(SUM(CASE WHEN data_scadenza >= CURRENT_DATE() THEN importo_residuo ELSE 0 END), 2) AS futuro
    FROM `{F_PARTITE_APERTE_FORNITORI}`
    WHERE societa_id = @societa
    GROUP BY codice_fornitore, nome_fornitore, mese
    ORDER BY codice_fornitore, mese
    """
    job_config = bq_mod.QueryJobConfig(
        query_parameters=[bq_mod.ScalarQueryParameter("societa", "STRING", societa)]
    )
    rows_raw = list(bq_client.query(sql, job_config=job_config))

    by_supplier: dict[int, dict] = {}
    for r in rows_raw:
        code = r.codice_fornitore
        if code not in by_supplier:
            by_supplier[code] = {
                "codice_fornitore": code,
                "nome": r.nome,
                "totale": 0,
                "scaduto": 0,
                "buckets": {},
            }
        by_supplier[code]["totale"] += float(r.totale)
        by_supplier[code]["scaduto"] += float(r.scaduto)
        if r.futuro and r.futuro != 0:
            by_supplier[code]["buckets"][r.mese] = float(r.futuro)

    return list(by_supplier.values())


# ── Runner ─────────────────────────────────────────────────────────────────


# ── Voce → PF Sheet mapping ──────────────────────────────────────────────

VOCE_TO_SHEET = {
    "USCITE_MATERIE_PRIME": "Materie Prime-Consumo ",
    "USCITE_UTENZE": "Utenze",
    "USCITE_SALARI": "Salari e Stipendi",
    "USCITE_TASSE": "Tasse e Imposte",
    "USCITE_COMMISSIONI": "Commisisoni Portali",
    "USCITE_MUTUI": "Mutui e Finaziamenti",
    "USCITE_CONSULENZE": "Consulenze",
    "USCITE_CANONE_PASSIVO": "Godimento Beni di Terzi",
    "USCITE_VARIE_EXT": " Varie ed Eventuali",
    "USCITE_SERVIZI_PRODUZIONE": "Canoni e servizi",
}

MONTH_NAMES_IT = {
    "GENNAIO": 1,
    "FEBBRAIO": 2,
    "MARZO": 3,
    "APRILE": 4,
    "MAGGIO": 5,
    "GIUGNO": 6,
    "LUGLIO": 7,
    "AGOSTO": 8,
    "SETTEMBRE": 9,
    "OTTOBRE": 10,
    "NOVEMBRE": 11,
    "DICEMBRE": 12,
}


def load_fornitori_map_full(csv_path: Path = FORNITORI_CSV) -> dict[int, dict]:
    """Load d_fornitori CSV, return {codice_fornitore: {voce_id, nome_pf}}."""
    result = {}
    with open(csv_path, newline="", encoding="utf-8") as f:
        for row in csv.DictReader(f):
            result[int(row["codice_fornitore"])] = {
                "voce_id": row["voce_id"],
                "nome_pf": row.get("nome_pf", "").strip(),
            }
    return result


def _build_month_col_map(ws) -> dict[int, int]:
    """Scan row 2 of a detail sheet, return {calendar_month: column}."""
    col_map = {}
    for col in range(1, ws.max_column + 1):
        val = ws.cell(row=2, column=col).value
        if val and str(val).strip().upper() in MONTH_NAMES_IT:
            month = MONTH_NAMES_IT[str(val).strip().upper()]
            # Take the rightmost occurrence (2026 for duplicated Sep-Dec)
            if month not in col_map or col > col_map[month]:
                col_map[month] = col
    return col_map


def _find_supplier_row(ws, nome_pf: str, max_row: int = 200) -> int | None:
    """Find the row in a detail sheet where column B matches nome_pf."""
    target = nome_pf.strip().lower()
    for r in range(3, min(ws.max_row + 1, max_row)):
        val = ws.cell(row=r, column=2).value
        if val and str(val).strip().lower() == target:
            return r
    return None


def write_back_to_pf(
    pf_path: Path,
    mapped: dict[str, list[dict]],
    fornitori_map_full: dict[int, dict],
    output_path: Path | None = None,
) -> tuple[Path, dict]:
    """Write cascaded scadenzario amounts into Rosa's PF Excel.

    Returns (output_path, summary) where summary tracks what was written.
    """
    wb = openpyxl.load_workbook(str(pf_path))
    out = output_path or pf_path

    summary: dict[str, list[dict]] = {}

    for voce_id, suppliers in mapped.items():
        sheet_name = VOCE_TO_SHEET.get(voce_id)
        if not sheet_name or sheet_name not in wb.sheetnames:
            continue

        ws = wb[sheet_name]
        month_col = _build_month_col_map(ws)
        voce_label = VOCE_LABELS.get(voce_id, voce_id)
        written = []

        for s in suppliers:
            info = fornitori_map_full.get(s["codice_fornitore"], {})
            nome_pf = info.get("nome_pf", "")
            if not nome_pf:
                continue

            row = _find_supplier_row(ws, nome_pf)
            if row is None:
                continue

            months_written = {}
            for month, amount in s["buckets"].items():
                col = month_col.get(month)
                if col and amount:
                    ws.cell(row=row, column=col, value=round(amount, 2))
                    months_written[month] = round(amount, 2)

            if months_written:
                written.append({"nome": nome_pf, "months": months_written})

        if written:
            summary[voce_label] = written

    wb.save(out)
    return out, summary


# ── Runner ─────────────────────────────────────────────────────────────────


def run(
    file: Path | None = None,
    pf: Path | None = None,
    societa: str = "ORTI",
    output: Path | None = None,
) -> Path:
    """Run the full pipeline: parse → map → generate."""
    if file:
        partite, bucket_months = parse_sintetica_scadenze(file)
    else:
        from core.bq.client import get_client

        bq = get_client()
        partite = load_from_bq(bq, societa)
        bucket_months = None

    fornitori_map = load_fornitori_map()
    mapped, unmapped = map_to_voci(partite, fornitori_map)

    forecasts = load_pf_forecasts(pf) if pf else None

    out_dir = output or Path(".")
    out_path = out_dir / f"Scadenzario_PF_{societa}_{date.today().isoformat()}.xlsx"

    generate_excel(mapped, forecasts, unmapped, out_path, bucket_months)
    return out_path
