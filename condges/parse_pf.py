"""Parse Rosa's Piano Finanziario Excel files.

Supports both ORTI and INTUR variants. Returns a PFData object with
structured data ready for BigQuery ingestion.

Usage:
    from condges.parse_pf import parse_pf

    with open("PF_ORTI.xlsx", "rb") as f:
        pf = parse_pf(f)
    print(pf.societa, pf.anno, pf.saldo_totale)
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import date
from io import BytesIO
from typing import BinaryIO

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

# ---------------------------------------------------------------------------
# Data structures
# ---------------------------------------------------------------------------


@dataclass
class VoceRow:
    voce_id: str
    excel_label: str
    importi: dict[int, float]  # {1..12: amount}


@dataclass
class PFData:
    societa: str
    anno: int
    data_saldo: date | None
    saldi_banca: dict[str, float]  # {banca_label: saldo}
    saldo_totale: float
    voci: dict[str, VoceRow]  # {voce_id: VoceRow}
    warnings: list[str] = field(default_factory=list)


# ---------------------------------------------------------------------------
# Label → voce_id mapping (normalised lowercase, stripped)
# ---------------------------------------------------------------------------

# Exact-match mapping (after lowercasing and stripping)
_LABEL_MAP: dict[str, str] = {
    # ORTI entrate
    "entrate hotel": "ENTRATE_HOTEL",
    "entrate residence": "ENTRATE_RESIDENCE",
    "entrate cvm": "ENTRATE_CVM",
    "entrate supermercato": "ENTRATE_SUPERMERCATO",
    "rientro sospesi": "ENTRATE_RIENTRO_SOSPESI",
    "caparre intur": "ENTRATE_CAPARRE",
    # INTUR entrate
    "fitto hotel": "ENTRATE_AFFITTI_INTUR",
    "fitto ar": "ENTRATE_AFFITTI_INTUR",
    "ribaltamento costi a orti": "ENTRATE_RIENTRO_SOSPESI",
    "entrate farmacia": "ENTRATE_AFFITTI_MINORI",
    "entrate spiaggia": "ENTRATE_SPIAGGIA",
    # Shared uscite
    "salari e stipendi": "USCITE_SALARI",
    "utenze": "USCITE_UTENZE",
    "materie prime/consumo": "USCITE_MATERIE_PRIME",
    "materie prime / consumo": "USCITE_MATERIE_PRIME",
    "tasse e imposte": "USCITE_TASSE",
    "commissioni portali": "USCITE_COMMISSIONI",
    "mutui e finaziamenti": "USCITE_MUTUI",       # Rosa's typo
    "mutui e finanziamenti": "USCITE_MUTUI",
    "consulenze": "USCITE_CONSULENZE",
    "godimento beni di terzi": "USCITE_GODIMENTO_BENI",
    "godimento benidi terzi": "USCITE_GODIMENTO_BENI",  # INTUR typo
    "varie ed eventuali": "USCITE_VARIE",
    "canoni e servizi": "USCITE_CANONI",
    "deposito fitto": "USCITE_DEPOSITO_FITTO",
    "caparre da girocantare aorti": "ENTRATE_CAPARRE_INTUR",
}

# Labels that are structural rows, not data voci — skip silently
_SKIP_PATTERNS: list[re.Pattern] = [
    re.compile(r"totale (entrate|uscite|banche)", re.I),
    re.compile(r"^={3,}$"),
    re.compile(r"saldo mese preced", re.I),
    re.compile(r"cash flow", re.I),
    re.compile(r"^saldo\b", re.I),            # e.g. "Saldo MPS" rows
    re.compile(r"^totale\b", re.I),
]


def _normalise(label: str) -> str:
    return label.strip().lower()


def _is_skip(label: str) -> bool:
    """Return True if this row should be silently skipped (structural row)."""
    norm = label.strip()
    for pat in _SKIP_PATTERNS:
        if pat.search(norm):
            return True
    return False


# ---------------------------------------------------------------------------
# Month detection
# ---------------------------------------------------------------------------

_MESI_IT = [
    "GENNAIO", "FEBBRAIO", "MARZO", "APRILE", "MAGGIO", "GIUGNO",
    "LUGLIO", "AGOSTO", "SETTEMBRE", "OTTOBRE", "NOVEMBRE", "DICEMBRE",
]
_MESE_INDEX: dict[str, int] = {m: i + 1 for i, m in enumerate(_MESI_IT)}


def _find_month_columns(ws: Worksheet, header_row: int) -> dict[int, int]:
    """Return {month_number: column_index} by scanning a row for Italian month names."""
    mapping: dict[int, int] = {}
    for cell in ws[header_row]:
        val = cell.value
        if isinstance(val, str):
            upper = val.strip().upper()
            if upper in _MESE_INDEX:
                mapping[_MESE_INDEX[upper]] = cell.column
    return mapping


# ---------------------------------------------------------------------------
# Anno detection
# ---------------------------------------------------------------------------

def _find_anno(ws: Worksheet) -> int:
    """Scan first 5 rows looking for an integer in the 2020-2030 range."""
    for row in ws.iter_rows(min_row=1, max_row=5, values_only=True):
        for cell_val in row:
            if isinstance(cell_val, (int, float)):
                v = int(cell_val)
                if 2020 <= v <= 2030:
                    return v
    return date.today().year


# ---------------------------------------------------------------------------
# Societa detection
# ---------------------------------------------------------------------------

def _find_societa(ws: Worksheet) -> str:
    """Look in A1 and A2 for ORTI / INTUR."""
    for row in (1, 2):
        val = ws.cell(row=row, column=1).value
        if isinstance(val, str):
            s = val.strip().upper()
            if s in ("ORTI", "INTUR"):
                return s
    return "UNKNOWN"


# ---------------------------------------------------------------------------
# Data saldo detection
# ---------------------------------------------------------------------------

def _to_date(val: object) -> date | None:
    """Coerce a datetime or date value to date, return None otherwise."""
    from datetime import datetime as _dt
    if isinstance(val, _dt):
        return val.date()
    if isinstance(val, date):
        return val
    return None


def _find_data_saldo(ws: Worksheet, societa: str) -> date | None:
    """Scan the first 5 rows for a date cell (not a string, not an int year).

    ORTI typically has the date in B2; INTUR in C3.
    """
    for row in range(1, 6):
        for col in range(1, ws.max_column + 1):
            val = ws.cell(row=row, column=col).value
            d = _to_date(val)
            if d is not None and 2020 <= d.year <= 2030:
                return d
    return None


# ---------------------------------------------------------------------------
# Saldi banca detection
# ---------------------------------------------------------------------------

def _find_saldi_banca(ws: Worksheet) -> tuple[dict[str, float], float]:
    """Scan rows 30-45 for 'Saldo XXX' labels with an adjacent numeric value.

    Returns ({label: amount}, totale).
    ORTI stores the value in col B (2); INTUR in col C (3).
    """
    saldi: dict[str, float] = {}
    totale = 0.0

    for row in range(30, 46):
        label_cell = ws.cell(row=row, column=1).value
        if not isinstance(label_cell, str):
            continue
        label_stripped = label_cell.strip()

        # "TOTALE BANCHE" row
        if re.search(r"totale\s+banche", label_stripped, re.I):
            for col in (2, 3):
                val = ws.cell(row=row, column=col).value
                if isinstance(val, (int, float)):
                    totale = float(val)
                    break
            continue

        # "Saldo XXX" rows — skip non-banca labels
        if re.match(r"saldo\b", label_stripped, re.I):
            lower = label_stripped.lower()
            if any(skip in lower for skip in ("periodo", "proiettato", "mese", "cumulat")):
                continue
            # Extract banca name: "Saldo MPS" -> "MPS", "Saldo Banca Sella" -> "Banca Sella"
            banca = re.sub(r"^saldo\s+", "", label_stripped, flags=re.I).strip()
            if not banca:
                continue
            for col in (2, 3):
                val = ws.cell(row=row, column=col).value
                if isinstance(val, (int, float)):
                    saldi[banca] = float(val)
                    break

    # If no explicit TOTALE row, sum what we found
    if totale == 0.0 and saldi:
        totale = sum(saldi.values())

    return saldi, totale


# ---------------------------------------------------------------------------
# Main parser
# ---------------------------------------------------------------------------

def parse_pf(source: BinaryIO | BytesIO | bytes) -> PFData:
    """Parse a Rosa Piano Finanziario Excel file.

    Args:
        source: file-like object (BytesIO), bytes, or any binary file handle.

    Returns:
        PFData with societa, anno, data_saldo, saldi_banca, saldo_totale, voci, warnings.
    """
    if isinstance(source, bytes):
        source = BytesIO(source)

    wb = load_workbook(source, data_only=True)

    # Find the right sheet
    sheet_name = "Piano Finanziario"
    if sheet_name not in wb.sheetnames:
        sheet_name = wb.sheetnames[0]
    ws: Worksheet = wb[sheet_name]

    # ---- Societa & anno ------------------------------------------------
    societa = _find_societa(ws)
    anno = _find_anno(ws)

    # ---- Data saldo ----------------------------------------------------
    data_saldo = _find_data_saldo(ws, societa)

    # ---- Month column map (try row 2, then row 3) ----------------------
    month_cols: dict[int, int] = {}
    for try_row in (2, 3):
        month_cols = _find_month_columns(ws, try_row)
        if month_cols:
            header_row = try_row
            break
    else:
        header_row = 2

    # ---- Voci ----------------------------------------------------------
    warnings: list[str] = []
    voci: dict[str, VoceRow] = {}

    for row in ws.iter_rows(min_row=header_row + 1, max_row=ws.max_row):
        label_cell = row[0]  # column A
        label_val = label_cell.value
        if label_val is None:
            continue
        if not isinstance(label_val, str):
            continue
        label = label_val.strip()
        if not label:
            continue

        # Skip structural rows
        if _is_skip(label):
            continue

        # Resolve voce_id
        voce_id = _LABEL_MAP.get(_normalise(label))
        if voce_id is None:
            # Only warn if the row has material non-zero data (>1€)
            nz = [
                ws.cell(row=label_cell.row, column=c).value
                for c in month_cols.values()
                if isinstance(ws.cell(row=label_cell.row, column=c).value, (int, float))
                and abs(ws.cell(row=label_cell.row, column=c).value) > 1
            ]
            if nz:
                warnings.append(f"Unknown voce label: '{label}' (values: {nz[:3]}{'…' if len(nz) > 3 else ''})")
            continue

        # Read amounts for each month
        importi: dict[int, float] = {}
        for month_num, col_idx in month_cols.items():
            cell = ws.cell(row=label_cell.row, column=col_idx)
            val = cell.value
            if isinstance(val, (int, float)):
                importi[month_num] = float(val)
            else:
                importi[month_num] = 0.0

        # Merge with existing row if voce_id already seen (e.g. Fitto Hotel + Fitto AR)
        if voce_id in voci:
            existing = voci[voce_id]
            for m, v in importi.items():
                existing.importi[m] = existing.importi.get(m, 0.0) + v
        else:
            voci[voce_id] = VoceRow(
                voce_id=voce_id,
                excel_label=label,
                importi=importi,
            )

    # ---- Saldi banca ---------------------------------------------------
    saldi_banca, saldo_totale = _find_saldi_banca(ws)

    return PFData(
        societa=societa,
        anno=anno,
        data_saldo=data_saldo,
        saldi_banca=saldi_banca,
        saldo_totale=saldo_totale,
        voci=voci,
        warnings=warnings,
    )


# ---------------------------------------------------------------------------
# Scadenzario Fornitori parser
# ---------------------------------------------------------------------------

# Italian month abbreviation → month number (first 3 chars, lowercase)
_MESI_ABBR: dict[str, int] = {
    "gen": 1, "feb": 2, "mar": 3, "apr": 4, "mag": 5, "giu": 6,
    "lug": 7, "ago": 8, "set": 9, "ott": 10, "nov": 11, "dic": 12,
}


@dataclass
class ScadenzarioData:
    fornitori: list[dict]          # [{fornitore, totale, scaduto, mese_4: x, ...}]
    totale_per_mese: dict[int, float]  # {mese: totale_uscite}
    scaduto_totale: float


def _parse_month_header(header: str) -> int | None:
    """Return month number (1-12) from a header like 'apr-26', or None."""
    s = header.strip().lower()
    prefix = s[:3]
    return _MESI_ABBR.get(prefix)


def parse_scadenzario(source: BinaryIO | BytesIO | bytes) -> ScadenzarioData:
    """Parse an interrogazionesituazionesinteticascadenze.XLSX scadenzario file.

    Args:
        source: file-like object (BytesIO), bytes, or any binary file handle.

    Returns:
        ScadenzarioData with fornitori list, totale_per_mese dict, and scaduto_totale.
    """
    if isinstance(source, bytes):
        source = BytesIO(source)

    wb = load_workbook(source, data_only=True)
    ws: Worksheet = wb.active

    # ---- Parse header row (row 1) -----------------------------------------
    header_row = list(ws.iter_rows(min_row=1, max_row=1, values_only=True))[0]

    fornitore_col: int | None = None
    totale_col: int | None = None
    scaduto_col: int | None = None
    month_cols: dict[int, int] = {}  # {month_number: col_index (0-based)}

    for col_idx, header_val in enumerate(header_row):
        if header_val is None:
            continue
        h = str(header_val).strip().lower()
        if fornitore_col is None and ("fornitore" in h or "ragione" in h):
            fornitore_col = col_idx
        elif totale_col is None and ("totale" in h):
            totale_col = col_idx
        elif scaduto_col is None and "scaduto" in h:
            scaduto_col = col_idx
        else:
            # Handle "Scadenze - In scadenza al DD/MM/YYYY" format
            month_num = _parse_month_header(h)
            if month_num is None:
                # Try extracting month from date pattern DD/MM/YYYY in header
                m = re.search(r"in scadenza al\s+\d{1,2}/(\d{1,2})/\d{4}", h)
                if m:
                    month_num = int(m.group(1))
            if month_num is not None:
                month_cols[month_num] = col_idx

    if fornitore_col is None:
        fornitore_col = 0

    # ---- Read data rows ---------------------------------------------------
    fornitori: list[dict] = []
    totale_per_mese: dict[int, float] = {m: 0.0 for m in month_cols}
    scaduto_totale: float = 0.0

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
        fornitore_val = row[fornitore_col] if fornitore_col < len(row) else None
        if fornitore_val is None:
            continue
        if not isinstance(fornitore_val, str) or not fornitore_val.strip():
            continue

        entry: dict = {"fornitore": fornitore_val.strip()}

        if totale_col is not None and totale_col < len(row):
            v = row[totale_col]
            entry["totale"] = float(v) if isinstance(v, (int, float)) else 0.0
        else:
            entry["totale"] = 0.0

        scaduto_val = 0.0
        if scaduto_col is not None and scaduto_col < len(row):
            v = row[scaduto_col]
            scaduto_val = float(v) if isinstance(v, (int, float)) else 0.0
        entry["scaduto"] = scaduto_val
        scaduto_totale += scaduto_val

        for month_num, col_idx in month_cols.items():
            v = row[col_idx] if col_idx < len(row) else None
            amount = float(v) if isinstance(v, (int, float)) else 0.0
            entry[f"mese_{month_num}"] = amount
            totale_per_mese[month_num] = totale_per_mese.get(month_num, 0.0) + amount

        fornitori.append(entry)

    return ScadenzarioData(
        fornitori=fornitori,
        totale_per_mese=totale_per_mese,
        scaduto_totale=scaduto_totale,
    )
