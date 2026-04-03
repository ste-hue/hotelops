#!/usr/bin/env python3
"""
Budget vs Consuntivo — App Interattiva (Gasparotto)

Dashboard BvA con slider crescita ricavi per BU e coefficienti stagionalità.
Legge da BigQuery, non scrive. Export Excel.

Usage:
    streamlit run condges/bva_app.py
"""

from __future__ import annotations

from datetime import date
from io import BytesIO

import pandas as pd
import streamlit as st

from core import config as cfg

MESI = ["Gen", "Feb", "Mar", "Apr", "Mag", "Giu", "Lug", "Ago", "Set", "Ott", "Nov", "Dic"]
ANNO_DEFAULT = 2026

# BU inferred from codice_conto prefix
BU_PREFIX = {
    "4791": "HOTEL",
    "4792": "RESIDENCE",
    "4793": "CVM",
    "4794": "LIDO",
    "4795": "HQ",
}

CATEGORIE_ORD = [
    "Ricavi",
    "Acquisti",
    "Costi Produttivi",
    "Costo del Personale",
    "Costi Amministrativi",
    "Costi Commerciali",
    "Oneri Tributari",
    "Oneri Finanziari",
]


@st.cache_resource
def get_bq():
    from google.cloud import bigquery
    return bigquery.Client(project=cfg.PROJECT)


def infer_bu(cod_conto: str) -> str:
    """Infer business unit from codice_conto prefix."""
    for prefix, bu in BU_PREFIX.items():
        if cod_conto.startswith(prefix):
            return bu
    return "HQ"


# ── Data Loading ──────────────────────────────────────────────────────────────


@st.cache_data(ttl=300, show_spinner=False)
def load_budget(societa: str, anno: int, _bq=None) -> pd.DataFrame:
    """Load budget from v_budget: one row per conto, 12 monthly columns."""
    bq = _bq or get_bq()
    sql = f"""
    SELECT cod_conto, codice_conto_display, descrizione,
           tipo_costo, categoria_ce, fonte,
           gen, feb, mar, apr, mag, giu, lug, ago, sett, ott, nov, dic,
           totale_annuo
    FROM `{cfg.V_BUDGET}`
    WHERE societa_id = '{societa}' AND anno = {anno}
    ORDER BY categoria_ce, cod_conto
    """
    df = bq.query(sql).to_dataframe()
    df = df.set_index("cod_conto")
    # Rename sett -> Set for consistency with MESI
    df = df.rename(columns={"sett": "Set"})
    # Capitalize month columns to match MESI
    month_map = {"gen": "Gen", "feb": "Feb", "mar": "Mar", "apr": "Apr",
                 "mag": "Mag", "giu": "Giu", "lug": "Lug", "ago": "Ago",
                 "ott": "Ott", "nov": "Nov", "dic": "Dic"}
    df = df.rename(columns=month_map)
    return df


@st.cache_data(ttl=300, show_spinner=False)
def load_consuntivo(societa: str, anno: int, _bq=None) -> pd.DataFrame:
    """Load consuntivo from f_movimenti_contabili, aggregated by conto × mese."""
    bq = _bq or get_bq()
    sql = f"""
    SELECT
      cod_conto,
      EXTRACT(MONTH FROM data_registrazione) AS mese,
      SUM(imp_dare - imp_avere) AS netto
    FROM `{cfg.F_MOVIMENTI_CONTABILI}`
    WHERE societa_id = '{societa}'
      AND EXTRACT(YEAR FROM data_registrazione) = {anno}
    GROUP BY cod_conto, EXTRACT(MONTH FROM data_registrazione)
    """
    df = bq.query(sql).to_dataframe()
    if df.empty:
        return pd.DataFrame(columns=["cod_conto"] + MESI).set_index("cod_conto")

    # Pivot: rows=cod_conto, cols=mesi
    pivot = df.pivot_table(index="cod_conto", columns="mese", values="netto", fill_value=0.0)
    for m in range(1, 13):
        if m not in pivot.columns:
            pivot[m] = 0.0
    pivot = pivot[[m for m in range(1, 13)]]
    pivot.columns = MESI
    return pivot


@st.cache_data(ttl=300, show_spinner=False)
def load_stagionalita(societa: str, _bq=None) -> pd.DataFrame:
    """Load seasonality coefficients per BU × mese."""
    bq = _bq or get_bq()
    sql = f"""
    SELECT business_unit_id AS bu, mese, coefficiente
    FROM `{cfg.D_COEFFICIENTI_STAGIONALITA}`
    WHERE societa_id = '{societa}'
    ORDER BY bu, mese
    """
    df = bq.query(sql).to_dataframe()
    return df.set_index(["bu", "mese"])


# ── Ricavi Ricalcolo ──────────────────────────────────────────────────────────


def recalc_ricavi(
    budget: pd.DataFrame,
    stag: pd.DataFrame,
    crescita: dict[str, float],
) -> pd.DataFrame:
    """
    Recalculate ricavi rows in budget using growth sliders + seasonality.

    For each ricavo conto:
      1. Infer BU from prefix (47.91=HOTEL, etc.)
      2. Get base annual total (from budget)
      3. Apply growth: annuo_new = annuo_base × (1 + crescita[bu])
      4. Distribute monthly: mese_new = annuo_new × coeff(bu, mese) / 12
    """
    df = budget.copy()
    ricavi_mask = df["categoria_ce"] == "Ricavi"

    for cod_conto in df[ricavi_mask].index:
        bu = infer_bu(cod_conto)
        growth = crescita.get(bu, 0.0)
        annuo_base = float(df.loc[cod_conto, "totale_annuo"])
        annuo_new = annuo_base * (1 + growth)

        for m_idx, nome in enumerate(MESI, 1):
            try:
                coeff = float(stag.loc[(bu, m_idx), "coefficiente"])
            except KeyError:
                coeff = 1.0  # fallback uniform
            df.loc[cod_conto, nome] = annuo_new * coeff / 12.0

        df.loc[cod_conto, "totale_annuo"] = df.loc[cod_conto, MESI].sum()

    return df


def recalc_variabili(
    budget: pd.DataFrame,
    ricavi_orig_totals: pd.Series,
    ricavi_new_totals: pd.Series,
) -> pd.DataFrame:
    """
    Scale variable costs proportionally to ricavi change per month.

    costo_var_mese_new = costo_var_mese_orig × (ricavi_new_mese / ricavi_orig_mese)
    """
    df = budget.copy()
    var_mask = df["tipo_costo"].isin(["V", "IP"]) & (df["categoria_ce"] != "Ricavi")

    for nome in MESI:
        orig = ricavi_orig_totals.get(nome, 0.0)
        new = ricavi_new_totals.get(nome, 0.0)
        if orig > 0:
            ratio = new / orig
        else:
            ratio = 1.0
        df.loc[var_mask, nome] = df.loc[var_mask, nome] * ratio

    # Recalc annual total
    df.loc[var_mask, "totale_annuo"] = df.loc[var_mask, MESI].sum(axis=1)
    return df


# ── Helpers ───────────────────────────────────────────────────────────────────


def fmt_eur(v) -> str:
    try:
        f = float(v)
        if f != f:
            return "—"
        return f"€ {f:,.0f}"
    except (TypeError, ValueError):
        return "—"


def past_month_count(anno: int) -> int:
    """Number of completed months (mesi chiusi) for the given year."""
    today = date.today()
    if anno < today.year:
        return 12
    if anno > today.year:
        return 0
    return today.month - 1


def build_categoria_summary(
    budget: pd.DataFrame,
    consuntivo: pd.DataFrame,
    mesi_chiusi: int,
) -> pd.DataFrame:
    """
    Build summary table: one row per categoria_ce, columns per mese.
    For closed months: budget, consuntivo, delta.
    For future months: budget only.
    """
    rows = []
    for cat in CATEGORIE_ORD:
        cat_budget = budget[budget["categoria_ce"] == cat]
        if cat_budget.empty:
            continue

        row = {"Categoria": cat}

        for m_idx, nome in enumerate(MESI, 1):
            bud_tot = cat_budget[nome].sum()

            if m_idx <= mesi_chiusi:
                # Consuntivo: sum matching cod_conto
                cons_tot = 0.0
                for cod in cat_budget.index:
                    if cod in consuntivo.index:
                        cons_tot += float(consuntivo.loc[cod, nome])
                delta = cons_tot - bud_tot
                delta_pct = (delta / abs(bud_tot) * 100) if bud_tot != 0 else 0
                row[f"{nome}_bud"] = round(bud_tot)
                row[f"{nome}_cons"] = round(cons_tot)
                row[f"{nome}_delta"] = f"{delta_pct:+.0f}%"
            else:
                row[f"{nome}_bud"] = round(bud_tot)

        rows.append(row)

    return pd.DataFrame(rows)


# ── Excel Export ──────────────────────────────────────────────────────────────


def generate_bva_excel(
    budget: pd.DataFrame,
    consuntivo: pd.DataFrame,
    societa: str,
    anno: int,
    mesi_chiusi: int,
    crescita: dict[str, float],
) -> bytes:
    """Generate BvA Excel with Dettaglio + Parametri sheets."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    wb.remove(wb.active)

    FILL_HEADER = PatternFill("solid", fgColor="D9E1F2")
    FILL_GREEN = PatternFill("solid", fgColor="E2EFDA")
    FILL_RED = PatternFill("solid", fgColor="FCE4D6")
    FONT_BOLD = Font(bold=True)
    EUR_FMT = '#,##0;(#,##0);"-"'

    # ── Sheet 1: Dettaglio ────────────────────────────────────────────────────
    ws = wb.create_sheet("Dettaglio")
    headers = ["Codice Conto", "Descrizione", "Categoria", "Fonte", "Tipo"]
    for nome in MESI:
        headers.append(f"{nome} Bud")
        if MESI.index(nome) < mesi_chiusi:
            headers += [f"{nome} Cons", f"{nome} Δ"]
    headers.append("Totale Bud")

    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(1, col_idx, h)
        cell.font = FONT_BOLD
        cell.fill = FILL_HEADER

    row = 2
    for cat in CATEGORIE_ORD:
        cat_df = budget[budget["categoria_ce"] == cat]
        for cod in cat_df.index:
            col = 1
            ws.cell(row, col, cat_df.loc[cod, "codice_conto_display"]); col += 1
            ws.cell(row, col, cat_df.loc[cod, "descrizione"] or ""); col += 1
            ws.cell(row, col, cat); col += 1
            ws.cell(row, col, cat_df.loc[cod, "fonte"]); col += 1
            ws.cell(row, col, cat_df.loc[cod, "tipo_costo"]); col += 1

            for m_idx, nome in enumerate(MESI):
                bud_val = round(float(cat_df.loc[cod, nome]))
                ws.cell(row, col, bud_val).number_format = EUR_FMT; col += 1

                if m_idx < mesi_chiusi:
                    cons_val = 0
                    if cod in consuntivo.index:
                        cons_val = round(float(consuntivo.loc[cod, nome]))
                    ws.cell(row, col, cons_val).number_format = EUR_FMT; col += 1
                    delta = cons_val - bud_val
                    cell = ws.cell(row, col, delta)
                    cell.number_format = EUR_FMT
                    cell.fill = FILL_GREEN if abs(delta) <= abs(bud_val) * 0.1 else FILL_RED
                    col += 1

            ws.cell(row, col, round(float(cat_df.loc[cod, "totale_annuo"]))).number_format = EUR_FMT
            row += 1

    # Auto-width first columns
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 40
    ws.column_dimensions["C"].width = 22

    # ── Sheet 2: Parametri ────────────────────────────────────────────────────
    ws_p = wb.create_sheet("Parametri")
    params = [
        ("Società", societa),
        ("Anno", anno),
        ("Generato", date.today().isoformat()),
        ("", ""),
        ("Crescita Hotel", f"{crescita.get('HOTEL', 0):.0%}"),
        ("Crescita Residence", f"{crescita.get('RESIDENCE', 0):.0%}"),
        ("Crescita CVM", f"{crescita.get('CVM', 0):.0%}"),
        ("Crescita Spiaggia", f"{crescita.get('LIDO', 0):.0%}"),
        ("Crescita HQ", f"{crescita.get('HQ', 0):.0%}"),
    ]
    for i, (k, v) in enumerate(params, 1):
        ws_p.cell(i, 1, k).font = FONT_BOLD
        ws_p.cell(i, 2, v)
    ws_p.column_dimensions["A"].width = 25

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── Main App ──────────────────────────────────────────────────────────────────


def main():
    st.set_page_config(
        page_title="Budget vs Consuntivo",
        page_icon="📊",
        layout="wide",
        initial_sidebar_state="expanded",
    )

    # ── Sidebar ───────────────────────────────────────────────────────────────
    with st.sidebar:
        st.header("📊 Budget vs Consuntivo")
        societa = st.selectbox("Società", ["ORTI", "INTUR"], key="bva_societa")
        anno = st.selectbox("Anno", [2025, 2026, 2027], index=1, key="bva_anno")

        st.divider()
        st.subheader("Crescita Ricavi")
        cr_hotel = st.slider("Hotel", -20, 50, 20, 1, format="%d%%", key="cr_hotel")
        cr_residence = st.slider("Residence", -20, 50, 5, 1, format="%d%%", key="cr_res")
        cr_cvm = st.slider("CVM", -20, 50, 5, 1, format="%d%%", key="cr_cvm")
        cr_lido = st.slider("Spiaggia", -20, 50, 8, 1, format="%d%%", key="cr_lido")
        cr_hq = st.slider("HQ/Supermercato", -20, 50, 0, 1, format="%d%%", key="cr_hq")

        crescita = {
            "HOTEL": cr_hotel / 100,
            "RESIDENCE": cr_residence / 100,
            "CVM": cr_cvm / 100,
            "LIDO": cr_lido / 100,
            "HQ": cr_hq / 100,
        }

        st.divider()
        if st.button("🔄 Ricarica da BQ", use_container_width=True):
            st.cache_data.clear()
            st.rerun()

    bq = get_bq()

    # ── Load data ─────────────────────────────────────────────────────────────
    with st.spinner("Caricamento dati da BigQuery..."):
        budget_orig = load_budget(societa, anno, _bq=bq)
        consuntivo = load_consuntivo(societa, anno, _bq=bq)
        stag = load_stagionalita(societa, _bq=bq)

    if budget_orig.empty:
        st.warning(f"Nessun budget per {societa} {anno}.")
        return

    # ── Recalc ricavi ─────────────────────────────────────────────────────────
    ricavi_orig_totals = budget_orig[budget_orig["categoria_ce"] == "Ricavi"][MESI].sum()
    budget = recalc_ricavi(budget_orig, stag, crescita)
    ricavi_new_totals = budget[budget["categoria_ce"] == "Ricavi"][MESI].sum()
    budget = recalc_variabili(budget, ricavi_orig_totals, ricavi_new_totals)

    mesi_chiusi = past_month_count(anno)

    # ── Header ────────────────────────────────────────────────────────────────
    tot_ricavi = budget[budget["categoria_ce"] == "Ricavi"]["totale_annuo"].sum()
    tot_costi = budget[budget["categoria_ce"] != "Ricavi"]["totale_annuo"].sum()
    margine = tot_ricavi - tot_costi

    col1, col2, col3 = st.columns(3)
    col1.metric("Ricavi Budget", fmt_eur(tot_ricavi))
    col2.metric("Costi Budget", fmt_eur(tot_costi))
    col3.metric("Margine", fmt_eur(margine))

    st.divider()

    # ── Sezione 1: Riepilogo per categoria ────────────────────────────────────
    st.subheader("Riepilogo per Categoria")
    summary = build_categoria_summary(budget, consuntivo, mesi_chiusi)
    if not summary.empty:
        st.dataframe(summary, hide_index=True, use_container_width=True)

    st.divider()

    # ── Sezione 2: Dettaglio per codice conto ─────────────────────────────────
    st.subheader("Dettaglio per Codice Conto")

    for cat in CATEGORIE_ORD:
        cat_budget = budget[budget["categoria_ce"] == cat].copy()
        if cat_budget.empty:
            continue

        cat_tot = cat_budget["totale_annuo"].sum()
        n_conti = len(cat_budget)

        with st.expander(f"**{cat}** — {n_conti} conti — Budget: {fmt_eur(cat_tot)}"):
            display_cols = ["codice_conto_display", "descrizione", "fonte"] + MESI + ["totale_annuo"]
            show = cat_budget[display_cols].copy()
            show = show.rename(columns={
                "codice_conto_display": "Conto",
                "descrizione": "Descrizione",
                "fonte": "Fonte",
                "totale_annuo": "Totale",
            })
            # Round month columns
            for nome in MESI:
                show[nome] = show[nome].round(0).astype(int)
            show["Totale"] = show["Totale"].round(0).astype(int)

            st.dataframe(show, hide_index=True, use_container_width=True)

            # If there's consuntivo data, show comparison for closed months
            if mesi_chiusi > 0:
                mesi_label = ", ".join(MESI[:mesi_chiusi])
                st.caption(f"Confronto consuntivo per: {mesi_label}")

                comp_rows = []
                for cod in cat_budget.index:
                    bud_sum = sum(float(cat_budget.loc[cod, MESI[m]]) for m in range(mesi_chiusi))
                    cons_sum = 0.0
                    if cod in consuntivo.index:
                        cons_sum = sum(float(consuntivo.loc[cod, MESI[m]]) for m in range(mesi_chiusi))
                    if bud_sum == 0 and cons_sum == 0:
                        continue
                    delta = cons_sum - bud_sum
                    delta_pct = (delta / abs(bud_sum) * 100) if bud_sum != 0 else 0
                    comp_rows.append({
                        "Conto": cat_budget.loc[cod, "codice_conto_display"],
                        "Descrizione": cat_budget.loc[cod, "descrizione"] or "",
                        f"Budget {mesi_label}": round(bud_sum),
                        f"Consuntivo {mesi_label}": round(cons_sum),
                        "Delta": round(delta),
                        "Delta %": f"{delta_pct:+.0f}%",
                    })

                if comp_rows:
                    comp_df = pd.DataFrame(comp_rows)
                    st.dataframe(comp_df, hide_index=True, use_container_width=True)

    # ── Sezione 3: Export ─────────────────────────────────────────────────────
    st.divider()
    col_exp, _ = st.columns([1, 3])
    with col_exp:
        if st.button("📥 Genera Excel", use_container_width=True):
            with st.spinner("Generazione Excel..."):
                xlsx = generate_bva_excel(
                    budget, consuntivo, societa, anno, mesi_chiusi, crescita
                )
            st.download_button(
                label="⬇️ Scarica BvA Excel",
                data=xlsx,
                file_name=f"BvA_{societa}_{anno}_{date.today().isoformat()}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True,
            )


if __name__ == "__main__":
    main()
