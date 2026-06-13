"""Estrazione dal PF precedente: previsioni (blocco B) + consuntivi chiusi.

Lettura TOLLERANTE: risolve i fogli sia coi nomi puliti del generatore sia
coi nomi legacy (typo storici) via VOCE_TO_SHEET_CANDIDATES.
"""

from __future__ import annotations

from io import BytesIO

import openpyxl

from verticals.condges.app_scadenzario import (
    VOCE_TO_SHEET_CANDIDATES,
    _build_month_col_map,
)
from verticals.condges.pf_generator.costanti import (
    LABEL_RETTIFICA,
    LABEL_SEZIONE_A,
    LABEL_SEZIONE_B,
    VOCE_SHEET_NAME,
)

_LABEL_SKIP = {
    LABEL_RETTIFICA.upper(),
    LABEL_SEZIONE_A.upper(),
    LABEL_SEZIONE_B.upper(),
    "PREVISIONALE",
}


def _risolvi_foglio(voce_id: str, sheetnames: list[str]) -> str | None:
    candidati = [VOCE_SHEET_NAME[voce_id]] + VOCE_TO_SHEET_CANDIDATES.get(voce_id, [])
    for nome in candidati:
        if nome in sheetnames:
            return nome
    return None


def _skip_label(nome: str) -> bool:
    up = nome.upper().strip()
    if not up:
        return True
    if any(lbl in up for lbl in _LABEL_SKIP):
        return True
    return up.startswith("TOTALE") or up.startswith("MATERIE PRIME")


def estrai_previsioni(
    pf_bytes: bytes,
    *,
    codici_partite: set[int],
    primo_mese_aperto: int,
) -> dict[str, dict]:
    """Per ogni voce: {'previsioni': [righe], 'consuntivi': {codice: {mese: v}}}.

    - previsione = riga con valori nei mesi APERTI il cui codice NON è nelle
      partite correnti (o senza codice). Le righe-fornitore con partite
      correnti appartengono al motore: i loro mesi aperti si rigenerano.
    - consuntivi = valori dei mesi CHIUSI delle righe con codice (qualunque),
      da ricopiare as-is nel nuovo file.
    Righe di servizio (TOTALE, RETTIFICA, header sezione, PREVISIONALE
    legacy) escluse per label.
    """
    wb = openpyxl.load_workbook(BytesIO(pf_bytes), data_only=True)
    out: dict[str, dict] = {}
    for voce_id in VOCE_SHEET_NAME:
        nome_foglio = _risolvi_foglio(voce_id, wb.sheetnames)
        if not nome_foglio:
            out[voce_id] = {"previsioni": [], "consuntivi": {}}
            continue
        ws = wb[nome_foglio]
        mc = _build_month_col_map(ws)
        previsioni: list[dict] = []
        consuntivi: dict[int, dict[int, float]] = {}
        for r in range(3, ws.max_row + 1):
            cod = ws.cell(row=r, column=1).value
            nome = str(ws.cell(row=r, column=2).value or "").strip()
            if _skip_label(nome) and not isinstance(cod, (int, float)):
                continue
            mesi_aperti: dict[int, float] = {}
            mesi_chiusi: dict[int, float] = {}
            for m, col in mc.items():
                v = ws.cell(row=r, column=col).value
                if not isinstance(v, (int, float)) or round(v, 2) == 0:
                    continue
                if m >= primo_mese_aperto:
                    mesi_aperti[m] = round(float(v), 2)
                else:
                    mesi_chiusi[m] = round(float(v), 2)
            codice = int(cod) if isinstance(cod, (int, float)) else None
            if codice is not None and mesi_chiusi:
                consuntivi[codice] = mesi_chiusi
            if not mesi_aperti:
                continue
            if codice is not None and codice in codici_partite:
                continue  # riga del motore: i mesi aperti si rigenerano
            if _skip_label(nome):
                continue
            previsioni.append({"codice": codice, "nome": nome, "mesi": mesi_aperti})
        out[voce_id] = {"previsioni": previsioni, "consuntivi": consuntivi}
    return out
