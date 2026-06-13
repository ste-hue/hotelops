"""Scrittura fogli del PF generato (layout standard restyled)."""
from __future__ import annotations

from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook import Workbook

from verticals.condges.pf_generator.costanti import (
    FILL_MESE_CHIUSO,
    FILL_SEZIONE_A,
    FILL_SEZIONE_B,
    FONT_PREVISIONE,
    LABEL_RETTIFICA,
    LABEL_SEZIONE_A,
    LABEL_SEZIONE_B,
    MESI,
    NUMFMT_CONTABILE,
    PRIMA_RIGA_BLOCCO_A,
    RIGA_HEADER_MESI,
    RIGA_SEZIONE_A,
    RIGA_TOTALE,
    VOCE_SHEET_NAME,
    col_mese,
)


def _intesta(ws, titolo: str, primo_mese_aperto: int) -> None:
    ws.cell(row=1, column=2, value=titolo).font = Font(bold=True, size=12)
    ws.cell(row=RIGA_HEADER_MESI, column=1, value="Cod").font = Font(bold=True)
    ws.cell(row=RIGA_HEADER_MESI, column=2, value="Fornitore / Voce").font = (
        Font(bold=True)
    )
    for m, nome in enumerate(MESI, start=1):
        c = ws.cell(row=RIGA_HEADER_MESI, column=col_mese(m), value=nome)
        c.font = Font(bold=True)
        if m < primo_mese_aperto:
            c.fill = PatternFill("solid", fgColor=FILL_MESE_CHIUSO)
    ws.freeze_panes = ws.cell(row=RIGA_HEADER_MESI + 1, column=3)
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 34
    for m in range(1, 13):
        ws.column_dimensions[get_column_letter(col_mese(m))].width = 12


def _scrivi_riga(ws, r: int, codice, nome: str, mesi: dict[int, float],
                 *, blu: bool = False) -> None:
    if codice is not None:
        ws.cell(row=r, column=1, value=int(codice))
    ws.cell(row=r, column=2, value=nome)
    for mese, val in sorted(mesi.items()):
        c = ws.cell(row=r, column=col_mese(mese), value=val)
        c.number_format = NUMFMT_CONTABILE
        if blu:
            c.font = Font(color=FONT_PREVISIONE)


def scrivi_foglio_voce(
    wb: Workbook,
    *,
    voce_id: str,
    blocco_a: list[dict],
    previsioni: list[dict],
    consuntivi: dict[int, dict[int, float]],
    rettifica: dict[int, float],
    primo_mese_aperto: int,
) -> str:
    """Crea il foglio della voce: blocco A + blocco B + rettifica + totale.

    Ritorna il nome del foglio creato.
    """
    nome_foglio = VOCE_SHEET_NAME[voce_id]
    ws = wb.create_sheet(nome_foglio)
    _intesta(ws, nome_foglio, primo_mese_aperto)

    sez_a = ws.cell(row=RIGA_SEZIONE_A, column=2, value=LABEL_SEZIONE_A)
    sez_a.fill = PatternFill("solid", fgColor=FILL_SEZIONE_A)
    sez_a.font = Font(bold=True)

    r = PRIMA_RIGA_BLOCCO_A
    for riga in blocco_a:
        mesi = dict(riga["mesi"])
        mesi.update(consuntivi.get(riga["codice"], {}))  # mesi chiusi as-is
        _scrivi_riga(ws, r, riga["codice"], riga["nome"], mesi)
        r += 1

    r += 1  # riga vuota di separazione
    sez_b = ws.cell(row=r, column=2, value=LABEL_SEZIONE_B)
    sez_b.fill = PatternFill("solid", fgColor=FILL_SEZIONE_B)
    sez_b.font = Font(bold=True)
    r += 1
    for riga in previsioni:
        _scrivi_riga(ws, r, riga.get("codice"), riga["nome"], riga["mesi"],
                     blu=True)
        r += 1
    if rettifica:
        _scrivi_riga(ws, r, None, LABEL_RETTIFICA, rettifica)
        for mese in rettifica:
            ws.cell(row=r, column=col_mese(mese)).font = Font(italic=True)
        r += 1
    fine_b = r - 1

    for m in range(1, 13):
        col = get_column_letter(col_mese(m))
        cella = ws.cell(
            row=RIGA_TOTALE, column=col_mese(m),
            value=f"=SUM({col}{PRIMA_RIGA_BLOCCO_A}:{col}{fine_b})",
        )
        cella.number_format = NUMFMT_CONTABILE
        cella.font = Font(bold=True)
    ws.cell(row=RIGA_TOTALE, column=2, value="TOTALE").font = Font(bold=True)
    return nome_foglio
