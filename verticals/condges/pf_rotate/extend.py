"""Estensione dell'orizzonte del template PF — spec 2026-08-10.

Aggiunge colonne-mese fino al periodo target copiando il pattern dell'ultima
colonna-mese esistente con traduzione dei riferimenti (Translator). Non tocca
mai i valori esistenti. La colonna TOTALI (se presente, master) viene spostata
a destra e le sue formule ripuntate al nuovo orizzonte.
"""

from __future__ import annotations

import re

from openpyxl import Workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.formula.translate import Translator
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from verticals.condges.pf_rotate.excel_model import (
    MESI_IT,
    find_month_periods,
    periodo_anno_mese,
    require_periodo,
)

_NOME_MESE = {v: k for k, v in MESI_IT.items()}
_HEADER_ROW = 2
_YEAR_ROW = 1

# Fogli-report generati da pf_generator.template (scrivi_da_mappare/scrivi_esclusi):
# 12 colonne mese-NUDO (no anno dichiarato in riga 1) — non sono griglie PF,
# find_month_periods ci solleverebbe sempre "Anno non dichiarato". Match per
# prefisso di titolo (case-sensitive, come i nomi reali): copre anche i
# duplicati auto-rinominati da openpyxl ("DA MAPPARE1", "ESCLUSI1").
_SERVICE_SHEET_PREFIXES = ("DA MAPPARE", "ESCLUSI")

# Riferimento cella A1-style, con eventuali $ (no supporto sheet-qualified: le
# formule della colonna TOTALI osservate nel template sono sempre same-sheet).
_CELL_REF_RE = re.compile(r"(\$?)([A-Z]{1,3})(\$?)(\d+)")


def _shift_totali_formula(formula: str, cols_to_shift: set[int], n_new: int) -> str:
    """Ripunta i riferimenti a `cols_to_shift` di `n_new` colonne a destra.

    Usato sulla colonna TOTALI dopo lo spostamento fisico: un riferimento
    all'ultimo mese vecchio (es. K37) o alla vecchia colonna TOTALI stessa
    (es. L12 dentro una formula su un'altra riga, tipo Cash Flow = TOTALE
    ENTRATE - TOTALE USCITE) deve seguire lo shift; i riferimenti al primo
    mese (es. C nel range di un SUM) restano fermi.
    """

    def repl(m: re.Match) -> str:
        dollar_col, col_letters, dollar_row, row = m.groups()
        col_idx = column_index_from_string(col_letters)
        if col_idx in cols_to_shift:
            col_letters = get_column_letter(col_idx + n_new)
        return f"{dollar_col}{col_letters}{dollar_row}{row}"

    return _CELL_REF_RE.sub(repl, formula)


def _extend_sheet(ws: Worksheet, target_periodo: int) -> list[str]:
    titolo = ws.title.strip()
    if titolo.startswith(_SERVICE_SHEET_PREFIXES):
        return [f"skip: {ws.title} (foglio di servizio)"]
    cols = find_month_periods(ws, header_row=_HEADER_ROW)
    if not cols:
        return []
    last_p = max(cols)
    if target_periodo <= last_p:
        return []
    src_col = cols[last_p]

    # TOTALI: colonna non-mese subito dopo l'ultimo mese (solo master).
    totali_col = src_col + 1
    has_totali = isinstance(ws.cell(_HEADER_ROW, totali_col).value, str)
    n_new = target_periodo - last_p
    if has_totali:
        ws.insert_cols(
            totali_col, n_new
        )  # sposta TOTALI a destra, formule NON tradotte
        # NB: le formule TOTALI restano testualmente identiche (=SUM(C6:K6), =K37):
        # i range esistenti non cambiano posizione, quindi restano corrette; si
        # estendono/ripuntano sotto.

    changed: list[str] = []
    for i in range(1, n_new + 1):
        p = last_p + i
        dst_col = src_col + i
        anno, mese = periodo_anno_mese(p)
        ws.cell(_HEADER_ROW, dst_col, _NOME_MESE[mese])
        if mese == 1:
            # marker anno sul primo mese dell'anno nuovo — decorativo (find_month_periods
            # deriva l'anno dal wrap, non da questa cella): se la riga 1 è merged in quel
            # punto (es. ORTI ' Varie ed Eventuali', I1:O1) la cella target è una
            # MergedCell read-only — salta il marker, non l'estensione delle colonne.
            if isinstance(ws.cell(_YEAR_ROW, dst_col), MergedCell):
                changed.append(
                    f"warn: marker anno saltato su '{ws.title}' (riga 1 merged)"
                )
            else:
                ws.cell(_YEAR_ROW, dst_col, anno)
        for r in range(1, ws.max_row + 1):
            v = ws.cell(r, src_col).value
            if isinstance(v, str) and v.startswith("="):
                src_ref = f"{get_column_letter(src_col)}{r}"
                dst_ref = f"{get_column_letter(dst_col)}{r}"
                ws.cell(r, dst_col).value = Translator(
                    v, origin=src_ref
                ).translate_formula(dst_ref)
        changed.append(f"{ws.title}!{get_column_letter(dst_col)} ({anno}-{mese:02d})")

    if has_totali:
        new_totali_col = totali_col + n_new
        # Righe di formula che vivevano su src_col (vecchio ultimo mese) o su
        # totali_col (vecchia colonna TOTALI, riferita da un'altra riga TOTALI,
        # es. Cash Flow = TOTALE ENTRATE - TOTALE USCITE) devono ripuntare di
        # n_new colonne: la prima diventa il nuovo ultimo mese, la seconda
        # diventa il nuovo TOTALI.
        cols_to_shift = {src_col, totali_col}
        for r in range(1, ws.max_row + 1):
            v = ws.cell(r, new_totali_col).value
            if not isinstance(v, str) or not v.startswith("="):
                continue
            v2 = _shift_totali_formula(v, cols_to_shift, n_new)
            if v2 != v:
                ws.cell(r, new_totali_col).value = v2
                changed.append(
                    f"{ws.title}!{get_column_letter(new_totali_col)}{r} (TOTALI)"
                )
    return changed


def extend_to(wb: Workbook, target_periodo: int) -> list[str]:
    """Estende ogni foglio con colonne-mese fino a target_periodo. In-place."""
    require_periodo(target_periodo, "target_periodo")
    changed: list[str] = []
    for name in wb.sheetnames:
        changed += _extend_sheet(wb[name], target_periodo)
    return changed


# ── trim_before (potatura annuale) ──────────────────────────────────────
#
# openpyxl.Worksheet.delete_cols sposta fisicamente il CONTENUTO delle colonne
# a destra dell'eliminazione, ma NON traduce un solo riferimento nelle formule
# (verificato empiricamente: `=A1+B1` in C1, dopo delete_cols(1,2), resta
# testualmente `=A1+B1` ma è finito in A1 — un self-ref rotto). Questo vale
# anche per le formule che NON referenziano affatto le colonne eliminate: si
# spostano fisicamente ma il testo resta quello vecchio, quindi anche loro
# vanno ri-tradotte. Strategia in due passate, prima di/dopo `delete_cols`:
#   1. (pre-delete) CONGELA a valore ogni formula in una colonna superstite
#      che referenzia una colonna in eliminazione — non esiste una traduzione
#      sensata per un riferimento che sta per sparire (vale per la cascata di
#      riga 4 sulla prima superstite, ma anche per un TOTALI il cui SUM copre
#      l'intero storico: il suo estremo sinistro non "trasla", va congelato).
#   2. (post-delete) Per tutto il resto (formule superstiti non congelate),
#      Translator dalla vecchia posizione (posizione attuale + n_del) alla
#      nuova: replica esattamente lo spostamento fisico che delete_cols ha
#      già fatto ai valori ma non alle formule.
# Alternativa scartata: ricostruire le colonne superstiti per traduzione dalla
# sola prima colonna congelata (proposta nel brief). Scartata perché la prima
# colonna, una volta congelata a valore, non è più un pattern di formula da
# cui il Translator possa derivare le altre — e comunque non risolve il caso
# TOTALI (colonna diversa da "prima superstite", con un SUM che copre sia
# storico eliminato sia superstite: né freeze-solo-primo né translate-tutto
# bastano da soli). Il presente approccio (freeze mirato + translate del resto)
# copre entrambi i casi con una sola passata di scansione.

_A1_TOKEN_RE = re.compile(
    r"(?:'[^']+'|[A-Za-z_][A-Za-z0-9_ ]*)?!?\$?([A-Z]{1,3})\$?\d+"
)


def _formula_refs_col_range(formula_body: str, lo: int, hi: int) -> bool:
    """True se `formula_body` contiene un riferimento NON qualificato da
    foglio (stesso foglio) a una colonna in [lo, hi]. I riferimenti
    sheet-qualified (`Utenze!C3`) sono ignorati: referenziano le colonne di
    UN ALTRO foglio, non quelle in eliminazione su questo foglio."""
    for m in _A1_TOKEN_RE.finditer(formula_body):
        if "!" in m.group(0):
            continue
        if lo <= column_index_from_string(m.group(1)) <= hi:
            return True
    return False


def _has_year_marker(ws: Worksheet, up_to_col: int) -> bool:
    for c in range(1, up_to_col + 1):
        v = ws.cell(_YEAR_ROW, c).value
        if isinstance(v, (int, float)) and 1900 < int(v) < 2100:
            return True
    return False


def _trim_sheet(ws: Worksheet, wb_values: Workbook, cutoff_periodo: int) -> list[str]:
    titolo = ws.title.strip()
    if titolo.startswith(_SERVICE_SHEET_PREFIXES):
        return []
    cols = find_month_periods(ws, header_row=_HEADER_ROW)
    to_delete_periods = [p for p in cols if p < cutoff_periodo]
    if not to_delete_periods:
        return []

    del_cols = sorted(cols[p] for p in to_delete_periods)
    first_del, last_del = del_cols[0], del_cols[-1]
    n_del = last_del - first_del + 1
    if del_cols != list(range(first_del, last_del + 1)):
        raise ValueError(
            f"{ws.title}: colonne mese da eliminare non contigue: {del_cols}"
        )
    c_first_old = min(c for p, c in cols.items() if p not in to_delete_periods)

    if ws.title not in wb_values.sheetnames:
        raise ValueError(
            f"{ws.title}: foglio assente nel gemello data_only — impossibile "
            "congelare i valori calcolati prima del trim"
        )
    wsv = wb_values[ws.title]

    changed: list[str] = []

    # Passata 1 (pre-delete): congela le formule superstiti che referenziano
    # una colonna in eliminazione — su TUTTE le colonne superstiti (non solo
    # la prima), perché una TOTALI più a destra può referenziare l'intero
    # storico eliminato con un SUM.
    max_col_pre = ws.max_column
    for r in range(1, ws.max_row + 1):
        for c in range(c_first_old, max_col_pre + 1):
            cell = ws.cell(r, c)
            v = cell.value
            if not (isinstance(v, str) and v.startswith("=")):
                continue
            if not _formula_refs_col_range(v[1:], first_del, last_del):
                continue
            cached = wsv.cell(r, c).value
            if cached is None:
                raise ValueError(
                    f"{ws.title}!{cell.coordinate}: valore calcolato assente — "
                    "apri e salva il file in Excel prima del trim"
                )
            cell.value = cached
            changed.append(f"freeze: {ws.title}!{cell.coordinate}")

    # Passata 2: elimina le colonne (contigue per costruzione).
    ws.delete_cols(first_del, n_del)
    c_first_new = first_del  # per costruzione: c_first_old - n_del == first_del

    # Passata 3: marker anno sulla prima superstite, se non più leggibile
    # (era su una colonna eliminata). MergedCell → warn e salta, come extend.
    if not _has_year_marker(ws, c_first_new):
        anno, _ = periodo_anno_mese(cutoff_periodo)
        target = ws.cell(_YEAR_ROW, c_first_new)
        if isinstance(target, MergedCell):
            changed.append(f"warn: marker anno saltato su '{ws.title}' (riga 1 merged)")
        else:
            ws.cell(_YEAR_ROW, c_first_new, anno)
            changed.append(
                f"marker anno: {ws.title}!{get_column_letter(c_first_new)}1 = {anno}"
            )

    # Passata 4 (post-delete): ri-traduce le formule superstiti NON congelate
    # dalla vecchia posizione (colonna attuale + n_del) alla nuova — replica
    # via Translator lo spostamento fisico che delete_cols ha fatto ai valori
    # ma non al testo delle formule.
    for r in range(1, ws.max_row + 1):
        for c in range(c_first_new, ws.max_column + 1):
            cell = ws.cell(r, c)
            v = cell.value
            if not (isinstance(v, str) and v.startswith("=")):
                continue
            old_ref = f"{get_column_letter(c + n_del)}{r}"
            new_ref = f"{get_column_letter(c)}{r}"
            translated = Translator(v, origin=old_ref).translate_formula(new_ref)
            if translated != v:
                cell.value = translated
                changed.append(f"shift: {ws.title}!{cell.coordinate}")

    # Guardia finale: mai output silenziosamente rotto.
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and "#REF!" in cell.value:
                raise ValueError(
                    f"{ws.title}!{cell.coordinate}: #REF! dopo il trim ({cell.value})"
                )

    return changed


def trim_before(wb: Workbook, wb_values: Workbook, cutoff_periodo: int) -> list[str]:
    """Elimina le colonne-mese con periodo < cutoff_periodo. In-place su wb.

    `wb_values` è il gemello data_only=True DELLO STESSO salvataggio di `wb`:
    serve a congelare a valore le formule che, dopo l'eliminazione,
    punterebbero a colonne non più esistenti. Se una cella del genere non ha
    un valore calcolato in cache (file mai aperto/salvato in Excel dopo
    l'ultima modifica), solleva ValueError — non esiste alternativa sicura,
    la formula andrebbe altrimenti persa senza sostituto.
    """
    require_periodo(cutoff_periodo, "cutoff_periodo")
    changed: list[str] = []
    for name in wb.sheetnames:
        changed += _trim_sheet(wb[name], wb_values, cutoff_periodo)
    return changed
