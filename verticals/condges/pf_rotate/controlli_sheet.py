"""Avanzamento dei controlli in-foglio al nuovo mese chiuso.

Le formule del foglio Controlli sono statiche: ancorate al mese chiuso di
quando furono scritte, dopo ogni rotation mostrano ERRORE finti sulla colonna
appena chiusa (thread "control-discrepancy"). Qui le riscriviamo spostando il
confine: mesi chiusi = hardcoded, primo mese aperto = catena dal saldo reale.

Le righe da riscrivere si trovano per pattern (non per posizione fissa) e le
righe saldo/periodo si estraggono dalle formule esistenti — layout-agnostico.
Se il foglio o i pattern non ci sono, no-op.
"""

from __future__ import annotations

import re

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from verticals.condges.pf_rotate.excel_model import MESI_IT, find_month_columns

PF_SHEET = "Piano Finanziario"
_SCAN_ROWS = 30
# 'Piano Finanziario'!D4='Piano Finanziario'!C37
_EQ_RE = re.compile(
    r"'Piano Finanziario'!([A-Z]{1,2})(\d+)='Piano Finanziario'!([A-Z]{1,2})(\d+)"
)

_MESE_NOME = {v: k.lower() for k, v in MESI_IT.items()}


def advance_controlli(wb: Workbook, mese_chiuso: int) -> list[str]:
    """Riscrive i check ancorati al mese nel foglio Controlli. Ritorna le celle toccate."""
    if "Controlli" not in wb.sheetnames or PF_SHEET not in wb.sheetnames:
        return []
    ct = wb["Controlli"]
    cols = find_month_columns(wb[PF_SHEET], header_row=2)
    if mese_chiuso not in cols:
        return []

    closed_col = cols[mese_chiuso]
    first_col = min(cols.values())
    last_col = max(cols.values())
    open_col = closed_col + 1
    if open_col > last_col:
        return []  # nessun mese aperto nel foglio: niente catena da riscrivere

    # Individua le 3 righe per pattern sulla formula in colonna B.
    hard_row = chain_start_row = chain_row = None
    saldo_row = periodo_row = None
    for r in range(1, _SCAN_ROWS + 1):
        v = ct.cell(r, 2).value
        if not isinstance(v, str) or not v.startswith("="):
            continue
        if "ISFORMULA" in v and PF_SHEET in v and hard_row is None:
            hard_row = r
            continue
        m = _EQ_RE.search(v)
        if m is None:
            continue
        if "AND(" in v:
            chain_row = chain_row or r
        else:
            chain_start_row = chain_start_row or r
            saldo_row = int(m.group(2))
            periodo_row = int(m.group(4))

    if saldo_row is None:
        # Senza il check-catena non conosciamo le righe saldo/periodo: no-op.
        return []

    closed_l = get_column_letter(closed_col)
    open_l = get_column_letter(open_col)
    first_l = get_column_letter(first_col)
    last_l = get_column_letter(last_col)
    changed: list[str] = []

    if hard_row is not None:
        # Estrai la riga del saldo iniziale dal check esistente (fallback: saldo_row).
        m = re.search(
            r"ISFORMULA\('Piano Finanziario'![A-Z]{1,2}(\d+)\)",
            ct.cell(hard_row, 2).value,
        )
        hard_target_row = int(m.group(1)) if m else saldo_row
        refs = ",".join(
            f"_xlfn.ISFORMULA('Piano Finanziario'!{get_column_letter(c)}{hard_target_row})"
            for c in range(first_col, closed_col + 1)
        )
        body = f"OR({refs})" if closed_col > first_col else refs
        ct.cell(hard_row, 1).value = (
            f"{first_l}{hard_target_row}:{closed_l}{hard_target_row} = "
            "saldi iniziali hardcoded (mesi chiusi)?"
        )
        ct.cell(hard_row, 2).value = f'=IF({body},"ERRORE","OK")'
        changed.append(f"B{hard_row}")

    def _eq(col: int) -> str:
        cl = get_column_letter(col)
        prev = get_column_letter(col - 1)
        return f"'Piano Finanziario'!{cl}{saldo_row}='Piano Finanziario'!{prev}{periodo_row}"

    if chain_start_row is not None:
        mese_aperto = mese_chiuso % 12 + 1
        ct.cell(chain_start_row, 1).value = (
            f"{open_l}{saldo_row} = {closed_l}{periodo_row} "
            f"({_MESE_NOME[mese_aperto]} parte da saldo reale {_MESE_NOME[mese_chiuso]})?"
        )
        ct.cell(chain_start_row, 2).value = f'=IF({_eq(open_col)},"OK","ERRORE")'
        changed.append(f"B{chain_start_row}")

    if chain_row is not None:
        eqs = ",".join(_eq(c) for c in range(open_col, last_col + 1))
        ct.cell(
            chain_row, 1
        ).value = f"Catena {open_l}{saldo_row}:{last_l}{saldo_row} = mese prec riga {periodo_row}?"
        ct.cell(chain_row, 2).value = f'=IF(AND({eqs}),"OK","ERRORE")'
        changed.append(f"B{chain_row}")

    return changed
