"""Motore di piazzamento scadenzario → fogli dettaglio del PF.

Unico writer condiviso (DRY) da CLI (`pf-rotate`, `scad`) e Streamlit
(`app_scadenzario.py`). Estratto da `app_scadenzario.py` per rimuovere la
dipendenza invertita (la pipeline `pf_rotate`, testata, importava dal modulo
UI Streamlit). Vedi spec `docs/superpowers/specs/2026-06-01-scadenzario-placement-unification.md`.

Layout-tollerante: la mappa colonna↔mese è derivata via
`excel_model.find_month_columns` (header riga 2).
"""

from __future__ import annotations

import re
from datetime import date
from io import BytesIO

import openpyxl
import pandas as pd

from core.schemas import FornitoreMapRow
from verticals.condges.pf_rotate.excel_model import find_month_columns

# -- Voce → foglio / etichette -------------------------------------------------

# Candidati di nome-foglio per voce_id: tollera i typo storici nei template PF
# di Rosa (es. "Materie Prime-Conumo ", spazi iniziali/finali).
VOCE_TO_SHEET_CANDIDATES: dict[str, list[str]] = {
    "USCITE_MATERIE_PRIME": ["Materie Prime-Consumo ", "Materie Prime-Conumo "],
    "USCITE_UTENZE": ["Utenze"],
    "USCITE_SALARI": ["Salari e Stipendi"],
    "USCITE_TASSE": ["Tasse e Imposte"],
    "USCITE_COMMISSIONI": ["Commisisoni Portali"],
    "USCITE_MUTUI": ["Mutui e Finaziamenti"],
    "USCITE_CONSULENZE": ["Consulenze"],
    "USCITE_CANONE_PASSIVO": ["Godimento Beni di Terzi"],
    "USCITE_VARIE_EXT": [" Varie ed Eventuali"],
    "USCITE_SERVIZI_PRODUZIONE": ["Canoni e servizi"],
}

# Mapping piatto per display (usa il primo candidato).
VOCE_TO_SHEET: dict[str, str] = {
    k: v[0] for k, v in VOCE_TO_SHEET_CANDIDATES.items()
}

VOCE_LABELS: dict[str, str] = {
    "USCITE_MATERIE_PRIME": "Materie Prime",
    "USCITE_UTENZE": "Utenze",
    "USCITE_SALARI": "Salari e Stipendi",
    "USCITE_TASSE": "Tasse e Imposte",
    "USCITE_COMMISSIONI": "Commissioni",
    "USCITE_MUTUI": "Mutui e Finanziamenti",
    "USCITE_CONSULENZE": "Consulenze",
    "USCITE_CANONE_PASSIVO": "Godimento Beni di Terzi",
    "USCITE_VARIE_EXT": "Varie ed Eventuali",
    "USCITE_SERVIZI_PRODUZIONE": "Canoni e servizi",
}

MESI_NOMI: list[str] = [
    "Gen",
    "Feb",
    "Mar",
    "Apr",
    "Mag",
    "Giu",
    "Lug",
    "Ago",
    "Set",
    "Ott",
    "Nov",
    "Dic",
]


# -- Helper di ricerca righe/colonne nel foglio dettaglio ----------------------


def resolve_sheet_name(voce_id: str, available_sheets: list[str]) -> str | None:
    """Trova il nome-foglio reale per un voce_id, gestendo i typo tra file PF."""
    candidates = VOCE_TO_SHEET_CANDIDATES.get(voce_id, [])
    for name in candidates:
        if name in available_sheets:
            return name
    return None


def _find_previsionale_row(ws, max_row: int = 200) -> int | None:
    """Trova la riga PREVISIONALE (col B contiene 'previsional')."""
    for r in range(3, min(ws.max_row + 1, max_row)):
        val = ws.cell(row=r, column=2).value
        if val and "previsional" in str(val).strip().lower():
            return r
    return None


def _find_total_row_and_range(
    ws,
    month_col: dict[int, int],
    max_row: int = 10,
) -> tuple[int | None, int, int]:
    """Trova la riga totale con SUM verticale in una colonna mese.

    Ritorna (total_row, sum_start_row, sum_end_row).
    """
    for r in range(3, min(ws.max_row + 1, max_row)):
        for month, col in month_col.items():
            val = ws.cell(row=r, column=col).value
            if not val or not isinstance(val, str):
                continue
            m = re.search(r"SUM\([A-Z]+(\d+):[A-Z]+(\d+)\)", val)
            if m:
                start = int(m.group(1))
                end = int(m.group(2))
                if end - start > 5:  # SUM verticale copre molte righe
                    return r, start, end
    return None, 0, 0


def _find_supplier_row_by_name(ws, nome_pf: str, max_row: int = 200) -> int | None:
    """Trova la riga dove col B == nome_pf (esatto, poi containment ≥4 char)."""
    if not nome_pf or not nome_pf.strip():
        return None
    target = nome_pf.strip().lower()
    # Match esatto
    for r in range(3, min(ws.max_row + 1, max_row)):
        val = ws.cell(row=r, column=2).value
        if val and str(val).strip().lower() == target:
            return r
    # Containment: in entrambe le direzioni, min 4 char per evitare falsi positivi
    if len(target) >= 4:
        for r in range(3, min(ws.max_row + 1, max_row)):
            val = ws.cell(row=r, column=2).value
            if not val:
                continue
            pf_name = str(val).strip().lower()
            if len(pf_name) < 4:
                continue
            if target in pf_name or pf_name in target:
                return r
    return None


def _find_supplier_row_by_codice(ws, codice: int, max_row: int = 200) -> int | None:
    """Trova la riga dove col A == codice_fornitore."""
    for r in range(3, min(ws.max_row + 1, max_row)):
        val = ws.cell(row=r, column=1).value
        if val and isinstance(val, (int, float)) and int(val) == codice:
            return r
    return None


# -- Writer principale ---------------------------------------------------------


def write_pf(
    pf_bytes: bytes,
    scad_df: pd.DataFrame,
    bucket_months: list[int],
    fornitori_map: dict[int, FornitoreMapRow],
    excluded: set[int] | None = None,
) -> tuple[bytes, dict[str, list]]:
    """Scrivi le scadenze in TUTTI i fogli dettaglio del PF, ritorna (bytes, summary).

    Gestisce l'aggiustamento PREVISIONALE: se la riga PREVISIONALE è dentro
    il range del SUM totale, la riduce dell'importo scadenzario in modo che
    il SUM resti corretto (= MAX(previsionale, scadenzario)).

    `fornitori_map` mappa codice_fornitore → FornitoreMapRow (loader canonico
    `pf_rotate.fornitori_map.load_fornitori`). `excluded` salta i codici dati.
    """
    wb = openpyxl.load_workbook(BytesIO(pf_bytes))
    wb_values = openpyxl.load_workbook(BytesIO(pf_bytes), data_only=True)
    wb_formulas = openpyxl.load_workbook(BytesIO(pf_bytes), data_only=False)

    current_month = date.today().month
    summary: dict[str, list] = {}

    # Raggruppa i fornitori dello scadenzario per voce_id
    excluded = excluded or set()
    scad_by_voce: dict[str, list[dict]] = {}
    for _, row in scad_df.iterrows():
        codice = int(row["codice_fornitore"])
        if codice in excluded:
            continue
        info = fornitori_map.get(codice)
        if not info:
            continue
        voce_id = info.voce_id
        nome_pf = info.nome_pf
        scad_by_voce.setdefault(voce_id, []).append(
            {
                "codice_fornitore": codice,
                "nome": row["nome"],
                "nome_pf": nome_pf,
                "scaduto": float(row.get("scaduto", 0) or 0),
                **{
                    f"mese_{m}": float(row.get(f"mese_{m}", 0) or 0)
                    for m in bucket_months
                },
            }
        )

    for voce_id, suppliers in scad_by_voce.items():
        sheet_name = resolve_sheet_name(voce_id, wb.sheetnames)
        if not sheet_name:
            continue

        ws = wb[sheet_name]
        ws_vals = wb_values[sheet_name]
        ws_form = wb_formulas[sheet_name]

        month_col = find_month_columns(ws_vals)
        prev_row = _find_previsionale_row(ws_vals)
        month_col_form = find_month_columns(ws_form)
        total_row, sum_start, sum_end = _find_total_row_and_range(
            ws_form, month_col_form
        )
        prev_in_sum = False
        if prev_row and total_row:
            prev_in_sum = sum_start <= prev_row <= sum_end

        written: list[dict] = []
        # Totale scadenzario scritto per mese (per aggiustamento PREVISIONALE)
        scad_totals: dict[int, float] = {}

        # Righe vuote dentro il range SUM per nuovi fornitori (evita insert_rows
        # che corrompe le formule e crea riferimenti circolari)
        empty_rows: list[int] = []
        search_start = sum_start if sum_start else 4
        search_end = sum_end if sum_end else (ws.max_row - 1)
        skip_rows = set()
        if prev_row:
            skip_rows.add(prev_row)
        if total_row:
            skip_rows.add(total_row)
        for r in range(search_start, search_end + 1):
            if r in skip_rows:
                continue
            a = ws_vals.cell(row=r, column=1).value
            b = ws_vals.cell(row=r, column=2).value
            if not a and not b:
                empty_rows.append(r)
        empty_row_idx = 0

        for s in suppliers:
            # Trova la riga: prima per codice, poi per nome
            pf_row = _find_supplier_row_by_codice(ws_vals, s["codice_fornitore"])
            if pf_row is None and s["nome_pf"]:
                pf_row = _find_supplier_row_by_name(ws_vals, s["nome_pf"])
            if pf_row is None:
                pf_row = _find_supplier_row_by_name(ws_vals, s["nome"])
            if pf_row is None and empty_row_idx < len(empty_rows):
                # Usa una riga vuota esistente invece di inserire
                pf_row = empty_rows[empty_row_idx]
                empty_row_idx += 1
                ws.cell(row=pf_row, column=1, value=s["codice_fornitore"])
                ws.cell(row=pf_row, column=2, value=s["nome_pf"] or s["nome"])
            if pf_row is None:
                continue

            # Importi: scaduto → mese corrente, bucket → loro mesi
            amounts: dict[int, float] = {}
            scaduto = s["scaduto"]
            if scaduto:
                amounts[current_month] = amounts.get(current_month, 0) + scaduto

            for month in bucket_months:
                val = s[f"mese_{month}"]
                if val:
                    amounts[month] = amounts.get(month, 0) + val

            # Scrivi: flip segno (scadenze negative = debito, PF positivo = uscita)
            months_written: dict[int, float] = {}
            for month, amount in amounts.items():
                col = month_col.get(month)
                if col:
                    pf_val = round(abs(amount), 2)
                    ws.cell(row=pf_row, column=col, value=pf_val)
                    months_written[month] = pf_val
                    scad_totals[month] = scad_totals.get(month, 0) + pf_val

            if months_written:
                written.append(
                    {
                        "codice": s["codice_fornitore"],
                        "nome": s["nome_pf"] or s["nome"],
                        "months": months_written,
                    }
                )

        # Aggiusta PREVISIONALE se è dentro il range SUM.
        # Dopo aver scritto le celle fornitore, calcola il totale di TUTTE le
        # righe non-prev nel range SUM, poi imposta PREVISIONALE così che:
        #   SUM = MAX(orig_previsionale, supplier_total)
        if prev_in_sum and prev_row and scad_totals:
            for month in scad_totals:
                col = month_col.get(month)
                if not col:
                    continue
                orig_prev = ws_vals.cell(row=prev_row, column=col).value
                if not orig_prev or not isinstance(orig_prev, (int, float)):
                    continue
                orig_prev_f = float(orig_prev)
                # Somma tutte le righe non-previsionale nel range SUM
                supplier_total = 0.0
                for r in range(sum_start, sum_end + 1):
                    if r == prev_row:
                        continue
                    v = ws.cell(row=r, column=col).value
                    if v and isinstance(v, (int, float)):
                        supplier_total += float(v)
                # Imposta PREV così che SUM = MAX(orig_prev, supplier_total)
                new_prev = max(0.0, orig_prev_f - supplier_total)
                ws.cell(row=prev_row, column=col, value=round(new_prev, 2))

        if written:
            summary[VOCE_LABELS.get(voce_id, voce_id)] = written

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue(), summary
