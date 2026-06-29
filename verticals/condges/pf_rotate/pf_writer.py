"""Motore di scrittura PF — estratto da app_scadenzario, casa engine accanto a pf_rotate.

Funzioni engine (read + write) che operano sui fogli dettaglio del Piano Finanziario.
Spostate verbatim da verticals/condges/app_scadenzario.py (deprecato).
"""

from __future__ import annotations

import csv
import logging
import re
from datetime import date
from io import BytesIO
from pathlib import Path

import openpyxl
import pandas as pd

from verticals.condges.pf_generator.blocchi import cascata_nc
from verticals.condges.pf_rotate.fornitori_map import VOCE_LABELS

log = logging.getLogger(__name__)

# -- Config --------------------------------------------------------------------

FORNITORI_CSV = (
    Path(__file__).resolve().parents[3]
    / "core"
    / "bq"
    / "dimensioni"
    / "d_fornitori.csv"
)

VOCE_TO_SHEET_CANDIDATES = {
    "USCITE_MATERIE_PRIME": ["Materie Prime-Consumo ", "Materie Prime-Conumo "],
    "USCITE_UTENZE": ["Utenze"],
    "USCITE_SALARI": ["Salari e Stipendi"],
    "USCITE_TASSE": ["Tasse e Imposte"],
    "USCITE_COMMISSIONI": ["Commisisoni Portali"],
    "USCITE_MUTUI": ["Mutui e Finaziamenti"],
    "USCITE_CONSULENZE": ["Consulenze"],
    "USCITE_CANONE_PASSIVO": ["Godimento Beni di Terzi"],
    "USCITE_VARIE_EXT": [" Varie ed Eventuali"],
    "USCITE_SERVIZI_PRODUZIONE": ["Canoni e servizi"],
}

MONTH_NAMES_IT = {
    "GENNAIO": 1,
    "FEBBRAIO": 2,
    "MARZO": 3,
    "APRILE": 4,
    "MAGGIO": 5,
    "GIUGNO": 6,
    "LUGLIO": 7,
    "AGOSTO": 8,
    "SETTEMBRE": 9,
    "OTTOBRE": 10,
    "NOVEMBRE": 11,
    "DICEMBRE": 12,
}

MESI_NOMI = [
    "Gen",
    "Feb",
    "Mar",
    "Apr",
    "Mag",
    "Giu",
    "Lug",
    "Ago",
    "Set",
    "Ott",
    "Nov",
    "Dic",
]


# -- Fornitori map -------------------------------------------------------------


def load_fornitori_map() -> dict[int, dict]:
    """Load d_fornitori CSV -> {codice_fornitore: {voce_id, nome_pf}}."""
    result = {}
    with open(FORNITORI_CSV, newline="", encoding="utf-8") as f:
        for row in csv.DictReader(f):
            result[int(row["codice_fornitore"])] = {
                "voce_id": row["voce_id"],
                "nome_pf": row.get("nome_pf", "").strip(),
            }
    return result


# -- Engine functions ----------------------------------------------------------


def resolve_sheet_name(voce_id: str, available_sheets: list[str]) -> str | None:
    """Find the actual sheet name for a voce_id, handling typos across PF files."""
    candidates = VOCE_TO_SHEET_CANDIDATES.get(voce_id, [])
    for name in candidates:
        if name in available_sheets:
            return name
    return None


def _build_month_col_map(ws) -> dict[int, int]:
    """Scan row 2 of a detail sheet, return {calendar_month: column}."""
    col_map: dict[int, int] = {}
    for col in range(1, ws.max_column + 1):
        val = ws.cell(row=2, column=col).value
        if val and str(val).strip().upper() in MONTH_NAMES_IT:
            month = MONTH_NAMES_IT[str(val).strip().upper()]
            if month not in col_map or col > col_map[month]:
                col_map[month] = col
    return col_map


def _find_previsionale_row(ws, max_row: int = 200) -> int | None:
    """Find the PREVISIONALE row (col B contains 'PREVISIONALE')."""
    for r in range(3, min(ws.max_row + 1, max_row)):
        val = ws.cell(row=r, column=2).value
        if val and "previsional" in str(val).strip().lower():
            return r
    return None


def _find_total_row_and_range(
    ws,
    month_col: dict[int, int],
    max_row: int = 10,
) -> tuple[int | None, int, int]:
    """Find the total row with a vertical SUM formula in a month column.

    Returns (total_row, sum_start_row, sum_end_row).
    """
    # Check month columns for vertical SUM formulas like =SUM(L5:L180)
    for r in range(3, min(ws.max_row + 1, max_row)):
        for month, col in month_col.items():
            val = ws.cell(row=r, column=col).value
            if not val or not isinstance(val, str):
                continue
            m = re.search(r"SUM\([A-Z]+(\d+):[A-Z]+(\d+)\)", val)
            if m:
                start = int(m.group(1))
                end = int(m.group(2))
                if end - start > 5:  # vertical SUM spans many rows
                    return r, start, end
    return None, 0, 0


def _find_supplier_row_by_name(ws, nome_pf: str, max_row: int = 200) -> int | None:
    """Find row in detail sheet where column B matches nome_pf.

    Tries exact match first, then containment in either direction
    (PF name contains search term, or search term contains PF name).
    """
    if not nome_pf or not nome_pf.strip():
        return None
    target = nome_pf.strip().lower()
    # Exact match first
    for r in range(3, min(ws.max_row + 1, max_row)):
        val = ws.cell(row=r, column=2).value
        if val and str(val).strip().lower() == target:
            return r
    # Containment match: either direction, min 4 chars to avoid false positives
    if len(target) >= 4:
        for r in range(3, min(ws.max_row + 1, max_row)):
            val = ws.cell(row=r, column=2).value
            if not val:
                continue
            pf_name = str(val).strip().lower()
            if len(pf_name) < 4:
                continue
            if target in pf_name or pf_name in target:
                return r
    return None


def _find_supplier_row_by_codice(ws, codice: int, max_row: int = 200) -> int | None:
    """Find row in detail sheet where column A matches codice_fornitore."""
    for r in range(3, min(ws.max_row + 1, max_row)):
        val = ws.cell(row=r, column=1).value
        if val and isinstance(val, (int, float)) and int(val) == codice:
            return r
    return None


def read_pf_sheet(
    wb_values,
    wb_formulas,
    sheet_name: str,
) -> tuple[pd.DataFrame, dict[int, int], int | None, bool]:
    """Read a PF detail sheet.

    Returns: (df, month_col, previsionale_row, prev_in_sum)
    """
    ws = wb_values[sheet_name]
    ws_cod = wb_formulas[sheet_name]

    month_col = _build_month_col_map(ws)
    prev_row = _find_previsionale_row(ws)
    month_col_form = _build_month_col_map(ws_cod)
    total_row, sum_start, sum_end = _find_total_row_and_range(ws_cod, month_col_form)

    prev_in_sum = False
    if prev_row and total_row:
        prev_in_sum = sum_start <= prev_row <= sum_end

    # Read supplier rows (match by codice in col 1 or name in col 2)
    rows = []
    for r in range(3, ws.max_row + 1):
        codice = ws_cod.cell(row=r, column=1).value
        nome = ws.cell(row=r, column=2).value
        if r == prev_row or r == total_row:
            continue
        if not nome:
            continue
        row_data = {"nome_pf": str(nome).strip(), "pf_row": r}
        if codice and isinstance(codice, (int, float)):
            row_data["codice_fornitore"] = int(codice)
        for month, col in month_col.items():
            val = ws.cell(row=r, column=col).value
            row_data[f"pf_mese_{month}"] = (
                float(val) if isinstance(val, (int, float)) else 0.0
            )
        rows.append(row_data)

    df = pd.DataFrame(rows)
    if df.empty:
        df = pd.DataFrame(columns=["codice_fornitore", "nome_pf", "pf_row"])
    return df, month_col, prev_row, prev_in_sum


def write_pf(
    pf_bytes: bytes,
    scad_df: pd.DataFrame,
    bucket_months: list[int],
    fornitori_map: dict[int, dict],
    excluded: set[int] | None = None,
    scaduto_month: int | None = None,
    clear_codici: set[int] | None = None,
) -> tuple[bytes, dict[str, list]]:
    """Write scadenze into ALL PF detail sheets, return (bytes, summary).

    ``scaduto_month``: mese in cui scrivere il bucket 'scaduto'. Per la
    rotation passare il primo mese aperto (mese_chiuso+1), così il risultato
    non dipende dal giorno del run. Default: mese di oggi (path app legacy).

    ``clear_codici``: codici fornitore le cui righe vengono ripulite nei mesi
    aperti (>= scaduto_month) PRIMA di riscrivere — rende il run idempotente
    e rimuove scritture stantie di run precedenti. Tocca solo righe con
    codice fornitore in col A; righe manuali (codici conto PF) restano intatte.

    Handles the PREVISIONALE adjustment: if the PREVISIONALE row is inside
    the SUM range of the total row, reduce it by the scadenzario total so
    the overall SUM stays correct (= MAX(previsionale, scadenzario)).
    """
    wb = openpyxl.load_workbook(BytesIO(pf_bytes))
    wb_values = openpyxl.load_workbook(BytesIO(pf_bytes), data_only=True)
    wb_formulas = openpyxl.load_workbook(BytesIO(pf_bytes), data_only=False)

    current_month = scaduto_month if scaduto_month is not None else date.today().month
    summary: dict[str, list] = {}

    # Group scadenzario suppliers by voce_id
    excluded = excluded or set()
    scad_by_voce: dict[str, list[dict]] = {}
    for _, row in scad_df.iterrows():
        codice = int(row["codice_fornitore"])
        if codice in excluded:
            continue
        info = fornitori_map.get(codice)
        if not info:
            continue
        voce_id = info["voce_id"]
        nome_pf = info["nome_pf"]
        scad_by_voce.setdefault(voce_id, []).append(
            {
                "codice_fornitore": codice,
                "nome": row["nome"],
                "nome_pf": nome_pf,
                "scaduto": float(row.get("scaduto", 0) or 0),
                **{
                    f"mese_{m}": float(row.get(f"mese_{m}", 0) or 0)
                    for m in bucket_months
                },
            }
        )

    # Pulizia scritture stantie su TUTTI i fogli voce, anche quelli che non
    # ricevono scritture in questo run (es. fornitore rimappato a un'altra
    # voce: la riga nel vecchio foglio va comunque azzerata nei mesi aperti).
    if clear_codici:
        for voce_id in VOCE_TO_SHEET_CANDIDATES:
            sheet_name = resolve_sheet_name(voce_id, wb.sheetnames)
            if not sheet_name:
                continue
            ws_c = wb[sheet_name]
            ws_cv = wb_values[sheet_name]
            mc = _build_month_col_map(ws_cv)
            open_cols = [c for m, c in mc.items() if m >= current_month]
            if not open_cols:
                continue
            for r in range(4, ws_c.max_row + 1):
                a = ws_cv.cell(row=r, column=1).value
                if isinstance(a, (int, float)) and int(a) in clear_codici:
                    for col in open_cols:
                        ws_c.cell(row=r, column=col).value = None

    for voce_id, suppliers in scad_by_voce.items():
        sheet_name = resolve_sheet_name(voce_id, wb.sheetnames)
        if not sheet_name:
            log.warning(
                "voce_id '%s' non ha un foglio dettaglio nel PF (%d fornitori skippati). "
                "Aggiungere il foglio per includere questi importi (punto B/#51).",
                voce_id,
                len(suppliers),
            )
            continue

        ws = wb[sheet_name]
        ws_vals = wb_values[sheet_name]
        ws_form = wb_formulas[sheet_name]

        month_col = _build_month_col_map(ws_vals)
        prev_row = _find_previsionale_row(ws_vals)
        month_col_form = _build_month_col_map(ws_form)
        total_row, sum_start, sum_end = _find_total_row_and_range(
            ws_form, month_col_form
        )
        prev_in_sum = False
        if prev_row and total_row:
            prev_in_sum = sum_start <= prev_row <= sum_end

        voce_label = VOCE_LABELS.get(voce_id, voce_id)
        written: list[dict] = []
        # Track total scadenzario written per month for PREVISIONALE adjustment
        scad_totals: dict[int, float] = {}

        # Find empty rows inside SUM range for new suppliers (avoid insert_rows
        # which corrupts formulas and creates circular references)
        empty_rows: list[int] = []
        search_start = sum_start if sum_start else 4
        search_end = sum_end if sum_end else (ws.max_row - 1)
        # Skip previsionale and total rows
        skip_rows = set()
        if prev_row:
            skip_rows.add(prev_row)
        if total_row:
            skip_rows.add(total_row)
        for r in range(search_start, search_end + 1):
            if r in skip_rows:
                continue
            a = ws_vals.cell(row=r, column=1).value
            b = ws_vals.cell(row=r, column=2).value
            if not a and not b:
                empty_rows.append(r)
        empty_row_idx = 0

        for s in suppliers:
            # Find the supplier row: try codice first, then name
            pf_row = _find_supplier_row_by_codice(ws_vals, s["codice_fornitore"])
            if pf_row is None and s["nome_pf"]:
                pf_row = _find_supplier_row_by_name(ws_vals, s["nome_pf"])
            if pf_row is None:
                pf_row = _find_supplier_row_by_name(ws_vals, s["nome"])
            if pf_row is None and empty_row_idx < len(empty_rows):
                # Use an existing empty row instead of inserting
                pf_row = empty_rows[empty_row_idx]
                empty_row_idx += 1
                ws.cell(row=pf_row, column=1, value=s["codice_fornitore"])
                ws.cell(row=pf_row, column=2, value=s["nome_pf"] or s["nome"])
            if pf_row is None:
                continue

            # Collect amounts: scaduto -> current month, buckets -> their months
            amounts: dict[int, float] = {}
            scaduto = s["scaduto"]
            if scaduto:
                amounts[current_month] = amounts.get(current_month, 0) + scaduto

            for month in bucket_months:
                val = s[f"mese_{month}"]
                if val:
                    amounts[month] = amounts.get(month, 0) + val

            # Note credito scalate in cascata (debiti negativi, NC positiva).
            netted = cascata_nc(amounts)

            # Write: flip sign (scadenze negative = debito, PF positive = uscita)
            months_written: dict[int, float] = {}
            for month, amount in netted.items():
                col = month_col.get(month)
                if col:
                    pf_val = round(abs(amount), 2)
                    ws.cell(row=pf_row, column=col, value=pf_val)
                    months_written[month] = pf_val
                    scad_totals[month] = scad_totals.get(month, 0) + pf_val

            if months_written:
                written.append(
                    {
                        "codice": s["codice_fornitore"],
                        "nome": s["nome_pf"] or s["nome"],
                        "months": months_written,
                    }
                )

        # Adjust PREVISIONALE if it's inside the SUM range.
        # After writing supplier cells, compute the total of ALL non-prev
        # rows in the SUM range, then set PREVISIONALE so that:
        #   SUM = MAX(orig_previsionale, supplier_total)
        if prev_in_sum and prev_row and scad_totals:
            for month in scad_totals:
                col = month_col.get(month)
                if not col:
                    continue
                orig_prev = ws_vals.cell(row=prev_row, column=col).value
                if not orig_prev or not isinstance(orig_prev, (int, float)):
                    continue
                orig_prev_f = float(orig_prev)
                # Sum all non-previsionale rows in the SUM range
                supplier_total = 0.0
                for r in range(sum_start, sum_end + 1):
                    if r == prev_row:
                        continue
                    v = ws.cell(row=r, column=col).value
                    if v and isinstance(v, (int, float)):
                        supplier_total += float(v)
                # Set PREV so SUM = MAX(orig_prev, supplier_total)
                new_prev = max(0.0, orig_prev_f - supplier_total)
                ws.cell(row=prev_row, column=col, value=round(new_prev, 2))

        if written:
            summary[voce_label] = written

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue(), summary
