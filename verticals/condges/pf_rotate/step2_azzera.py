"""Step 2 della rotation — azzera la colonna del mese chiuso.

Regola: solo celle-valore → None. Mai toccare formule con riferimenti.
"""

from __future__ import annotations

from dataclasses import dataclass

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from verticals.condges.pf_rotate.excel_model import (
    find_layout,
    find_month_columns,
    is_value_cell,
)

# Fogli da scansionare (master + dettagli canonici). I dettagli usano gli alias
# typo-tolerant — vengono filtrati a quelli effettivamente presenti.
DETAIL_SHEET_ALIASES = [
    "Salari e Stipendi",
    "Utenze",
    ["Materie Prime-Consumo ", "Materie Prime-Conumo "],
    "Tasse e Imposte",
    "Commisisoni Portali",
    "Mutui e Finaziamenti",
    "Consulenze",
    "Godimento Beni di Terzi",
    " Varie ed Eventuali",
    "Canoni e servizi",
]

# Detail: dalla r3 in giù. Anatomia REALE dei fogli dettaglio: r3 = costante
# manuale (previsione del mese, es. Materie Prime 40k), r4 = =SUM(r5:...),
# r5+ = fornitori. La vecchia assunzione "r3 = riga totale" era invertita e
# lasciava vive le previsioni manuali del mese chiuso (leak nella cascata via
# formule master protette). is_value_cell salta comunque le formule.
DETAIL_VALUE_START_ROW = 3


@dataclass(frozen=True)
class CellChange:
    sheet: str
    ref: str
    old: object
    new: object


def _resolve(wb: Workbook, alias_spec: str | list[str]) -> str | None:
    """Return existing sheet name from a string or list-of-aliases spec."""
    if isinstance(alias_spec, str):
        return alias_spec if alias_spec in wb.sheetnames else None
    for name in alias_spec:
        if name in wb.sheetnames:
            return name
    return None


def azzera_mese(wb: Workbook, mese_chiuso: int) -> list[CellChange]:
    """Svuota le celle-valore della colonna del mese chiuso in master + dettagli.

    Returns list of CellChange recording every overwrite.
    """
    changes: list[CellChange] = []

    # ── Master ─────────────────────────────────────────────────────────
    if "Piano Finanziario" not in wb.sheetnames:
        raise ValueError("Foglio 'Piano Finanziario' assente — input non valido.")
    pf = wb["Piano Finanziario"]
    mese_cols_master = find_month_columns(pf, header_row=2)
    if mese_chiuso not in mese_cols_master:
        raise ValueError(f"Mese chiuso {mese_chiuso} non trovato nel master.")
    col_master = mese_cols_master[mese_chiuso]
    col_letter = get_column_letter(col_master)

    layout = find_layout(wb)
    master_value_rows = list(layout.entrate_rows)

    for r in master_value_rows:
        cell = pf.cell(r, col_master)
        if is_value_cell(cell):
            old = cell.value
            cell.value = None
            changes.append(
                CellChange("Piano Finanziario", f"{col_letter}{r}", old, None)
            )

    # ── Dettagli ───────────────────────────────────────────────────────
    for spec in DETAIL_SHEET_ALIASES:
        name = _resolve(wb, spec)
        if name is None:
            continue
        ws = wb[name]
        mese_cols = find_month_columns(ws, header_row=2)
        if mese_chiuso not in mese_cols:
            continue
        col = mese_cols[mese_chiuso]
        cl = get_column_letter(col)
        for r in range(DETAIL_VALUE_START_ROW, ws.max_row + 1):
            cell = ws.cell(r, col)
            if is_value_cell(cell):
                old = cell.value
                cell.value = None
                changes.append(CellChange(name, f"{cl}{r}", old, None))

    return changes
