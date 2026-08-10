"""Excel model primitives for PF rotation.

Pure functions on openpyxl Workbook/Worksheet. No I/O, no BQ.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from typing import Iterable

from openpyxl.worksheet.worksheet import Worksheet

MESI_IT = {
    "GENNAIO": 1,
    "FEBBRAIO": 2,
    "MARZO": 3,
    "APRILE": 4,
    "MAGGIO": 5,
    "GIUGNO": 6,
    "LUGLIO": 7,
    "AGOSTO": 8,
    "SETTEMBRE": 9,
    "OTTOBRE": 10,
    "NOVEMBRE": 11,
    "DICEMBRE": 12,
}

PERIODO_MIN = 2000 * 12  # nessun periodo reale è sotto il 2000


def periodo(anno: int, mese: int) -> int:
    """Chiave periodo ordinale: gen 2026 e gen 2027 sono chiavi diverse."""
    return anno * 12 + (mese - 1)


def periodo_anno_mese(p: int) -> tuple[int, int]:
    anno, m0 = divmod(p, 12)
    return anno, m0 + 1


def require_periodo(p: int, nome: str = "periodo") -> None:
    """Guardia: periodo e mese nudo sono entrambi int — qui si separano."""
    if p < PERIODO_MIN:
        raise ValueError(
            f"{nome}={p} sembra un mese nudo (1-12) o un anno: serve un periodo "
            f"ordinale anno*12+(mese-1), es. periodo(2026, 6) = {2026 * 12 + 5}"
        )


def find_month_periods(
    ws: Worksheet, header_row: int = 2, year_row: int = 1
) -> dict[int, int]:
    """Scan header_row per mesi ITA, datati per anno. Return {periodo: col_idx}.

    L'anno base è l'ultima cella numerica 1900<v<2100 in year_row fino alla
    prima colonna-mese inclusa (ORTI master: C1; fogli dettaglio: D1).
    L'anno incrementa a ogni wrap (numero mese che scende, es. DIC→GEN).
    Niente inferenze dall'orologio: senza anno dichiarato → ValueError.
    """
    months: list[tuple[int, int]] = []  # (col, mese 1-12) in ordine di colonna
    first_month_col: int | None = None
    for c in range(1, ws.max_column + 1):
        v = ws.cell(header_row, c).value
        if isinstance(v, str) and v.strip().upper() in MESI_IT:
            if first_month_col is None:
                first_month_col = c
            months.append((c, MESI_IT[v.strip().upper()]))
    if not months:
        return {}
    base_year: int | None = None
    for c in range(1, first_month_col + 1):
        v = ws.cell(year_row, c).value
        if isinstance(v, (int, float)) and 1900 < int(v) < 2100:
            base_year = int(v)
    if base_year is None:
        raise ValueError(
            f"Anno non dichiarato in riga {year_row} del foglio '{ws.title}': "
            "impossibile datare le colonne-mese."
        )
    result: dict[int, int] = {}
    year = base_year
    prev_mese: int | None = None
    for col, mese in months:
        if prev_mese is not None and mese < prev_mese:
            year += 1
        prev_mese = mese
        result[periodo(year, mese)] = col
    return result


def find_month_columns(ws: Worksheet, header_row: int = 2) -> dict[int, int]:
    """DEPRECATA (compat transitoria): mese nudo → col. Latest col wins.

    Cancellare quando tutti i caller usano find_month_periods (Task 6).
    """
    out: dict[int, int] = {}
    for p, col in find_month_periods(ws, header_row=header_row).items():
        out[periodo_anno_mese(p)[1]] = col  # latest wins per costruzione (ordine col)
    return out


# Matches an A1-style reference (e.g. A1, $B$5, 'Sheet Name'!D3, Utenze!D3).
# Avoid matching function names by requiring a digit somewhere in the token.
_A1_REF_RE = re.compile(
    r"(?:'[^']+'|[A-Z][A-Za-z0-9_ ]*)?!?\$?[A-Z]{1,3}\$?\d+",
)


def is_formula_with_refs(cell) -> bool:
    """True if cell value is a formula string containing A1-style references."""
    v = cell.value
    if not isinstance(v, str) or not v.startswith("="):
        return False
    # Strip the leading '='
    body = v[1:]
    return _A1_REF_RE.search(body) is not None


def is_value_cell(cell) -> bool:
    """True if cell holds a value (number, string, formula-of-constants).

    False for None and formulas that reference other cells.
    """
    v = cell.value
    if v is None:
        return False
    if isinstance(v, (int, float)):
        return True
    if isinstance(v, str):
        if not v.startswith("="):
            return True
        # Formula: only-constants → value-equivalent; with refs → not value.
        return not is_formula_with_refs(cell)
    return False


def resolve_sheet_name(wb, candidates: Iterable[str]) -> str | None:
    """Return the first candidate that exists in wb.sheetnames, else None."""
    names = set(wb.sheetnames)
    for c in candidates:
        if c in names:
            return c
    return None


@dataclass(frozen=True)
class PFLayout:
    saldo_iniziale_row: int  # row of "SALDO MESE PREC..." label
    totale_entrate_row: int  # row of "TOTALE ENTRATE"
    totale_uscite_row: int  # row of "TOTALE USCITE"
    cashflow_row: int  # row of "Cash Flow" or "CASH FLOW"
    saldi_banca_rows: list[int]  # rows of "Saldo <banca>" labels
    totale_banche_row: int  # row of "TOTALE BANCHE"
    saldo_proiettato_row: int  # row of "Saldo di Periodo" / "Saldo Proiettato"
    entrate_rows: list[
        int
    ]  # data rows between saldo_iniziale and totale_entrate (exclusive)
    uscite_rows: list[
        int
    ]  # data rows between totale_entrate and totale_uscite (exclusive)
    snapshot_kind: str = (
        "month-closed"  # "month-closed" | "fixed-snapshot" (C is DATA RILEVAZ)
    )


def find_layout(wb) -> PFLayout:
    """Scan 'Piano Finanziario' col A for canonical labels, derive row positions.

    Raises ValueError if any required anchor not found.
    """
    if "Piano Finanziario" not in wb.sheetnames:
        raise ValueError("Foglio 'Piano Finanziario' assente")
    pf = wb["Piano Finanziario"]

    labels: dict[int, str] = {}  # row -> normalized label
    for r in range(1, min(pf.max_row + 1, 60)):
        v = pf.cell(r, 1).value
        if isinstance(v, str):
            labels[r] = v.strip().upper()

    def _find(needle: str) -> int:
        needle_up = needle.upper()
        for r, lbl in labels.items():
            if needle_up in lbl:
                return r
        raise ValueError(f"Riga con label contenente '{needle}' non trovata in col A")

    saldo_iniziale_row = _find("SALDO MESE PREC")
    totale_entrate_row = _find("TOTALE ENTRATE")
    totale_uscite_row = _find("TOTALE USCITE")
    cashflow_row = _find("CASH FLOW")
    totale_banche_row = _find("TOTALE BANCHE")
    saldo_proiettato_row = _find("SALDO DI PERIODO")

    # Saldi banca: righe tra cashflow_row+1 e totale_banche_row-1 con label "SALDO <banca>"
    saldi_banca_rows = []
    for r in range(cashflow_row + 1, totale_banche_row):
        lbl = labels.get(r, "")
        if lbl.startswith("SALDO ") and "MESE" not in lbl and "PERIODO" not in lbl:
            saldi_banca_rows.append(r)

    entrate_rows = [
        r for r in range(saldo_iniziale_row + 1, totale_entrate_row) if labels.get(r)
    ]
    uscite_rows = [
        r for r in range(totale_entrate_row + 1, totale_uscite_row) if labels.get(r)
    ]

    # snapshot_kind: INTUR-style files have C1 = "DATA RILEVAZ" (col C is a fixed
    # snapshot of saldi banche, independent of which month is closed).
    c1_val = pf["C1"].value
    snapshot_kind = "month-closed"
    if isinstance(c1_val, str) and "data rilevaz" in c1_val.strip().lower():
        snapshot_kind = "fixed-snapshot"

    return PFLayout(
        saldo_iniziale_row=saldo_iniziale_row,
        totale_entrate_row=totale_entrate_row,
        totale_uscite_row=totale_uscite_row,
        cashflow_row=cashflow_row,
        saldi_banca_rows=saldi_banca_rows,
        totale_banche_row=totale_banche_row,
        saldo_proiettato_row=saldo_proiettato_row,
        entrate_rows=entrate_rows,
        uscite_rows=uscite_rows,
        snapshot_kind=snapshot_kind,
    )


_SUM_RANGE_RE = re.compile(r"SUM\(\$?[A-Z]+\$?(\d+):\$?[A-Z]+\$?(\d+)\)")


def find_total_row_and_sum_range(
    ws: Worksheet,
    scan_max_row: int = 10,
) -> tuple[int | None, int, int]:
    """Locate the total row + vertical SUM range in a PF detail sheet.

    Strategy: scan first `scan_max_row` rows for a vertical SUM formula in any
    column. Return (total_row, sum_start, sum_end). If none found, (None, 0, 0).
    """
    for r in range(1, scan_max_row + 1):
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            if not isinstance(v, str):
                continue
            m = _SUM_RANGE_RE.search(v)
            if not m:
                continue
            start, end = int(m.group(1)), int(m.group(2))
            if end - start >= 5:  # vertical = many rows
                return r, start, end
    return None, 0, 0
