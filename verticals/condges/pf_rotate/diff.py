"""Semantic xlsx diff per i test golden — ignora metadata / order / calc-state."""

from __future__ import annotations

import math
from io import BytesIO

import openpyxl


def _val(cell):
    v = cell.value
    if isinstance(v, str) and v.startswith("="):
        # Normalize formula: lowercase refs
        return ("FORMULA", v.lower())
    if isinstance(v, float) and math.isnan(v):
        return ("NaN",)
    return v


def semantic_diff(
    a_bytes: bytes,
    b_bytes: bytes,
    *,
    target_cells: dict[str, list[str]],
    tol: float = 1e-6,
) -> list[str]:
    """Return list of human-readable diffs. Empty list = equivalent."""
    wa = openpyxl.load_workbook(BytesIO(a_bytes), data_only=False)
    wb = openpyxl.load_workbook(BytesIO(b_bytes), data_only=False)
    diffs: list[str] = []
    for sheet, refs in target_cells.items():
        if sheet not in wa.sheetnames or sheet not in wb.sheetnames:
            diffs.append(f"sheet missing: {sheet}")
            continue
        sa, sb = wa[sheet], wb[sheet]
        for ref in refs:
            va, vb = _val(sa[ref]), _val(sb[ref])
            if isinstance(va, (int, float)) and isinstance(vb, (int, float)):
                if not math.isclose(va, vb, abs_tol=tol, rel_tol=tol):
                    diffs.append(f"{sheet}!{ref}: {va} vs {vb}")
            elif va != vb:
                diffs.append(f"{sheet}!{ref}: {va!r} vs {vb!r}")
    return diffs
