"""Step 0 — normalizza i fogli dettaglio INTUR: col A → codice fornitore."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

import openpyxl

from verticals.condges.pf_rotate.fornitori_map import load_fornitori


DETAIL_SHEETS = [
    "Salari e Stipendi",
    "Utenze",
    "Materie Prime-Consumo ",
    "Materie Prime-Conumo ",
    "Tasse e Imposte",
    "Commisisoni Portali",
    "Mutui e Finaziamenti",
    "Consulenze",
    "Godimento Beni di Terzi",
    " Varie ed Eventuali",
    "Canoni e servizi",
]


def _find_codice_by_name(target: str, fornitori_by_codice: dict) -> int | None:
    """Match a fornitore name to a codice. Exact first, then containment ≥4 char."""
    if not target:
        return None
    t = target.strip().lower()
    # exact
    for cod, row in fornitori_by_codice.items():
        for cand in (row.nome_esolver, row.nome_pf):
            if cand and cand.strip().lower() == t:
                return cod
    # containment
    if len(t) >= 4:
        for cod, row in fornitori_by_codice.items():
            for cand in (row.nome_esolver, row.nome_pf):
                if not cand or len(cand) < 4:
                    continue
                cl = cand.strip().lower()
                if t in cl or cl in t:
                    return cod
    return None


def normalize_intur(
    *,
    pf_bytes: bytes,
    fornitori_csv: Path,
) -> tuple[bytes, list[dict], list[dict]]:
    """Normalizza col A dei fogli dettaglio INTUR.

    Scansiona i fogli dettaglio da riga 4 in poi (riga 3 è il totale/label voce).
    Per ogni riga con col B non vuota, cerca il fornitore in d_fornitori[INTUR].
    Se trovato, sovrascrive col A col codice_fornitore e registra l'overwrite.

    Returns:
        - bytes del nuovo workbook
        - overwrites: lista {foglio, riga, vecchio, nuovo_codice, nome_b}
        - unmapped: lista {foglio, riga, nome_b}  (col B presente ma codice non trovato)
    """
    wb = openpyxl.load_workbook(BytesIO(pf_bytes))
    fornitori = load_fornitori(fornitori_csv, societa="INTUR")
    overwrites: list[dict] = []
    unmapped: list[dict] = []

    for sheet_name in DETAIL_SHEETS:
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        # Start at row 4: row 3 is the TOTALE/label row (col B = voce name, not fornitore)
        for r in range(4, ws.max_row + 1):
            nome = ws.cell(r, 2).value
            if not nome or not isinstance(nome, str):
                continue
            current_a = ws.cell(r, 1).value
            codice = _find_codice_by_name(nome, fornitori)
            if codice is None:
                unmapped.append({"foglio": sheet_name, "riga": r, "nome_b": nome})
                continue
            # Sovrascrivi col A
            if current_a != codice:
                overwrites.append(
                    {
                        "foglio": sheet_name,
                        "riga": r,
                        "vecchio": current_a,
                        "nuovo_codice": codice,
                        "nome_b": nome,
                    }
                )
                ws.cell(r, 1, codice)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue(), overwrites, unmapped
