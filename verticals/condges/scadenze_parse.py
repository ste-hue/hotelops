"""Shared parsers for supplier scadenzario Excel exports.

- ``parse_scadenze``: Esolver 'Situazione partite sintetica per fornitori' (invoice rows).
- ``partite_df_to_scadenzario_data`` / ``sintetica_list_to_scadenzario_data``: normalize to
  :class:`parse_pf.ScadenzarioData` for tesoreria cashflow + Excel export.

``parse_scadenzario`` (flexible header matrix) stays in :mod:`condges.parse_pf`.
"""

from __future__ import annotations

import re
from datetime import date
from typing import BinaryIO

import openpyxl
import pandas as pd

from verticals.condges.parse_pf import ScadenzarioData


def parse_scadenze(
    file_bytes: BinaryIO,
    *,
    primo_mese_aperto: tuple[int, int] | None = None,
) -> tuple[pd.DataFrame, list[int]]:
    """Parse Esolver 'Situazione partite sintetica/dettagliata per fornitori'.

    Layout sintetica (one row per invoice):
      col 11: codice_fornitore, col 12+13: nome (Ragione sociale 1+2),
      col 20: importo, col 23: data_scadenza.
    Layout dettagliata (one row per partita; codice/nome same columns):
      col 34: data_scadenza, col 37: saldo scadenza in UDC (residuo aperto).

    Returns (df, bucket_months) where df has one row per supplier with columns:
      codice_fornitore, nome, totale, scaduto, mese_4, mese_5, ...
    Invoices with scadenza before ``primo_mese_aperto`` (anno, mese) are
    aggregated into 'scaduto'. Default: il mese corrente di oggi — per la
    rotation passare SEMPRE il primo mese aperto, così il risultato non
    dipende dal giorno in cui gira lo script.
    """
    wb = openpyxl.load_workbook(file_bytes, data_only=True)
    ws = wb.active

    if primo_mese_aperto is not None:
        current_year, current_month = primo_mese_aperto
    else:
        current_month = date.today().month
        current_year = date.today().year

    invoices: list[dict] = []
    for row_idx in range(1, ws.max_row + 1):
        codice = None
        nome = ""
        scad = None
        importo_col = 20

        c11 = ws.cell(row=row_idx, column=11).value
        if c11 and isinstance(c11, (int, float)):
            codice = int(c11)
            nome1 = str(ws.cell(row=row_idx, column=12).value or "").strip()
            v13 = ws.cell(row=row_idx, column=13).value
            nome2 = v13.strip() if isinstance(v13, str) else ""
            nome = f"{nome1} {nome2}".strip()
            scad = ws.cell(row=row_idx, column=23).value
            if not hasattr(scad, "month"):
                # layout dettagliata: scadenza in col 34, residuo in col 37
                scad_dett = ws.cell(row=row_idx, column=34).value
                if hasattr(scad_dett, "month"):
                    scad = scad_dett
                    importo_col = 37

        if codice is None:
            c14 = ws.cell(row=row_idx, column=14).value
            if c14:
                m = re.match(r"^(\d+)\s+(.+)", str(c14).strip())
                if m:
                    codice = int(m.group(1))
                    nome = m.group(2).strip()
                    scad = ws.cell(row=row_idx, column=13).value

        if codice is None or scad is None:
            continue
        if not hasattr(scad, "month"):
            continue

        importo = ws.cell(row=row_idx, column=importo_col).value or 0
        invoices.append(
            {
                "codice": int(codice),
                "nome": nome,
                "importo": float(importo),
                "scad_month": scad.month,
                "scad_year": scad.year,
            }
        )

    if not invoices:
        return pd.DataFrame(columns=["codice_fornitore", "nome", "totale", "scaduto"]), []

    supplier_data: dict[int, dict] = {}
    all_months: set[int] = set()

    for inv in invoices:
        cod = inv["codice"]
        if cod not in supplier_data:
            supplier_data[cod] = {"nome": inv["nome"], "scaduto": 0.0, "totale": 0.0}
        amt = inv["importo"]
        supplier_data[cod]["totale"] += amt

        if inv["scad_year"] < current_year or (
            inv["scad_year"] == current_year and inv["scad_month"] < current_month
        ):
            supplier_data[cod]["scaduto"] += amt
        else:
            m = inv["scad_month"]
            all_months.add(m)
            key = f"mese_{m}"
            supplier_data[cod][key] = supplier_data[cod].get(key, 0.0) + amt

    rows = []
    for cod, data in supplier_data.items():
        row = {
            "codice_fornitore": cod,
            "nome": data["nome"],
            "totale": data["totale"],
            "scaduto": data["scaduto"],
        }
        for m in all_months:
            row[f"mese_{m}"] = data.get(f"mese_{m}", 0.0)
        rows.append(row)

    df = pd.DataFrame(rows)
    bucket_months = sorted(all_months)
    return df, bucket_months


def partite_df_to_scadenzario_data(
    df: pd.DataFrame, bucket_months: list[int]
) -> ScadenzarioData:
    """Convert ``parse_scadenze`` output to :class:`ScadenzarioData` for tesoreria."""
    if df.empty:
        return ScadenzarioData(fornitori=[], totale_per_mese={}, scaduto_totale=0.0)

    totale_per_mese: dict[int, float] = {m: 0.0 for m in bucket_months}
    scaduto_totale = 0.0
    fornitori: list[dict] = []

    for _, row in df.iterrows():
        cod = int(row["codice_fornitore"])
        nome = str(row.get("nome", "") or "")
        entry: dict = {
            "fornitore": f"{cod} {nome}".strip(),
            "totale": float(row.get("totale", 0) or 0),
            "scaduto": float(row.get("scaduto", 0) or 0),
        }
        scaduto_totale += entry["scaduto"]
        for m in bucket_months:
            col = f"mese_{m}"
            if col in row.index and pd.notna(row[col]):
                v = float(row[col] or 0)
                entry[col] = v
                totale_per_mese[m] = totale_per_mese.get(m, 0.0) + v
        fornitori.append(entry)

    return ScadenzarioData(
        fornitori=fornitori,
        totale_per_mese=totale_per_mese,
        scaduto_totale=scaduto_totale,
    )


def sintetica_list_to_scadenzario_data(
    suppliers: list[dict], bucket_months: list[int]
) -> ScadenzarioData:
    """Convert output of :func:`condges.scadenzario_excel.parse_sintetica_scadenze`."""
    totale_per_mese: dict[int, float] = {m: 0.0 for m in bucket_months}
    scaduto_totale = 0.0
    fornitori: list[dict] = []

    for s in suppliers:
        entry: dict = {
            "fornitore": f"{s['codice_fornitore']} {s['nome']}".strip(),
            "totale": float(s.get("totale", 0) or 0),
            "scaduto": float(s.get("scaduto", 0) or 0),
        }
        scaduto_totale += entry["scaduto"]
        for m, amt in s.get("buckets", {}).items():
            key = f"mese_{m}"
            entry[key] = float(amt)
            totale_per_mese[m] = totale_per_mese.get(m, 0.0) + float(amt)
        fornitori.append(entry)

    return ScadenzarioData(
        fornitori=fornitori,
        totale_per_mese=totale_per_mese,
        scaduto_totale=scaduto_totale,
    )


