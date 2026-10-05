"""Budget camere per driver — il foglio dove Stefano scrive gli aumenti.

Valori = base 2026 da BigQuery. Formule = tutto ciò che dipende dagli input
(celle gialle): cambiando un aumento, ricavo ed effetti si aggiornano in Excel.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

INPUT = PatternFill("solid", fgColor="FFF2CC")
GRASSETTO = Font(bold=True)
MESI = {
    1: "gennaio",
    2: "febbraio",
    3: "marzo",
    4: "aprile",
    5: "maggio",
    6: "giugno",
    7: "luglio",
    8: "agosto",
    9: "settembre",
    10: "ottobre",
    11: "novembre",
    12: "dicembre",
}
COL_PREZZI = [
    "Mese",
    "Nome mese",
    "Categoria",
    "Camere 2026",
    "Camere 2027",
    "Stato base",
    "Notti 2026 (calendario 2027)",
    "Prezzo medio 2026",
    "Notti 2027",
    "Aumento %",
    "Prezzo 2027",
    "Ricavo base",
    "Ricavo 2027 a prezzi 2026",
    "Ricavo 2027",
    "Prezzo base Lybra",
    "Ragione",
]
COL_MESI = [
    "Mese",
    "Nome mese",
    "Stato",
    "Giorni 2026",
    "Giorni 2027",
    "Notti 2026 reali",
    "Ricavo camere 2026 reale (01ROOM)",
    "Ricavo 2026 sul calendario 2027 (01ROOM)",
    "Effetto calendario",
    "Notti base",
    "Ricavo base (gestionale)",
    "Notti 2027",
    "Ricavo 2027 (gestionale)",
    "Effetto occupazione (gestionale)",
    "Effetto mix (gestionale)",
    "Effetto prezzo (gestionale)",
    "Raccordo 01ROOM",
    "Budget 2027 su base 01ROOM",
    "Occupazione 2027",
    "Effetto occupazione (01ROOM)",
    "Effetto mix (01ROOM)",
    "Effetto prezzo (01ROOM)",
]


def _intesta(ws, colonne: list[str]) -> None:
    ws.append(colonne)
    for cella in ws[1]:
        cella.font = GRASSETTO
    ws.freeze_panes = "A2"


def _prezzi(ws, righe: list[dict]) -> None:
    _intesta(ws, COL_PREZZI)
    for i, r in enumerate(righe, start=2):
        ws.append(
            [
                r["mese"],
                MESI[r["mese"]],
                r["categoria"],
                r["camere_2026"],
                r["camere_2027"],
                r["stato"],
                r["notti_base"],
                r["prezzo_base"],
                r["notti_2027"],
                0,
                f'=IF(H{i}="","",H{i}*(1+J{i}))',
                f'=IF(H{i}="",0,G{i}*H{i})',
                f'=IF(H{i}="",0,I{i}*H{i})',
                f'=IF(K{i}="",0,I{i}*K{i})',
                None,
                None,
            ]
        )
        # senza prezzo 2026 non c'è nulla da aumentare: il prezzo si scrive in K
        for col in "IJOP" if r["prezzo_base"] is not None else "IJKOP":
            ws[f"{col}{i}"].fill = INPUT
        for col in "HKO":
            ws[f"{col}{i}"].number_format = "#,##0.00"
        for col in "LMN":
            ws[f"{col}{i}"].number_format = "#,##0"
        ws[f"J{i}"].number_format = "0.0%"


def _mesi(ws, mesi: list[dict], ultima_prezzi: int) -> None:
    _intesta(ws, COL_MESI)

    def somma(col: str, i: int) -> str:
        return (
            f"SUMIF(Prezzi!$A$2:$A${ultima_prezzi},A{i},"
            f"Prezzi!${col}$2:${col}${ultima_prezzi})"
        )

    for i, x in enumerate(mesi, start=2):
        ws.append(
            [
                x["mese"],
                MESI[x["mese"]],
                x["stato"],
                x["giorni_2026"],
                x["giorni_2027"],
                x["notti_reali"],
                x["ricavo_reale"],
                x["ricavo_calendario"],
                f'=IF(OR(G{i}="",H{i}=""),"n.d.",H{i}-G{i})',
                f"={somma('G', i)}",
                f"={somma('L', i)}",
                f"={somma('I', i)}",
                f"={somma('N', i)}",
                f"=IF(J{i}=0,0,(L{i}-J{i})*K{i}/J{i})",
                f"={somma('M', i)}-IF(J{i}=0,0,L{i}*K{i}/J{i})",
                f"=M{i}-{somma('M', i)}",
                x["raccordo"],
                f'=IF(Q{i}="","n.d.",M{i}*Q{i})',
                f"=L{i}/({x['camere']}*E{i})",
            ]
            + [f'=IF(Q{i}="","n.d.",{col}{i}*Q{i})' for col in "NOP"]
        )
    fine = len(mesi) + 1
    t = fine + 1
    camere = mesi[0]["camere"]
    # un totale con un mese mancante o in stima è "n.d.", mai una somma parziale
    stima = f'COUNTIF(C2:C{fine},"stima")>0'
    senza_raccordo = f"COUNTBLANK(Q2:Q{fine})>0"
    ws.append(
        ["Totale", None, None]
        + [f"=SUM({col}2:{col}{fine})" for col in "DEF"]
        + [
            f'=IF(COUNTBLANK({col}2:{col}{fine})>0,"n.d.",SUM({col}2:{col}{fine}))'
            for col in "GH"
        ]
        + [f'=IF(COUNTBLANK(G2:H{fine})>0,"n.d.",SUM(I2:I{fine}))']
        + [f'=IF({stima},"n.d.",SUM({col}2:{col}{fine}))' for col in "JKLMNOP"]
        + [
            None,
            f'=IF({senza_raccordo},"n.d.",SUM(R2:R{fine}))',
            f'=IF(ISNUMBER(L{t}),L{t}/({camere}*E{t}),"n.d.")',
        ]
        + [f'=IF({senza_raccordo},"n.d.",SUM({col}2:{col}{fine}))' for col in "TUV"]
    )
    for cella in ws[t]:
        cella.font = GRASSETTO
    for riga in ws.iter_rows(min_row=2, max_row=t):
        for cella in riga:
            lettera = cella.column_letter
            if lettera in "GHIKMNOPRTUV":
                cella.number_format = "#,##0"
            elif lettera == "Q":
                cella.number_format = "0.0000"
            elif lettera == "S":
                cella.number_format = "0.0%"


def _leggimi(ws, dati: dict) -> None:
    (am, ag), (cm, cg) = dati["apertura"], dati["chiusura"]
    base, budget = dati["anno_base"], dati["anno_base"] + 1
    stime = [MESI[x["mese"]] for x in dati["mesi"] if x["stato"] == "stima"]
    righe = [
        f"Budget camere Hotel Panorama {budget} — base {base}",
        f"Calendario {budget}: dal {ag} {MESI[am]} al {cg} {MESI[cm]} (estremi inclusi).",
        "Foglio Prezzi: scrivi solo nelle celle gialle. Aumento % sul prezzo medio "
        f"{base} della categoria in quel mese; Ragione = perché.",
        f"Notti {budget} precompilate: le notti {base} di ogni categoria, scalate col "
        f"rapporto camere {budget}/{base}, a parità di notti totali del mese. "
        "Correggile se sai di più.",
        "Categoria senza prezzo 2026 in un mese (nessuna vendita): la cella Prezzo 2027 "
        "è gialla, scrivi lì il prezzo.",
        "Foglio Mesi: effetto calendario, occupazione, mix e prezzo si aggiornano da soli.",
        "Prezzi ed effetti (gestionale) sono sulla base del rapporto per tipologia venduta. "
        "Le colonne (01ROOM) li riportano al ricavo camere canonico col Raccordo: "
        "ricavo 2026 reale + calendario + occupazione + mix + prezzo = budget 2027.",
        f"Ricavi {base} vuoti = ricavo camere non caricato per tutti i giorni del mese.",
    ]
    if stime:
        righe.append(
            "Mesi in stato stima ("
            + ", ".join(stime)
            + "): notti e prezzi sono le prenotazioni alla data dell'esportazione, "
            "quindi sottostimati. Raccordo e totali restano n.d. finché non arriva "
            "l'esportazione nuova."
        )
    for testo in righe:
        ws.append([testo])
    ws["A1"].font = GRASSETTO
    ws.column_dimensions["A"].width = 120


def scrivi(dati: dict, path: Path) -> Path:
    wb = Workbook()
    _leggimi(wb.active, dati)
    wb.active.title = "Leggimi"
    _prezzi(wb.create_sheet("Prezzi"), dati["righe"])
    _mesi(wb.create_sheet("Mesi"), dati["mesi"], ultima_prezzi=len(dati["righe"]) + 1)
    wb.save(path)
    return path
