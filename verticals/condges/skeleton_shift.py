"""Monthly skeleton shift for the PF cashflow rollover.

Shifts VALUE cells one column left inside the "forward months" range so the
file represents month N+1..N+9 instead of N..N+8. FORMULA cells are left in
place because they reference fixed cells (=C37, =SUM(D5:D178), …) that
remain correct. New last column starts empty for the wraparound month.

Only operates on the PF skeleton: month labels (row 2), year markers (row 1),
and value cells in the shift range. Saldi banca, SALDO MESE PRECED and the
data riferimento are owned by the bank-balance update step.
"""

from __future__ import annotations

from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

MONTHS_IT = [
    "GENNAIO", "FEBBRAIO", "MARZO", "APRILE", "MAGGIO", "GIUGNO",
    "LUGLIO", "AGOSTO", "SETTEMBRE", "OTTOBRE", "NOVEMBRE", "DICEMBRE",
]

# Master "Piano Finanziario" sheet: forward months occupy cols 3..11.
MASTER_FIRST_COL = 3
MASTER_LAST_COL = 11
# Detail per-voce sheets: forward months occupy cols 4..12 (col 3 holds the
# annual aggregate formula =SUM(D..L..)).
DETAIL_FIRST_COL = 4
DETAIL_LAST_COL = 12

# Master rows whose values get shifted. Row 4 col 3 (SALDO MESE PRECED) is
# excluded — its value is replaced by the bank-balance update.
MASTER_VALUE_ROWS = list(range(6, 12)) + list(range(14, 25))


def _is_formula(value) -> bool:
    return isinstance(value, str) and value.startswith("=")


def rollover_month_labels(
    current_month: int, current_year: int, n_forward: int = 9
) -> tuple[list[str], list[tuple[int, int]]]:
    """Compute new month labels + year-change slots for the next-month skeleton.

    Args:
        current_month: the month currently in the leftmost forward column
            (e.g., 4 for an April PF that we are rolling to May).
        current_year: the year matching ``current_month``.
        n_forward: how many forward months the file shows (default 9).

    Returns:
        (labels, year_changes)
        labels: ``n_forward`` ITALIAN month names starting from ``current_month + 1``.
        year_changes: list of ``(slot_index, new_year)`` for each transition
            from December to January inside the new label range. slot_index is
            0-based into ``labels``.
    """
    labels: list[str] = []
    year_changes: list[tuple[int, int]] = []
    year = current_year
    for i in range(n_forward):
        # next month after the current one, then advance with i
        idx = (current_month + i) % 12  # 0-based into MONTHS_IT
        if i > 0 and idx == 0:
            year += 1
            year_changes.append((i, year))
        labels.append(MONTHS_IT[idx])
    return labels, year_changes


def _detect_current_month_from_header(
    ws: Worksheet, first_col: int, last_col: int
) -> int:
    """Read the row 2 month header in ``first_col`` and return its month int."""
    label = ws.cell(row=2, column=first_col).value
    if not label:
        raise ValueError(
            f"No month label in row 2 col {first_col} of sheet '{ws.title}'"
        )
    name = str(label).strip().upper()
    try:
        return MONTHS_IT.index(name) + 1
    except ValueError as e:
        raise ValueError(
            f"Unknown month label '{label}' in row 2 col {first_col} of "
            f"sheet '{ws.title}'"
        ) from e


def _shift_value_row(
    ws: Worksheet, row: int, first_col: int, last_col: int
) -> None:
    """Shift VALUE cells one column left in ``row[first_col..last_col]``.

    Behavior per destination cell at (row, c) for c in [first_col, last_col - 1]:
      - if dst is a formula: skip (formula stays).
      - else: dst.value = src.value, where src = (row, c + 1), unless src is a
        formula in which case dst is cleared (None).
    The last column (last_col) value is cleared (unless it holds a formula).
    """
    for c in range(first_col, last_col):
        dst = ws.cell(row=row, column=c)
        if _is_formula(dst.value):
            continue
        src_val = ws.cell(row=row, column=c + 1).value
        dst.value = None if _is_formula(src_val) else src_val
    last = ws.cell(row=row, column=last_col)
    if not _is_formula(last.value):
        last.value = None


def _unmerge_overlapping(
    ws: Worksheet, row: int, first_col: int, last_col: int
) -> None:
    """Unmerge any single-row ranges on ``row`` that overlap [first_col, last_col]."""
    to_unmerge = []
    for r in list(ws.merged_cells.ranges):
        if r.min_row == row and r.max_row == row:
            if not (r.max_col < first_col or r.min_col > last_col):
                to_unmerge.append(str(r))
    for r in to_unmerge:
        ws.unmerge_cells(r)


def _apply_headers(
    ws: Worksheet,
    first_col: int,
    last_col: int,
    labels: list[str],
    current_year: int,
    year_changes: list[tuple[int, int]],
) -> None:
    """Write new month labels (row 2) and year markers (row 1).

    Row 1 may contain merged ranges spanning the year. Unmerge first, write the
    new year markers (one per year-run), then re-merge each run.
    """
    # Row 2: month labels
    for i, label in enumerate(labels):
        ws.cell(row=2, column=first_col + i, value=label)

    # Compute year per slot index
    years_per_slot = [current_year] * len(labels)
    for slot_idx, new_year in year_changes:
        for j in range(slot_idx, len(labels)):
            years_per_slot[j] = max(years_per_slot[j], new_year)

    # Unmerge any row-1 range overlapping our cols, then clear and rewrite
    _unmerge_overlapping(ws, row=1, first_col=first_col, last_col=last_col)
    for c in range(first_col, last_col + 1):
        ws.cell(row=1, column=c, value=None)

    # Write year markers + re-merge runs of equal years
    i = 0
    while i < len(labels):
        j = i
        while j + 1 < len(labels) and years_per_slot[j + 1] == years_per_slot[i]:
            j += 1
        ws.cell(row=1, column=first_col + i, value=float(years_per_slot[i]))
        if j > i:
            c1 = get_column_letter(first_col + i)
            c2 = get_column_letter(first_col + j)
            ws.merge_cells(f"{c1}1:{c2}1")
        i = j + 1


def hide_past_months(
    ws: Worksheet, first_col: int, last_col: int, primo_mese_aperto: int
) -> None:
    """Nasconde (NON cancella) le colonne dei mesi prima del primo mese aperto.

    Cerca nella riga 2 (header mesi) la prima colonna etichettata col primo mese
    aperto e marca ``hidden=True`` tutte le colonne mese che la precedono. Dati e
    formule restano intatti — la bussola non si perde, il file parte visivamente
    dal mese aperto. No-op se l'etichetta non c'è.
    """
    target = MONTHS_IT[(primo_mese_aperto - 1) % 12]
    target_col = None
    for c in range(first_col, last_col + 1):
        v = ws.cell(row=2, column=c).value
        if v and str(v).strip().upper() == target:
            target_col = c
            break
    if target_col is None:
        return
    for c in range(first_col, target_col):
        ws.column_dimensions[get_column_letter(c)].hidden = True


def hide_past_columns_rotation(wb, primo_mese_aperto: int) -> None:
    """Applica hide_past_months a tutti i fogli del PF post-rotation.

    Master 'Piano Finanziario': mesi in C..R. Fogli voce: mesi in D..S.
    Fogli non-mensili (Controlli/DA MAPPARE/ESCLUSI) saltati.
    """
    for ws in wb.worksheets:
        if ws.title in ("Controlli", "DA MAPPARE", "ESCLUSI"):
            continue
        if ws.title.strip() == "Piano Finanziario":
            hide_past_months(ws, 3, 18, primo_mese_aperto)
        else:
            hide_past_months(ws, 4, 19, primo_mese_aperto)


def shift_master_sheet(ws: Worksheet, current_year: int) -> None:
    """Shift the master 'Piano Finanziario' sheet for monthly rollover.

    Reads the current month from row 2 col 3, computes new labels for
    cols 3..11, shifts VALUES in rows ``MASTER_VALUE_ROWS`` left by 1, and
    updates row 1 / row 2 headers.

    Does NOT touch row 4 (SALDO MESE PRECED), rows 31-35 (saldi banca),
    or the date reference. Those belong to the bank-balance update step.
    """
    current_month = _detect_current_month_from_header(
        ws, MASTER_FIRST_COL, MASTER_LAST_COL
    )
    labels, year_changes = rollover_month_labels(
        current_month, current_year, n_forward=MASTER_LAST_COL - MASTER_FIRST_COL + 1
    )
    for row in MASTER_VALUE_ROWS:
        _shift_value_row(ws, row, MASTER_FIRST_COL, MASTER_LAST_COL)
    _apply_headers(
        ws, MASTER_FIRST_COL, MASTER_LAST_COL, labels, current_year, year_changes
    )


def shift_detail_sheet(ws: Worksheet, current_year: int) -> None:
    """Shift a per-voce detail sheet for monthly rollover.

    Reads the current month from row 2 col 4, computes new labels for
    cols 4..12, and shifts VALUES for every row from 3 to ``ws.max_row``.
    Formula cells (e.g., =SUM(...) on row 4 totale voce and =SUM(D{r}:L{r})
    on annual-aggregate col 3) are preserved.
    """
    current_month = _detect_current_month_from_header(
        ws, DETAIL_FIRST_COL, DETAIL_LAST_COL
    )
    labels, year_changes = rollover_month_labels(
        current_month, current_year, n_forward=DETAIL_LAST_COL - DETAIL_FIRST_COL + 1
    )
    for row in range(3, ws.max_row + 1):
        _shift_value_row(ws, row, DETAIL_FIRST_COL, DETAIL_LAST_COL)
    _apply_headers(
        ws, DETAIL_FIRST_COL, DETAIL_LAST_COL, labels, current_year, year_changes
    )
