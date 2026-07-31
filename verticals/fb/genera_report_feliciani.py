#!/usr/bin/env python3
"""Modello Excel F&B per il consulente Feliciani (8 fogli).

Cruscotto (reale vs target per outlet, base NETTO, gating mese chiuso) ·
Giornaliero Servizio (POS, LORDO) · Coperti Completi (paganti/non paganti) ·
Breakfast · Consumi · Menu Engineering (quadranti) · Definizioni · Trend Settimanale.
Spec: docs/superpowers/specs/2026-07-30-modello-fb-feliciani-design.md

Usage:
    python -m verticals.fb.genera_report_feliciani [--output X.xlsx] [--days 30] [--months 13]
"""

from __future__ import annotations

import argparse
import logging
from datetime import date

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from core.config import DATASET, PROJECT

log = logging.getLogger("fb.report_feliciani")

SERVIZIO_ORDER = {"pranzo": 0, "cena": 1, "bar": 2, "altro": 3}

FMT_EURO = "#,##0.00 €"
FMT_RATIO = "0.00"
FMT_INT = "#,##0"
FMT_PCT = "+0.0%;-0.0%"


def query_giornaliero(client, days: int) -> pd.DataFrame:
    q = f"""
    SELECT * FROM `{PROJECT}.{DATASET}.v_fb_giornaliero`
    WHERE data >= DATE_SUB(
      (SELECT MAX(data) FROM `{PROJECT}.{DATASET}.v_fb_giornaliero`),
      INTERVAL @days - 1 DAY)
    ORDER BY data, servizio
    """
    from google.cloud import bigquery

    job = client.query(
        q,
        job_config=bigquery.QueryJobConfig(
            query_parameters=[bigquery.ScalarQueryParameter("days", "INT64", days)]
        ),
    )
    df = pd.DataFrame([dict(r) for r in job.result()])
    if df.empty:
        return df
    df["_ord"] = df["servizio"].map(SERVIZIO_ORDER).fillna(9)
    return df.sort_values(["data", "_ord"]).drop(columns="_ord").reset_index(drop=True)


PASTO_ORDER = {"BRK": 0, "LUNCH": 1, "DINNER": 2}


def query_coperti_completi(client, days: int) -> pd.DataFrame:
    from google.cloud import bigquery

    q = f"""
    SELECT * FROM `{PROJECT}.{DATASET}.v_fb_coperti_giornaliero`
    WHERE data_servizio >= DATE_SUB(
      (SELECT MAX(data_servizio) FROM `{PROJECT}.{DATASET}.v_fb_coperti_giornaliero`),
      INTERVAL @days - 1 DAY)
    """
    job = client.query(
        q,
        job_config=bigquery.QueryJobConfig(
            query_parameters=[bigquery.ScalarQueryParameter("days", "INT64", days)]
        ),
    )
    df = pd.DataFrame([dict(r) for r in job.result()])
    if df.empty:
        return df
    df["_ord"] = df["tipo_pasto"].map(PASTO_ORDER).fillna(9)
    return (
        df.sort_values(["data_servizio", "_ord"])
        .drop(columns="_ord")
        .reset_index(drop=True)
    )


def add_sheet_coperti(wb: Workbook, df: pd.DataFrame) -> None:
    ws = wb.create_sheet("Coperti Completi")
    rows = [
        [
            r["data_servizio"],
            r["tipo_pasto"],
            r["hotel"],
            r["residence"],
            r["cvm"],
            r["esterni"],
            r["paganti"],
            r["dipendenti"],
            r["courtesy_pm"],
            r["totale"],
        ]
        for _, r in df.iterrows()
    ]
    _write_sheet(
        ws,
        [
            "Data",
            "Pasto",
            "Hotel",
            "Residence",
            "CVM",
            "Esterni",
            "Paganti",
            "Dipendenti",
            "Courtesy/PM",
            "Totale",
        ],
        rows,
        dict.fromkeys(range(2, 10), FMT_INT),
    )


def build_trend_settimanale(df: pd.DataFrame) -> pd.DataFrame:
    """Aggrega il giornaliero per settimana ISO (lun-dom) con Δ% vs settimana precedente."""
    if df.empty:
        return pd.DataFrame()
    g = df.copy()
    g["data"] = pd.to_datetime(g["data"])
    g["settimana"] = g["data"].dt.to_period("W-SUN")
    wk = (
        g.groupby("settimana", sort=True)
        .agg(
            coperti=("coperti", "sum"),
            ricavi=("ricavi", "sum"),
            giorni=("data", "nunique"),
        )
        .reset_index()
    )
    wk["coperto_medio"] = wk["ricavi"] / wk["coperti"].replace(0, pd.NA)
    wk["delta_coperti_pct"] = wk["coperti"].pct_change()
    wk["delta_ricavi_pct"] = wk["ricavi"].pct_change()
    # Settimane incomplete: il confronto col totale della settimana piena è fuorviante
    parziale = wk["giorni"] < 7
    wk.loc[parziale, ["delta_coperti_pct", "delta_ricavi_pct"]] = pd.NA
    wk.loc[
        parziale.shift(1, fill_value=False), ["delta_coperti_pct", "delta_ricavi_pct"]
    ] = pd.NA
    # Settimane non consecutive (buco stagionale): Δ% oltre il buco non ha senso
    ordinali = wk["settimana"].map(lambda p: p.ordinal)
    non_consecutive = ordinali.diff() != 1
    non_consecutive.iloc[0] = False  # la prima ha già Δ% NA da pct_change
    wk.loc[non_consecutive, ["delta_coperti_pct", "delta_ricavi_pct"]] = pd.NA
    wk["settimana"] = wk["settimana"].astype(str) + parziale.map(
        {True: " (parziale)", False: ""}
    )
    return (
        wk.drop(columns="giorni").iloc[::-1].reset_index(drop=True)
    )  # più recente in alto


OUTLET_ORDER = {"BREAKFAST": 0, "RISTORANTE": 1, "BAR": 2, "MENSA_STAFF": 3}


def query_modello_mensile(client, months: int = 13) -> pd.DataFrame:
    from google.cloud import bigquery

    q = f"""
    SELECT * FROM `{PROJECT}.{DATASET}.v_fb_modello_mensile`
    WHERE periodo >= DATE_SUB(
      (SELECT MAX(periodo) FROM `{PROJECT}.{DATASET}.v_fb_modello_mensile`),
      INTERVAL @months MONTH)
    ORDER BY periodo, outlet
    """
    job = client.query(
        q,
        job_config=bigquery.QueryJobConfig(
            query_parameters=[bigquery.ScalarQueryParameter("months", "INT64", months)]
        ),
    )
    df = pd.DataFrame([dict(r) for r in job.result()])
    if df.empty:
        return df
    df["_ord"] = df["outlet"].map(OUTLET_ORDER).fillna(9)
    return (
        df.sort_values(["periodo", "_ord"]).drop(columns="_ord").reset_index(drop=True)
    )


def _label_mese(r) -> str:
    label = f"{r['anno']}-{r['mese']:02d}"
    # MENSA_STAFF non ha mai costi separabili (escono da CUCINA): niente
    # etichetta "in corso", che lì significherebbe la cosa sbagliata.
    if not r["has_costi"] and r.get("outlet") != "MENSA_STAFF":
        label += " (in corso — senza costi)"
    return label


def add_sheet_cruscotto(wb: Workbook, df: pd.DataFrame) -> None:
    ws = wb.create_sheet("Cruscotto")
    rows = []
    for _, r in df.iterrows():
        chiuso = bool(r["has_costi"])
        rows.append(
            [
                _label_mese(r),
                r["outlet"],
                r["coperti_paganti"],
                r["coperti_non_paganti"],
                r["ricavo_per_coperto"] if r["has_ricavi"] else None,
                r["costo_per_coperto"] if chiuso else None,
                r["food_cost_pct"] if chiuso else None,
                r["margine_per_coperto"] if chiuso else None,
                None,
                None,  # Target €/cop, Δ vs Target — compila Feliciani
            ]
        )
    _write_sheet(
        ws,
        [
            "Mese",
            "Outlet",
            "Coperti Paganti",
            "Coperti Non Paganti",
            "Ricavo/Cop",
            "Costo/Cop",
            "Food Cost %",
            "Margine/Cop",
            "Target Margine/Cop",
            "Δ vs Target",
        ],
        rows,
        {
            2: FMT_INT,
            3: FMT_INT,
            4: FMT_EURO,
            5: FMT_EURO,
            6: "0.0%",
            7: FMT_EURO,
            8: FMT_EURO,
            9: FMT_EURO,
        },
    )
    ws.column_dimensions["A"].width = 26


def add_sheet_breakfast(wb: Workbook, df: pd.DataFrame) -> None:
    ws = wb.create_sheet("Breakfast")
    brk = df[df["outlet"] == "BREAKFAST"]
    rows = [
        [
            _label_mese(r),
            r["coperti_paganti"],
            r["costo_netto"] if r["has_costi"] else None,
            r["costo_per_coperto"] if r["has_costi"] else None,
            r["ricavo_netto"] if r["has_ricavi"] else None,
        ]
        for _, r in brk.iterrows()
    ]
    _write_sheet(
        ws,
        [
            "Mese",
            "Coperti",
            "Costo Economato",
            "Costo/Coperto",
            "Ricavo Esplicito (il grosso è in tariffa camera)",
        ],
        rows,
        {1: FMT_INT, 2: FMT_EURO, 3: FMT_EURO, 4: FMT_EURO},
    )
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["E"].width = 42


def query_consumi_reparto(client, months: int = 13) -> pd.DataFrame:
    from google.cloud import bigquery

    q = f"""
    SELECT anno, mese,
      IF(reparto_id IN ('BRK','CUCINA','CANTINA'), reparto_id, 'ALTRI') AS reparto,
      CASE reparto_id WHEN 'BRK' THEN 'BREAKFAST' WHEN 'CUCINA' THEN 'RISTORANTE'
                      WHEN 'CANTINA' THEN 'BAR' END AS outlet,
      SUM(importo) AS importo
    FROM `{PROJECT}.{DATASET}.f_consumi_economato`
    WHERE codice_prodotto NOT IN (
      'BEV.CAF.00014','FOO.FRS.00013','FOO.FRS.00011','FOO.FRS.00012','FOO.FRS.00009',
      'BEV.BOL.00013','FOO.FRS.00008','FOO.FAR.00003','FOO.BUR.00001')
      AND DATE(anno, mese, 1) >= DATE_SUB(
        (SELECT DATE(MAX(anno), MAX(mese), 1) FROM `{PROJECT}.{DATASET}.f_consumi_economato`),
        INTERVAL @months MONTH)
    GROUP BY 1, 2, 3, 4
    ORDER BY 1, 2, 3
    """
    job = client.query(
        q,
        job_config=bigquery.QueryJobConfig(
            query_parameters=[bigquery.ScalarQueryParameter("months", "INT64", months)]
        ),
    )
    return pd.DataFrame([dict(r) for r in job.result()])


def add_sheet_consumi(
    wb: Workbook, df_reparti: pd.DataFrame, df_modello: pd.DataFrame
) -> None:
    ws = wb.create_sheet("Consumi")
    cop = {
        (r["anno"], r["mese"], r["outlet"]): r["coperti_paganti"]
        for _, r in df_modello.iterrows()
    }
    rows = []
    for _, r in df_reparti.iterrows():
        denominatore = cop.get((r["anno"], r["mese"], r["outlet"]))
        rows.append(
            [
                f"{r['anno']}-{r['mese']:02d}",
                r["reparto"],
                r["outlet"] or "—",
                r["importo"],
                (r["importo"] / denominatore) if denominatore else None,
            ]
        )
    _write_sheet(
        ws,
        ["Mese", "Reparto", "Outlet", "Consumo €", "€/Coperto Pagante"],
        rows,
        {3: FMT_EURO, 4: FMT_EURO},
    )


def query_menu_engineering(client, days: int) -> pd.DataFrame:
    from google.cloud import bigquery

    q = f"""
    WITH ultima_foto AS (
      SELECT sala, piatto, costo_unitario, tipo, descrizione
      FROM `{PROJECT}.{DATASET}.f_menu_engineering`
      WHERE snapshot_date = (SELECT MAX(snapshot_date)
                             FROM `{PROJECT}.{DATASET}.f_menu_engineering`)
        AND costo_unitario IS NOT NULL AND costo_unitario > 0
        AND NOT STARTS_WITH(piatto, 'COP')
    ),
    vendite AS (
      SELECT codice_articolo AS piatto, SUM(quantita) AS qty,
             SAFE_DIVIDE(SUM(importo_netto), SUM(quantita)) AS prezzo_medio_netto
      FROM `{PROJECT}.{DATASET}.f_vendite_fb`
      WHERE data_servizio >= DATE_SUB(
        (SELECT MAX(data_servizio) FROM `{PROJECT}.{DATASET}.f_vendite_fb`),
        INTERVAL @days - 1 DAY)
      GROUP BY 1
      HAVING qty > 0
    )
    SELECT
      f.piatto, ANY_VALUE(f.descrizione) AS descrizione, ANY_VALUE(f.tipo) AS tipo,
      ANY_VALUE(f.costo_unitario) AS costo_unitario,
      v.qty, v.prezzo_medio_netto,
      v.prezzo_medio_netto - ANY_VALUE(f.costo_unitario) AS margine_unitario,
      (v.prezzo_medio_netto - ANY_VALUE(f.costo_unitario)) * v.qty AS margine_totale
    FROM ultima_foto f
    JOIN vendite v USING (piatto)
    GROUP BY f.piatto, v.qty, v.prezzo_medio_netto
    ORDER BY margine_totale DESC
    """
    job = client.query(
        q,
        job_config=bigquery.QueryJobConfig(
            query_parameters=[bigquery.ScalarQueryParameter("days", "INT64", days)]
        ),
    )
    return pd.DataFrame([dict(r) for r in job.result()])


def classifica_quadranti(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    med_qty = out["qty"].median()
    med_margine = out["margine_unitario"].median()

    def _q(r):
        alta = r["qty"] >= med_qty
        alto = r["margine_unitario"] >= med_margine
        return {
            (True, True): "Star",
            (True, False): "Cavallo",
            (False, True): "Enigma",
            (False, False): "Cane",
        }[(alta, alto)]

    out["quadrante"] = out.apply(_q, axis=1)
    return out


def add_sheet_menu_engineering(wb: Workbook, df: pd.DataFrame) -> None:
    ws = wb.create_sheet("Menu Engineering")
    rows = [
        [
            r["descrizione"] or r["piatto"],
            r["tipo"],
            r["qty"],
            r["prezzo_medio_netto"],
            r["costo_unitario"],
            r["margine_unitario"],
            r["margine_totale"],
            r["quadrante"],
        ]
        for _, r in df.iterrows()
    ]
    _write_sheet(
        ws,
        [
            "Piatto",
            "Categoria",
            "Qty Periodo",
            "Prezzo Medio Netto",
            "Costo Unitario",
            "Margine Unitario",
            "Margine Totale",
            "Quadrante",
        ],
        rows,
        {2: FMT_INT, 3: FMT_EURO, 4: FMT_EURO, 5: FMT_EURO, 6: FMT_EURO},
    )
    ws.column_dimensions["A"].width = 38


def _style_header(ws, ncols: int) -> None:
    fill = PatternFill("solid", fgColor="1F4E5F")
    for col in range(1, ncols + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"


def _write_sheet(
    ws, headers: list[str], rows: list[list], formats: dict[int, str]
) -> None:
    """headers, poi righe; formats = {col_index_0based: number_format}."""
    thin = Border(bottom=Side(style="thin", color="D9D9D9"))
    ws.append(headers)
    for row in rows:
        ws.append(row)
    for r in range(2, ws.max_row + 1):
        for c0, fmt in formats.items():
            ws.cell(row=r, column=c0 + 1).number_format = fmt
        for c in range(1, len(headers) + 1):
            ws.cell(row=r, column=c).border = thin
    for c, h in enumerate(headers, start=1):
        width = max(len(h) + 2, 12)
        ws.column_dimensions[get_column_letter(c)].width = width
    _style_header(ws, len(headers))


DEFINIZIONI = [
    (
        "Base valore",
        "Cruscotto/Breakfast/Consumi/Menu Engineering = NETTO (imponibile). "
        "Giornaliero Servizio = LORDO IVA (POS cassa). Mai confrontare le due basi.",
    ),
    (
        "Denominatore per-coperto",
        "Solo coperti PAGANTI (esclusi DIPENDENTI, COURTESY, PM). "
        "La mensa staff è un costo, non un cliente.",
    ),
    (
        "Mese (in corso — senza costi)",
        "I consumi economato arrivano a mese chiuso: "
        "margine e food cost compaiono solo dopo l'ingest consumi.",
    ),
    (
        "Breakfast",
        "Il ricavo colazione è quasi tutto DENTRO la tariffa camera: la metrica "
        "di governo è il costo/coperto, non il margine esplicito.",
    ),
    (
        "Bar (mensile)",
        "Il PMS non conta i coperti bar: il per-coperto bar vive nel foglio "
        "Giornaliero Servizio (fonte POS).",
    ),
    (
        "Menu Engineering",
        "Costi unitari dall'ultima foto RistoCube; quantità e prezzi medi "
        "netti dal venduto del periodo del report. Quadranti su mediane.",
    ),
    (
        "Fonti e freschezza",
        "POS comande e vendite: a export. Coperti: giornalieri (Hoxell). "
        "Consumi e ricavi PMS: mensili.",
    ),
]


def add_sheet_definizioni(wb: Workbook) -> None:
    ws = wb.create_sheet("Definizioni")
    _write_sheet(ws, ["Voce", "Definizione"], [list(t) for t in DEFINIZIONI], {})
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 100


def add_sheet_giornaliero(wb: Workbook, df_giorni: pd.DataFrame) -> None:
    ws = wb.create_sheet("Giornaliero Servizio")
    rows = [
        [
            r["data"],
            r["servizio"].capitalize(),
            r["n_comande"],
            r["coperti"],
            r["ricavi"],
            r["coperto_medio"],
            r["articoli"],
            r["articoli_per_coperto"],
            r["bevande_per_coperto"],
            r["food_per_coperto"],
            r["antipasti_per_coperto"],
            r["primi_per_coperto"],
            r["secondi_per_coperto"],
            r["dessert_per_coperto"],
            r["bottiglie_vino_per_10_coperti"],
        ]
        for _, r in df_giorni.iterrows()
    ]
    _write_sheet(
        ws,
        [
            "Data",
            "Servizio",
            "Comande",
            "Coperti",
            "Ricavi Lordi",
            "Coperto Medio",
            "Articoli",
            "Art/Cop",
            "Bevande/Cop",
            "Food/Cop",
            "Antipasti/Cop",
            "Primi/Cop",
            "Secondi/Cop",
            "Dessert/Cop",
            "Vino/10Cop",
        ],
        rows,
        {
            2: FMT_INT,
            3: FMT_INT,
            4: FMT_EURO,
            5: FMT_EURO,
            6: FMT_INT,
            **dict.fromkeys(range(7, 15), FMT_RATIO),
        },
    )


def add_sheet_trend(wb: Workbook, df_trend: pd.DataFrame) -> None:
    ws = wb.create_sheet("Trend Settimanale")
    rows = [
        [
            r["settimana"],
            r["coperti"],
            r["ricavi"],
            r["coperto_medio"],
            r["delta_coperti_pct"],
            r["delta_ricavi_pct"],
        ]
        for _, r in df_trend.iterrows()
    ]
    _write_sheet(
        ws,
        [
            "Settimana (lun-dom)",
            "Coperti",
            "Ricavi",
            "Coperto Medio",
            "Δ% Coperti",
            "Δ% Ricavi",
        ],
        rows,
        {1: FMT_INT, 2: FMT_EURO, 3: FMT_EURO, 4: FMT_PCT, 5: FMT_PCT},
    )
    ws.column_dimensions["A"].width = 32


def build_workbook(
    *,
    df_giorni: pd.DataFrame,
    df_trend: pd.DataFrame,
    df_coperti: pd.DataFrame,
    df_modello: pd.DataFrame,
    df_reparti: pd.DataFrame,
    df_menu: pd.DataFrame,
) -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)  # i fogli si creano tutti via add_sheet_*
    add_sheet_cruscotto(wb, df_modello)
    add_sheet_giornaliero(wb, df_giorni)
    add_sheet_coperti(wb, df_coperti)
    add_sheet_breakfast(wb, df_modello)
    add_sheet_consumi(wb, df_reparti, df_modello)
    add_sheet_menu_engineering(wb, df_menu)
    add_sheet_definizioni(wb)
    add_sheet_trend(wb, df_trend)
    return wb


def genera_report(output_path: str, days: int = 30, months: int = 13) -> str:
    from google.cloud import bigquery

    client = bigquery.Client(project=PROJECT)
    df_giorni = query_giornaliero(client, days)
    if df_giorni.empty:
        raise SystemExit("Nessun dato in v_fb_giornaliero — verifica la vista/ingest.")
    wb = build_workbook(
        df_giorni=df_giorni,
        df_trend=build_trend_settimanale(df_giorni),
        df_coperti=query_coperti_completi(client, days),
        df_modello=query_modello_mensile(client, months),
        df_reparti=query_consumi_reparto(client, months),
        df_menu=classifica_quadranti(query_menu_engineering(client, days)),
    )
    wb.save(output_path)
    log.info(
        "OK %s — giornaliero %s → %s (%d righe), modello mensile %d mesi",
        output_path,
        df_giorni["data"].min(),
        df_giorni["data"].max(),
        len(df_giorni),
        months,
    )
    return output_path


def main() -> None:
    ap = argparse.ArgumentParser(
        description="Modello Excel F&B per Feliciani (8 fogli)"
    )
    ap.add_argument(
        "--output",
        default=f"modello_fb_feliciani_{date.today():%Y%m%d}.xlsx",
        help="path Excel di output",
    )
    ap.add_argument(
        "--days",
        type=int,
        default=30,
        help="giorni per fogli giornalieri e menu engineering (default 30)",
    )
    ap.add_argument(
        "--months",
        type=int,
        default=13,
        help="mesi per Cruscotto/Breakfast/Consumi (default 13)",
    )
    args = ap.parse_args()
    logging.basicConfig(level=logging.INFO, format="%(levelname)s: %(message)s")
    genera_report(args.output, days=args.days, months=args.months)


if __name__ == "__main__":
    main()
