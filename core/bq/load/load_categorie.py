#!/usr/bin/env python3
"""
CATEGORIE ingest — loads canonical conto→(TipoCosto, Categoria_CE) mapping.

Source: "CATEGORIE" sheet from "Costi Ricavi 2025-2026 Budget.xlsx"
Target: d_categorie_conti in BigQuery (WRITE_TRUNCATE — dimension table)

This is the authoritative mapping for:
  - tipo_costo (IP, F, V, P, X)
  - categoria_ce (Ricavi, Acquisti, Costi Produttivi, Costo del Personale,
                   Costi Commerciali, Costi Amministrativi, Oneri Tributari,
                   Oneri Finanziari)

Usage:
    python -m core.bq.load.load_categorie --dry-run
    python -m core.bq.load.load_categorie \\
        --file "/path/to/Costi Ricavi 2025-2026 Budget.xlsx"
"""

from __future__ import annotations

import argparse
import csv
import logging
import sys
from pathlib import Path

try:
    import openpyxl

    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

try:
    from google.cloud import bigquery

    HAS_BQ = True
except ImportError:
    HAS_BQ = False

from core.bq.client import get_client
from core.config import PROJECT

BQ_TABLE = f"{PROJECT}.hotelops.d_categorie_conti"

DEFAULT_FILE = Path(
    "/Users/stefanodellapietra/Desktop/WORK/artifacts/"
    "Costi Ricavi 2025-2026 Budget.xlsx"
)

SHEET_NAME = "CATEGORIE"


def setup_logger() -> logging.Logger:
    log = logging.getLogger("ingest_categorie")
    log.setLevel(logging.DEBUG)
    if not log.handlers:
        h = logging.StreamHandler(sys.stdout)
        h.setFormatter(logging.Formatter("%(levelname)s  %(message)s"))
        log.addHandler(h)
    return log


def parse_categorie(filepath: Path, logger: logging.Logger) -> list[dict]:
    """Parse CATEGORIE sheet → list of dicts."""
    if not HAS_OPENPYXL:
        logger.error("openpyxl non installato: pip install openpyxl")
        sys.exit(1)

    wb = openpyxl.load_workbook(str(filepath), read_only=True, data_only=True)

    if SHEET_NAME not in wb.sheetnames:
        logger.error(f"Foglio '{SHEET_NAME}' non trovato in {filepath.name}")
        logger.info(f"  Fogli disponibili: {wb.sheetnames}")
        wb.close()
        return []

    ws = wb[SHEET_NAME]
    rows_raw = list(ws.iter_rows(min_row=2, values_only=True))  # Skip header
    wb.close()

    records = []
    skipped = 0

    for row in rows_raw:
        if not row or not row[0]:
            skipped += 1
            continue

        conto = str(row[0]).strip()
        if not conto or conto == "None":
            skipped += 1
            continue

        descrizione = str(row[1]).strip() if row[1] else ""
        tipo_costo = str(row[2]).strip() if row[2] else ""
        categoria_ce = str(row[3]).strip() if row[3] else ""
        anno_inizio = int(row[4]) if row[4] else None
        anno_fine = int(row[5]) if row[5] else None

        records.append(
            {
                "codice_conto": conto,
                "descrizione": descrizione,
                "tipo_costo": tipo_costo,
                "categoria_ce": categoria_ce,
                "anno_inizio": anno_inizio,
                "anno_fine": anno_fine,
            }
        )

    logger.info(f"  CATEGORIE: {len(records)} righe parsate (skip {skipped})")
    return records


def quality_summary(rows: list[dict], logger: logging.Logger) -> None:
    """Print summary by categoria_ce."""
    by_cat: dict[str, int] = {}
    by_tipo: dict[str, int] = {}
    for r in rows:
        cat = r["categoria_ce"] or "?"
        tipo = r["tipo_costo"] or "?"
        by_cat[cat] = by_cat.get(cat, 0) + 1
        by_tipo[tipo] = by_tipo.get(tipo, 0) + 1

    logger.info("═" * 50)
    logger.info("  CATEGORIE SUMMARY")
    logger.info("═" * 50)
    for cat, n in sorted(by_cat.items()):
        logger.info(f"  {cat:<28s}: {n:>4d} conti")
    logger.info(f"  {'─' * 38}")
    for tipo, n in sorted(by_tipo.items()):
        logger.info(f"  TipoCosto {tipo:<4s}: {n:>4d}")
    logger.info("═" * 50)


FIELDS = [
    "codice_conto",
    "descrizione",
    "tipo_costo",
    "categoria_ce",
    "anno_inizio",
    "anno_fine",
]


def dump_csv(rows: list[dict], path: Path, logger: logging.Logger) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=FIELDS, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)
    logger.info(f"  CSV: {path}  ({len(rows)} righe)")


BQ_SCHEMA = (
    [
        bigquery.SchemaField("codice_conto", "STRING", mode="REQUIRED"),
        bigquery.SchemaField("descrizione", "STRING"),
        bigquery.SchemaField("tipo_costo", "STRING"),
        bigquery.SchemaField("categoria_ce", "STRING"),
        bigquery.SchemaField("anno_inizio", "INTEGER"),
        bigquery.SchemaField("anno_fine", "INTEGER"),
    ]
    if HAS_BQ
    else []
)


def load_to_bq(rows: list[dict], logger: logging.Logger) -> None:
    if not rows:
        logger.warning("Nessuna riga da caricare")
        return

    bq_client = get_client()
    job = bq_client.load_table_from_json(
        rows,
        BQ_TABLE,
        job_config=bigquery.LoadJobConfig(
            schema=BQ_SCHEMA,
            write_disposition=bigquery.WriteDisposition.WRITE_TRUNCATE,
        ),
    )
    job.result()
    if job.errors:
        logger.error(f"BQ errors: {job.errors}")
        sys.exit(1)
    logger.info(f"  {BQ_TABLE}: {len(rows)} righe caricate (WRITE_TRUNCATE)")


def main() -> None:
    parser = argparse.ArgumentParser(description="CATEGORIE → d_categorie_conti")
    parser.add_argument(
        "--file",
        default=str(DEFAULT_FILE),
        help="Path al file XLSX con foglio CATEGORIE",
    )
    parser.add_argument(
        "--dry-run", action="store_true", help="Parse + CSV, no BQ write"
    )
    parser.add_argument(
        "--output-dir", default="output", help="Directory per CSV output in dry-run"
    )
    args = parser.parse_args()

    logger = setup_logger()
    filepath = Path(args.file)
    if not filepath.exists():
        logger.error(f"File non trovato: {filepath}")
        sys.exit(1)

    logger.info(f"File: {filepath.name}")
    rows = parse_categorie(filepath, logger)
    if not rows:
        logger.error("Nessuna riga estratta")
        sys.exit(1)

    quality_summary(rows, logger)

    if args.dry_run:
        dump_csv(rows, Path(args.output_dir) / "d_categorie_conti.csv", logger)
        logger.info("DRY RUN — nessuna scrittura su BQ.")
        return

    if not HAS_BQ:
        logger.error("google-cloud-bigquery non installato")
        sys.exit(1)

    load_to_bq(rows, logger)
    logger.info("✓ DONE")


if __name__ == "__main__":
    main()
